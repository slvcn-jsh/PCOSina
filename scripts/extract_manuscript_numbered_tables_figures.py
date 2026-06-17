from __future__ import annotations

from io import BytesIO
import re
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "[Manuscript] Quadrant_Revised_Ch 1-6_S2526T3 (3).docx"
OUT_DIR = ROOT / "docs" / "defense"
FIGURE_ASSET_DIR = OUT_DIR / "PCOSINA_Manuscript_Figures_1_to_26_Study_Copy_2026-06-09.assets"
TABLES_DOCX = OUT_DIR / "PCOSINA_Manuscript_Tables_1_to_58_Study_Copy_2026-06-09.docx"
FIGURES_DOCX = OUT_DIR / "PCOSINA_Manuscript_Figures_1_to_26_Study_Copy_2026-06-09.docx"
DOWNLOADS = Path.home() / "Downloads"

TABLE_CAPTION_RE = re.compile(r"^\s*Table\s+(\d+)\.\s*(.*)$", re.IGNORECASE)
FIGURE_CAPTION_RE = re.compile(r"^\s*Figure\s+(\d+)\.\s*(.*)$", re.IGNORECASE)


def paragraph_text_from_element(element: Any) -> str:
    return " ".join("".join(element.xpath(".//w:t/text()")).split())


def source_blocks(document: Document) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph_index += 1
            blocks.append(
                {
                    "type": "paragraph",
                    "paragraph_index": paragraph_index,
                    "element": child,
                    "text": paragraph_text_from_element(child),
                    "blips": child.xpath('.//*[local-name()="blip"]'),
                }
            )
        elif child.tag == qn("w:tbl"):
            table_index += 1
            blocks.append(
                {
                    "type": "table",
                    "table_index": table_index,
                    "element": child,
                    "text": "",
                    "blips": child.xpath('.//*[local-name()="blip"]'),
                }
            )
    return blocks


def set_landscape(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)


def setup_styles(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(12)


def add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(subtitle)
    run.bold = True
    run.font.size = Pt(11)
    p = document.add_paragraph()
    p.add_run("Source manuscript: ").bold = True
    p.add_run(str(SOURCE.relative_to(ROOT)).replace("\\", "/"))
    p = document.add_paragraph()
    p.add_run("Scope: ").bold = True
    p.add_run("Only numbered manuscript items are included. Appendix images and extra Word table objects used for equations are excluded.")


def source_cell_text(cell: Any) -> str:
    paragraphs = [" ".join(paragraph.text.split()) for paragraph in cell.paragraphs]
    return "\n".join(text for text in paragraphs if text)


def copy_table_as_visible_grid(target_document: Document, source_table: Table) -> None:
    row_count = len(source_table.rows)
    col_count = max((len(row.cells) for row in source_table.rows), default=0)
    if row_count <= 0 or col_count <= 0:
        target_document.add_paragraph("[Empty table object in source manuscript]")
        return

    target_table = target_document.add_table(rows=row_count, cols=col_count)
    target_table.style = "Table Grid"
    target_table.autofit = True
    for row_index, source_row in enumerate(source_table.rows):
        for col_index in range(col_count):
            target_cell = target_table.cell(row_index, col_index)
            if col_index < len(source_row.cells):
                target_cell.text = source_cell_text(source_row.cells[col_index])
            else:
                target_cell.text = ""
            for paragraph in target_cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)


def table_entries(source_document: Document, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    used_table_indexes: set[int] = set()
    for block_index, block in enumerate(blocks):
        if block["type"] != "paragraph":
            continue
        match = TABLE_CAPTION_RE.match(str(block["text"]))
        if not match:
            continue
        number = int(match.group(1))
        if number < 1 or number > 58:
            continue
        for next_index in range(block_index + 1, len(blocks)):
            next_block = blocks[next_index]
            if next_block["type"] == "paragraph":
                next_match = TABLE_CAPTION_RE.match(str(next_block["text"]))
                if next_match and int(next_match.group(1)) != number:
                    break
            if next_block["type"] == "table" and next_block["table_index"] not in used_table_indexes:
                used_table_indexes.add(next_block["table_index"])
                entries.append(
                    {
                        "number": number,
                        "caption": str(block["text"]),
                        "caption_paragraph_index": block["paragraph_index"],
                        "source_table_index": next_block["table_index"],
                        "table": Table(next_block["element"], source_document),
                    }
                )
                break
    return sorted(entries, key=lambda item: item["number"])


def build_tables_doc(source_document: Document, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    entries = table_entries(source_document, blocks)
    document = Document()
    set_landscape(document)
    setup_styles(document)
    add_title(
        document,
        "PCOSina Manuscript Tables 1 to 58",
        "Numbered manuscript tables only, organized for final-defense study",
    )
    p = document.add_paragraph()
    p.add_run("Total numbered tables included: ").bold = True
    p.add_run(str(len(entries)))
    document.add_heading("Table Index", level=1)
    for entry in entries:
        document.add_paragraph(f"Table {entry['number']}. {entry['caption']}")

    for entry in entries:
        document.add_page_break()
        document.add_heading(entry["caption"], level=1)
        p = document.add_paragraph()
        p.add_run("Source: ").bold = True
        p.add_run(
            f"caption paragraph {entry['caption_paragraph_index']}; "
            f"source Word table object {entry['source_table_index']}"
        )
        copy_table_as_visible_grid(document, entry["table"])

    document.save(TABLES_DOCX)
    return entries


def figure_image_blocks(blocks: list[dict[str, Any]], caption_block_index: int) -> list[dict[str, Any]]:
    caption_block = blocks[caption_block_index]
    if caption_block.get("blips"):
        return [caption_block]

    # Most manuscript figures place the image immediately before the caption.
    # Figure 10 has several blank paragraphs between the image and caption, so
    # the window is intentionally wider while still stopping at the previous
    # numbered figure caption.
    for index in range(caption_block_index - 1, max(-1, caption_block_index - 65), -1):
        block = blocks[index]
        if block["type"] == "paragraph" and FIGURE_CAPTION_RE.match(str(block["text"])):
            break
        if block["type"] == "paragraph" and block.get("blips"):
            return [block]

    for index in range(caption_block_index + 1, min(len(blocks), caption_block_index + 65)):
        block = blocks[index]
        if block["type"] == "paragraph" and FIGURE_CAPTION_RE.match(str(block["text"])):
            break
        if block["type"] == "paragraph" and block.get("blips"):
            return [block]
    return []


def figure_table_block(blocks: list[dict[str, Any]], caption_block_index: int) -> dict[str, Any] | None:
    # Figure 15 is a high-level pseudocode figure stored as a Word table object.
    for index in range(caption_block_index - 1, max(-1, caption_block_index - 12), -1):
        block = blocks[index]
        if block["type"] == "paragraph" and FIGURE_CAPTION_RE.match(str(block["text"])):
            break
        if block["type"] == "table":
            return block
    for index in range(caption_block_index + 1, min(len(blocks), caption_block_index + 12)):
        block = blocks[index]
        if block["type"] == "paragraph" and FIGURE_CAPTION_RE.match(str(block["text"])):
            break
        if block["type"] == "table":
            return block
    return None


def media_extension(part: Any) -> str:
    suffix = Path(str(getattr(part, "partname", ""))).suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg"}:
        return ".jpg" if suffix == ".jpeg" else suffix
    content_type = str(getattr(part, "content_type", "")).lower()
    if "jpeg" in content_type or "jpg" in content_type:
        return ".jpg"
    if "png" in content_type:
        return ".png"
    return ".bin"


def figure_entries(source_document: Document, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    FIGURE_ASSET_DIR.mkdir(parents=True, exist_ok=True)
    for existing in FIGURE_ASSET_DIR.glob("*"):
        if existing.is_file():
            existing.unlink()

    entries: list[dict[str, Any]] = []
    for block_index, block in enumerate(blocks):
        if block["type"] != "paragraph":
            continue
        match = FIGURE_CAPTION_RE.match(str(block["text"]))
        if not match:
            continue
        number = int(match.group(1))
        if number < 1 or number > 26:
            continue
        image_paths: list[Path] = []
        related_blocks = figure_image_blocks(blocks, block_index)
        image_count = 0
        for image_block in related_blocks:
            for blip in image_block.get("blips") or []:
                rid = blip.get(qn("r:embed"))
                if not rid:
                    continue
                part = source_document.part.related_parts.get(rid)
                if part is None:
                    continue
                image_count += 1
                suffix = media_extension(part)
                image_path = FIGURE_ASSET_DIR / f"figure_{number:02d}_{image_count:02d}{suffix}"
                image_path.write_bytes(part.blob)
                image_paths.append(image_path)
        source_table = None
        source_table_index = None
        if not image_paths:
            table_block = figure_table_block(blocks, block_index)
            if table_block is not None:
                source_table = Table(table_block["element"], source_document)
                source_table_index = table_block["table_index"]
        entries.append(
            {
                "number": number,
                "caption": str(block["text"]),
                "caption_paragraph_index": block["paragraph_index"],
                "image_paths": image_paths,
                "source_table": source_table,
                "source_table_index": source_table_index,
            }
        )
    return sorted(entries, key=lambda item: item["number"])


def build_figures_doc(source_document: Document, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    entries = figure_entries(source_document, blocks)
    document = Document()
    set_landscape(document)
    setup_styles(document)
    add_title(
        document,
        "PCOSina Manuscript Figures 1 to 26",
        "Numbered manuscript figures only, organized for final-defense study",
    )
    p = document.add_paragraph()
    p.add_run("Total numbered figures included: ").bold = True
    p.add_run(str(len(entries)))
    document.add_heading("Figure Index", level=1)
    for entry in entries:
        document.add_paragraph(f"Figure {entry['number']}. {entry['caption']}")

    for entry in entries:
        document.add_page_break()
        document.add_heading(entry["caption"], level=1)
        p = document.add_paragraph()
        p.add_run("Source: ").bold = True
        p.add_run(f"caption paragraph {entry['caption_paragraph_index']}")
        if not entry["image_paths"]:
            if entry.get("source_table") is not None:
                p = document.add_paragraph()
                p.add_run("Stored form: ").bold = True
                p.add_run(f"Word table object {entry['source_table_index']} in the source manuscript")
                copy_table_as_visible_grid(document, entry["source_table"])
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.add_run(entry["caption"]).italic = True
            else:
                document.add_paragraph("[No nearby embedded image or table-form figure found for this numbered figure caption.]")
            continue
        for image_path in entry["image_paths"]:
            document.add_picture(str(image_path), width=Inches(9.2))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(entry["caption"]).italic = True

    document.save(FIGURES_DOCX)
    return entries


def copy_to_downloads(paths: list[Path]) -> list[Path]:
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    copied: list[Path] = []
    for path in paths:
        destination = DOWNLOADS / path.name
        shutil.copy2(path, destination)
        copied.append(destination)
    return copied


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    source_document = Document(SOURCE)
    blocks = source_blocks(source_document)
    tables = build_tables_doc(source_document, blocks)
    figures = build_figures_doc(source_document, blocks)
    copied = copy_to_downloads([TABLES_DOCX, FIGURES_DOCX])
    missing_tables = [number for number in range(1, 59) if number not in {entry["number"] for entry in tables}]
    missing_figures = [number for number in range(1, 27) if number not in {entry["number"] for entry in figures}]
    figures_without_images = [entry["number"] for entry in figures if not entry["image_paths"]]
    figures_without_visual = [
        entry["number"]
        for entry in figures
        if not entry["image_paths"] and entry.get("source_table") is None
    ]
    print(f"source={SOURCE}")
    print(f"tables_doc={TABLES_DOCX}")
    print(f"figures_doc={FIGURES_DOCX}")
    print(f"tables_included={len(tables)}")
    print(f"figures_included={len(figures)}")
    print(f"missing_tables={missing_tables}")
    print(f"missing_figures={missing_figures}")
    print(f"figures_without_images={figures_without_images}")
    print(f"figures_without_visual={figures_without_visual}")
    print("downloads=" + ";".join(str(path) for path in copied))


if __name__ == "__main__":
    main()
