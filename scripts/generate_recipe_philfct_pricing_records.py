"""Generate PhilFCT and pricing DOCX records for every active recipe.

The output mirrors the worksheet process in app/PCOSINA_PhilFCT_TEMPLATE.docx:
recipe summary, ingredient serving list, PhilFCT source values, nutrition
computation, pricing worksheet, and totals cross-check.
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = ROOT / "backend"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from scripts.generate_current_catalog_evidence_doc import (
    NUTRIENT_KEYS,
    OFFICIAL_SOURCES,
    all_evidence_rows,
    float_or_zero,
    philfct_code_index,
    priced_meal_index,
    read_json,
    recipe_evidence_rows,
)
from scripts.export_budget_support_runtime_catalog import PHP_PER_KG


GENERATED_DATE = "2026-08-11"
TEMPLATE_PATH = ROOT / "app" / "PCOSINA_PhilFCT_TEMPLATE.docx"
RUNTIME_PATH = BACKEND_ROOT / "recipes.json"
PHILFCT_CACHE_PATH = BACKEND_ROOT / "seed_data" / "philfct_reference_cache.json"
PRICED_DRAFT_PATH = BACKEND_ROOT / "seed_data" / "pcosina_philfct_portioned_priced_catalog_v2_draft.json"
OUTPUT_DIR = ROOT / "docs" / "thesis_validation" / "03_ACTUAL_SYSTEM_DATA_EXPORTS" / "recipe_records"
MASTER_DOCX = OUTPUT_DIR / "PCOSINA_PhilFCT_Pricing_Recipe_Records_Master_Index.docx"
JSON_PATH = OUTPUT_DIR / "pcosina_recipe_philfct_pricing_records.json"
CSV_PATH = OUTPUT_DIR / "pcosina_recipe_philfct_pricing_evidence_rows.csv"
INDEX_CSV_PATH = OUTPUT_DIR / "pcosina_recipe_philfct_pricing_record_index.csv"


MEAL_ORDER = {"Breakfast": 1, "Lunch": 2, "Dinner": 3, "Universal": 4}
VOLUME_NAMES = {
    "Breakfast": "PCOSINA_PhilFCT_Pricing_Recipe_Records_Volume_01_Breakfast.docx",
    "Lunch": "PCOSINA_PhilFCT_Pricing_Recipe_Records_Volume_02_Lunch.docx",
    "Dinner": "PCOSINA_PhilFCT_Pricing_Recipe_Records_Volume_03_Dinner.docx",
    "Universal": "PCOSINA_PhilFCT_Pricing_Recipe_Records_Volume_04_Universal.docx",
}

COOKED_TO_MARKET_WEIGHT_FACTOR = {
    "cooked white rice": 0.34,
    "cooked munggo": 0.40,
    "cooked chickpeas": 0.45,
    "red kidney beans boiled": 0.45,
}

PANTRY_ALLOCATION_RULES = {
    "water": {
        "unit": "l assumed tap",
        "unit_price": 0.0,
        "confidence": "not_applicable",
        "source": "PCOSina assumption: household/tap water companion; no purchased bottled water charged.",
        "source_url": "",
        "rule_id": "pcosina_water_zero_cost",
    },
    "salt": {
        "unit": "kg pantry allocation",
        "unit_price": 34.44,
        "confidence": "medium",
        "source": "DA proxy: RFO I Daily Price Index, July 16, 2026; salt PHP 34.44/kg; validate against local store price.",
        "source_url": "https://ilocos.da.gov.ph/wp-content/uploads/July-16-2026-DPI-Ilocos-Region.pdf",
        "rule_id": "da_rfo1_salt_proxy_allocation",
        "default_portion_g": 1.0,
        "tsp_g": 6.0,
        "tbsp_g": 18.0,
    },
    "pepper": {
        "unit": "kg pantry allocation",
        "unit_price": 650.0,
        "confidence": "low",
        "source": "Non-DA audit estimate: black pepper pantry allocation; validate against local store price.",
        "source_url": "https://www.dti.gov.ph/konsyumer/e-presyo/",
        "rule_id": "reviewed_pepper_allocation",
        "default_portion_g": 0.25,
        "tsp_g": 2.3,
        "tbsp_g": 6.9,
    },
    "bay_leaf": {
        "unit": "kg pantry allocation",
        "unit_price": 650.0,
        "confidence": "low",
        "source": "Non-DA audit estimate: bay leaf pantry allocation; validate against local store price.",
        "source_url": "https://www.dti.gov.ph/konsyumer/e-presyo/",
        "rule_id": "reviewed_bay_leaf_allocation",
        "default_portion_g": 0.125,
        "piece_g": 0.5,
    },
    "chili_spice": {
        "unit": "kg pantry allocation",
        "unit_price": 650.0,
        "confidence": "low",
        "source": "PCOSina pantry spice fallback pending local validation.",
        "source_url": "https://www.dti.gov.ph/konsyumer/e-presyo/",
        "rule_id": "pcosina_chili_spice_allocation",
        "default_portion_g": 0.5,
        "tsp_g": 2.0,
        "tbsp_g": 6.0,
        "piece_g": 1.0,
    },
    "seasoning_mix": {
        "unit": "kg pantry allocation",
        "unit_price": 1000.0,
        "confidence": "low",
        "source": "PCOSina seasoning fallback pending local validation.",
        "source_url": "https://www.dti.gov.ph/konsyumer/e-presyo/",
        "rule_id": "pcosina_seasoning_mix_allocation",
        "default_portion_g": 2.0,
        "tsp_g": 4.0,
        "tbsp_g": 12.0,
    },
    "frying_oil": {
        "unit": "kg pantry allocation",
        "unit_price": 150.0,
        "confidence": "low",
        "source": "DA proxy: RFO I Daily Price Index, July 16, 2026; palm olein PHP 138/L converted with 0.92 kg/L assumption; validate frying absorption.",
        "source_url": "https://ilocos.da.gov.ph/wp-content/uploads/July-16-2026-DPI-Ilocos-Region.pdf",
        "rule_id": "pcosina_frying_oil_allocation",
        "default_portion_g": 10.0,
        "tsp_g": 4.5,
        "tbsp_g": 13.5,
    },
}


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    return " ".join(text.split())


def short_text(value: Any, limit: int = 80) -> str:
    text = safe_text(value)
    if len(text) <= limit:
        return text
    return text[: limit - 3].rstrip() + "..."


def number(value: Any, digits: int = 2) -> str:
    try:
        raw = float(value)
    except (TypeError, ValueError):
        return ""
    if raw.is_integer():
        return str(int(raw))
    return f"{raw:.{digits}f}"


def money(value: Any) -> str:
    try:
        return f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return ""


def nutrition(recipe: dict[str, Any], key: str) -> Any:
    nutrition_map = recipe.get("nutrition") if isinstance(recipe.get("nutrition"), dict) else {}
    aliases = {
        "calories": ("calories", "kcal"),
        "protein_g": ("protein_g", "proteinGrams"),
        "carbs_g": ("carbs_g", "carbsGrams"),
        "fat_g": ("fat_g", "fatsGrams"),
        "fiber_g": ("fiber_g", "fiberGrams"),
    }
    for candidate in aliases.get(key, (key,)):
        if candidate in nutrition_map:
            return nutrition_map.get(candidate)
        if candidate in recipe:
            return recipe.get(candidate)
    return ""


def recipe_id(recipe: dict[str, Any]) -> str:
    return safe_text(recipe.get("id"))


def recipe_name(recipe: dict[str, Any]) -> str:
    return safe_text(recipe.get("name") or recipe.get("title") or recipe_id(recipe))


def meal_type(recipe: dict[str, Any]) -> str:
    return safe_text(recipe.get("mealType") or "Universal") or "Universal"


def git_commit_short() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"],
            cwd=str(ROOT),
            text=True,
            stderr=subprocess.DEVNULL,
        ).strip()
    except Exception:
        return "unknown"


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell: Any, width_inches: float) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.33)
    section.right_margin = Inches(0.33)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(7.4)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(16)
    styles["Title"].font.bold = True
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = RGBColor(31, 78, 121)


def add_table(
    doc: Document,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    *,
    widths: list[float] | None = None,
    font_size: float = 6.2,
    header_fill: str = "1F4E79",
) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, heading in enumerate(headers):
        cell = header_row.cells[idx]
        cell.text = safe_text(heading)
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[idx])
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)

    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.text = safe_text(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                set_cell_width(cell, widths[idx])
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
    doc.add_paragraph()
    return table


def add_title(doc: Document, title_text: str, subtitle: str, meta: str) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(title_text)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.add_run(subtitle)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(meta)


def row_status(delta: float, tolerance: float) -> str:
    return "OK" if abs(delta) <= tolerance else "Review"


def is_rnd_reviewed(recipe: dict[str, Any]) -> bool:
    return safe_text(recipe.get("nutritionReviewStatus")) == "rnd_representative_meal_reviewed"


def is_nonnutritive_row(row: dict[str, Any]) -> bool:
    ingredient = safe_text(row.get("ingredient")).lower()
    note = safe_text(row.get("nutrition_note")).lower()
    return ingredient in {"water"} or "ignored non-nutritive" in note or "excluded from nutrition totals" in note


def is_pcosina_fallback_price(row: dict[str, Any]) -> bool:
    status = safe_text(row.get("pricing_status"))
    source = safe_text(row.get("market_source"))
    return "PCOSina" in status or "PCOSina" in source


def row_text(row: dict[str, Any]) -> str:
    return f"{safe_text(row.get('ingredient'))} {safe_text(row.get('source_text'))}".lower()


def pantry_rule_key(row: dict[str, Any]) -> str:
    text = row_text(row)
    if "water spinach" in text or "kangkong" in text:
        text_without_water_spinach = text.replace("water spinach", "").replace("kangkong", "")
    else:
        text_without_water_spinach = text
    if re.search(r"\bwater\b|tubig", text_without_water_spinach):
        return "water"
    if any(term in text for term in ("oil for frying", "oil for deep", "oil as needed", "peanut oil for frying")):
        return "frying_oil"
    if any(term in text for term in ("bay leaf", "bay leaves", "laurel", "dahon ng laurel")):
        return "bay_leaf"
    has_salt = bool(re.search(r"\bsalt\b|\basin\b", text))
    has_pepper = any(term in text for term in ("ground black pepper", "black pepper", "peppercorn", "paminta")) and "bell pepper" not in text
    if has_salt and has_pepper:
        return "salt_pepper"
    if has_pepper:
        return "pepper"
    if has_salt:
        return "salt"
    if any(term in text for term in ("chili flakes", "dried chili", "dried chilies", "sili", "paprika", "five-spice", "five spice")):
        return "chili_spice"
    if any(term in text for term in ("seasoning", "magic sarap", "bouillon", "pork cube", "beef cube", "chicken cube")):
        return "seasoning_mix"
    return ""


def source_servings(row: dict[str, Any]) -> float:
    value = float_or_zero(row.get("source_servings"))
    return value if value > 0 else 1.0


def _parse_first_amount(text: str) -> float | None:
    normalized = safe_text(text).lower()
    normalized = normalized.replace("½", " 1/2 ").replace("¼", " 1/4 ").replace("¾", " 3/4 ")
    match = re.search(r"(\d+(?:\.\d+)?\s+\d+/\d+|\d+/\d+|\d+(?:\.\d+)?)", normalized)
    if not match:
        return None
    raw = match.group(1).strip()
    if " " in raw and "/" in raw:
        whole, fraction = raw.split(maxsplit=1)
        num, den = fraction.split("/")
        return float(whole) + float(num) / float(den)
    if "/" in raw:
        num, den = raw.split("/")
        return float(num) / float(den)
    return float(raw)


def estimated_pantry_portion_g(row: dict[str, Any], rule: dict[str, Any]) -> float:
    text = row_text(row)
    amount = _parse_first_amount(text)
    servings = source_servings(row)
    if amount is not None:
        if re.search(r"\b(tsp|teaspoon|teaspoons)\b", text):
            return round(amount * float(rule.get("tsp_g") or rule.get("default_portion_g") or 0) / servings, 4)
        if re.search(r"\b(tbsp|tablespoon|tablespoons)\b", text):
            return round(amount * float(rule.get("tbsp_g") or rule.get("default_portion_g") or 0) / servings, 4)
        if re.search(r"\b(piece|pieces|pc|pcs|leaf|leaves)\b", text):
            return round(amount * float(rule.get("piece_g") or rule.get("default_portion_g") or 0) / servings, 4)
        if re.search(r"\b(g|gram|grams)\b", text):
            return round(amount / servings, 4)
    return float(rule.get("default_portion_g") or 0.0)


def allocated_price_parts(row: dict[str, Any]) -> list[dict[str, Any]]:
    if float_or_zero(row.get("cost_php")) > 0:
        return []
    key = pantry_rule_key(row)
    if not key:
        return []
    if key == "salt_pepper":
        return [
            {"key": "salt", **PANTRY_ALLOCATION_RULES["salt"]},
            {"key": "pepper", **PANTRY_ALLOCATION_RULES["pepper"]},
        ]
    return [{"key": key, **PANTRY_ALLOCATION_RULES[key]}]


def allocated_price(row: dict[str, Any]) -> dict[str, Any]:
    parts = allocated_price_parts(row)
    if not parts:
        return {}
    line_items = []
    total = 0.0
    for part in parts:
        if part["key"] == "water":
            portion_g = 0.0
            cost = 0.0
        else:
            portion_g = estimated_pantry_portion_g(row, part)
            cost = portion_g / 1000 * float(part["unit_price"])
        total += cost
        line_items.append(
            {
                "name": part["key"],
                "portion_g": round(portion_g, 4),
                "unit": part["unit"],
                "unit_price": part["unit_price"],
                "confidence": part["confidence"],
                "source": part["source"],
                "source_url": part["source_url"],
                "rule_id": part["rule_id"],
                "cost_php": round(cost, 4),
            }
        )
    confidence_order = {"not_applicable": 0, "low": 1, "medium": 2, "high": 3}
    confidence = min(
        (item["confidence"] for item in line_items),
        key=lambda item: confidence_order.get(item, 1),
    )
    unit_price_text = " + ".join(
        f"{money(item['unit_price'])}/kg" if "kg" in item["unit"] else f"{money(item['unit_price'])}/{item['unit']}"
        for item in line_items
    )
    unit_text = " + ".join(dict.fromkeys(item["unit"] for item in line_items))
    source_text = " | ".join(dict.fromkeys(item["source"] for item in line_items))
    source_url = " | ".join(dict.fromkeys(item["source_url"] for item in line_items if item["source_url"]))
    rule_id = " + ".join(item["rule_id"] for item in line_items)
    note = "; ".join(
        f"{item['name']}: {item['portion_g']:g} g x PHP {item['unit_price']:g}/kg = PHP {item['cost_php']:.2f}"
        if item["cost_php"] > 0
        else f"{item['name']}: PHP 0 assumption"
        for item in line_items
    )
    return {
        "cost_php": round(total, 2),
        "market_unit": unit_text,
        "unit_price_php": unit_price_text,
        "price_confidence": confidence,
        "market_source": source_text,
        "source_url": source_url,
        "price_rule_id": rule_id,
        "pricing_status": "pantry_spice_portion_allocation" if total > 0 else "zero_cost_companion",
        "note": f"Allocated pantry/non-nutritive pricing for purchasing clarity: {note}. Validate or replace during review.",
        "line_items": line_items,
    }


def effective_cost_php(row: dict[str, Any]) -> float:
    cost = float_or_zero(row.get("cost_php"))
    if cost > 0:
        return cost
    return float_or_zero(allocated_price(row).get("cost_php"))


def display_nutrition_value(row: dict[str, Any], section: str, key: str) -> Any:
    values = row.get(section) or {}
    value = values.get(key)
    if value in (None, "") and is_nonnutritive_row(row):
        return 0.0
    return value


def fallback_price_basis(row: dict[str, Any]) -> dict[str, str]:
    """Expose embedded PCOSina fallback pricing so reviewers can replace it."""
    if not is_pcosina_fallback_price(row) or float_or_zero(row.get("cost_php")) <= 0:
        return {"unit": "", "unit_price": "", "confidence": "", "note": ""}
    ingredient = safe_text(row.get("ingredient")).lower()
    portion_g = float_or_zero(row.get("portion_g"))
    unit_price = PHP_PER_KG.get(ingredient)
    if unit_price is None:
        if portion_g <= 0:
            return {
                "unit": "fallback estimate",
                "unit_price": "",
                "confidence": "low",
                "note": "Embedded PCOSina fallback cost; market unit/source still needs reviewer input.",
            }
        derived = float_or_zero(row.get("cost_php")) / (portion_g / 1000)
        return {
            "unit": "kg consumed fallback",
            "unit_price": money(derived),
            "confidence": "low",
            "note": "Derived from embedded consumed-portion cost; replace with sourced market unit and price.",
        }

    factor = COOKED_TO_MARKET_WEIGHT_FACTOR.get(ingredient, 1.0)
    if factor != 1.0:
        note = (
            f"Fallback basis: {portion_g:g} g consumed x {factor:g} market-weight factor "
            f"x PHP {unit_price:g}/kg; replace with sourced market price."
        )
    else:
        note = f"Fallback basis: consumed portion x PHP {unit_price:g}/kg; replace with sourced market price."
    return {
        "unit": "kg fallback",
        "unit_price": money(unit_price),
        "confidence": "low",
        "note": note,
    }


def display_market_unit(row: dict[str, Any]) -> str:
    allocated = allocated_price(row)
    return safe_text(row.get("market_unit")) or fallback_price_basis(row)["unit"] or safe_text(allocated.get("market_unit"))


def display_unit_price(row: dict[str, Any]) -> str:
    allocated = allocated_price(row)
    return safe_text(row.get("unit_price_php")) or fallback_price_basis(row)["unit_price"] or safe_text(allocated.get("unit_price_php"))


def display_price_confidence(row: dict[str, Any]) -> str:
    allocated = allocated_price(row)
    return safe_text(row.get("price_confidence")) or fallback_price_basis(row)["confidence"] or safe_text(allocated.get("price_confidence"))


def price_basis_note(row: dict[str, Any]) -> str:
    return fallback_price_basis(row)["note"] or safe_text(allocated_price(row).get("note"))


def display_price_rule_id(row: dict[str, Any]) -> str:
    return safe_text(row.get("price_rule_id")) or safe_text(allocated_price(row).get("price_rule_id"))


def display_market_source(row: dict[str, Any]) -> str:
    return safe_text(row.get("market_source")) or safe_text(allocated_price(row).get("market_source"))


def display_price_source_url(row: dict[str, Any]) -> str:
    return safe_text(row.get("price_source_url")) or safe_text(allocated_price(row).get("source_url"))


def display_pricing_status(row: dict[str, Any]) -> str:
    return safe_text(row.get("pricing_status")) or safe_text(allocated_price(row).get("pricing_status"))


def has_cost_value(row: dict[str, Any]) -> bool:
    try:
        float(row.get("cost_php"))
        return True
    except (TypeError, ValueError):
        return False


def has_effective_cost_value(row: dict[str, Any]) -> bool:
    return has_cost_value(row) or bool(allocated_price(row))


def evidence_issue_flags(row: dict[str, Any]) -> list[str]:
    flags: list[str] = []
    portion = row.get("portion_g")
    included = safe_text(row.get("nutrition_included"))
    code_or_name = source_name_code(row)
    cost_value = effective_cost_php(row)
    nonnutritive = is_nonnutritive_row(row)
    allocated = allocated_price(row)

    if portion in (None, "") and not nonnutritive:
        flags.append("Missing portion")
    if included != "Yes" and not nonnutritive:
        flags.append("Nutrition not included")
    if included == "Yes" and not code_or_name:
        flags.append("Missing PhilFCT source")
    if included == "Yes" and not safe_text(row.get("philfct_code")):
        flags.append("Confirm PhilFCT code")
    if safe_text(row.get("row_type")) in {"Generated/RND component", "Complete-plate add-on"}:
        flags.append("Verify generated component")
    if safe_text(row.get("parse_status")) in {
        "estimated_default_portion",
        "estimated_garnish_portion",
        "parsed_range_estimate",
        "parsed_range_piece_estimate",
        "parsed_embedded_weight_estimate",
        "parsed_embedded_weight_range",
        "parsed_embedded_each_weight",
    }:
        flags.append("Validate inferred portion")
    if not has_cost_value(row):
        flags.append("Missing cost")
    if cost_value > 0 and not (
        display_price_rule_id(row)
        or display_price_source_url(row)
        or is_pcosina_fallback_price(row)
    ):
        flags.append("Missing price source")
    if cost_value > 0 and (is_pcosina_fallback_price(row) or allocated) and not (
        safe_text(row.get("price_rule_id")) or safe_text(row.get("price_source_url"))
    ):
        flags.append("Validate fallback price source")
    if allocated and cost_value > 0:
        flags.append("Validate pantry price allocation")
    if nonnutritive and not allocated and safe_text(row.get("ingredient")).lower() != "water" and float_or_zero(row.get("cost_php")) == 0:
        flags.append("Review ignored row classification")
    if cost_value > 0 and display_price_confidence(row).lower() not in {"high", "medium"}:
        flags.append("Check price confidence")

    return sorted(set(flags))


def evidence_check_status(row: dict[str, Any]) -> str:
    flags = evidence_issue_flags(row)
    if not flags:
        return "Ready for reviewer confirmation"
    if any(flag in flags for flag in ("Missing portion", "Nutrition not included", "Missing PhilFCT source", "Missing cost", "Missing price source")):
        return "Needs data correction/check"
    return "Needs cross-check"


def evidence_action_needed(row: dict[str, Any]) -> str:
    flags = evidence_issue_flags(row)
    if not flags:
        return "Confirm row and fill reviewer/date."
    return "Review/complete: " + "; ".join(flag.lower() for flag in flags) + "."


def recipe_issue_flags(recipe: dict[str, Any], rows: list[dict[str, Any]]) -> list[str]:
    flags: Counter[str] = Counter()
    if not is_rnd_reviewed(recipe):
        flags["Needs final recipe sign-off"] += 1
    if safe_text(recipe.get("nutritionConfidence")).lower() != "high":
        flags["Medium nutrition confidence"] += 1
    if safe_text(recipe.get("nutritionReviewStatus")) == "source_mapped_needs_final_review":
        flags["Source mapping needs final review"] += 1
    if any(check[-1] != "OK" for check in cross_check_rows(recipe, rows)):
        flags["Computed total mismatch"] += 1
    for row in rows:
        for flag in evidence_issue_flags(row):
            flags[flag] += 1
    return [f"{flag} ({count})" if count > 1 else flag for flag, count in flags.most_common()]


def recipe_priority(recipe: dict[str, Any], rows: list[dict[str, Any]]) -> str:
    flat_flags = "; ".join(recipe_issue_flags(recipe, rows))
    high_markers = (
        "Computed total mismatch",
        "Missing portion",
        "Nutrition not included",
        "Missing PhilFCT source",
        "Missing cost",
        "Missing price source",
        "Review ignored row classification",
    )
    if any(marker in flat_flags for marker in high_markers):
        return "High"
    if "Medium nutrition confidence" in flat_flags or "Needs final recipe sign-off" in flat_flags:
        return "Medium"
    return "Low"


def recipe_team_status(recipe: dict[str, Any], rows: list[dict[str, Any]]) -> str:
    priority = recipe_priority(recipe, rows)
    if priority == "High":
        return "Needs correction/check"
    if is_rnd_reviewed(recipe):
        return "RND-reviewed sample; spot-check only"
    return "Needs final team/RND confirmation"


def recipe_next_action(recipe: dict[str, Any], rows: list[dict[str, Any]]) -> str:
    flags = recipe_issue_flags(recipe, rows)
    if not flags:
        return "Fill reviewer/date and mark confirmed."
    if recipe_priority(recipe, rows) == "High":
        return "Resolve high-priority ingredient/source gaps, then re-run generator."
    return "Verify source mapping, pricing basis, and serving realism; fill reviewer/date."


def recipe_review_tracker_rows(recipe: dict[str, Any], rows: list[dict[str, Any]]) -> list[list[Any]]:
    checks = cross_check_rows(recipe, rows)
    calculation_status = "OK within rounding" if all(check[-1] == "OK" for check in checks) else "Review computed deltas"
    price_flags = Counter()
    philfct_flags = Counter()
    serving_flags = Counter()
    for row in rows:
        row_flags = evidence_issue_flags(row)
        for flag in row_flags:
            if "price" in flag.lower() or "cost" in flag.lower():
                price_flags[flag] += 1
            if "PhilFCT" in flag or "Nutrition" in flag:
                philfct_flags[flag] += 1
            if "portion" in flag.lower() or "component" in flag.lower():
                serving_flags[flag] += 1
    return [
        [
            "Overall recipe status",
            recipe_team_status(recipe, rows),
            recipe_next_action(recipe, rows),
            "",
            "",
            "",
        ],
        [
            "Serving/portion realism",
            "Needs check" if serving_flags else "Ready for confirmation",
            "Confirm edible one-person portions and update notes if adjusted.",
            "",
            "",
            "",
        ],
        [
            "PhilFCT source mapping",
            "Needs check" if philfct_flags or not is_rnd_reviewed(recipe) else "RND reviewed",
            "Confirm exact PhilFCT food name/code and per-100g values.",
            "",
            "",
            "",
        ],
        [
            "Nutrition arithmetic",
            calculation_status,
            "If values are edited, recompute totals and compare against runtime values.",
            "",
            "",
            "",
        ],
        [
            "Pricing source/basis",
            "Needs check" if price_flags else "Ready for confirmation",
            "Confirm source URL, unit price, market unit, and portion-cost formula.",
            "",
            "",
            "",
        ],
        [
            "Final sign-off",
            "Signed/validated" if is_rnd_reviewed(recipe) else "Pending",
            "Fill reviewer, date, and correction notes before importing back to source data.",
            "",
            "",
            "",
        ],
    ]


def totals_for_rows(rows: list[dict[str, Any]]) -> dict[str, float]:
    sums = {key: 0.0 for key in NUTRIENT_KEYS}
    sums["cost_php"] = 0.0
    for row in rows:
        for key in NUTRIENT_KEYS:
            sums[key] += float_or_zero(display_nutrition_value(row, "computed", key))
        sums["cost_php"] += effective_cost_php(row)
    return sums


def cross_check_rows(recipe: dict[str, Any], rows: list[dict[str, Any]]) -> list[list[Any]]:
    sums = totals_for_rows(rows)
    output = []
    labels = [
        ("Energy", "calories", "kcal", 0.55),
        ("Protein", "protein_g", "g", 0.55),
        ("Carbohydrates", "carbs_g", "g", 0.55),
        ("Fat", "fat_g", "g", 0.55),
        ("Fiber", "fiber_g", "g", 0.55),
    ]
    for label, key, unit, tolerance in labels:
        runtime_total = float_or_zero(nutrition(recipe, key))
        computed_total = float_or_zero(sums.get(key))
        delta = round(computed_total - runtime_total, 2)
        output.append(
            [
                label,
                number(runtime_total, 2),
                number(computed_total, 2),
                number(delta, 2),
                unit,
                row_status(delta, tolerance),
            ]
        )
    runtime_cost = float_or_zero(recipe.get("estimatedCostPhp"))
    computed_cost = float_or_zero(sums.get("cost_php"))
    cost_delta = round(computed_cost - runtime_cost, 2)
    output.append(
        [
            "Cost",
            money(runtime_cost),
            money(computed_cost),
            money(cost_delta),
            "PHP",
            row_status(cost_delta, 0.05),
        ]
    )
    return output


def record_review_status(recipe: dict[str, Any]) -> str:
    status = safe_text(recipe.get("nutritionReviewStatus"))
    confidence = safe_text(recipe.get("nutritionConfidence"))
    if status == "rnd_representative_meal_reviewed":
        return "RND representative meal reviewed"
    if confidence == "high":
        return "Source mapped; final human review still required"
    return "Source mapped with medium confidence; final human review required"


def component_note(recipe: dict[str, Any]) -> str:
    names = [safe_text(item.get("name")) for item in recipe.get("ingredients") or [] if safe_text(item.get("name"))]
    lower_names = [name.lower() for name in names]
    go = [name for name in names if any(token in name.lower() for token in ("rice", "noodle", "pasta", "bread", "potato", "corn", "oat"))]
    grow = [
        name
        for name in names
        if any(
            token in name.lower()
            for token in ("chicken", "fish", "bangus", "tilapia", "tuna", "sardine", "shrimp", "egg", "tofu", "munggo", "bean", "pork", "beef", "peanut")
        )
    ]
    fat = [name for name in names if any(token in name.lower() for token in ("oil", "coconut milk", "peanut"))]
    glow = [name for name in names if name not in set(go + grow + fat)]
    pieces = []
    if go:
        pieces.append("Go: " + ", ".join(go[:4]))
    if grow:
        pieces.append("Grow: " + ", ".join(grow[:4]))
    if glow:
        pieces.append("Glow/other: " + ", ".join(glow[:6]))
    if fat:
        pieces.append("Fat: " + ", ".join(fat[:3]))
    if not pieces and lower_names:
        pieces.append("Ingredients listed for RND review.")
    return "; ".join(pieces) or "No ingredient rows found."


def source_name_code(row: dict[str, Any]) -> str:
    code = safe_text(row.get("philfct_code"))
    name = safe_text(row.get("philfct_name"))
    if code and name:
        return f"{name} / {code}"
    return name or code


def source_date_text(row: dict[str, Any]) -> str:
    source = display_market_source(row)
    url = display_price_source_url(row)
    status = display_pricing_status(row)
    basis = price_basis_note(row)
    if source and url:
        return short_text(f"{source} | {url}", 100)
    if basis:
        return short_text(f"{source or status} | {basis}", 100)
    return short_text(source or url or status, 100)


def recipe_record_rows(recipe: dict[str, Any], rows: list[dict[str, Any]]) -> dict[str, list[list[Any]]]:
    ingredient_tracker_rows = []
    ingredient_rows = []
    philfct_rows = []
    computation_rows = []
    pricing_rows = []
    for row in rows:
        portion = row.get("portion_g")
        ingredient = safe_text(row.get("ingredient"))
        source_text = safe_text(row.get("source_text"))
        ingredient_tracker_rows.append(
            [
                ingredient,
                evidence_check_status(row),
                evidence_action_needed(row),
                "",
                "",
                "",
            ]
        )
        ingredient_rows.append(
            [
                ingredient,
                f"{number(portion, 2)} g" if portion not in (None, "") else "",
                short_text(source_text, 80),
                row.get("row_type"),
            ]
        )
        philfct_rows.append(
            [
                meal_type(recipe),
                ingredient,
                number(portion, 2),
                source_name_code(row),
                number(display_nutrition_value(row, "per_100g", "calories"), 2),
                number(display_nutrition_value(row, "per_100g", "protein_g"), 2),
                number(display_nutrition_value(row, "per_100g", "carbs_g"), 2),
                number(display_nutrition_value(row, "per_100g", "fat_g"), 2),
                number(display_nutrition_value(row, "per_100g", "fiber_g"), 2),
                short_text(row.get("nutrition_note"), 85),
            ]
        )
        computation_rows.append(
            [
                ingredient,
                number(portion, 2),
                "portion_g / 100 * per_100g",
                number(display_nutrition_value(row, "computed", "calories"), 2),
                number(display_nutrition_value(row, "computed", "protein_g"), 2),
                number(display_nutrition_value(row, "computed", "carbs_g"), 2),
                number(display_nutrition_value(row, "computed", "fat_g"), 2),
                number(display_nutrition_value(row, "computed", "fiber_g"), 2),
                row.get("nutrition_included"),
            ]
        )
        pricing_rows.append(
            [
                ingredient,
                f"{number(portion, 2)} g" if portion not in (None, "") else "",
                display_market_unit(row),
                display_unit_price(row),
                money(effective_cost_php(row)),
                source_date_text(row),
                "",
                short_text(evidence_action_needed(row), 70),
            ]
        )
    return {
        "ingredient_tracker": ingredient_tracker_rows,
        "ingredients": ingredient_rows,
        "philfct": philfct_rows,
        "computation": computation_rows,
        "pricing": pricing_rows,
        "cross_check": cross_check_rows(recipe, rows),
    }


def add_recipe_record(doc: Document, recipe: dict[str, Any], rows: list[dict[str, Any]], record_number: int, first: bool) -> None:
    if not first:
        doc.add_page_break()
    doc.add_heading(f"Recipe Record {record_number:03d}: {recipe_id(recipe)} - {recipe_name(recipe)}", level=1)
    generated_rows = recipe_record_rows(recipe, rows)
    source_count = sum(1 for row in rows if safe_text(row.get("philfct_code")) or safe_text(row.get("philfct_name")))
    priced_count = sum(1 for row in rows if has_effective_cost_value(row))

    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Recipe ID", recipe_id(recipe)],
            ["Recipe name", recipe_name(recipe)],
            ["Meal", meal_type(recipe)],
            ["Runtime dataset", safe_text(recipe.get("sourceDataset"))],
            ["Nutrition data source", safe_text(recipe.get("nutritionDataSource"))],
            ["Nutrition confidence", safe_text(recipe.get("nutritionConfidence"))],
            ["Review status", record_review_status(recipe)],
            ["Team review status", recipe_team_status(recipe, rows)],
            ["Review priority", recipe_priority(recipe, rows)],
            ["Main items to check", short_text("; ".join(recipe_issue_flags(recipe, rows)), 180)],
            ["PhilFCT coverage", number(recipe.get("philfctCoverage"), 2)],
            ["Estimated cost PHP", money(recipe.get("estimatedCostPhp"))],
            ["Cooking minutes", recipe.get("minutes")],
            ["Evidence rows", f"{len(rows)} total; {source_count} with PhilFCT/source names; {priced_count} with cost values"],
            ["Tags", short_text(", ".join(recipe.get("tags") or []), 140)],
        ],
        widths=[1.7, 8.0],
        font_size=7.1,
        header_fill="5B9BD5",
    )

    add_table(
        doc,
        ["Review item", "Current status", "Action needed", "Reviewer", "Date checked", "Correction notes"],
        recipe_review_tracker_rows(recipe, rows),
        widths=[1.35, 1.55, 3.4, 0.9, 0.85, 1.75],
        font_size=6.2,
        header_fill="C00000",
    )

    add_table(
        doc,
        ["Meal", "Recipe/plate", "kcal", "Protein", "Carbs", "Fat", "Fiber", "Pinggang Pinoy / completeness note"],
        [
            [
                meal_type(recipe),
                recipe_name(recipe),
                number(nutrition(recipe, "calories"), 0),
                f"{number(nutrition(recipe, 'protein_g'), 1)} g",
                f"{number(nutrition(recipe, 'carbs_g'), 1)} g",
                f"{number(nutrition(recipe, 'fat_g'), 1)} g",
                f"{number(nutrition(recipe, 'fiber_g'), 1)} g",
                component_note(recipe),
            ]
        ],
        widths=[0.7, 2.0, 0.45, 0.55, 0.55, 0.45, 0.45, 4.5],
        font_size=6.7,
    )

    add_table(
        doc,
        ["Ingredient", "Check status", "Action needed", "Reviewer", "Date checked", "Correction notes"],
        generated_rows["ingredient_tracker"],
        widths=[1.8, 1.55, 3.35, 0.9, 0.85, 1.25],
        font_size=5.95,
        header_fill="C0504D",
    )

    add_table(
        doc,
        ["Ingredient", "Edible portion for one person", "Source text / quantity basis", "Row type"],
        generated_rows["ingredients"],
        widths=[2.4, 1.2, 4.3, 1.8],
        font_size=6.3,
    )

    add_table(
        doc,
        ["Meal", "Ingredient", "Portion (g)", "PhilFCT food name/code", "kcal/100g", "Protein/100g", "Carbs/100g", "Fat/100g", "Fiber/100g", "Remarks"],
        generated_rows["philfct"],
        widths=[0.65, 1.45, 0.58, 2.25, 0.52, 0.58, 0.58, 0.52, 0.52, 2.1],
        font_size=5.75,
    )

    add_table(
        doc,
        ["Ingredient", "Portion (g)", "Formula", "Computed kcal", "Computed protein", "Computed carbs", "Computed fat", "Computed fiber", "Included"],
        generated_rows["computation"],
        widths=[1.8, 0.62, 1.55, 0.7, 0.78, 0.72, 0.65, 0.68, 0.55],
        font_size=5.9,
        header_fill="385723",
    )

    add_table(
        doc,
        ["Ingredient", "Portion used (g)", "Market unit", "Price per market unit", "Computed cost for portion", "Source/date", "Checked by", "Remarks"],
        generated_rows["pricing"],
        widths=[1.65, 0.75, 0.7, 0.75, 0.85, 2.6, 1.1, 1.35],
        font_size=5.7,
        header_fill="7F6000",
    )

    add_table(
        doc,
        ["Item", "Runtime recipe total", "Ingredient-row computed total", "Delta", "Unit", "Status"],
        generated_rows["cross_check"],
        widths=[1.4, 1.2, 1.55, 0.75, 0.55, 1.0],
        font_size=6.4,
        header_fill="8064A2",
    )


def build_volume_doc(meal: str, recipes: list[dict[str, Any]], rows_by_recipe: dict[str, list[dict[str, Any]]], git_short: str) -> Path:
    output_path = OUTPUT_DIR / VOLUME_NAMES[meal]
    doc = Document()
    configure_document(doc)
    add_title(
        doc,
        f"PCOSina PhilFCT and Pricing Recipe Records - {meal}",
        "Template-based recipe records for active runtime catalog cross-checking.",
        f"Generated: {GENERATED_DATE} | Git commit: {git_short} | Template inspected: {TEMPLATE_PATH.relative_to(ROOT)} | Records: {len(recipes)}",
    )
    doc.add_heading("Volume Scope", level=1)
    add_table(
        doc,
        ["Question", "Answer"],
        [
            ["Which recipes are included?", f"All active runtime recipes with mealType = {meal}."],
            ["What is repeated per recipe?", "Recipe summary, serving rows, PhilFCT worksheet, nutrition computation, pricing worksheet, and totals cross-check."],
            ["How should the team use the tracker?", "Fill Reviewer, Date checked, and Correction notes directly in the red review tables. Edit data cells only when a source value is being corrected."],
            ["Are these final RND-certified values?", "No. Rows marked source_mapped_needs_final_review remain reviewable records until final human/RND verification."],
        ],
        widths=[2.4, 7.2],
        font_size=7.6,
        header_fill="5B9BD5",
    )
    for index, recipe in enumerate(recipes, start=1):
        add_recipe_record(doc, recipe, rows_by_recipe.get(recipe_id(recipe), []), index, first=index == 1)
    doc.save(output_path)
    return output_path


def source_reference_rows() -> list[list[Any]]:
    rows = []
    for source_name, role, url, use in OFFICIAL_SOURCES:
        rows.append([source_name, role, url, use])
    return rows


def build_master_doc(
    recipes: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    rows_by_recipe: dict[str, list[dict[str, Any]]],
    volume_paths: dict[str, Path],
    git_short: str,
) -> None:
    doc = Document()
    configure_document(doc)
    add_title(
        doc,
        "PCOSina PhilFCT and Pricing Recipe Records - Master Index",
        "Index and methodology for the per-recipe nutrition and pricing evidence volumes.",
        f"Generated: {GENERATED_DATE} | Git commit: {git_short} | Active runtime catalog: {len(recipes)} recipes",
    )

    meal_counts = Counter(meal_type(recipe) for recipe in recipes)
    dataset_counts = Counter(safe_text(recipe.get("sourceDataset")) for recipe in recipes)
    confidence_counts = Counter(safe_text(recipe.get("nutritionConfidence")) for recipe in recipes)
    review_counts = Counter(safe_text(recipe.get("nutritionReviewStatus")) for recipe in recipes)
    cross_statuses = Counter()
    priority_counts: Counter[str] = Counter()
    team_status_counts: Counter[str] = Counter()
    issue_counts: Counter[str] = Counter()
    ingredient_rows_needing_check = 0
    for recipe in recipes:
        rows = rows_by_recipe.get(recipe_id(recipe), [])
        checks = cross_check_rows(recipe, rows)
        if all(row[-1] == "OK" for row in checks):
            cross_statuses["OK"] += 1
        else:
            cross_statuses["Review"] += 1
        priority_counts[recipe_priority(recipe, rows)] += 1
        team_status_counts[recipe_team_status(recipe, rows)] += 1
        for flag in recipe_issue_flags(recipe, rows):
            issue_counts[flag.split(" (", 1)[0]] += 1
        for row in rows:
            if evidence_issue_flags(row):
                ingredient_rows_needing_check += 1

    doc.add_heading("1. Scope and Summary", level=1)
    add_table(
        doc,
        ["Question", "Answer"],
        [
            ["Template basis", str(TEMPLATE_PATH.relative_to(ROOT))],
            ["Active recipe catalog", str(RUNTIME_PATH.relative_to(ROOT))],
            ["Active recipe count documented", len(recipes)],
            ["Meal type counts", "; ".join(f"{key}: {value}" for key, value in sorted(meal_counts.items(), key=lambda item: MEAL_ORDER.get(item[0], 99)))],
            ["Ingredient-level evidence rows", len(evidence_rows)],
            ["Recipe cross-check status", "; ".join(f"{key}: {value}" for key, value in sorted(cross_statuses.items()))],
            ["Nutrition confidence counts", "; ".join(f"{key}: {value}" for key, value in confidence_counts.most_common())],
            ["Review status counts", "; ".join(f"{key}: {value}" for key, value in review_counts.most_common())],
            ["Review priority counts", "; ".join(f"{key}: {value}" for key, value in priority_counts.most_common())],
            ["Team status counts", "; ".join(f"{short_text(key, 48)}: {value}" for key, value in team_status_counts.most_common())],
            ["Ingredient rows needing check", f"{ingredient_rows_needing_check} of {len(evidence_rows)}"],
            ["Dataset counts", "; ".join(f"{short_text(key, 55)}: {value}" for key, value in dataset_counts.most_common())],
        ],
        widths=[2.2, 7.5],
        font_size=7.4,
        header_fill="5B9BD5",
    )

    doc.add_heading("2. Team Review Tracking Guide", level=1)
    add_table(
        doc,
        ["Tracker field", "Meaning", "How the team should use it"],
        [
            ["Review priority", "High means a source/cost/portion gap needs correction; Medium means final source/RND sign-off is still pending; Low means mostly spot-check.", "Start with High, then Medium, then Low."],
            ["Team review status", "Plain-language status for whether the recipe still needs checking.", "Use this before opening a volume so work can be assigned quickly."],
            ["Issue flags", "Specific reason the record needs attention.", "Use this to know whether to check PhilFCT, portion, pricing, or final sign-off."],
            ["Reviewer / Date checked", "Blank fields intended for team editing.", "Fill after checking the row or recipe."],
            ["Correction notes", "Blank field for exact changes or source notes.", "Write the corrected value/source here before feeding the reviewed document back for import."],
        ],
        widths=[1.35, 3.1, 5.25],
        font_size=6.8,
        header_fill="C00000",
    )
    add_table(
        doc,
        ["Issue flag", "Recipe count"],
        [[flag, count] for flag, count in issue_counts.most_common()],
        widths=[4.2, 1.2],
        font_size=6.8,
        header_fill="C0504D",
    )

    doc.add_heading("3. Computation Method", level=1)
    add_table(
        doc,
        ["Work item", "Formula / method", "Review note"],
        [
            ["Ingredient nutrition", "ingredient contribution = portion_g / 100 * PhilFCT per-100g nutrient value", "This is the same arithmetic pattern requested in the PhilFCT worksheet."],
            ["Meal nutrition total", "meal total = sum of included ingredient contributions", "Totals are compared against backend/recipes.json values."],
            ["Ingredient pricing", "portion cost follows the linked pricing rule basis, usually portion_g / 1000 * PHP per kg", "Other units are retained from the pricing rule where available."],
            ["Meal cost", "meal cost = sum of ingredient consumed-portion costs", "Compared against estimatedCostPhp in the active runtime recipe."],
            ["Tolerance", "nutrition delta <= 0.55 and cost delta <= PHP 0.05 is marked OK", "Tolerance is for rounding only; source mapping still needs final review where status says so."],
        ],
        widths=[1.5, 4.4, 3.8],
        font_size=7.0,
        header_fill="385723",
    )

    doc.add_heading("4. Source References Used by the Records", level=1)
    add_table(
        doc,
        ["Source", "Role", "URL", "How used"],
        source_reference_rows(),
        widths=[2.0, 1.4, 2.4, 3.9],
        font_size=6.0,
        header_fill="7030A0",
    )

    doc.add_heading("5. Record Volumes", level=1)
    add_table(
        doc,
        ["Volume", "Meal type", "Recipe count", "File"],
        [
            [idx, meal, meal_counts.get(meal, 0), volume_paths[meal].name]
            for idx, meal in enumerate(VOLUME_NAMES, start=1)
        ],
        widths=[0.6, 1.2, 1.0, 5.5],
        font_size=7.5,
        header_fill="7F6000",
    )

    doc.add_heading("6. Recipe Record Index", level=1)
    index_rows = []
    for idx, recipe in enumerate(sorted_recipes(recipes), start=1):
        meal = meal_type(recipe)
        rows = rows_by_recipe.get(recipe_id(recipe), [])
        checks = cross_check_rows(recipe, rows)
        index_rows.append(
            [
                idx,
                meal,
                recipe_id(recipe),
                short_text(recipe_name(recipe), 42),
                number(nutrition(recipe, "calories"), 0),
                money(recipe.get("estimatedCostPhp")),
                safe_text(recipe.get("nutritionConfidence")),
                short_text(recipe.get("nutritionReviewStatus"), 30),
                "OK" if all(row[-1] == "OK" for row in checks) else "Review",
                recipe_priority(recipe, rows),
                recipe_team_status(recipe, rows),
                short_text("; ".join(recipe_issue_flags(recipe, rows)), 80),
                short_text(recipe_next_action(recipe, rows), 90),
                volume_paths[meal].name,
                "",
                "",
                "",
            ]
        )
    add_table(
        doc,
        ["#", "Meal", "Recipe ID", "Recipe name", "kcal", "Cost PHP", "Conf", "Review status", "Calc", "Priority", "Team status", "Issue flags", "Next action", "Volume", "Reviewer", "Date", "Notes"],
        index_rows,
        widths=[0.22, 0.42, 0.52, 1.0, 0.28, 0.38, 0.3, 0.65, 0.3, 0.35, 0.85, 0.95, 1.0, 0.9, 0.42, 0.38, 0.7],
        font_size=4.65,
    )
    doc.save(MASTER_DOCX)


def sorted_recipes(recipes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(recipes, key=lambda recipe: (MEAL_ORDER.get(meal_type(recipe), 99), recipe_name(recipe), recipe_id(recipe)))


def write_json(
    recipes: list[dict[str, Any]],
    rows_by_recipe: dict[str, list[dict[str, Any]]],
    volume_paths: dict[str, Path],
    git_short: str,
) -> None:
    records = []
    for recipe in sorted_recipes(recipes):
        rid = recipe_id(recipe)
        rows = rows_by_recipe.get(rid, [])
        enriched_rows = []
        for row in rows:
            enriched = dict(row)
            allocated = allocated_price(row)
            enriched["display_per_100g"] = {
                key: display_nutrition_value(row, "per_100g", key)
                for key in NUTRIENT_KEYS
            }
            enriched["display_computed"] = {
                key: display_nutrition_value(row, "computed", key)
                for key in NUTRIENT_KEYS
            }
            enriched["display_cost_php"] = effective_cost_php(row)
            enriched["display_market_unit"] = display_market_unit(row)
            enriched["display_unit_price_php"] = display_unit_price(row)
            enriched["display_price_confidence"] = display_price_confidence(row)
            enriched["display_price_rule_id"] = display_price_rule_id(row)
            enriched["display_market_source"] = display_market_source(row)
            enriched["display_price_source_url"] = display_price_source_url(row)
            enriched["display_pricing_status"] = display_pricing_status(row)
            enriched["priceBasisNote"] = price_basis_note(row)
            if allocated:
                enriched["pantryPriceAllocation"] = allocated
            enriched["reviewStatus"] = evidence_check_status(row)
            enriched["issueFlags"] = evidence_issue_flags(row)
            enriched["actionNeeded"] = evidence_action_needed(row)
            enriched["reviewer"] = ""
            enriched["dateChecked"] = ""
            enriched["correctionNotes"] = ""
            enriched_rows.append(enriched)
        records.append(
            {
                "recipeId": rid,
                "recipeName": recipe_name(recipe),
                "mealType": meal_type(recipe),
                "sourceDataset": recipe.get("sourceDataset"),
                "nutritionDataSource": recipe.get("nutritionDataSource"),
                "nutritionConfidence": recipe.get("nutritionConfidence"),
                "nutritionReviewStatus": recipe.get("nutritionReviewStatus"),
                "philfctCoverage": recipe.get("philfctCoverage"),
                "estimatedCostPhp": recipe.get("estimatedCostPhp"),
                "teamReviewStatus": recipe_team_status(recipe, rows),
                "reviewPriority": recipe_priority(recipe, rows),
                "issueFlags": recipe_issue_flags(recipe, rows),
                "nextAction": recipe_next_action(recipe, rows),
                "reviewer": "",
                "dateChecked": "",
                "correctionNotes": "",
                "runtimeNutrition": {key: nutrition(recipe, key) for key in NUTRIENT_KEYS},
                "computedTotals": totals_for_rows(rows),
                "crossCheckRows": cross_check_rows(recipe, rows),
                "evidenceRows": enriched_rows,
                "volumeFile": volume_paths[meal_type(recipe)].name,
            }
        )
    payload = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "generatedDate": GENERATED_DATE,
        "gitCommitShort": git_short,
        "templatePath": str(TEMPLATE_PATH.relative_to(ROOT)),
        "runtimeCatalog": str(RUNTIME_PATH.relative_to(ROOT)),
        "recipeCount": len(recipes),
        "records": records,
    }
    JSON_PATH.write_text(json.dumps(payload, indent=2, ensure_ascii=True), encoding="utf-8")


def write_csvs(recipes: list[dict[str, Any]], rows_by_recipe: dict[str, list[dict[str, Any]]], volume_paths: dict[str, Path]) -> None:
    index_rows = []
    evidence_rows = []
    for index, recipe in enumerate(sorted_recipes(recipes), start=1):
        rid = recipe_id(recipe)
        meal = meal_type(recipe)
        checks = cross_check_rows(recipe, rows_by_recipe.get(rid, []))
        index_rows.append(
            {
                "Record Number": index,
                "Recipe ID": rid,
                "Recipe Name": recipe_name(recipe),
                "Meal Type": meal,
                "Calories": nutrition(recipe, "calories"),
                "Protein g": nutrition(recipe, "protein_g"),
                "Carbs g": nutrition(recipe, "carbs_g"),
                "Fat g": nutrition(recipe, "fat_g"),
                "Fiber g": nutrition(recipe, "fiber_g"),
                "Estimated Cost PHP": recipe.get("estimatedCostPhp"),
                "PhilFCT Coverage": recipe.get("philfctCoverage"),
                "Nutrition Confidence": recipe.get("nutritionConfidence"),
                "Nutrition Review Status": recipe.get("nutritionReviewStatus"),
                "Calculation Status": "OK" if all(row[-1] == "OK" for row in checks) else "Review",
                "Team Review Status": recipe_team_status(recipe, rows_by_recipe.get(rid, [])),
                "Review Priority": recipe_priority(recipe, rows_by_recipe.get(rid, [])),
                "Issue Flags": "; ".join(recipe_issue_flags(recipe, rows_by_recipe.get(rid, []))),
                "Next Action": recipe_next_action(recipe, rows_by_recipe.get(rid, [])),
                "Reviewer": "",
                "Date Checked": "",
                "Correction Notes": "",
                "Volume File": volume_paths[meal].name,
            }
        )
        for row_index, row in enumerate(rows_by_recipe.get(rid, []), start=1):
            evidence_rows.append(
                {
                    "Recipe ID": rid,
                    "Recipe Name": recipe_name(recipe),
                    "Meal Type": meal,
                    "Row Number": row_index,
                    "Row Type": row.get("row_type"),
                    "Ingredient": row.get("ingredient"),
                    "Source Text": row.get("source_text"),
                    "Portion g": row.get("portion_g"),
                    "Parse Status": row.get("parse_status"),
                    "PhilFCT Code": row.get("philfct_code"),
                    "PhilFCT Name": row.get("philfct_name"),
                    "kcal per 100g": display_nutrition_value(row, "per_100g", "calories"),
                    "Protein per 100g": display_nutrition_value(row, "per_100g", "protein_g"),
                    "Carbs per 100g": display_nutrition_value(row, "per_100g", "carbs_g"),
                    "Fat per 100g": display_nutrition_value(row, "per_100g", "fat_g"),
                    "Fiber per 100g": display_nutrition_value(row, "per_100g", "fiber_g"),
                    "Computed kcal": display_nutrition_value(row, "computed", "calories"),
                    "Computed Protein": display_nutrition_value(row, "computed", "protein_g"),
                    "Computed Carbs": display_nutrition_value(row, "computed", "carbs_g"),
                    "Computed Fat": display_nutrition_value(row, "computed", "fat_g"),
                    "Computed Fiber": display_nutrition_value(row, "computed", "fiber_g"),
                    "Nutrition Included": row.get("nutrition_included"),
                    "Nutrition Note": row.get("nutrition_note"),
                    "Cost PHP": effective_cost_php(row),
                    "Pricing Status": display_pricing_status(row),
                    "Price Rule ID": display_price_rule_id(row),
                    "Unit Price PHP": display_unit_price(row),
                    "Market Unit": display_market_unit(row),
                    "Price Confidence": display_price_confidence(row),
                    "Price Basis Note": price_basis_note(row),
                    "Market Source": display_market_source(row),
                    "Price Source URL": display_price_source_url(row),
                    "Review Status": evidence_check_status(row),
                    "Issue Flags": "; ".join(evidence_issue_flags(row)),
                    "Action Needed": evidence_action_needed(row),
                    "Reviewer": "",
                    "Date Checked": "",
                    "Correction Notes": "",
                    "Volume File": volume_paths[meal].name,
                }
            )

    with INDEX_CSV_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(index_rows[0].keys()))
        writer.writeheader()
        writer.writerows(index_rows)

    with CSV_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(evidence_rows[0].keys()))
        writer.writeheader()
        writer.writerows(evidence_rows)


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    recipes = read_json(RUNTIME_PATH)
    philfct = read_json(PHILFCT_CACHE_PATH)
    priced = read_json(PRICED_DRAFT_PATH)
    if not isinstance(recipes, list):
        raise ValueError("backend/recipes.json must be a JSON array")
    if not isinstance(philfct, list):
        raise ValueError("PhilFCT reference cache must be a JSON array")

    priced_by_recipe = priced_meal_index(priced)
    philfct_by_code = philfct_code_index(philfct)
    rows_by_recipe: dict[str, list[dict[str, Any]]] = {}
    for recipe in recipes:
        rows_by_recipe[recipe_id(recipe)] = recipe_evidence_rows(recipe, priced_by_recipe, philfct_by_code)
    evidence_rows = all_evidence_rows(recipes, priced, philfct)
    git_short = git_commit_short()

    by_meal: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for recipe in sorted_recipes(recipes):
        by_meal[meal_type(recipe)].append(recipe)

    volume_paths: dict[str, Path] = {}
    for meal in VOLUME_NAMES:
        volume_paths[meal] = build_volume_doc(meal, by_meal.get(meal, []), rows_by_recipe, git_short)

    build_master_doc(recipes, evidence_rows, rows_by_recipe, volume_paths, git_short)
    write_json(recipes, rows_by_recipe, volume_paths, git_short)
    write_csvs(recipes, rows_by_recipe, volume_paths)

    cross_ok = 0
    for recipe in recipes:
        if all(row[-1] == "OK" for row in cross_check_rows(recipe, rows_by_recipe.get(recipe_id(recipe), []))):
            cross_ok += 1

    summary = {
        "masterDocx": str(MASTER_DOCX),
        "volumeDocx": {meal: str(path) for meal, path in volume_paths.items()},
        "json": str(JSON_PATH),
        "evidenceCsv": str(CSV_PATH),
        "indexCsv": str(INDEX_CSV_PATH),
        "recipeCount": len(recipes),
        "evidenceRowCount": len(evidence_rows),
        "crossCheckOkRecipeCount": cross_ok,
        "crossCheckReviewRecipeCount": len(recipes) - cross_ok,
        "mealCounts": dict(Counter(meal_type(recipe) for recipe in recipes)),
    }
    print(json.dumps(summary, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
