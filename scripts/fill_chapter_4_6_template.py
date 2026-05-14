from __future__ import annotations

import csv
import json
import re
import subprocess
from pathlib import Path
from statistics import mean

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


REPO = Path(__file__).resolve().parents[1]
DOCX = REPO / "docs" / "PCOSina_Chapter_4-6_Thesis_Template.docx"


def git_output(*args: str) -> str:
    return subprocess.check_output(["git", *args], cwd=REPO, text=True).strip()


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def load_json(path: Path) -> dict:
    return json.loads(read_text(path))


def load_android_version() -> tuple[str, str]:
    text = read_text(REPO / "app" / "build.gradle.kts")
    code = re.search(r"versionCode\s*=\s*(\d+)", text)
    name = re.search(r'versionName\s*=\s*"([^"]+)"', text)
    return (name.group(1) if name else "not found", code.group(1) if code else "not found")


def csv_count(path: Path) -> int:
    with path.open(newline="", encoding="utf-8") as handle:
        return sum(1 for _ in csv.DictReader(handle))


def validation_category_counts(path: Path) -> dict[str, int]:
    counts: dict[str, int] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            category = row["category"]
            counts[category] = counts.get(category, 0) + 1
    return counts


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def set_first_starting(doc: Document, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            paragraph.text = text
            return


def set_exact(doc: Document, old: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text == old:
            paragraph.text = new
            return


def cell_set(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def trim_table(table, desired_rows: int) -> None:
    while len(table.rows) > desired_rows:
        row = table.rows[-1]
        row._tr.getparent().remove(row._tr)


def fill_table(table, headers: list[str], rows: list[list[str]]) -> None:
    desired_rows = len(rows) + 1
    while len(table.rows) < desired_rows:
        table.add_row()
    trim_table(table, desired_rows)
    for column_index, header in enumerate(headers):
        cell_set(table.rows[0].cells[column_index], header)
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row):
            cell_set(table.rows[row_index].cells[column_index], str(value))


def fill_one_cell(table, text: str) -> None:
    cell_set(table.rows[0].cells[0], text)


def main() -> None:
    branch = git_output("branch", "--show-current")
    commit = git_output("rev-parse", "HEAD")
    version_name, version_code = load_android_version()
    validation_case_path = REPO / "docs" / "thesis_validation" / "05_VALIDATION_TEST_CASES.csv"
    validation_cases = csv_count(validation_case_path)
    category_counts = validation_category_counts(validation_case_path)
    training = load_json(REPO / "ml" / "offline_training" / "artifacts" / "model_v1" / "training_metrics.json")
    report = load_json(REPO / "backend" / "comment5_algorithm_tradeoff_report.json")
    quality = training["dataset_manifest"]["quality"]
    splits = training["dataset_manifest"]["splitRows"]
    agg_cpsat = report["aggregate"]["cpsat_two_stage"]
    agg_greedy = report["aggregate"]["greedy_heuristic"]
    agg_proxy = report["aggregate"]["gen_proxy"]

    scenario_runtime_rows = [
        ("S1", "Base", 0.2852, 1.9052, "100.00%", "100.00%", "3.00"),
        ("S2", "BudgetNoPork", 0.2240, 1.1701, "100.00%", "100.00%", "3.00"),
        ("S3", "VegLactose", 0.1366, 3.9361, "100.00%", "100.00%", "3.19"),
        ("S4", "PantryConstrained", 0.2898, 2.4698, "100.00%", "100.00%", "3.00"),
        ("S5", "ResilienceProxy300", 6.9306, 2.2258, "100.00%", "100.00%", "n/a"),
    ]
    scenario_quality_rows = [
        ("S1", 531.57, 530.14, 33.27, 33.10, "0 / 0"),
        ("S2", 540.43, 539.00, 33.27, 33.02, "0 / 0"),
        ("S3", 605.43, 601.86, 36.31, 36.28, "0 / 0"),
        ("S4", 565.71, 557.14, 33.97, 33.10, "0 / 0"),
    ]
    before_runs_ms = [207816, 188249]
    after_runs_ms = [9803, 13674, 6890, 6486, 9593]
    before_avg = mean(before_runs_ms)
    after_avg = mean(after_runs_ms)
    improvement = before_avg / after_avg

    doc = Document(DOCX)
    if doc.paragraphs:
        doc.paragraphs[0].text = "PCOSina Chapters 4-6 Thesis Draft"

    guide_texts = {
        "Template Use Notes",
        "Delete this guide page before final submission if your school format does not allow template notes.",
        "Use Chapter 4 for actual results, tables, figures, test outputs, survey findings, and interpretation.",
        "Use Chapter 5 to summarize the findings according to the general objective, specific objectives, and research questions.",
        "Use Chapter 6 for conclusions, recommendations, future work, and limitations.",
        "Keep PCOSina framed as a wellness and meal-planning decision-support system, not a diagnostic or treatment application.",
        "Do not claim ISO certification. State that ISO/IEC 25010 is used as a software quality evaluation framework.",
        "For every table and figure, include a short interpretation paragraph explaining what the result means for the study.",
        "Use exact release/version evidence: final APK/version number, GitHub branch, commit hash, and test date.",
        "Mark missing values as placeholders only while drafting. Remove all placeholders before final binding/submission.",
    }
    for paragraph in list(doc.paragraphs):
        if paragraph.text in guide_texts:
            delete_paragraph(paragraph)

    category_summary = ", ".join(f"{key} {value}" for key, value in sorted(category_counts.items()))

    set_first_starting(
        doc,
        "Template with placeholders.",
        "Repository-backed Chapter 4-6 draft updated with PCOSina implementation evidence, computation exports, "
        "benchmark metrics, operational reports, and ML artifacts. Human respondent, CS/IT expert, and health "
        "consultation results remain marked pending where no raw data was found in the repository.",
    )
    set_exact(doc, "Template Evidence Checklist", "Evidence Checklist")

    replacements = [
        (
            "[WRITE OVERVIEW HERE]",
            "This chapter presents the available evaluation evidence for PCOSina using only repository-backed data "
            "from the Android application, backend planner, validation exports, benchmark report, operational "
            "roadmap reports, and ML artifacts. The evidence is strongest for implemented functionality, formula "
            "correctness, deterministic planner behavior, pantry and budget computations, test coverage, and guarded "
            "ML ranking. Sections that require actual human respondents, CS/IT validators, or OBGYN or nutrition "
            "consultation are intentionally marked pending because no raw response matrix or signed validation "
            "summary was found in the current workspace.",
        ),
        (
            "[VERSION STATEMENT]",
            f"The PCOSina build described in this draft is Android version {version_name} with versionCode "
            f"{version_code}, from Git branch {branch} at commit {commit}. APK artifacts are present under "
            "app/build/outputs/apk, including debug, staging, release, and androidTest outputs. The thesis evidence "
            "pack at docs/thesis_validation was generated from repository artifacts and should be regenerated after "
            "any final release build or respondent-testing APK change.",
        ),
        (
            "[SYSTEM OUTPUT INTRODUCTION]",
            "PCOSina is implemented as an Android wellness decision-support application for Filipino individuals "
            "with PCOS-related meal-planning needs. Its user workflow covers startup and legal acceptance, profile "
            "setup, goal selection, health and calorie computation, weekly meal-plan generation, recipe details, "
            "grocery guidance, pantry behavior, progress logging, reminders, and support or settings screens. The "
            "system must be described as a wellness and meal-planning tool, not as a diagnostic engine, treatment "
            "system, or medical device.",
        ),
        (
            "[INTERPRETATION]",
            "The implemented modules satisfy the expected workflow from Chapter 3 because user input flows into "
            "health computations, deterministic safety filtering, weekly optimization, grocery output, and progress "
            "support. The available evidence also confirms the important boundary that ML is optional and assistive "
            "only. It may rank already-safe candidates, but it does not override allergies, restrictions, budget "
            "ceilings, repetition rules, or the CP-SAT final selection authority.",
        ),
        (
            "[CONSULTATION INTRODUCTION]",
            "This section is reserved for actual CS/IT expert validation and OBGYN, nutritionist, or health "
            "consultation results. The repository contains instrument drafts and implementation evidence, but it "
            "does not contain signed validator forms, respondent matrices, expert ratings, or dated consultation "
            "summaries. Therefore, the tables below are filled only with readiness notes and safe wording boundaries "
            "until actual expert evidence is inserted.",
        ),
        (
            "[VALIDATION INTERPRETATION]",
            "No final expert-validation conclusion should be claimed yet. What can be stated from the system scan is "
            "that the manuscript must use accurate technical wording: PCOSina uses deterministic filtering plus "
            "OR-Tools CP-SAT integer optimization, not a pure MILP implementation; pantry behavior is a soft reward "
            "and grocery-support feature, not a full hard pantry-feasibility guarantee; and PCOSina remains a "
            "wellness decision-support application rather than a clinical diagnostic system.",
        ),
        (
            "[TESTING INTRODUCTION]",
            f"System testing evidence was gathered from the repository test inventory and thesis validation pack. "
            f"The current validation case list contains {validation_cases} cases, including health computations, "
            "safety filters, planner constraints, planner contracts, no-safe-plan handling, grocery aggregation, "
            "offline reliability, input validation, and API contract behavior. The strongest implemented tests "
            "include BMI and calorie target checks, allergy-family exclusion, macro ratios, budget enforcement, "
            "household cost scaling, repetition limits, no-safe-plan responses, grocery aggregation, meal swaps, "
            "profile validation, and schema handling.",
        ),
        (
            "[TESTING INTERPRETATION]",
            "The testing evidence supports the claim that the core software functions are implemented and that many "
            "critical planner and Android flows are covered. The test scan also identifies gaps that should be "
            "addressed before final defense, including additional Android parity tests for BMR and all activity "
            "multipliers, more grocery pantry-name matching tests, and a full planner execution run on a machine "
            "where OR-Tools is not blocked by application-control policy.",
        ),
        (
            "[BENCHMARK INTRODUCTION]",
            "The planner benchmark compares a simpler rule-filter plus greedy baseline with the proposed PCOSina "
            "two-stage planner. PCOSina first filters unsafe or incompatible recipes, then formulates the weekly "
            "assignment as a 0-1 CP-SAT optimization model using OR-Tools. The modeled decision variable is whether "
            "recipe r is assigned to meal slot s. The model enforces exactly one recipe per slot, meal-slot "
            "compatibility, allergy and restriction exclusions, repeat limits, adjacency rules, and budget caps when "
            "a budget is resolved. The objective minimizes weighted nutrition deviations, cost pressure, preparation "
            "burden, repetition pressure, and other soft penalties while rewarding pantry overlap and vegetable "
            "diversity. OR-Tools supplies the CP-SAT solving algorithm; the study's defensible formula is the "
            "explicit model formulation and objective terms, not a hidden library-internal closed-form equation.",
        ),
        (
            "Note. Define feasibility rate",
            "Note. In Table 4.8, feasibility is interpreted as hard-safe plan completion for the benchmark rows. "
            "Strict nutrition feasibility is discussed separately through calorie and macro deviation values because "
            "the benchmark report records nonzero nutrition deviation even when hard constraints are preserved.",
        ),
        (
            "[BENCHMARK INTERPRETATION]",
            f"The proposed CP-SAT planner is slower than the greedy baseline in matched scenarios because it solves "
            f"a richer weekly planning problem. Across the benchmark aggregate report, the greedy baseline averaged "
            f"{agg_greedy['avg_runtime_s']:.4f} seconds, while the CP-SAT planner averaged "
            f"{agg_cpsat['avg_runtime_s']:.4f} seconds. However, the generated-proxy comparison demonstrates why "
            f"constraint enforcement matters: its hard violation rate was {agg_proxy['hard_violation_rate_pct']:.1f}%, "
            f"while the CP-SAT planner's hard violation rate was {agg_cpsat['hard_violation_rate_pct']:.1f}%. "
            "The benchmark therefore supports a constraint-first interpretation: PCOSina trades some runtime for "
            "deterministic safety and more explainable planning behavior.",
        ),
        (
            "[PANTRY/BUDGET INTRODUCTION]",
            "Pantry and budget computations demonstrate PCOSina's practical planning value. The backend planner "
            "normalizes pantry entries and rewards pantry overlap during ranking and optimization, while the Android "
            "grocery screen supports pantry-aware user review. Budget handling is stronger than pantry handling: when "
            "a weekly budget is resolved, the backend enforces it as a hard cap and may also use cost in the soft "
            "objective depending on policy. Pantry should be described as pantry-aware support and reward, not as a "
            "hard guarantee that every ingredient will already exist in the user's pantry.",
        ),
        (
            "[PANTRY/BUDGET INTERPRETATION]",
            "The sample pantry computation shows why the system can prioritize recipes that use available "
            "ingredients, while the grocery computation shows the remaining practical limitation: backend token "
            "overlap and Android grocery-name matching are not identical. Budget estimates are useful for planning "
            "and comparison, but they should be presented as estimates or ranges because market prices vary by store, "
            "location, season, and data-source freshness.",
        ),
        (
            "[OPERATIONAL PERFORMANCE INTRODUCTION]",
            "Operational performance evidence comes from the April 2026 roadmap reports and the Chapter 4-6 "
            "generator constants. These records captured a planner latency regression and the later recovery "
            "measurements. This section reports those observed values as engineering evidence, not as a final "
            "large-sample performance study.",
        ),
        (
            "[PERFORMANCE INTERPRETATION]",
            f"Observed generation latency improved from an average of {before_avg / 1000:.2f} seconds before the "
            f"recovery work to {after_avg / 1000:.2f} seconds after the fix, or about {improvement:.2f} times faster. "
            "This supports the claim that the planner became more usable for respondent testing, although formal "
            "performance benchmarking should still be repeated on the final release build and device set.",
        ),
        (
            "[ML INTRODUCTION]",
            "PCOSina includes a guarded, system-managed ML-assist ranking layer for Stage 1 candidate ordering. The "
            "ML model uses structured candidate features such as nutrition values, estimated cost, pantry overlap, "
            "meal type, cooking time, restriction counts, allergy counts, budget context, and feedback-derived reason "
            "tags. It operates only after rule-based safety filtering and before final CP-SAT selection. If the "
            "ranker is unavailable, invalid, or outside the allowed rollout path, deterministic ranking remains the "
            "fallback.",
        ),
        (
            "[ML INTERPRETATION]",
            "The ML artifact metrics support the claim that learned ranking can improve candidate ordering, but they "
            "do not justify giving ML final decision authority. The model is best described as an assistive ranker "
            "over already-safe candidates. CP-SAT optimization and hard constraints remain the final authority for "
            "weekly meal selection.",
        ),
        (
            "[USER EVALUATION INTRODUCTION]",
            "This section is reserved for respondent distribution, TAM results, ISO/IEC 25010 results, Cronbach "
            "alpha reliability, and open-ended feedback. The repository contains questionnaire drafts and deployment "
            "support files, but no completed respondent dataset was found. Final weighted means, percentages, and "
            "reliability coefficients must therefore be inserted only after actual responses are collected and cleaned.",
        ),
        (
            "[USER EVALUATION INTERPRETATION]",
            "No final user-acceptance or software-quality conclusion should be claimed from this draft alone. The "
            "correct current interpretation is that the system is ready for respondent testing from an "
            "implementation-evidence standpoint, while TAM, ISO/IEC 25010, and Cronbach alpha results remain pending "
            "until the response matrix is available.",
        ),
        (
            "[DISCUSSION - IMPLEMENTATION]",
            "Implementation evidence shows that PCOSina achieved the main software workflow from profile capture to "
            "meal planning, grocery support, and progress logging. The implementation is not merely a visual "
            "prototype; it contains Android domain computations, repository logic, backend planner services, policy "
            "configuration, test artifacts, and ML training artifacts.",
        ),
        (
            "[DISCUSSION - VALIDATION]",
            "Validation evidence currently supports software behavior and computation correctness more strongly than "
            "human evaluation. Expert consultation and respondent findings must still be inserted from actual forms "
            "or survey files. The manuscript should explicitly separate repository-backed system validation from "
            "pending human-subject evaluation.",
        ),
        (
            "[DISCUSSION - BENCHMARK]",
            "The algorithm evidence shows a clear tradeoff. A greedy planner can run faster, but the CP-SAT planner "
            "gives PCOSina a more defensible constraint-first model for weekly assignment, budget handling, variety, "
            "nutrition deviation, and no-safe-plan behavior. This is the stronger technical basis for a thesis "
            "defense than claiming speed alone.",
        ),
        (
            "[DISCUSSION - ML]",
            "ML is appropriately limited to ranking support. The current evidence is credible because the model has "
            "dataset, split, and metric artifacts, and because the architecture prevents ML from overriding safety "
            "constraints or final solver feasibility.",
        ),
        (
            "[DISCUSSION - USER EVALUATION]",
            "User evaluation remains the major missing evidence stream. The paper can discuss readiness, instruments, "
            "and expected analysis formulas, but final acceptance claims must wait for respondent counts, weighted "
            "means, reliability analysis, and open-ended feedback themes.",
        ),
        (
            "[CHAPTER 4 CLOSING]",
            "Overall, Chapter 4 evidence supports implementation completeness, formula-backed health computations, "
            "deterministic planner safety boundaries, benchmarked algorithm behavior, pantry and budget computation "
            "support, operational recovery, and guarded ML ranking. It does not yet support final claims about "
            "expert approval, clinical appropriateness, or respondent acceptance because those datasets were not "
            "found in the repository.",
        ),
        (
            "[STUDY SUMMARY]",
            "PCOSina was developed as an offline-first Android meal-planning decision-support system for Filipino "
            "individuals with PCOS-related dietary planning needs. Its scope is wellness support: profile-based "
            "computations, preference-aware filtering, deterministic weekly optimization, pantry-aware grocery "
            "guidance, budget awareness, and progress support. It does not diagnose PCOS, verify clinical symptoms, "
            "prescribe treatment, or replace professional consultation.",
        ),
        (
            "[EVALUATION SUMMARY]",
            "The available evaluation evidence consists of implementation inspection, formula validation, validation "
            "test cases, benchmark and operational performance records, policy exports, recipe inventory summaries, "
            "and ML training artifacts. Expert validation and user evaluation sections remain incomplete until actual "
            "CS/IT, health consultation, and respondent survey data are supplied.",
        ),
        (
            "[GENERAL OBJECTIVE FINDING]",
            "General objective: to develop and evaluate PCOSina, an offline-first Filipino-PCOS meal-planning "
            "application that uses deterministic filtering and optimization to generate practical weekly meal plans "
            "with pantry and grocery support. Based on repository-backed evidence, the development component is "
            "substantially achieved. The evaluation component is partially achieved because system and algorithm "
            "evidence exists, while human respondent and expert-validation evidence is still pending.",
        ),
        (
            "Replace the bullets below",
            "The overall findings below are limited to evidence available in the repository as of this draft and "
            "should be updated after respondent and expert-validation data are inserted.",
        ),
        (
            "[Finding 1:",
            "Finding 1: PCOSina implemented the required user workflow from onboarding and profile setup to meal "
            "generation, grocery guidance, progress logging, reminders, settings, and support screens.",
        ),
        (
            "[Finding 2:",
            "Finding 2: The proposed planner is implemented as deterministic filtering plus OR-Tools CP-SAT integer "
            "optimization, preserving hard safety constraints while minimizing weighted nutrition, cost, variety, "
            "and practicality penalties.",
        ),
        (
            "[Finding 3:",
            "Finding 3: Pantry and budget computations support practical meal planning, but the current pantry "
            "behavior is a soft reward and grocery-support mechanism rather than a hard pantry-feasibility constraint.",
        ),
        (
            "[Finding 4:",
            "Finding 4: Operational performance evidence shows a substantial recovery from minute-level planner "
            "latency to second-level generation times, making respondent testing more practical.",
        ),
        (
            "[Finding 5:",
            "Finding 5: Guarded ML-assist ranking has credible artifact evidence and improves candidate ordering "
            "metrics, but final user acceptance, CS/IT validation, and health consultation conclusions remain pending.",
        ),
        (
            "Write conclusions directly",
            "The following conclusions are drawn only from the findings above and avoid clinical outcome claims.",
        ),
        (
            "[Conclusion 1:",
            "Conclusion 1: PCOSina achieved the core system-development objective by implementing a local-first "
            "Android meal-planning workflow with profile capture, meal planning, grocery support, progress logging, "
            "and support screens.",
        ),
        (
            "[Conclusion 2:",
            "Conclusion 2: The two-stage planner is technically feasible when described accurately as deterministic "
            "filtering followed by OR-Tools CP-SAT integer optimization, not as a pure MILP solver.",
        ),
        (
            "[Conclusion 3:",
            "Conclusion 3: Pantry-aware grocery guidance and budget estimation support practical planning decisions, "
            "but pantry use must be presented as a reward and guidance mechanism unless future work implements full "
            "hard pantry feasibility.",
        ),
        (
            "[Conclusion 4:",
            "Conclusion 4: Local-first continuity is supported for saved plans, pantry data, grocery outputs, profile "
            "information, and progress logs, although new plan generation may still depend on backend availability in "
            "the current architecture.",
        ),
        (
            "[Conclusion 5:",
            "Conclusion 5: Guarded ML-assist can improve Stage 1 candidate ranking when bounded by deterministic "
            "filters and CP-SAT final selection authority.",
        ),
        (
            "[Conclusion 6:",
            "Conclusion 6: PCOSina should be understood only as a wellness decision-support system and not as a "
            "diagnostic, treatment, prescription, or medical-device application.",
        ),
        (
            "[Conclusion 7:",
            "Conclusion 7: Final conclusions about user acceptance, ISO/IEC 25010 quality ratings, and health expert "
            "approval must wait for actual respondent and expert-validation data.",
        ),
        (
            "These are near-term improvements",
            "The following recommendations are near-term system improvements based on verified implementation gaps "
            "and thesis-defense risks.",
        ),
        (
            "These are larger improvements",
            "The following recommendations are larger improvements beyond the current thesis scope.",
        ),
        (
            "Limitations protect",
            "The following limitations are retained to prevent overclaiming and to keep the study aligned with the "
            "implemented PCOSina system.",
        ),
    ]
    for prefix, text in replacements:
        set_first_starting(doc, prefix, text)

    set_exact(doc, "Suggested Appendices Checklist", "Appendix Support Checklist")
    set_exact(
        doc,
        "This page may be moved to your actual appendix section or deleted after use.",
        "The appendix list below identifies support files that should be attached or referenced to defend the Chapter 4-6 evidence.",
    )

    figure_notes = {
        2: "Pending actual screenshot capture from the final Android release build: login, legal acceptance, and startup flow.",
        3: "Pending actual screenshot capture from the final Android release build: profile setup, PCOS-related inputs, and goal selection.",
        4: "Pending actual screenshot capture from the final Android release build: meal plan generation and weekly output.",
        5: "Pending actual screenshot capture from the final Android release build: pantry-aware grocery list and recipe details.",
        6: "Pending actual screenshot capture from the final Android release build: progress logging, reminders, settings, and support screens.",
        11: "Pending functional screenshot: profile setup validation evidence.",
        12: "Pending functional screenshot: successful meal-plan generation evidence.",
        13: "Pending functional screenshot: grocery and pantry output evidence.",
        18: "Benchmark chart asset available at output/doc/assets/comment12/figure_4_5_scenario_runtime_and_quality_comparison.png. Numeric values are supplied in Tables 4.8 and 4.9.",
        19: "Benchmark acceptance/safety chart asset available at output/doc/assets/comment12/figure_4_6_acceptance_criteria_matrix.png. Numeric hard-violation values are supplied in Table 4.9.",
        22: "Pending actual screenshot capture: pantry input and pantry-aware grocery output from final release build.",
    }
    for index, note in figure_notes.items():
        fill_one_cell(doc.tables[index], note)

    fill_table(
        doc.tables[0],
        ["Evidence", "Where to Insert", "Status"],
        [
            ["Final app release/version/commit hash", "4.1 and Appendix", f"Completed: Android {version_name} ({version_code}), branch {branch}, commit {commit}"],
            ["System screenshots", "4.2 figures and Appendix", "Pending final release screenshot capture"],
            ["CS/IT expert validation summary", "4.3", "Pending actual validator forms or ratings"],
            ["OBGYN/health consultation summary", "4.3", "Pending actual consultation evidence"],
            ["Functional test case results", "4.4", f"Completed/partially completed: {validation_cases} validation cases listed"],
            ["Algorithm benchmark tables", "4.5", "Completed for available benchmark report; rerun after final release if needed"],
            ["Pantry/grocery/budget evidence", "4.6", "Completed for sample computation and policy exports"],
            ["Operational loading/performance evidence", "4.7", "Completed for April 2026 recovery log; final device benchmark recommended"],
            ["Guarded ML ranking evidence", "4.8", "Completed: model artifact and training metrics available"],
            ["Survey weighted means and Cronbach alpha", "4.9", "Pending raw respondent dataset"],
        ],
    )

    fill_table(
        doc.tables[1],
        ["Evidence Source", "Data Collected", "Purpose in Chapter 4", "Status/Location"],
        [
            ["Version control and Android build config", f"Branch {branch}, commit {commit}, version {version_name} ({version_code})", "Defines the evaluated build", "git metadata and app/build.gradle.kts"],
            ["Thesis validation evidence pack", f"{validation_cases} validation cases; category counts: {category_summary}", "Supports formula, planner, safety, and output validation", "docs/thesis_validation"],
            ["Recipe inventory export", "1,114 bundled recipes; avg 357.55 kcal; avg 10.33 g protein; avg 24.58 min", "Shows actual dataset size and planning pool", "docs/thesis_validation/03_ACTUAL_SYSTEM_DATA_EXPORTS/recipe_inventory.csv"],
            ["Policy and constraint exports", "Nutrition ranges, 7-day/3-meal horizon, hard and soft planner constraints", "Grounds the optimization model in source-backed policy", "policy_values.csv and planner_constraints.csv"],
            ["Benchmark report", "Greedy, CP-SAT, and generated-proxy runtime, violation, and deviation metrics", "Supports algorithm comparison", "backend/comment5_algorithm_tradeoff_report.json"],
            ["Operational roadmap reports", "Before/after planner latency observations", "Supports loading and performance discussion", "docs/roadmap/team_system_evolution_report_2026-04-14.md"],
            ["ML training artifacts", "74,112 rows, 46 features, AUC/NDCG/MAP metrics", "Supports guarded ML-assist ranking discussion", "ml/offline_training/artifacts/model_v1"],
        ],
    )

    fill_table(
        doc.tables[7],
        ["Module/Screen", "Implemented Function", "Evidence", "Status"],
        [
            ["Login, legal gate, and startup", "Account entry, disclaimer/legal boundary, startup flow", "LoginScreen, SignUpScreen, SplashScreen, auth tests", "Implemented; screenshot pending"],
            ["Profile setup", "Age, height, weight, activity, restrictions, allergies, symptoms, pantry and preferences", "UserProfile model, profile screens, ProfileStepOneValidationUiTest", "Implemented"],
            ["Health computation", "BMI, BMI category, BMR, TDEE, target calories", "HealthMetrics.kt and HealthMetricsTest", "Implemented with parity notes"],
            ["Meal plan generation", "7-day, 3-meal weekly plan request and response handling", "MealPlanRepository, MealPlanViewModel, backend planner", "Implemented"],
            ["Planner safety filtering", "Allergy, restriction, cooking-time, budget, meal-slot, repeat, and no-safe-plan behavior", "backend/services/meal_planner.py and tests", "Implemented"],
            ["Grocery and pantry", "Aggregated grocery list, price estimates, pantry-aware display", "GroceryAggregation.kt, PriceCatalog.kt, GroceryRefinedScreen", "Implemented with known pantry-matching caveat"],
            ["Progress logging", "Meal check-ins, reflections, progress recovery", "ProgressViewModel, progress tests", "Implemented"],
            ["Settings, reminders, support", "Notification controls, profile maintenance, support flow", "SettingsScreen, notification/support tests", "Implemented"],
            ["ML-assist ranking", "Optional Stage 1 candidate ranking after safety filtering", "ml_ranker.py and model_v1 training metrics", "Implemented as guarded assistive layer"],
        ],
    )

    fill_table(
        doc.tables[8],
        ["Area Evaluated", "Validator Comment/Concern", "Action Taken or Planned", "Defense Impact", "Status"],
        [
            ["CS/IT validation dataset", "No completed CS/IT validator forms found in repository", "Collect signed/dated validation summary", "Prevents unsupported expert-validation claim", "Pending"],
            ["ISO/IEC 25010 usage", "Instrument draft exists but no final ratings found", "Use as evaluation framework, not certification", "Avoids false ISO-compliance claim", "Pending results"],
            ["Algorithm wording", "Repo evidence shows OR-Tools CP-SAT, not pure MILP", "Use deterministic filtering plus CP-SAT integer optimization", "Improves technical accuracy", "Done"],
            ["Offline-first wording", "Android uses DataStore and encrypted local artifacts", "Avoid claiming Room/SQLite if not implemented", "Keeps architecture description truthful", "Done"],
            ["Planner hard constraints", "Hard allergy/restriction/budget/repeat behavior is source-backed", "Attach test and constraint evidence", "Supports defense of safety boundaries", "Ready for review"],
            ["Pantry feasibility", "Pantry is rewarded, not hard-enforced", "Use pantry-aware support wording", "Avoids overstating planner capability", "Done"],
            ["Maintainability", "Separation exists across UI, repository, backend planner, policy, and ML artifacts", "Ask CS/IT validator to review module boundaries", "Supports maintainability discussion", "Pending expert signoff"],
        ],
    )

    fill_table(
        doc.tables[9],
        ["Health/Medical Area", "Consultation Finding", "Implication for PCOSina", "Action Taken or Thesis Boundary"],
        [
            ["Health consultation dataset", "No completed OBGYN/nutritionist consultation form found", "Cannot claim expert health approval yet", "Pending actual consultation evidence"],
            ["System scope", "Repository identity frames PCOSina as wellness decision-support", "No diagnosis, treatment, prescription, or medical-device claim", "Retain disclaimer and scope boundary"],
            ["BMI/BMR/TDEE", "Formula evidence exists in Android and backend code", "Useful for planning context only", "State that values are estimates, not prescriptions"],
            ["PCOS-oriented macro ratios", "Backend implements mild/moderate/severe ratio mapping", "Supports planner targeting but needs expert review", "Mark nutrition validation pending"],
            ["Allergy/restriction safety", "Hard exclusion logic is source-backed", "Supports safety filtering for declared user inputs", "Do not claim clinical allergy diagnosis"],
            ["Nutrient limitations", "Sodium and sugar penalties exist in backend policy, but Android DTO parity is incomplete", "Avoid overclaiming full nutrient display coverage", "Recommend nutritionist review and DTO parity work"],
        ],
    )

    fill_table(
        doc.tables[10],
        ["Test Case", "Feature Tested", "Expected Result", "Actual Result", "Status", "Evidence"],
        [
            ["TC-BMI-001", "BMI calculation", "60 kg, 165 cm -> 22.04 Normal", "HealthMetrics unit test exists", "Implemented", "HealthMetricsTest"],
            ["TC-CAL-001", "Weight-loss calorie target", "TDEE minus 500 kcal/day", "HealthMetrics unit test exists", "Implemented", "targetCalories_weightLoss_usesDeficit"],
            ["TC-MACRO-002", "Moderate insulin resistance macros", "0.28 protein, 0.35 carbs, 0.37 fat", "Backend test exists", "Implemented", "test_macro_ratios_moderate"],
            ["TC-ALLERGY-001", "Fish allergy filter", "Bangus/tilapia/galunggong/salmon/tuna excluded", "Backend allergen-family test exists", "Implemented", "test_allergy_filter_blocks_descendant_ingredients"],
            ["TC-BUDGET-001", "Weekly budget hard cap", "No solution if all combinations exceed budget", "Backend unit test exists", "Implemented", "test_solve_meal_plan_enforces_budget_as_hard_cap"],
            ["TC-HOUSEHOLD-001", "Household scaling", "Cost and grocery estimates increase monotonically", "Backend tests exist", "Implemented", "test_estimate_cost_scales_monotonically_with_household_size"],
            ["TC-REPEAT-001", "Recipe repetition limit", "Repeat overflow blocked", "Swap and planner tests exist", "Implemented", "test_build_swap_candidates_blocks_current_recipe_and_repetition_overflow"],
            ["TC-NOSAFE-001", "No-safe-plan response", "Structured no-safe-plan response with guidance", "Integration tests exist", "Implemented", "test_no_safe_plan_response_contract"],
            ["TC-GROCERY-001", "Grocery aggregation", "Matching grocery names merge and prices sum", "Android unit test exists", "Implemented", "GroceryAggregationTest"],
            ["TC-SCHEMA-001", "API schema version", "Schema mismatch path is surfaced", "Backend/API contract test exists", "Implemented", "test_schema_contract"],
        ],
    )

    fill_table(
        doc.tables[14],
        ["Issue", "Observed Effect", "Corrective Action", "Current Status"],
        [
            ["Android/backend calorie rounding", "Android preview can differ from backend target by 1 kcal", "Document round() vs int() behavior and use backend as planner authority", "Documented"],
            ["Pantry token vs grocery-name matching", "Backend pantry overlap may not match Android grocery item names exactly", "Describe pantry as guidance/reward and add UI pantry matching tests", "Known gap"],
            ["Separate price catalogs", "Android and backend estimates may differ", "State price estimates as ranges and identify layer-specific source", "Known gap"],
            ["OR-Tools import blocked on this machine", "Full local solve could not be rerun in this environment", "Use static source evidence and rerun on an unblocked machine", "Pending rerun"],
            ["Human evaluation data absent", "Cannot compute weighted means or Cronbach alpha", "Collect and clean respondent/expert datasets before final defense", "Pending"],
        ],
    )

    fill_table(
        doc.tables[15],
        ["Scenario", "Label", "Description", "Main Constraint Tested"],
        [
            ["S1", "Base", "Standard profile with no extra restrictions", "General hard-safe feasibility and nutrition deviation"],
            ["S2", "BudgetNoPork", "No Pork restriction with PHP 1,000 weekly budget", "Budget cap plus restriction filtering"],
            ["S3", "VegLactose", "Vegetarian and Lactose Intolerant restrictions", "Restrictive candidate pool and meal-slot compatibility"],
            ["S4", "PantryConstrained", "No Beef profile with pantry items egg, garlic, onion, tomato, chicken, rice, tofu", "Pantry reward and recipe practicality"],
            ["S5", "ResilienceProxy300", "Stress/proxy scenario for recovery behavior", "Runtime and non-silent fallback behavior"],
        ],
    )

    fill_table(
        doc.tables[16],
        ["Scenario", "Time Baseline (s)", "Time Proposed (s)", "Hard-Safe Baseline", "Hard-Safe Proposed", "Diversity / Pantry U"],
        [[f"{sid} - {label}", f"{tb:.4f}", f"{tp:.4f}", fb, fp, div] for sid, label, tb, tp, fb, fp, div in scenario_runtime_rows],
    )

    fill_table(
        doc.tables[17],
        ["Scenario", "Calorie Dev. Baseline", "Calorie Dev. Proposed", "Macro Dev. Baseline", "Macro Dev. Proposed", "Hard Violations"],
        [[sid, f"{cb:.2f}", f"{cp:.2f}", f"{mb:.2f}%", f"{mp:.2f}%", hv] for sid, cb, cp, mb, mp, hv in scenario_quality_rows],
    )

    fill_table(
        doc.tables[20],
        ["Ingredient", "Required Amount", "Available Pantry Amount", "Grocery Deficit", "Interpretation"],
        [
            ["egg", "2 pieces", "2 pieces if pantry entry is exact", "0 pieces", "Covered when exact pantry item exists"],
            ["white rice cooked", "2/3 cup", "generic pantry entry: rice", "UI may still show item", "Backend token overlap counts rice, but UI name matching may require manual check-off"],
            ["tomato chopped", "1/3 cup", "generic pantry entry: tomato", "UI may still show item", "Backend token overlap counts tomato; grocery UI may need a more specific pantry name"],
            ["onion chopped", "2 tbsp", "generic pantry entry: onion", "UI may still show item", "Backend token overlap counts onion; grocery UI may need a more specific pantry name"],
            ["talong roasted", "1 medium", "none in sample pantry", "1 medium", "Buy or add exact pantry item"],
        ],
    )

    fill_table(
        doc.tables[21],
        ["Category/Item", "Estimated Price Range", "Quantity", "Estimated Subtotal", "Budget Effect"],
        [
            ["Egg price rule", "PHP 7 per piece", "2 pieces", "PHP 14", "Small part of sample PHP 1,500 weekly budget"],
            ["Sample recipe ph_qk_052", "Backend estimate PHP 55", "1 household-size serving", "PHP 55", "Within sample weekly budget if repeated reasonably"],
            ["Backend recipe-cost policy", "Recipe total * 0.75, clamped PHP 30 to PHP 450", "per recipe estimate", "varies by recipe", "Prevents extreme estimates but remains approximate"],
            ["Weekly budget cap", "PHP 1,500 sample budget", "21 meal slots", "sum of selected recipe costs", "Backend enforces cap when budget is resolved"],
            ["Market-price caveat", "Depends on location/date/store", "all ingredients", "not exact", "Use estimates for planning, not accounting"],
        ],
    )

    fill_table(
        doc.tables[23],
        ["Metric", "Before Fix", "After Fix", "Interpretation"],
        [
            ["Observed slow runs", "207,816 ms and 188,249 ms", "6,486 ms to 13,674 ms", "Recent successful runs returned to practical waiting time"],
            ["Average generation time", f"{before_avg:.2f} ms ({before_avg / 1000:.2f} s)", f"{after_avg:.2f} ms ({after_avg / 1000:.2f} s)", "Average latency dropped from minutes to seconds"],
            ["Best successful generation", "Not below 188,249 ms in cited window", f"{min(after_runs_ms):,} ms", "Best recent run is under seven seconds"],
            ["Worst successful generation", "207,816 ms", f"{max(after_runs_ms):,} ms", "Upper bound dropped sharply"],
            ["Improvement factor", "baseline", f"{improvement:.2f}x faster", "Performance issue was materially reduced"],
        ],
    )

    fill_table(
        doc.tables[24],
        ["Metric", "Observed Value", "Interpretation"],
        [
            ["Dataset version", training["dataset_manifest"]["datasetVersion"], "Offline training data snapshot"],
            ["Model", training["model_name"], "Stage 1 candidate ranker"],
            ["Rows", f"{quality['rows']:,}", "Candidate-feature rows used for training/evaluation"],
            ["Positive / negative labels", f"{quality['positives']:,} / {quality['negatives']:,}", f"Positive rate {quality['positive_rate']:.4f}"],
            ["Unique requests / users", f"{quality['unique_requests']} / {quality['unique_users']}", "Request-level and user-level coverage"],
            ["Train / val / test rows", f"{splits['train']:,} / {splits['val']:,} / {splits['test']:,}", "Leakage-aware split by request id"],
            ["Feature count", str(training["feature_count"]), "Structured features only; no free-text generation"],
            ["Test AUC / log loss", f"{training['test_binary_metrics']['auc']:.4f} / {training['test_binary_metrics']['logloss']:.4f}", "Binary ranking-signal quality"],
            ["Test NDCG@10 / MAP@10", f"{training['test_ranking_metrics']['ndcg@10']:.4f} / {training['test_ranking_metrics']['map@10']:.4f}", "Top-list ranking quality"],
            ["Heuristic baseline NDCG@10 / MAP@10", f"{training['test_baseline_ranking_metrics']['ndcg@10']:.4f} / {training['test_baseline_ranking_metrics']['map@10']:.4f}", "ML improves candidate ordering but remains assistive"],
        ],
    )

    fill_table(
        doc.tables[25],
        ["Respondent Group", "Target Count", "Valid Responses", "Percentage"],
        [
            ["Primary users", "To be set by sampling plan", "Pending", "Pending"],
            ["CS/IT evaluators", "To be set by validation plan", "Pending", "Pending"],
            ["Health/nutrition consultant", "To be set by consultation plan", "Pending", "Pending"],
            ["Total", "Pending", "Pending", "Pending"],
        ],
    )

    fill_table(
        doc.tables[26],
        ["Construct/Section", "Number of Items", "Cronbach's Alpha", "Interpretation"],
        [
            ["TAM - Perceived Usefulness", "Pending", "Pending", "Requires respondent matrix"],
            ["TAM - Perceived Ease of Use", "Pending", "Pending", "Requires respondent matrix"],
            ["TAM - Behavioral Intention", "Pending", "Pending", "Requires respondent matrix"],
            ["ISO/IEC 25010", "Pending", "Pending", "Requires respondent or expert ratings"],
            ["Overall instrument", "Pending", "Pending", "Compute only after final data cleaning"],
        ],
    )

    fill_table(
        doc.tables[27],
        ["TAM Construct", "Weighted Mean", "Verbal Interpretation", "Result Summary"],
        [
            ["Perceived Usefulness", "Pending", "Pending", "No raw respondent data found"],
            ["Perceived Ease of Use", "Pending", "Pending", "No raw respondent data found"],
            ["Attitude Toward Use", "Pending", "Pending", "No raw respondent data found"],
            ["Behavioral Intention", "Pending", "Pending", "No raw respondent data found"],
            ["Overall TAM", "Pending", "Pending", "Compute after survey collection"],
        ],
    )

    fill_table(
        doc.tables[28],
        ["Quality Characteristic", "Weighted Mean", "Verbal Interpretation", "Evidence/Comment"],
        [
            ["Functional Suitability", "Pending", "Pending", "Implementation evidence exists; respondent score pending"],
            ["Performance Efficiency", "Pending", "Pending", "Operational latency evidence exists; respondent score pending"],
            ["Compatibility / Offline Support", "Pending", "Pending", "Saved local data evidence exists; respondent score pending"],
            ["Usability", "Pending", "Pending", "Screens and UI tests exist; respondent score pending"],
            ["Reliability", "Pending", "Pending", "No-safe-plan and recovery tests exist; respondent score pending"],
            ["Security / Privacy", "Pending", "Pending", "Privacy-safe logging evidence exists; respondent score pending"],
            ["Maintainability", "Pending", "Pending", "Architecture evidence exists; CS/IT expert score pending"],
        ],
    )

    fill_table(
        doc.tables[29],
        ["Feedback Theme", "Summary of Comments", "System Implication", "Action/Recommendation"],
        [
            ["Usability", "Pending open-ended responses", "Cannot infer theme yet", "Insert coded themes after survey"],
            ["Meal relevance", "Pending open-ended responses", "Cannot infer theme yet", "Insert coded themes after survey"],
            ["Grocery usefulness", "Pending open-ended responses", "Cannot infer theme yet", "Insert coded themes after survey"],
            ["Performance/loading", "Pending open-ended responses", "Cannot infer theme yet", "Insert coded themes after survey"],
            ["Privacy/trust", "Pending open-ended responses", "Cannot infer theme yet", "Insert coded themes after survey"],
        ],
    )

    fill_table(
        doc.tables[30],
        ["General Objective Component", "Chapter 4 Evidence", "Finding"],
        [
            ["Android mobile application", "Implemented workflow, screens, models, repositories, and tests", "Substantially achieved"],
            ["Deterministic meal planner", "Filtering rules, CP-SAT model, constraints, objective summary, benchmark report", "Achieved with accurate CP-SAT wording"],
            ["Pantry and grocery support", "Pantry overlap, grocery aggregation, sample manual computation", "Achieved as pantry-aware support; not hard pantry feasibility"],
            ["Offline-first behavior", "DataStore/local artifact evidence and offline-read guarantees", "Partially achieved; new generation may require backend"],
            ["Evaluation", "Tests, validation cases, benchmark, ML artifacts, performance report", "Partially achieved; human evaluation pending"],
        ],
    )

    fill_table(
        doc.tables[31],
        ["Specific Objective", "Evidence from Chapter 4", "Finding", "Interpretation"],
        [
            ["Objective 1: Implement the two-stage meal-planning algorithm", "Planner constraints, objective terms, benchmark report", "Achieved", "Use deterministic filtering plus CP-SAT wording"],
            ["Objective 2: Personalize based on restrictions, preferences, and PCOS-oriented nutrition targets", "Profile fields, allergy/restriction tests, BMI/BMR/TDEE/macros", "Achieved with nutrition-review caveat", "Formula estimates require expert review"],
            ["Objective 3: Integrate pantry inventory and grocery recommendations", "Pantry overlap count, grocery aggregation, price catalog", "Partially achieved", "Pantry is soft/rewarded and UI matching has caveats"],
            ["Objective 4: Support offline-first mobile use", "Saved local data, preferences, recovery policy, offline-read guarantees", "Partially achieved", "Saved data works offline; new plans may need backend"],
            ["Objective 5: Evaluate performance, usability, cultural relevance, and computational efficiency", "Benchmark, performance logs, test inventory, recipe dataset", "Partially achieved", "Respondent usability and expert ratings pending"],
        ],
    )

    fill_table(
        doc.tables[32],
        ["Research Question", "Summary Finding"],
        [
            ["RQ1: How can PCOSina generate nutritionally guided weekly meal plans?", "By computing profile-based targets, filtering unsafe candidates, and assigning recipes through CP-SAT optimization over 7 days and 3 meals per day."],
            ["RQ2: How does the system handle allergies, restrictions, and preferences?", "Allergy and restriction rules are hard filters before ranking and solving; ML cannot override them."],
            ["RQ3: How does pantry and budget information affect planning?", "Budget can be a hard cap when resolved; pantry overlap is rewarded and informs grocery guidance but is not hard-enforced."],
            ["RQ4: How effective and usable is the application?", "Implementation, tests, and performance evidence are available, but final usability conclusions require respondent data."],
            ["RQ5: What is the role of ML in the system?", "ML is limited to guarded Stage 1 ranking over already-safe candidates; CP-SAT remains final authority."],
        ],
    )

    fill_table(
        doc.tables[33],
        ["Appendix", "Suggested Content"],
        [
            ["Appendix A", "Version evidence: branch, commit, app version, APK build record"],
            ["Appendix B", "System screenshots for Figures 4.1 to 4.5 and 4.11"],
            ["Appendix C", "docs/thesis_validation/05_VALIDATION_TEST_CASES.csv"],
            ["Appendix D", "docs/thesis_validation/06_SAMPLE_MANUAL_COMPUTATION.md"],
            ["Appendix E", "docs/thesis_validation/14_SYSTEM_COMPUTATIONS_AND_FORMULAS_CENTRALIZED.docx plus planner constraints, policy values, and objective summary exports"],
            ["Appendix F", "backend/comment5_algorithm_tradeoff_report.json benchmark report"],
            ["Appendix G", "ML training metrics and reason-feedback artifacts"],
            ["Appendix H", "Operational performance roadmap reports"],
            ["Appendix I", "CS/IT expert validation forms and signed summaries when collected"],
            ["Appendix J", "Respondent survey matrix, weighted means, and Cronbach alpha computation when collected"],
        ],
    )

    doc.save(DOCX)
    print(f"UPDATED {DOCX}")


if __name__ == "__main__":
    main()
