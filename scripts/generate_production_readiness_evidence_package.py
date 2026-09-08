#!/usr/bin/env python
from __future__ import annotations

import json
import os
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
DATE_LABEL = "2026-09-07"
EVIDENCE_DIR = ROOT / "docs" / "production_readiness" / f"evidence_{DATE_LABEL}"
LOG_DIR = EVIDENCE_DIR / "logs"
LOCAL_ARTIFACT_DIR = EVIDENCE_DIR / "artifacts"
BENCHMARK_DIR = ROOT / "benchmarks" / "reports" / f"production_readiness_{DATE_LABEL}"
OUTPUT_DIR = ROOT / "output" / "doc"


def _read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _latest(pattern: str) -> Path | None:
    matches = sorted(BENCHMARK_DIR.glob(pattern), key=lambda p: p.stat().st_mtime if p.exists() else 0)
    return matches[-1] if matches else None


def _git_status() -> str:
    proc = subprocess.run(
        ["git", "status", "--short", "--branch"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    return (proc.stdout or "") + (proc.stderr or "")


def _status_from_log(path: Path, success_text: str) -> str:
    text = _read_text(path)
    if success_text in text:
        return "Pass"
    if text.strip():
        return "Fail/Blocked"
    return "Not run"


def _rel(path: Path | str | None) -> str:
    if not path:
        return ""
    candidate = Path(path)
    try:
        if candidate.is_absolute():
            return str(candidate.relative_to(ROOT)).replace("\\", "/")
        return str(candidate).replace("\\", "/")
    except ValueError:
        return str(candidate).replace("\\", "/")


def _build_manifest() -> dict[str, Any]:
    planner_report = BENCHMARK_DIR / "planner_profiles.local_host.json"
    planner = _read_json(planner_report)
    suite = planner.get("suite") or {}
    local_metrics = BENCHMARK_DIR / "go_live_metrics.local_host.json"
    nutrition = _read_json(LOCAL_ARTIFACT_DIR / "catalog_nutrition_readiness.json")
    staging_report_path = _latest("rollout_readiness_batch.staging.*.json")
    production_report_path = _latest("rollout_readiness_batch.production.*.json")
    staging_report = _read_json(staging_report_path) if staging_report_path else {}
    production_report = _read_json(production_report_path) if production_report_path else {}

    pass_items = [
        {
            "area": "Backend release/security tests",
            "status": _status_from_log(LOG_DIR / "backend_release_security_pytest.txt", "50 passed"),
            "evidence": _rel(LOG_DIR / "backend_release_security_pytest.txt"),
            "summary": "Runtime readiness, admin feedback security, operator auth unit policy, async ownership, and ops access tests passed.",
        },
        {
            "area": "Backend planner/ML regression tests",
            "status": _status_from_log(LOG_DIR / "backend_planner_ml_pytest.txt", "120 passed"),
            "evidence": _rel(LOG_DIR / "backend_planner_ml_pytest.txt"),
            "summary": "Planner, ranker, dataset builder, reason-feedback builder, and ML readiness script unit tests passed.",
        },
        {
            "area": "No-safe-plan contract",
            "status": _status_from_log(LOG_DIR / "backend_no_safe_plan_contract_pytest.txt", "11 passed"),
            "evidence": _rel(LOG_DIR / "backend_no_safe_plan_contract_pytest.txt"),
            "summary": "Backend contract tests passed for no-safe-plan behavior and hard-constraint handling.",
        },
        {
            "area": "Backend sync replay and reinstall recovery",
            "status": _status_from_log(LOG_DIR / "backend_sync_replay_pytest.txt", "6 passed"),
            "evidence": _rel(LOG_DIR / "backend_sync_replay_pytest.txt"),
            "summary": "Local deterministic sync conflict/replay and reinstall recovery tests passed.",
        },
        {
            "area": "Android debug unit tests",
            "status": _status_from_log(LOG_DIR / "android_debug_unit_tests_final.txt", "BUILD SUCCESSFUL"),
            "evidence": _rel(LOG_DIR / "android_debug_unit_tests_final.txt"),
            "summary": "Full Android debug unit test suite passed after fixing the progress weekly-spend inline validation policy.",
        },
        {
            "area": "Mobile ML event coverage",
            "status": _status_from_log(LOG_DIR / "mobile_ml_event_coverage.txt", "MOBILE ML EVENT COVERAGE PASSED"),
            "evidence": _rel(LOG_DIR / "mobile_ml_event_coverage.txt"),
            "summary": "Mobile event coverage gate passed for the required ML event hooks.",
        },
        {
            "area": "Local planner benchmark",
            "status": "Pass" if str(suite.get("thresholdStatus") or "").lower() == "pass" else "Fail/Blocked",
            "evidence": _rel(planner_report),
            "summary": (
                f"{suite.get('passedRuns', 0)}/{suite.get('totalRuns', 0)} local runs passed; "
                f"p95={suite.get('p95RuntimeMs')} ms; max={suite.get('maxRuntimeMs')} ms; "
                f"failed={suite.get('failedRuns', 0)}."
            ),
        },
        {
            "area": "Local go-live metrics format gate",
            "status": _status_from_log(LOG_DIR / "go_live_gates_local_host.txt", "GO-LIVE GATE CHECK PASSED"),
            "evidence": _rel(local_metrics),
            "summary": "The explicit local-host metrics file passes the go-live gate checker. It is not production clearance.",
        },
        {
            "area": "Local canary drill",
            "status": _status_from_log(LOG_DIR / "canary_drill_local_receipt_check.txt", "CANARY DRILL RECEIPT CHECK PASSED"),
            "evidence": _rel(BENCHMARK_DIR / "canary_drill_receipt.local_host.json"),
            "summary": "Local canary drill and receipt contract passed without live webhook delivery requirements.",
        },
        {
            "area": "Local rollout-readiness batch",
            "status": "Pass" if staging_report.get("status") == "ok" else "Fail/Blocked",
            "evidence": _rel(staging_report_path),
            "summary": "Permissive local staging rollout batch passed; skipped external Redis/webhook receipts were not required in this mode.",
        },
        {
            "area": "Local queue throughput",
            "status": _status_from_log(LOG_DIR / "queue_memory_throughput_check.txt", "QUEUE BROKER THROUGHPUT REPORT CHECK PASSED"),
            "evidence": _rel(BENCHMARK_DIR / "queue_broker_throughput.memory.local.json"),
            "summary": "In-memory queue throughput probe passed and was validated by the checker.",
        },
        {
            "area": "Nutrition readiness",
            "status": "Pass" if nutrition.get("ok") is True else "Fail/Blocked",
            "evidence": _rel(LOCAL_ARTIFACT_DIR / "catalog_nutrition_readiness.json"),
            "summary": (
                f"{nutrition.get('completeNutritionProfileCount', 0)}/{nutrition.get('activeRecipeCount', 0)} active recipes have complete nutrition profiles; "
                f"{nutrition.get('trustedNutritionCount', 0)} trusted; warnings={len(nutrition.get('warnings') or [])}."
            ),
        },
        {
            "area": "DOCX structural validation",
            "status": _status_from_log(LOG_DIR / "docx_structural_validation.txt", "DOCX STRUCTURE OK"),
            "evidence": _rel(LOG_DIR / "docx_structural_validation.txt"),
            "summary": "The generated DOCX opens through python-docx and contains the expected paragraphs and tables.",
        },
    ]

    blocker_items = [
        {
            "area": "Connected Android/emulator proof",
            "importance": "Critical",
            "status": "Blocked",
            "evidence": _rel(LOG_DIR / "adb_devices.txt"),
            "effect_if_done": "Proves the APK launches and core flows work on a real Android runtime.",
            "risk_if_skipped": "Unit tests may pass while manifest, permissions, lifecycle, storage, or ANR issues still break on device.",
            "current_blocker": "ADB preflight passes, but `adb devices -l` lists no attached or authorized device.",
        },
        {
            "area": "Staged multi-device sync validation",
            "importance": "High",
            "status": "Blocked",
            "evidence": _rel(LOG_DIR / "backend_sync_replay_pytest.txt"),
            "effect_if_done": "Proves real staging payloads survive device switching and conflict resolution.",
            "risk_if_skipped": "Saved plans, groceries, profile, or progress can become stale, duplicated, overwritten, or inconsistent across devices.",
            "current_blocker": "Only deterministic local replay passed; no staging devices/accounts/receipts were available in this workspace.",
        },
        {
            "area": "Production Redis/load evidence",
            "importance": "Critical for backend release",
            "status": "Blocked",
            "evidence": _rel(LOG_DIR / "redis_environment_check.txt"),
            "effect_if_done": "Proves queued meal-plan jobs do not disappear, hang, or overload under production broker conditions.",
            "risk_if_skipped": "Users may hit timeouts or lost background plan jobs once real traffic uses Redis-backed queues.",
            "current_blocker": "`PCOSINA_REDIS_URL` is not set, so only memory queue throughput could be tested locally.",
        },
        {
            "area": "Live canary/webhook delivery",
            "importance": "Critical",
            "status": "Blocked",
            "evidence": _rel(LOG_DIR / "canary_webhook_secrets_staging.txt"),
            "effect_if_done": "Proves production alerts and dashboard notifications fire when canary guardrails run.",
            "risk_if_skipped": "A bad rollout could fail silently without operator alerts.",
            "current_blocker": "Staging webhook environment variables are missing in this workspace.",
        },
        {
            "area": "Strict production rollout batch",
            "importance": "Critical",
            "status": "Failed/Blocked",
            "evidence": _rel(production_report_path),
            "effect_if_done": "Gives one auditable production release pass that ties schema, auth, canary, Redis, queue, and sync evidence together.",
            "risk_if_skipped": "Production readiness remains a confidence claim instead of an evidence-backed release decision.",
            "current_blocker": "; ".join(production_report.get("blockers") or ["Strict production rollout report was not found."]),
        },
        {
            "area": "ML sustained canary/drift evidence",
            "importance": "Medium to High",
            "status": "Deferred",
            "evidence": _rel(LOG_DIR / "ml_readiness_shadow_final.txt"),
            "effect_if_done": "Proves ML ranking remains helpful and does not regress key cohorts before widening.",
            "risk_if_skipped": "ML should stay limited to shadow/canary or be disabled; deterministic planning can still release without ML widening.",
            "current_blocker": "Fresh reason-feedback artifact exists, but ML readiness fails because the dataset manifest is missing; attempted local telemetry regeneration did not complete cleanly, and its partial Stage 1 seed rows were cleared.",
        },
        {
            "area": "Clean commit/push state",
            "importance": "High for team delivery",
            "status": "Deferred",
            "evidence": _rel(LOG_DIR / "git_status_before_package.txt"),
            "effect_if_done": "Makes the fixed version reproducible for teammates and CI.",
            "risk_if_skipped": "Fixes and evidence remain local until intentionally staged, reviewed, committed, and pushed.",
            "current_blocker": "The worktree contains many pre-existing modified/untracked files, so a broad commit would mix unrelated work into this evidence package.",
        },
    ]

    return {
        "generatedAtUtc": datetime.now(timezone.utc).isoformat(),
        "dateLabel": DATE_LABEL,
        "decision": "Local evidence package complete; production readiness is not closed until external blockers are resolved.",
        "productionClearance": False,
        "passItems": pass_items,
        "blockerItems": blocker_items,
        "keyArtifacts": {
            "localMetrics": _rel(local_metrics),
            "plannerBenchmark": _rel(planner_report),
            "stagingRolloutBatch": _rel(staging_report_path),
            "productionRolloutBatch": _rel(production_report_path),
            "nutritionReadiness": _rel(LOCAL_ARTIFACT_DIR / "catalog_nutrition_readiness.json"),
            "localCanaryReceipt": _rel(BENCHMARK_DIR / "canary_drill_receipt.local_host.json"),
        },
    }


def _markdown(manifest: dict[str, Any]) -> str:
    lines: list[str] = []
    lines.append("# PCOSINA Production Readiness Evidence Packet")
    lines.append("")
    lines.append(f"Date: {DATE_LABEL}")
    lines.append(f"Generated: {manifest['generatedAtUtc']}")
    lines.append("")
    lines.append("## Release Decision")
    lines.append("")
    lines.append("Status: NOT production-ready yet.")
    lines.append("")
    lines.append(
        "The local app/backend fixes and local release tooling are mostly handled, and this packet contains saved proof files. "
        "However, production readiness is not closed because connected Android, staged multi-device sync, production Redis/load, live webhook/canary delivery, strict production rollout, and clean commit/push proof are still not complete."
    )
    lines.append("")
    lines.append("## What Changed In This Run")
    lines.append("")
    lines.append("- Fixed the Android progress weekly budget card so actual weekly spend uses inline validation copy instead of Toast-style feedback.")
    lines.append("- Preserved the release truth boundary: local-host go-live metrics were generated under a separate local path and the default production metrics path was not overwritten.")
    lines.append("- Generated this evidence folder and ZIP package so the proof is made of files, not just chat text.")
    lines.append("")
    lines.append("## Passing Local Evidence")
    lines.append("")
    lines.append("| Area | Status | Evidence file | Summary |")
    lines.append("|---|---:|---|---|")
    for item in manifest["passItems"]:
        lines.append(f"| {item['area']} | {item['status']} | `{item['evidence']}` | {item['summary']} |")
    lines.append("")
    lines.append("## Remaining Production Blockers")
    lines.append("")
    lines.append("| Area | Importance | Status | Effect if completed | Risk if skipped | Current blocker / evidence |")
    lines.append("|---|---|---:|---|---|---|")
    for item in manifest["blockerItems"]:
        evidence = f"`{item['evidence']}`" if item.get("evidence") else ""
        lines.append(
            f"| {item['area']} | {item['importance']} | {item['status']} | "
            f"{item['effect_if_done']} | {item['risk_if_skipped']} | {item['current_blocker']} Evidence: {evidence} |"
        )
    lines.append("")
    lines.append("## Meaning Of Proof Files")
    lines.append("")
    lines.append(
        "Yes: the proof is stored as files. The important proof files are logs, JSON reports, CSV benchmark outputs, this Markdown report, the DOCX report, and the ZIP archive. "
        "They can be attached to a mentor update, copied into thesis appendices, or used as release evidence."
    )
    lines.append("")
    lines.append(
        "DOCX QA: structural validation passed through python-docx. Visual rendering was attempted, but this machine is missing `soffice` and `pdftoppm`; see `docs/production_readiness/evidence_2026-09-07/logs/docx_render_check.txt`."
    )
    lines.append("")
    lines.append("## Production Closeout Commands Still Needed")
    lines.append("")
    lines.append("Run these only when the missing external environment is available:")
    lines.append("")
    lines.append("```powershell")
    lines.append(".\\scripts\\run_connected_android_tests.ps1")
    lines.append("python scripts\\run_rollout_readiness_batch.py --environment production --metrics benchmarks\\reports\\go_live_metrics.json --require-live-webhooks --require-redis --require-external-canary-receipt --require-external-queue-throughput-report")
    lines.append("```")
    lines.append("")
    lines.append("Before the strict production batch, configure real production/staging secrets and supply external canary and Redis throughput receipts.")
    lines.append("")
    lines.append("## Final Position")
    lines.append("")
    lines.append(
        "This repo now has a complete local evidence packet for the work that can be proven on this machine. "
        "It should be described as locally validated and release-candidate work, not fully production-ready, until the external blockers are resolved and archived."
    )
    lines.append("")
    return "\n".join(lines)


def _write_docx(markdown_path: Path, manifest: dict[str, Any]) -> Path | None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except Exception:
        return None

    out = OUTPUT_DIR / f"PCOSINA_Production_Readiness_Evidence_Packet_{DATE_LABEL}.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PCOSINA Production Readiness Evidence Packet")
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph(f"Date: {DATE_LABEL}")
    doc.add_paragraph(f"Generated: {manifest['generatedAtUtc']}")
    doc.add_paragraph("Status: NOT production-ready yet.")
    doc.add_paragraph(
        "Local fixes and local release evidence are packaged. Production readiness is still blocked by external device, staging, Redis, live webhook/canary, strict rollout, and clean commit/push evidence."
    )

    doc.add_heading("Passing Local Evidence", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Area", "Status", "Evidence file", "Summary"]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for item in manifest["passItems"]:
        cells = table.add_row().cells
        cells[0].text = item["area"]
        cells[1].text = item["status"]
        cells[2].text = item["evidence"]
        cells[3].text = item["summary"]

    doc.add_heading("Remaining Production Blockers", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Area", "Importance", "Status", "Effect if done", "Risk/current blocker"]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for item in manifest["blockerItems"]:
        cells = table.add_row().cells
        cells[0].text = item["area"]
        cells[1].text = item["importance"]
        cells[2].text = item["status"]
        cells[3].text = item["effect_if_done"]
        cells[4].text = f"{item['risk_if_skipped']} Current: {item['current_blocker']} Evidence: {item['evidence']}"

    doc.add_heading("Proof Files", level=1)
    doc.add_paragraph(
        "The proof is stored as files: test logs, JSON reports, benchmark CSV/JSON, this DOCX, the Markdown report, and the ZIP archive."
    )
    doc.add_paragraph(f"Markdown source: {_rel(markdown_path)}")
    doc.save(out)
    return out


def _write_zip(files: list[Path], zip_path: Path) -> None:
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    seen: set[str] = set()
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            if not path.exists() or not path.is_file():
                continue
            arcname = _rel(path)
            if arcname in seen:
                continue
            seen.add(arcname)
            archive.write(path, arcname=arcname)


def main() -> int:
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    LOCAL_ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    BENCHMARK_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    (LOG_DIR / "git_status_before_package.txt").write_text(_git_status(), encoding="utf-8")
    manifest = _build_manifest()

    manifest_path = LOCAL_ARTIFACT_DIR / "production_readiness_evidence_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")

    blocker_manifest = {
        "generatedAtUtc": manifest["generatedAtUtc"],
        "productionClearance": False,
        "blockers": manifest["blockerItems"],
    }
    blocker_path = LOCAL_ARTIFACT_DIR / "production_readiness_blockers.json"
    blocker_path.write_text(json.dumps(blocker_manifest, indent=2, sort_keys=True), encoding="utf-8")

    markdown_path = EVIDENCE_DIR / f"PCOSINA_Production_Readiness_Evidence_Packet_{DATE_LABEL}.md"
    markdown_path.write_text(_markdown(manifest), encoding="utf-8")
    docx_path = _write_docx(markdown_path, manifest)

    package_files: list[Path] = [
        markdown_path,
        manifest_path,
        blocker_path,
        ROOT / "RELEASE_CHECKLIST.md",
        ROOT / "docs" / "roadmap" / "progress_ledger.md",
        ROOT / "docs" / "roadmap" / "ml_progress_ledger.md",
    ]
    if docx_path is not None:
        package_files.append(docx_path)
    package_files.extend(sorted(LOG_DIR.glob("*")))
    package_files.extend(sorted(LOCAL_ARTIFACT_DIR.glob("*")))
    package_files.extend(sorted(BENCHMARK_DIR.glob("*")))

    zip_path = OUTPUT_DIR / f"PCOSINA_Production_Readiness_Evidence_Packet_{DATE_LABEL}.zip"
    _write_zip(package_files, zip_path)

    print(f"MARKDOWN={markdown_path}")
    if docx_path is None:
        print("DOCX=not generated; python-docx is not installed")
    else:
        print(f"DOCX={docx_path}")
    print(f"MANIFEST={manifest_path}")
    print(f"BLOCKERS={blocker_path}")
    print(f"ZIP={zip_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
