#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[3]
THESIS_DIR = REPO_ROOT / "docs" / "thesis_validation"
EXPORT_DIR = THESIS_DIR / "03_ACTUAL_SYSTEM_DATA_EXPORTS"
SOURCE_MD = THESIS_DIR / "14_SYSTEM_COMPUTATIONS_AND_FORMULAS_CENTRALIZED.md"
OUTPUT_DOCX = THESIS_DIR / "14_SYSTEM_COMPUTATIONS_AND_FORMULAS_CENTRALIZED_FINAL.docx"


PINK = "D2466E"
PINK_HEADER = "D85A7F"
PINK_DARK = "A6264B"
PINK_LIGHT = "FCE4EC"
PINK_PALE = "FDF2F6"
GRAY = "505050"
LIGHT_GRAY = "F4F4F4"
TEXT = "222222"
BORDER = "E7A1B4"
WHITE = "FFFFFF"


PURPOSES = {
    "01_SYSTEM_DATA_DICTIONARY.csv": "Input and output field inventory across Android, backend, API, and policy models.",
    "05_VALIDATION_TEST_CASES.csv": "Validation case inventory mapped to actual rules, code paths, and tests.",
    "allergen_family_tokens.csv": "Normalized allergen family tokens used for safety filtering.",
    "allergen_synonyms.csv": "Allergen alias mapping used before hard allergy exclusion.",
    "android_test_inventory.csv": "Android unit and instrumented test inventory.",
    "api_endpoints.csv": "Backend API endpoint inventory and contract surface.",
    "backend_test_inventory.csv": "Backend test inventory.",
    "ingredient_synonyms.csv": "Ingredient alias support for normalization.",
    "optimization_objective_summary.csv": "Live optimizer objective components and source locations.",
    "planner_constraints.csv": "Hard and soft planner constraints with source locations.",
    "policy_values.csv": "Planner policy constants, limits, and environment overrides.",
    "price_catalog_rules_android.csv": "Android-side price catalog and cost rules.",
    "price_catalog_rules_backend.csv": "Backend-side price catalog and cost rules.",
    "recipe_inventory.csv": "Bundled recipe inventory in tabular form.",
    "recipe_inventory.json": "Bundled recipe inventory in JSON form.",
    "stored_local_artifacts.csv": "Local storage artifacts and offline-first evidence.",
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_paragraph_border(paragraph, color: str = PINK_HEADER, size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, *, font: str = "Times New Roman", size: float = 11.0, color: str | None = None, bold=None, italic=None) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    style_specs = [
        ("Title", 20, True, PINK),
        ("Heading 1", 16, True, PINK),
        ("Heading 2", 13, True, GRAY),
        ("Heading 3", 12, True, GRAY),
    ]
    for style_name, size, bold, color in style_specs:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    try:
        code_style = styles.add_style("Formula Block", 1)
    except ValueError:
        code_style = styles["Formula Block"]
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.font.size = Pt(9.2)
    code_style.font.color.rgb = rgb(TEXT)
    code_style.paragraph_format.space_before = Pt(2)
    code_style.paragraph_format.space_after = Pt(2)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("PCOSina Thesis Validation | System Computations and Formulas | Final")
    style_run(footer_run, size=8.5, color=GRAY)


def add_text(paragraph, text: str, *, size: float = 11.0, color: str = TEXT) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            style_run(run, font="Consolas", size=max(8.8, size - 1.0), color=PINK_DARK)
        else:
            run = paragraph.add_run(part)
            style_run(run, size=size, color=color)


def add_banner(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PINK_HEADER)
    set_cell_borders(cell, PINK_HEADER)
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    style_run(run, size=10.5, color=WHITE, bold=True)


def add_info_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = ""
        cells[1].text = ""
        set_cell_shading(cells[0], PINK_LIGHT)
        set_cell_borders(cells[0])
        set_cell_borders(cells[1])
        set_cell_margins(cells[0])
        set_cell_margins(cells[1])
        p0 = cells[0].paragraphs[0]
        r0 = p0.add_run(label)
        style_run(r0, size=9.6, color=PINK_DARK, bold=True)
        p1 = cells[1].paragraphs[0]
        add_text(p1, value, size=9.6, color=TEXT)


def add_pink_table(doc: Document, headers: list[str], rows: list[list[str]], *, font_size: float = 8.8) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    set_table_repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, PINK_HEADER)
        set_cell_borders(cell, PINK_DARK)
        set_cell_margins(cell, top=90, start=80, bottom=90, end=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        style_run(run, size=font_size, color=WHITE, bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_shading(cell, PINK_PALE if row_index % 2 == 0 else WHITE)
            set_cell_borders(cell)
            set_cell_margins(cell, top=70, start=70, bottom=70, end=70)
            p = cell.paragraphs[0]
            add_text(p, str(value), size=font_size, color=TEXT)
    doc.add_paragraph()


def add_formula_block(doc: Document, formula: str, *, code_block: bool = False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PINK_PALE)
    set_cell_borders(cell, BORDER, size="8")
    set_cell_margins(cell, top=95, start=120, bottom=95, end=120)
    paragraph = cell.paragraphs[0]
    paragraph.style = doc.styles["Formula Block"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(formula)
    style_run(run, font="Consolas", size=8.6 if code_block else 9.3, color=TEXT)
    doc.add_paragraph()


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.25 + (0.25 * level))
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    paragraph.paragraph_format.space_after = Pt(2)
    add_text(paragraph, text, size=10.6, color=TEXT)


def add_heading(doc: Document, text: str, level: int) -> None:
    mapped = min(level, 3)
    paragraph = doc.add_heading(text, level=mapped)
    if mapped == 1:
        set_paragraph_border(paragraph)


def add_paragraph(doc: Document, text: str) -> None:
    stripped = text.strip()
    if is_inline_formula(stripped):
        add_formula_block(doc, stripped[1:-1])
        return
    paragraph = doc.add_paragraph()
    add_text(paragraph, stripped)


def is_inline_formula(text: str) -> bool:
    return text.startswith("`") and text.endswith("`") and text.count("`") == 2


def parse_markdown_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    def split_row(line: str) -> list[str]:
        return [cell.strip() for cell in line.strip().strip("|").split("|")]

    headers = split_row(lines[0])
    rows = [split_row(line) for line in lines[2:]]
    width = len(headers)
    clean_rows = []
    for row in rows:
        if len(row) < width:
            row = row + [""] * (width - len(row))
        clean_rows.append(row[:width])
    return headers, clean_rows


def flush_paragraph_buffer(doc: Document, buffer: list[str]) -> None:
    if not buffer:
        return
    add_paragraph(doc, " ".join(item.strip() for item in buffer))
    buffer.clear()


def render_markdown(doc: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]

    buffer: list[str] = []
    code_buffer: list[str] = []
    in_code = False
    table_buffer: list[str] = []

    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph_buffer(doc, buffer)
            if not in_code:
                in_code = True
                code_buffer = []
            else:
                in_code = False
                add_formula_block(doc, "\n".join(code_buffer), code_block=True)
                code_buffer = []
            index += 1
            continue

        if in_code:
            code_buffer.append(raw)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and "---" in lines[index + 1]:
            flush_paragraph_buffer(doc, buffer)
            table_buffer = [line]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_buffer.append(lines[index].rstrip())
                index += 1
            headers, rows = parse_markdown_table(table_buffer)
            add_pink_table(doc, headers, rows)
            continue

        if not stripped:
            flush_paragraph_buffer(doc, buffer)
            index += 1
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph_buffer(doc, buffer)
            level = len(heading.group(1)) - 1
            add_heading(doc, heading.group(2), level)
            index += 1
            continue

        bullet = re.match(r"^(\s*)-\s+(.+)$", line)
        if bullet:
            flush_paragraph_buffer(doc, buffer)
            level = max(0, len(bullet.group(1)) // 2)
            add_bullet(doc, bullet.group(2), level=level)
            index += 1
            continue

        buffer.append(stripped)
        index += 1

    flush_paragraph_buffer(doc, buffer)


def csv_count(path: Path) -> tuple[int, int]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    return max(0, len(rows) - 1), len(rows[0]) if rows else 0


def json_count(path: Path) -> int:
    data = json.loads(path.read_text(encoding="utf-8"))
    return len(data) if hasattr(data, "__len__") else 1


def build_data_coverage_rows() -> list[list[str]]:
    rows: list[list[str]] = []
    for path in [THESIS_DIR / "01_SYSTEM_DATA_DICTIONARY.csv", THESIS_DIR / "05_VALIDATION_TEST_CASES.csv"]:
        row_count, column_count = csv_count(path)
        rows.append([path.name, f"{row_count} rows / {column_count} columns", PURPOSES[path.name]])
    for path in sorted(EXPORT_DIR.iterdir()):
        if path.suffix.lower() == ".csv":
            row_count, column_count = csv_count(path)
            count_text = f"{row_count} rows / {column_count} columns"
        elif path.suffix.lower() == ".json":
            count_text = f"{json_count(path)} entries"
        else:
            continue
        rows.append([f"03_ACTUAL_SYSTEM_DATA_EXPORTS/{path.name}", count_text, PURPOSES.get(path.name, "Current validation export.")])
    return rows


def add_front_matter(doc: Document) -> None:
    add_banner(doc, "FINAL THESIS VALIDATION REFERENCE")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PCOSina System Computations and Formulas")
    style_run(run, size=20, color=PINK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Centralized computation, constraint, data-source, and formula inventory")
    style_run(run, size=11.5, color=GRAY, italic=True)

    try:
        generated = datetime.now(ZoneInfo("Asia/Manila")).strftime("%Y-%m-%d %H:%M %Z")
    except Exception:
        generated = datetime.now().strftime("%Y-%m-%d %H:%M")

    add_info_table(
        doc,
        [
            ("Document status", "Finalized defense copy generated from the centralized Markdown source."),
            ("Generated", generated),
            ("Primary source", "docs/thesis_validation/14_SYSTEM_COMPUTATIONS_AND_FORMULAS_CENTRALIZED.md"),
            ("Evidence basis", "docs/thesis_validation evidence pack, actual source files, policy exports, recipe inventory, and validation artifacts."),
            ("System scope", "Offline-first Filipino-PCOS wellness meal planning decision support."),
            ("Claim boundary", "Not a diagnosis engine, not a medical device, and not a clinical treatment recommendation system."),
        ],
    )

    add_heading(doc, "Defense Framing", 1)
    add_paragraph(
        doc,
        "Use this document as the single reference for explaining how PCOSina computes profile metrics, filters recipes, scores candidates, builds weekly plans, and validates outputs. The strongest defensible wording is that PCOSina uses deterministic profile normalization, deterministic safety and preference filtering, and deterministic 0-1 integer optimization implemented through OR-Tools CP-SAT.",
    )
    for item in [
        "The current live planner is OR-Tools CP-SAT, not a classical MILP solver. Binary decision variables are still shown because they make the optimization model explainable.",
        "ML is optional and assistive only. It may rank or boost candidates, but it must not override allergies, excluded ingredients, budget caps, nutrition bounds, pantry logic, or explicit hard preferences.",
        "Pantry overlap is currently a reward and thresholded shortlist signal, not a complete ingredient-by-ingredient feasibility proof.",
        "Survey and statistician formulas are validation methods. They should not be described as live runtime computations unless the application actually emits those values.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Formula Reading Guide", 1)
    add_pink_table(
        doc,
        ["Notation", "Meaning", "Defense-safe explanation"],
        [
            ["clamp(x, min, max)", "Restricts a value inside an allowed range.", "Used to keep calorie, macro, budget, and policy values inside explicit limits."],
            ["int(x)", "Converts a decimal value to an integer.", "Backend target formulas use integer conversion, which may differ from Android rounding by 1 kcal or gram."],
            ["round(x)", "Rounds to the nearest whole number.", "Android preview values use rounding in some health-metric displays."],
            ["abs(x)", "Absolute value.", "Used for deviation/error terms where only distance from target matters."],
            ["x_(r,s)", "Binary decision variable for assigning recipe r to slot s.", "A value of 1 means the recipe is selected for that meal slot; 0 means it is not selected."],
            ["y_i", "Binary usage variable for recipe i.", "Connects slot assignments to weekly recipe-use and repetition constraints."],
            ["Hard constraint", "A rule that cannot be violated.", "Allergies, restrictions, slot assignment, adjacent duplicates, repetition caps, and weekly budget caps are enforced before or inside optimization."],
            ["Soft term", "A penalty or reward in the objective.", "Nutrition deviation, pantry reward, prep-time burden, and diversity terms guide the optimizer without replacing hard safety rules."],
        ],
        font_size=8.4,
    )

    add_heading(doc, "Evidence Coverage Snapshot", 1)
    add_pink_table(doc, ["Artifact", "Coverage", "Use in defense"], build_data_coverage_rows(), font_size=7.8)

    add_heading(doc, "Main Computation Inventory", 1)


def add_back_matter(doc: Document) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Final Defense Checklist", 1)
    add_pink_table(
        doc,
        ["Defense claim", "Supported by", "Safe wording"],
        [
            ["Profile and nutrition metrics are deterministic.", "HealthMetrics, planner profile preparation, policy constants, and formula inventory.", "The same validated inputs produce the same computed targets except for documented seeded jitter in daily targets."],
            ["Safety filters run before optimization.", "Allergy/restriction normalization, planner constraints export, and restriction failure logic.", "Unsafe or disallowed recipes are removed before solver assignment."],
            ["Optimization is deterministic and explainable.", "CP-SAT model variables, constraints, objective components, and policy exports.", "The implementation is a deterministic 0-1 CP-SAT optimizer with binary assignment variables."],
            ["ML cannot override hard constraints.", "ML ranker guardrails and stage 1 boost formula.", "ML may affect ordering or bounded score boosts only after deterministic constraints are respected."],
            ["Validation artifacts are grounded in actual repo data.", "Evidence index, data dictionary, recipe inventory, test inventories, and source exports.", "The tables cite generated artifacts from the dev branch evidence pack instead of fabricated survey or runtime results."],
            ["Clinical claims are intentionally limited.", "Product identity and claim boundary.", "PCOSina is wellness decision support for meal planning, not a diagnostic or treatment system."],
        ],
        font_size=8.2,
    )


def main() -> None:
    if not SOURCE_MD.exists():
        raise FileNotFoundError(SOURCE_MD)
    doc = Document()
    configure_document(doc)
    add_front_matter(doc)
    render_markdown(doc, SOURCE_MD)
    add_back_matter(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
