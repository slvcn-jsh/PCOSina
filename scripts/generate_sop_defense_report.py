from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "defense" / "PCOSINA_SOP_Defense_Explanation_Report_2026-05-23.docx"


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, bold: bool = False, font_size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for index, header in enumerate(headers):
        cell = hdr.cells[index]
        set_cell_text(cell, header, bold=True, font_size=9)
        shade_cell(cell, "D9EAD3")
        if widths:
            cell.width = Inches(widths[index])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, font_size=8)
            if widths:
                cells[index].width = Inches(widths[index])
    document.add_paragraph()
    return table


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(31, 78, 55)


def add_para(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)


def add_bullets(document: Document, items: list[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.size = Pt(10)


def add_numbered(document: Document, items: list[str]) -> None:
    for index, item in enumerate(items, 1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.32)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(f"{index}. {item}")
        run.font.size = Pt(10)


def add_callout(document: Document, title: str, body: str, fill: str = "EAF4EF") -> None:
    table = document.add_table(rows=1, cols=1)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(31, 78, 55)
    body_p = cell.add_paragraph()
    body_p.paragraph_format.space_after = Pt(0)
    body_run = body_p.add_run(body)
    body_run.font.size = Pt(9.5)
    document.add_paragraph()


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Calibri"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 55)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(11)


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("PCOSINA")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(31, 78, 55)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Statement of the Problem Defense Explanation Report")
    run.bold = True
    run.font.size = Pt(18)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Revised SOP wording, algorithmic how-process, required evidence, and defense-safe claim boundaries")
    run.font.size = Pt(11)
    run.italic = True

    add_callout(
        document,
        "Defense Position",
        "PCOSina should be defended as an offline-first Filipino-PCOS wellness decision-support system. "
        "It is not a diagnosis engine, treatment engine, or medical device. The planning contract is deterministic "
        "filtering plus deterministic OR-Tools CP-SAT optimization. ML may assist ranking only and must not override hard constraints.",
        fill="EAF4EF",
    )

    meta = [
        ["Prepared for", "Final defense preparation"],
        ["Prepared on", date.today().isoformat()],
        ["Repository", str(ROOT)],
        ["Primary purpose", "Make every SOP question answerable as a concrete process with inputs, outputs, evidence, and metrics."],
    ]
    add_table(document, ["Field", "Value"], meta, widths=[1.6, 5.8])
    document.add_page_break()


def add_executive_summary(document: Document) -> None:
    add_heading(document, "Executive Summary", 1)
    add_para(
        document,
        "The safest final-defense strategy is to revise the SOPs so each question asks about what the system actually implements "
        "and what can be measured. The report below turns each question into a defendable pipeline: profile inputs, deterministic "
        "filtering, candidate scoring, CP-SAT optimization, local/offline persistence, pantry-aware grocery support, and evaluation.",
    )
    add_bullets(
        document,
        [
            "Remove insulin resistance from the SOP wording and from the final system narrative. Current code and older validation docs still contain insulinResistanceLevel, so this must be cleaned or explicitly deprecated before claiming it is removed.",
            "Replace food-waste language with help promote reduction or support reduction. The system can use pantry data and grocery lists, but actual waste reduction depends on user adherence and must be measured separately.",
            "Keep offline claims precise. Saved profile, plan, grocery, pantry, progress, and reminder data can be defended as local-first; new backend-assisted generation or swap may still require internet.",
            "For SOP 5, Cronbach's alpha supports questionnaire reliability only. It does not prove system quality by itself. Final conclusions also need weighted means, standard deviations, algorithmic metrics, expert validation, and actual respondent data.",
        ],
    )

    add_table(
        document,
        ["Claim Area", "Defense-Safe Wording"],
        [
            ["Medical scope", "Wellness decision-support for meal planning, not diagnosis or treatment."],
            ["Algorithm", "Deterministic filtering plus deterministic OR-Tools CP-SAT optimization; ML only ranks or personalizes candidates before final solving."],
            ["Nutrition", "Nutritionally adequate means measured against calorie, macro, fiber, and visible nutrition constraints available in the system data."],
            ["Pantry and grocery", "Pantry-aware planning can encourage use of available ingredients and reduce unnecessary purchases if the user keeps pantry data current and follows the plan."],
            ["Offline-first", "Saved artifacts remain available offline; online-only actions are blocked or clearly explained."],
            ["Evaluation", "Use algorithmic metrics, expert ratings, user-centered ISO 25010 ratings, and reliability/statistical treatment."],
        ],
        widths=[1.65, 5.75],
    )


def add_revised_sops(document: Document) -> None:
    add_heading(document, "Revised SOP Set", 1)
    add_para(
        document,
        "All five questions still start with How, but the wording is tightened so the answers can be demonstrated with the actual PCOSina pipeline and evidence.",
    )
    add_table(
        document,
        ["No.", "Revised Statement of the Problem", "Main Change"],
        [
            [
                "1",
                "How can a weekly meal-planning system be designed to generate nutritionally adequate meal plans for Filipinos with PCOS while considering weight-management goals and dietary constraints?",
                "Removed insulin resistance and focused on measurable nutrition adequacy, weight management, and constraints.",
            ],
            [
                "2",
                "How can user preferences, Filipino food practices, and dietary restrictions be effectively integrated into a personalized meal-planning algorithm for individuals with PCOS?",
                "Kept personalization and cultural relevance, but links them to recipe catalog, constraints, and scoring.",
            ],
            [
                "3",
                "How can pantry inventory data be integrated into meal planning to help promote reduced food waste and unnecessary grocery purchases while maintaining nutritional adequacy?",
                "Changed reduce food waste into help promote reduced food waste. This avoids a guarantee that depends on user behavior.",
            ],
            [
                "4",
                "How can an offline-first mobile application be designed to support saved meal planning, saved recipe access, grocery/pantry management, and progress tracking in environments with limited or no internet connectivity?",
                "Clarifies saved/local-first access while avoiding an unsupported claim that every new backend-assisted action works offline.",
            ],
            [
                "5",
                "How does the proposed system perform in terms of nutritional quality, usability, cultural relevance, and computational efficiency based on algorithmic, expert, and user-centered evaluation metrics?",
                "Keeps performance/evaluation broad but makes the evidence types explicit.",
            ],
        ],
        widths=[0.45, 4.8, 2.15],
    )


def sop_section(
    document: Document,
    number: int,
    title: str,
    defense_answer: str,
    process: list[str],
    inputs: list[str],
    evidence: list[str],
    metrics: list[str],
    risk: str,
) -> None:
    add_heading(document, f"SOP {number}: {title}", 1)
    add_callout(document, "Direct Defense Answer", defense_answer)
    add_heading(document, "How PCOSina Solves This", 2)
    add_numbered(document, process)
    add_heading(document, "Inputs Required", 2)
    add_bullets(document, inputs)
    add_heading(document, "Outputs and Evidence Required", 2)
    add_bullets(document, evidence)
    add_heading(document, "Metrics or Results Needed", 2)
    add_bullets(document, metrics)
    add_callout(document, "Claim Boundary", risk, fill="FFF2CC")


def add_sop_details(document: Document) -> None:
    sop_section(
        document,
        1,
        "Nutritionally Adequate Weekly Plans Without Insulin-Resistance Claiming",
        "PCOSina answers this by converting a user's profile and weight-management goal into calorie and nutrition targets, "
        "filtering unsafe or incompatible recipes, and using CP-SAT to assemble a complete weekly plan that minimizes nutrition deviation while preserving hard constraints.",
        [
            "Collect profile inputs: age, height, weight, activity level, goal, allergies, dietary restrictions, budget, household size, cooking-time limit, variety preference, and pantry items.",
            "Compute wellness planning targets: BMI category for feedback, BMR, TDEE, and daily calorie target. Weight loss is handled as a goal-based calorie adjustment, not as a clinical insulin-resistance treatment.",
            "Load the Filipino recipe catalog with calories, protein, carbohydrates, fats, fiber, cost estimates, prep time, meal type, and ingredient tokens.",
            "Apply deterministic hard filters before optimization: allergies, recognized restrictions, conflicting profile states, cooking-time limit, meal-slot eligibility, and budget ceiling when configured.",
            "Score safe candidates using nutrition, calorie distance, estimated cost, pantry overlap, prep-time behavior, and optional ML or shadow ranking as assistive Stage 1 signals.",
            "Solve the final weekly assignment using OR-Tools CP-SAT with one recipe per meal slot, repetition limits, budget cap, and nutrition deviation penalties.",
            "Return either a complete success plan with explanation metadata or a structured no-safe-plan response with safe suggested relaxations.",
        ],
        [
            "Validated user profile and stated planning goal.",
            "Recipe catalog with reviewed or at least traceable nutrition values.",
            "Policy values for calorie, macro, fiber, sodium/sugar advisory limits, budget, variety, and solver time.",
            "Manual formula worksheet for BMI, BMR, TDEE, target calories, macro grams, and grocery cost examples.",
            "No-safe-plan test cases for impossible or conflicting constraints.",
        ],
        [
            "Weekly meal plan, daily nutrition totals, and solver diagnostics.",
            "Manual-vs-system computation table for health and nutrition formulas.",
            "Planner success/no-safe-plan evidence for representative profiles.",
            "Expert review confirming that the meal-plan criteria are appropriate as wellness support.",
        ],
        [
            "Hard-constraint violation count should be zero for allergies, restrictions, meal-slot validity, and configured budget.",
            "Nutrition quality can be reported through calorie/macro/fiber deviation from target ranges.",
            "Planner success rate and no-safe-plan reason distribution should be reported for scenario tests.",
            "Nutrition expert ratings should support adequacy and safety wording.",
        ],
        "Do not say the system treats insulin resistance or chooses meal plans based on clinical insulin-resistance severity. If insulinResistanceLevel remains in the app or backend, the thesis wording and demo will conflict with the revised SOP.",
    )

    sop_section(
        document,
        2,
        "Preferences, Filipino Food Practices, and Dietary Restrictions",
        "PCOSina answers this by treating hard restrictions as non-negotiable filters and treating preferences/cultural fit as scoring and optimization signals. The system personalizes the plan without allowing soft preferences or ML ranking to override safety rules.",
        [
            "Represent food practices through a Filipino recipe catalog and ingredient names that include local terms and common substitutions.",
            "Normalize ingredients, pantry entries, and allergy terms so Filipino synonyms such as bawang/garlic or manok/chicken map to consistent tokens.",
            "Infer recipe tags such as meal type, protein group, dairy/meat/seafood presence, high-protein, high-fiber, and low-carb indicators.",
            "Filter hard constraints first: allergies, dietary restrictions, explicit excluded foods, conflict checks, and cooking-time limits.",
            "Rank remaining recipes by protein, calorie distance, cost estimate, pantry overlap, prep-time behavior, and planning priority.",
            "Use CP-SAT to balance the full week, not only individual meals, so preferences are integrated with variety, repetition, nutrition, budget, and pantry behavior.",
            "Explain the output through plan metadata, grocery output, and no-safe-plan guidance when preferences are too restrictive.",
        ],
        [
            "User preferences: meal count, variety preference, planning priority, max cooking time, budget, pantry items, and optional symptoms/goals.",
            "Dietary restrictions and allergies entered as hard constraints.",
            "Cultural recipe dataset and local ingredient synonym mappings.",
            "Preference and restriction test cases, including Filipino synonym tests.",
        ],
        [
            "Personalized weekly plan with Filipino recipes and readable meal slots.",
            "Restriction-compliance evidence from backend tests and manual validation.",
            "Cultural relevance evidence from user survey and nutrition/health expert validation.",
            "Explanation of which inputs were hard rules versus soft preferences.",
        ],
        [
            "Restriction compliance pass/fail.",
            "Preference match rate or qualitative preference-fit rating.",
            "Cultural relevance weighted mean from respondents or experts.",
            "Usability score for profile setup and plan readability.",
        ],
        "Do not describe the algorithm as a pure decision tree. A better defense phrase is decision flow plus optimization computation: Stage 1 filters/ranks; Stage 2 solves the weekly combination.",
    )

    sop_section(
        document,
        3,
        "Pantry-Aware Planning and Grocery Reduction Support",
        "PCOSina answers this by using pantry entries as planning signals and grocery-list context. Pantry data can encourage the planner and user to use available ingredients first, but actual waste reduction depends on whether the pantry is accurate and whether the user follows the plan.",
        [
            "Let the user enter pantry items during profile/grocery setup.",
            "Normalize pantry strings into tokens and compare them with recipe ingredient tokens.",
            "Compute pantry overlap count for each candidate recipe.",
            "Add pantry overlap as a Stage 1 ranking bonus and Stage 2 reward/slack signal, while still preserving nutrition and hard constraints.",
            "Generate or rebuild grocery lists from selected recipes and show pantry/bought/remaining categories where the UI supports it.",
            "Use estimated prices to help the user distinguish ingredients already available from ingredients still needed.",
            "Evaluate the real effect using adherence data, purchase logs, or user survey items, not by assuming the system automatically reduces waste.",
        ],
        [
            "Accurate pantry inventory from the user.",
            "Recipe ingredient lists and normalized pantry/ingredient token mapping.",
            "Current generated plan and grocery item aggregation.",
            "Purchase or pantry-adherence records if the thesis will claim real-world reduction.",
        ],
        [
            "Pantry overlap counts in candidate scoring and plan explanation.",
            "Grocery list showing what is needed versus pantry/bought items.",
            "Estimated grocery cost output.",
            "User or expert rating for grocery usefulness and pantry practicality.",
        ],
        [
            "Pantry overlap per plan or per meal.",
            "Number or estimated cost of grocery items already covered by pantry.",
            "User-reported reduction in unnecessary purchases.",
            "Optional pre/post waste or leftover log if real waste reduction is claimed.",
            "Nutritional adequacy must still be measured with calories/macros/fiber, not pantry use alone.",
        ],
        "Current evidence says pantry is not a hard all-ingredients-in-pantry feasibility constraint. Use wording like helps promote reduced waste or supports pantry-aware planning, not guarantees waste reduction.",
    )

    sop_section(
        document,
        4,
        "Offline-First Mobile Support",
        "PCOSina answers this by making local saved data the day-to-day source of continuity. The app should keep profile, saved plans, grocery/pantry data, progress logs, and reminders visible offline, while online-only actions provide clear feedback instead of failing silently.",
        [
            "Persist user profile, saved plan/history, grocery snapshots, pantry entries, progress logs, and reminder preferences locally.",
            "Treat local data as the first UI source for continuity, with cloud sync used for backup, reinstall recovery, or multi-device continuity.",
            "Observe connectivity state in meal plan, grocery, dashboard, and progress flows.",
            "Allow saved artifacts to remain readable without internet.",
            "Block or explain online-dependent actions such as new backend-assisted plan generation or meal swap when the device is offline.",
            "Queue or retry sync-supported actions when connectivity returns, depending on the feature.",
            "Validate with an offline checklist: generate/save online, disable internet, reopen, view saved data, attempt online-only action, confirm clear feedback, re-enable internet.",
        ],
        [
            "Local storage artifacts and repository boundaries.",
            "Connectivity state and offline/online UI states.",
            "Saved profile, plan, grocery, pantry, progress, and notification data.",
            "Offline-first test checklist and connected-device verification evidence.",
        ],
        [
            "Screens showing saved plan/grocery/pantry/progress available offline.",
            "Clear offline banners or internet-required messages for blocked actions.",
            "Sync/recovery behavior when internet returns.",
            "Test evidence from unit, instrumentation, or manual offline walkthrough.",
        ],
        [
            "Offline task completion rate for saved data access.",
            "Number of critical saved artifacts available offline.",
            "Crash/freeze count during offline scenario.",
            "User rating for offline support under ISO 25010 compatibility/offline items.",
        ],
        "Do not claim full offline generation unless the final system actually contains a local solver and local recipe/nutrition catalog sufficient for new planning. The current defense-safe claim is offline-first access to saved artifacts plus clear handling of online-dependent actions.",
    )

    sop_section(
        document,
        5,
        "Evaluation of Nutrition Quality, Usability, Cultural Relevance, and Efficiency",
        "PCOSina answers this by separating system evaluation into algorithmic evidence, expert validation, and user-centered survey results. Cronbach's alpha belongs to the survey instrument reliability layer, while algorithm quality and system performance need their own metrics.",
        [
            "Define evaluation dimensions: nutritional quality, hard-constraint safety, usability, cultural relevance, offline support, reliability, and computational efficiency.",
            "Run algorithmic scenarios and collect solver status, runtime, candidate counts, hard-constraint violations, nutrition deviation, pantry overlap, and no-safe-plan outcomes.",
            "Run manual computation validation for BMI, BMR, TDEE, calorie targets, pantry overlap, grocery cost examples, and planner constraints.",
            "Collect nutrition/health expert ratings for wellness appropriateness, Filipino meal relevance, allergy/restriction safety, and clarity of no-safe-plan guidance.",
            "Collect user/IT expert survey responses using the ISO 25010-aligned questionnaire.",
            "Compute frequency, percentage, weighted mean, standard deviation, and overall mean by category.",
            "Compute Cronbach's alpha if required to show questionnaire consistency, then interpret survey results only after reliability and response data are available.",
        ],
        [
            "Algorithm benchmark output or scenario logs.",
            "Raw ISO 25010 respondent response matrix.",
            "Nutrition/health expert validation matrix and comments.",
            "Cronbach's alpha worksheet or statistician-verified output.",
            "Manual computation validation table.",
        ],
        [
            "Algorithm result tables: success rate, runtime, hard-constraint violations, nutrition deviations, pantry overlap, and no-safe-plan cases.",
            "Survey result tables: weighted mean, standard deviation, interpretation, and Cronbach's alpha.",
            "Expert result tables: nutrition quality, PCOS wellness relevance, cultural relevance, grocery usefulness, and safety.",
            "Chapter 5 narrative that reports only collected data, not assumed results.",
        ],
        [
            "Nutritional quality: calorie/macro/fiber deviation and expert rating.",
            "Usability: ISO 25010 usability weighted mean and task feedback.",
            "Cultural relevance: Filipino meal relevance expert/user rating.",
            "Computational efficiency: average runtime, P95 runtime, max runtime, timeouts, and success rate.",
            "Reliability: Cronbach's alpha for the questionnaire, plus no-safe-plan handling and offline continuity evidence for the system.",
        ],
        "A passing Cronbach's alpha does not prove the system works. It only supports that the questionnaire responses are internally consistent. System performance must still be supported by actual algorithm logs, user ratings, and expert evaluation.",
    )


def add_insulin_cleanup(document: Document) -> None:
    add_heading(document, "Required Cleanup: Removing Insulin Resistance", 1)
    add_para(
        document,
        "Because the revised SOP removes insulin resistance, the system narrative and implementation should not keep presenting insulin resistance as a required planning input. The current repository still contains insulinResistanceLevel in app, backend, schema, policy, tests, and generated thesis validation files.",
    )
    add_table(
        document,
        ["Layer", "Current Trace", "Action Needed"],
        [
            ["Android model", "UserProfile.insulinResistanceLevel", "Remove, deprecate, or keep internally hidden with a migration default. Do not show it as a required input."],
            ["Android UI", "UserProfileScreen StepTwoMedical insulin dropdown and blocker text", "Remove the dropdown and replace with optional symptoms/goals or general wellness factors."],
            ["Android persistence", "UserPreferencesRepository insulin keys and profile payload fields", "Stop writing new insulin fields; migrate old saved values safely if needed."],
            ["Backend domain/API", "backend/domain/models.py, backend/main.py, backend/schema/pcosina_contract.json", "Remove from public request contract or accept only as legacy ignored input."],
            ["Planner logic", "macro_ratios(insulin_level) and solve_meal_plan macro target use", "Replace with goal/symptom/general nutrition policy that does not mention insulin-resistance severity."],
            ["Policy defaults", "stage1.cold_start_defaults.insulinResistanceLevel", "Remove from policy config and generated policy exports."],
            ["Tests", "test_macro_ratios_moderate/severe and profile UI tests", "Replace with tests for revised macro policy and no clinical insulin-resistance claim."],
            ["Docs/evidence pack", "thesis_validation formula/data dictionary/test cases and two-stage planner writeup", "Regenerate or revise docs so Chapter 1/3/4 do not contradict the revised SOP."],
        ],
        widths=[1.55, 2.45, 3.4],
    )
    add_callout(
        document,
        "Recommended Replacement Logic",
        "For final defense, describe nutrition targeting through weight-management goal, activity level, visible nutrition bounds, symptoms/goals if retained, and expert-reviewed wellness criteria. This avoids clinical severity labels while preserving a measurable planning pipeline.",
        fill="EAF4EF",
    )


def add_evidence_matrix(document: Document) -> None:
    add_heading(document, "Evidence Checklist", 1)
    add_para(
        document,
        "Use this checklist as the defense preparation board. A SOP is strongest when it has input data, process evidence, output evidence, and evaluation metrics.",
    )
    add_table(
        document,
        ["Evidence Needed", "Why It Matters", "Current Source or Required Artifact"],
        [
            ["Revised SOP wording", "Prevents overclaiming and removes insulin resistance from the research question.", "Chapter 1 SOP section to be edited."],
            ["Planner process diagram", "Shows the How: inputs -> filters -> scoring -> CP-SAT -> plan/no-safe-plan.", "docs/thesis/pcosina_two_stage_planner_algorithm.md, then revise insulin wording."],
            ["Nutrition formula table", "Defends nutritional adequacy computation.", "docs/thesis_validation/02_FORMULAS_AND_COMPUTATIONS.md, after insulin cleanup."],
            ["Recipe catalog summary", "Supports Filipino meal and nutrition-data claims.", "docs/production_readiness/local_catalog_nutrition_readiness.json and backend/recipes.json."],
            ["Algorithm benchmark", "Supports computational efficiency and feasibility.", "docs/roadmap/progress_ledger.md and benchmark JSON/CSV outputs."],
            ["Pantry/grocery evidence", "Supports pantry-aware planning and purchase-reduction support.", "GroceryAggregation tests, pantry overlap outputs, grocery screenshots, user adherence logs if claiming reduction."],
            ["Offline-first walkthrough", "Supports SOP 4.", "ADR-002, offline validation checklist, screenshots/manual test evidence."],
            ["User survey matrix", "Required for SOP 5 user-centered metrics.", "Raw respondent spreadsheet, weighted means, SD, category means."],
            ["Cronbach alpha", "Shows questionnaire reliability if required.", "Statistician-verified alpha computation from actual respondent item matrix."],
            ["Nutrition expert validation", "Supports nutrition quality and cultural relevance.", "Nutrition expert validation form and scored results."],
            ["Risk/gap table", "Keeps defense honest.", "docs/thesis_validation/12_IMPLEMENTATION_GAPS_AND_ACTION_ITEMS.md updated after changes."],
        ],
        widths=[1.9, 2.55, 2.95],
    )

    add_heading(document, "Current Evidence Highlights From Repository", 2)
    add_bullets(
        document,
        [
            "The current planner is documented as deterministic Stage 1 filtering/ranking plus OR-Tools CP-SAT Stage 2 optimization.",
            "The current validation pack lists 1,130 active recipes with complete active nutrition profiles, but also warns that 953 of 1,130 active recipes still need reviewed nutrition provenance.",
            "The roadmap benchmark update reports 100/100 local planner runs with average 257 ms, P95 313 ms, max 357 ms, zero hard-constraint violations, and no planner timeouts after optimization passes.",
            "The ML evaluation report shows LightGBM canary readiness, but ML remains assistive; Stage 2 hard-rule authority is retained.",
            "The offline-first ADR states that local-first data is the default source of day-to-day continuity and cloud sync is supportive.",
            "The statistician packet and ISO 25010 questionnaire draft are prepared, but final respondent datasets must be collected before final Chapter 5 claims.",
        ],
    )


def add_panel_scripts(document: Document) -> None:
    add_heading(document, "Panel Answer Scripts", 1)
    scripts = [
        [
            "If asked why insulin resistance was removed",
            "We removed it because our final defense should stay within what we can responsibly validate. PCOSina supports PCOS-related wellness meal planning through weight-management goals, nutrition targets, dietary restrictions, pantry data, and expert review. It does not diagnose or clinically manage insulin resistance.",
        ],
        [
            "If asked how the algorithm works",
            "The algorithm has two major stages. First, it normalizes and filters recipes using hard rules such as allergies, restrictions, cooking time, and budget. Then it scores safe candidates. Second, CP-SAT chooses the final weekly combination while balancing nutrition, repetition, variety, pantry use, and cost. If it cannot create a safe full plan, it returns no-safe-plan guidance instead of forcing an unsafe plan.",
        ],
        [
            "If asked whether the app reduces food waste",
            "The system supports that goal, but it does not guarantee it. It uses pantry data to prioritize recipes that overlap with available ingredients and to help the grocery list distinguish what may still need to be bought. Actual waste reduction depends on pantry accuracy and user adherence, so we would measure it through user logs or survey results.",
        ],
        [
            "If asked what offline-first means",
            "Offline-first means saved data remains useful when internet is limited. Users can still view saved profile, meal plan, grocery/pantry information, progress logs, and reminders. New backend-assisted planning may require internet, but the app should explain that clearly and preserve local artifacts.",
        ],
        [
            "If asked about Cronbach's alpha",
            "Cronbach's alpha is used to check whether the survey instrument is internally consistent. It does not measure the algorithm directly. For system performance, we also use algorithmic metrics such as success rate, runtime, nutrition deviation, hard-constraint compliance, expert validation, and user weighted means.",
        ],
    ]
    add_table(document, ["Panel Question", "Defense-Safe Answer"], scripts, widths=[2.1, 5.3])


def add_final_action_plan(document: Document) -> None:
    add_heading(document, "Action Plan Before Final Defense", 1)
    add_numbered(
        document,
        [
            "Revise Chapter 1 SOP wording using the five revised questions in this report.",
            "Remove or deprecate insulinResistanceLevel from the app/backend request path, profile UI, planner macro policy, tests, and generated evidence docs.",
            "Regenerate the thesis validation pack after code cleanup so the data dictionary, formulas, test cases, and tables no longer contradict the SOP.",
            "Prepare a one-page algorithm diagram showing Stage 1 filtering/ranking and Stage 2 CP-SAT solving.",
            "Prepare a pantry/grocery slide that uses helps promote reduced waste, not guarantees waste reduction.",
            "Run or archive the latest planner benchmark and hard-constraint validation results.",
            "Collect respondent and expert validation data before writing final Chapter 5 claims.",
            "Compute weighted means, standard deviations, and Cronbach's alpha from the actual response matrix only.",
            "Prepare offline-first screenshots or a manual test record showing saved artifacts available without internet.",
            "Keep every defense statement inside wellness decision-support scope.",
        ],
    )


def add_appendix_sources(document: Document) -> None:
    add_heading(document, "Repository Source Anchors Used", 1)
    add_table(
        document,
        ["Topic", "Source Anchor"],
        [
            ["Planner algorithm", "backend/services/meal_planner.py; docs/thesis/pcosina_two_stage_planner_algorithm.md"],
            ["No-safe-plan contract", "docs/safety/no_safe_plan_contract.md; backend/services/plan_response_builder.py; app/src/main/java/com/pcosina/app/ui/MealPlanViewModel.kt"],
            ["Offline-first contract", "docs/adr/ADR-002-offline-first-sync-contract.md"],
            ["Validation/evidence pack", "docs/thesis_validation/00_EVIDENCE_INDEX.md through 13_GENERATED_ARTIFACTS_SUMMARY.md"],
            ["Implementation gaps", "docs/thesis_validation/12_IMPLEMENTATION_GAPS_AND_ACTION_ITEMS.md"],
            ["ISO 25010 survey draft", "docs/thesis_validation/09_ISO_25010_QUESTIONNAIRE_DRAFT.md"],
            ["Nutrition expert form", "docs/thesis_validation/10_NUTRITION_EXPERT_VALIDATION_FORM.md"],
            ["ML evaluation", "docs/ml/evaluation_report_lightgbm_v1.md"],
            ["Roadmap benchmarks/status", "docs/roadmap/progress_ledger.md; docs/roadmap/ml_progress_ledger.md"],
            ["Nutrition readiness", "docs/production_readiness/local_catalog_nutrition_readiness.json"],
        ],
        widths=[2.2, 5.2],
    )


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)
    add_executive_summary(document)
    add_revised_sops(document)
    add_sop_details(document)
    add_insulin_cleanup(document)
    add_evidence_matrix(document)
    add_panel_scripts(document)
    add_final_action_plan(document)
    add_appendix_sources(document)

    core = document.core_properties
    core.title = "PCOSINA SOP Defense Explanation Report"
    core.subject = "Statement of the Problem defense preparation"
    core.author = "OpenAI Codex"
    core.keywords = "PCOSina, SOP, defense, meal planning, CP-SAT, offline-first, pantry"

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
