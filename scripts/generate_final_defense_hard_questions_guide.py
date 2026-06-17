from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "defense" / f"PCOSINA_Final_Defense_Hard_Questions_Master_QA_{date.today().isoformat()}.docx"


def set_font(run, size=10, bold=False, italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 16, bold=True)


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 9, italic=True)


def add_h1(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_font(r, 12, bold=True)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, 9.5, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, 9.5)
    else:
        r = p.add_run(text)
        set_font(r, 9.5)


def add_quote(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, 9.5, italic=True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_font(r, 8.5, bold=True)
        if widths:
            cell.width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(value))
            set_font(r, 8.2)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()


def build_doc() -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    add_title(doc, "PCOSina Final Defense Hard Questions Master Guide")
    add_subtitle(doc, "Prepared from repository evidence and current defense discussions, June 2026")
    add_body(
        doc,
        "Scope boundary: PCOSina is an offline-first Filipino-PCOS meal-planning wellness decision-support system. "
        "It is not a diagnosis engine, treatment system, or medical device. The safest defense style is to explain "
        "what the software actually computes and validates, then state limitations honestly.",
        bold_prefix="Scope boundary:",
    )

    add_h1(doc, "1. One-Liners To Memorize")
    add_table(
        doc,
        ["Panel concern", "Best short answer"],
        [
            ["Greedy is faster", "Greedy wins raw speed; PCOSina wins whole-week constraint-aware planning and explainability."],
            ["Why greedy baseline", "It is the natural simple reference for a scheduling/recommendation task: rule-filter, then pick locally good meals slot by slot."],
            ["Why generative proxy", "It tests the risk of fluent generated plans without deterministic constraint authority; constrained generation would need the same rule and solver guardrails."],
            ["Why LightGBM", "It is a small, fast, tabular ranker for structured recipe/profile features; it assists Stage 1 ranking only and never overrides hard constraints."],
            ["ML accuracy", "This is a ranking task, not a diagnosis classifier; we use AUC, log loss, NDCG@10, and MAP@10, while CP-SAT remains final authority."],
            ["Time complexity", "Stage 1 is roughly linear-plus-sorting; CP-SAT is combinatorial in worst case, controlled by filtering, pruning, capping, and a fixed 21-slot horizon."],
            ["CP-SAT vs MILP", "Our math is a MILP-style 0-1 linear assignment formulation; OR-Tools CP-SAT is the solver that searches for a feasible or best assignment."],
            ["40% carbs", "It is a bounded moderate-carb wellness policy, not a medical prescription and not a rice-only rule."],
            ["Offline-first", "Previously saved profile, plan, grocery, pantry, and progress artifacts remain readable locally; new optimization may need backend connectivity."],
        ],
        [1.8, 5.9],
    )

    add_h1(doc, "2. Why Greedy Heuristic Was Chosen")
    add_body(
        doc,
        "Defense logic: Greedy heuristic is the simplest credible algorithmic baseline for meal scheduling. It reflects "
        "what many simple recommenders do: filter obviously invalid recipes, then fill each meal slot with the locally "
        "best available candidate. It is fast, transparent, and easy to reproduce, so it is a fair baseline.",
        bold_prefix="Defense logic:",
    )
    add_body(
        doc,
        "Why not compare to binary search, sorting, or graph algorithms: those are lower-level techniques, not complete "
        "weekly meal-planning methods. Sorting can rank recipes, but it cannot by itself enforce a 7-day meal plan with "
        "budget, nutrition, allergies, meal slots, repetition, pantry, and variety. Graph or metaheuristic approaches "
        "could be future work, but they are less directly aligned with the defended 0-1 assignment formulation.",
    )
    add_quote(
        doc,
        "Defense script: The greedy baseline was not chosen because it is weak. It was chosen because it is the clearest "
        "simple alternative: it can generate a plan quickly using local slot decisions. PCOSina was then evaluated against "
        "that baseline to show what is gained by solving the whole week as one constrained optimization problem.",
    )

    add_h1(doc, "3. Greedy Baseline vs PCOSina Two-Stage")
    add_table(
        doc,
        ["Area", "Greedy baseline", "PCOSina two-stage"],
        [
            ["Candidate safety", "Rule-filtered candidates under the same scenarios; checked for violations.", "Production Stage 1 filtering, scoring, optional ML ranking, pruning, and capping."],
            ["Selection style", "Local: choose acceptable meals slot by slot.", "Global: CP-SAT chooses one 21-meal weekly combination."],
            ["ML", "No evidence that LightGBM was part of the greedy reference.", "LightGBM can boost Stage 1 ranking, but only as assistive support."],
            ["Solver", "No CP-SAT; no global Minimize Z.", "Binary x[s,r] variables, linear constraints, weighted objective, CP-SAT solve."],
            ["Result interpretation", "Faster and can be safe when rule-filtered.", "Slower than greedy in the tradeoff table but more defensible for whole-week balancing."],
        ],
        [1.5, 3.0, 3.2],
    )
    add_body(
        doc,
        "Table 55/56 result boundary: In the Comment #5 comparison, greedy and PCOSina both had zero hard violations in the "
        "matched rows. That does not mean the algorithms are the same. It means the comparison was fair: both were tested "
        "under the same scenarios and validation metrics. The stronger PCOSina claim is whole-week optimization, not that "
        "the greedy baseline was unsafe.",
        bold_prefix="Table 55/56 result boundary:",
    )
    add_body(
        doc,
        "Exact aggregate evidence: CP-SAT two-stage averaged 2.2806 s with objective surrogate J 0.2236. Greedy averaged "
        "0.2359 s with J 0.2248. Generative proxy averaged 0.0061 s but had 100% hard violation rate. The safe statement "
        "is: greedy is much faster; CP-SAT is slightly better on the aggregate quality surrogate and has stronger global "
        "planning structure; unconstrained generation is fast but unsafe under implemented hard-rule validation.",
    )

    add_h1(doc, "4. Why The Generative Proxy Was Unconstrained")
    add_body(
        doc,
        "The purpose of the generative proxy was not to defeat a best possible AI planner. It was to test a realistic risk: "
        "a fluent generated meal list can look culturally relevant but fail hard constraints such as restrictions, repeats, "
        "or budget. A constrained generative system would require deterministic filters, validators, and probably a solver; "
        "at that point it becomes a hybrid architecture similar to PCOSina's safety contract.",
    )
    add_quote(
        doc,
        "Defense script: We used the unconstrained proxy to show that natural-looking generated output is not enough for a "
        "safety-sensitive meal-planning task. A constrained generative planner is valid future work, but it would still need "
        "the same deterministic hard-rule layer and validation contract.",
    )

    add_h1(doc, "5. Why LightGBM Was Chosen")
    add_body(
        doc,
        "LightGBM fits PCOSina because the ML problem is tabular ranking: recipe calories, protein, carbs, fats, fiber, cost, "
        "cooking time, meal-slot flags, pantry overlap, restrictions, and feedback tags. A deep neural model or language model "
        "would be heavier and harder to explain for this small structured ranking layer. LightGBM is fast enough for Stage 1 "
        "candidate scoring and supports feature-importance style inspection.",
    )
    add_table(
        doc,
        ["Metric", "Current artifact value", "Defense meaning"],
        [
            ["Rows / positive rate", "74,112 rows / 6.27%", "Many safe candidate rows; selected-by-solver rows are a minority."],
            ["Feature count", "46", "Structured recipe/profile/feedback features, not free-text generation."],
            ["Test AUC", "0.9373", "Selected candidates generally score above unselected safe candidates."],
            ["Test log loss", "0.1024", "Probability confidence is reasonably controlled; lower is better."],
            ["Test NDCG@10", "0.9211 vs heuristic 0.8838", "Better top-10 ranking order than deterministic heuristic alone."],
            ["Test MAP@10", "0.6779 vs heuristic 0.6275", "More selected candidates appear early in the top-10 list."],
        ],
        [1.5, 1.7, 4.5],
    )
    add_body(
        doc,
        "Why 15% ML weight: The ML score is intentionally a small boost, not a replacement for deterministic scoring. At 15%, "
        "it can change ordering when candidates are close, but it should not dominate safety, nutrition anchors, cost, cooking "
        "time, or hard filters. The 30% cap is the guardrail; the current 15% is conservative because ML is assistive.",
        bold_prefix="Why 15% ML weight:",
    )
    add_body(
        doc,
        "Synthetic 60 users: The 60 unique users in Table 39 are synthetic/offline scenario identities used for reproducible "
        "training telemetry, not real respondent claims. This does not destroy credibility if stated honestly. It means the ML "
        "artifact validates ranking behavior under controlled scenario coverage, while live user generalization remains a "
        "rollout limitation and recommendation.",
        bold_prefix="Synthetic 60 users:",
    )

    add_h1(doc, "6. ML Accuracy Defense")
    add_body(
        doc,
        "Do not claim one overall ML accuracy percentage as the main metric. The model is not diagnosing PCOS and is not "
        "classifying users. It ranks already-safe candidate recipes. Accuracy can be computed by thresholding selected vs "
        "unselected labels, but it would be misleading because only 6.27% of candidate rows are positive. A model could get "
        "high accuracy by mostly predicting 'not selected' and still be poor at ranking useful meals.",
    )
    add_quote(
        doc,
        "Defense script: If the panel requires a classifier-style metric, we can provide AUC and log loss as binary evaluation "
        "metrics, but we should explain that NDCG@10 and MAP@10 are more appropriate because the app uses ML to rank safe "
        "candidate recipes, not to make a medical diagnosis or final meal-plan decision.",
    )

    add_h1(doc, "7. Time Complexity Answer")
    add_table(
        doc,
        ["Part", "Approximate complexity", "Defense explanation"],
        [
            ["Filtering", "O(R x I)", "R recipes checked against I average ingredient/restriction tokens."],
            ["Scoring", "O(C)", "Each safe candidate receives deterministic score signals."],
            ["Sorting/ranking", "O(C log C)", "Candidates are ordered before pruning/capping."],
            ["Similarity dedup", "Up to O(C^2)", "Pairwise similarity can be quadratic, so PCOSina prunes and caps the pool."],
            ["LightGBM inference", "Approximately O(C x T x L)", "C candidates through a fixed number of trees/leaves; practically near-linear for fixed model size."],
            ["CP-SAT", "Combinatorial worst case", "The solver searches binary assignments over slots and candidates; Stage 1 controls practical size."],
        ],
        [1.5, 1.7, 4.5],
    )
    add_quote(
        doc,
        "Defense script: PCOSina does have time complexity, but not one single Big-O for the whole system. Filtering is roughly "
        "linear, ranking is dominated by sorting, deduplication can be quadratic, and CP-SAT is combinatorial in worst case. "
        "That is exactly why Stage 1 filtering, pruning, pool capping, and a fixed 21-slot horizon are necessary.",
    )

    add_h1(doc, "8. CP-SAT and MILP-Style Formulation")
    add_body(
        doc,
        "Best wording: PCOSina uses a MILP-style 0-1 linear assignment formulation implemented with OR-Tools CP-SAT. It is "
        "MILP-style because the model has binary decision variables, linear constraints, and a weighted linear objective. It "
        "is not solved by a classical LP/MIP simplex/branch-and-bound MPSolver path; the actual solver is CP-SAT.",
        bold_prefix="Best wording:",
    )
    add_table(
        doc,
        ["Evidence", "Code location", "What it proves"],
        [
            ["OR-Tools import", "backend/services/meal_planner.py:10", "Uses OR-Tools CP-SAT API."],
            ["Model creation", "backend/services/meal_planner.py:2893", "Creates CpModel."],
            ["Binary variable", "backend/services/meal_planner.py:2898", "x[s,i] is 0/1 recipe-slot selection."],
            ["Slot coverage", "backend/services/meal_planner.py:2905", "Each slot gets exactly one recipe."],
            ["Repeat limit", "backend/services/meal_planner.py:2937", "Recipe use is bounded per plan attempt."],
            ["Budget cap", "backend/services/meal_planner.py:3002", "Weekly cost cap enforced when active."],
            ["Daily nutrition", "backend/services/meal_planner.py:3017-3038", "Calories, macros, and fiber constraints."],
            ["Objective", "backend/services/meal_planner.py:3093", "Weighted minimization objective."],
            ["Solve call", "backend/services/meal_planner.py:3161", "Solver actually solves the model."],
        ],
        [1.5, 2.2, 4.0],
    )
    add_body(
        doc,
        "If we only had CP-SAT without a formulation, nothing meaningful would be solved: the solver needs variables, "
        "constraints, and an objective. If we only had a formulation without a solver, we would only have math on paper. "
        "OR-Tools CP-SAT is the engine that searches the feasible space and returns OPTIMAL, FEASIBLE, INFEASIBLE, or UNKNOWN.",
    )
    add_body(
        doc,
        "Why OR-Tools CP-SAT: It supports integer/binary constraints, assignment-style problems, feasibility statuses, time "
        "limits, multiple workers, and production Python integration. It is appropriate for a 0-1 weekly meal assignment problem "
        "where all selections are discrete.",
        bold_prefix="Why OR-Tools CP-SAT:",
    )

    add_h1(doc, "9. Macro, Calorie, and Goal Policy")
    add_table(
        doc,
        ["Policy item", "Current system behavior", "Defense wording"],
        [
            ["BMR", "(10 x kg) + (6.25 x cm) - (5 x age) - 161", "Uses a recognized female BMR estimation equation as planning input, not diagnosis."],
            ["TDEE", "BMR x activity multiplier", "Estimates maintenance energy before goal adjustment."],
            ["Weight loss", "TDEE - 500 kcal, then clamp", "Moderate planning deficit; not a prescription."],
            ["General health", "Keeps balanced default targets", "No forced deficit; favors balanced meals."],
            ["Symptom management", "+4g fiber min, -8g sugar max, -10g carb target; boosts high-fiber/steady-carb candidates", "Nutrition nudges, not medical treatment."],
            ["Macro ratio", "25% protein / 40% carbs / 35% fat", "Moderated-carb PCOS wellness policy with guardrails."],
            ["Calorie floor/ceiling", "1200 / 3200 kcal", "Prevents extreme targets in the planner."],
        ],
        [1.5, 2.5, 3.7],
    )
    add_body(
        doc,
        "Better defense for 40% carbohydrates: Do not say it is only because there is no rice control. Say: PCOSina controls "
        "total carbohydrate exposure rather than only rice because rice is only one carbohydrate source. The 40% policy is a "
        "moderated-carb wellness target: lower than a high-carb pattern, not extreme low-carb or keto, and protected by minimum "
        "grams, calorie floors, fiber targets, and solver bounds. Rice portion control is a valid future enhancement, but total "
        "carb control is a broader first implementation.",
        bold_prefix="Better defense for 40% carbohydrates:",
    )
    add_body(
        doc,
        "Do not overclaim: These values are not personalized medical prescriptions and do not prove clinical outcomes. They are "
        "planning parameters for a wellness decision-support tool and should be reviewed by nutrition professionals before clinical use.",
        bold_prefix="Do not overclaim:",
    )

    add_h1(doc, "10. Benchmark Defense")
    add_body(
        doc,
        "Benchmark means controlled automated evaluation. In PCOSina, it means predefined scenarios/profiles were run through "
        "Python benchmark scripts and archived artifacts. The benchmark measured runtime, completion, hard-rule violations, nutrition "
        "gaps, repeat violations, budget overage, and objective surrogate values. It is not AI guessing and not manual timing.",
    )
    add_quote(
        doc,
        "Defense script: We used automated benchmark testing. The inputs were predefined to control the experiment, the actual planner "
        "or reference implementation produced the outputs, and the script computed metrics such as runtime, pass/fail, hard violations, "
        "calorie gap, macro gap, P95 runtime, and objective surrogate. Predefined data improves reproducibility; it does not mean the "
        "results were invented.",
    )

    add_h1(doc, "11. Why PCOS and Not Diabetes")
    add_body(
        doc,
        "PCOS is a valid scope because it is a common hormonal and metabolic condition affecting reproductive-aged women, with "
        "nutrition, weight management, insulin resistance risk, symptom burden, fertility concerns, and mental well-being implications. "
        "Diabetes is also important, but it has different clinical monitoring needs and higher medical-device risk. PCOSina deliberately "
        "limits itself to wellness meal planning for a specific underserved group.",
    )
    add_quote(
        doc,
        "Defense script: Diabetes apps usually focus on blood glucose tracking and medical monitoring. PCOSina focuses on PCOS-oriented "
        "meal planning, Filipino food context, pantry-aware groceries, offline access, restrictions/allergies, and weekly planning. "
        "That focus is the edge: it is not a generic calorie tracker and it is not a diabetes medical device.",
    )

    add_h1(doc, "12. Adherence and Progress Features")
    add_body(
        doc,
        "The adherence defense is not that the app guarantees behavior change. The defense is that it reduces friction: saved plans, "
        "grocery guidance, reminders, progress logging, meal feedback, swap options, and local trends make it easier for users to continue. "
        "Actual long-term adherence is a limitation and recommendation for future longitudinal evaluation.",
    )

    add_h1(doc, "13. Admin and Operator Side")
    add_body(
        doc,
        "The admin side is for controlled maintenance: recipe/content correction, nutrition corrections, policy or operations review, and "
        "support workflows. It is a strength because datasets change and recipe/nutrition data need governance. It is also a risk: careless "
        "admin edits can damage credibility. The defense is role-based access, review workflow, auditability, and the limitation that unvalidated "
        "new meals should not be claimed as nutritionist-approved.",
    )

    add_h1(doc, "14. Render, Backend, and Cloud Backup")
    add_body(
        doc,
        "Why backend/Render is needed: The mobile app is Kotlin/Android; the authoritative optimizer uses Python OR-Tools CP-SAT and ML tooling. "
        "Running the solver backend-side avoids shipping heavy solver/ML dependencies into the mobile client and keeps policy, telemetry, and "
        "no-safe-plan handling centralized. If there is no network, the app should still show previously saved local artifacts, but generating a "
        "fresh optimized plan may require backend access.",
    )
    add_body(
        doc,
        "Cloud backup wording: Use careful wording. Cloud sync is optional/supplementary for account continuity, backup, reinstall, or multi-device "
        "support. Local app storage remains the main offline source for day-to-day use. Do not claim that every feature is fully cloud-backed unless "
        "you show the specific sync evidence.",
        bold_prefix="Cloud backup wording:",
    )

    add_h1(doc, "15. Offline-First and DataStore")
    add_body(
        doc,
        "DataStore is Jetpack's local persistence library for storing app preferences and small structured local state inside the PCOSina app's "
        "private app storage. It is not a separate external database app. It contributes to offline-first behavior because saved profile/session/"
        "preference artifacts can be read without internet. PCOSina also uses encrypted shared preferences/local artifacts for sensitive or local "
        "content. Android does not currently use Room/SQLite as the mobile database.",
    )
    add_quote(
        doc,
        "Defense script: Offline-first does not mean every action works without internet. It means the app is still usable with previously saved "
        "local state: profile, saved plan, grocery guidance, pantry/progress artifacts, and reminders. New optimization may need the backend, but "
        "the user does not lose the current saved experience in airplane mode.",
    )

    add_h1(doc, "16. Evidence and Source Map")
    add_table(
        doc,
        ["Question area", "Repository evidence"],
        [
            ["Planner CP-SAT", "backend/services/meal_planner.py"],
            ["Policy values", "backend/policy_config.py"],
            ["20-profile benchmark", "benchmarks/reports/planner_realistic_profiles.local.lightgbm.final.json"],
            ["Tradeoff comparison", "backend/comment5_algorithm_tradeoff_summary.csv and records.csv"],
            ["ML dataset/metrics", "ml/offline_training/artifacts/dataset_v1/dataset_manifest.json and model_v1/training_metrics.json"],
            ["ML safety role", "docs/ml/model_card_lightgbm_v1.md and ml/offline_training/README.md"],
            ["Architecture/offline", "ARCHITECTURE.md, README.md, docs/adr/ADR-002-offline-first-sync-contract.md"],
            ["No-safe-plan contract", "docs/safety/no_safe_plan_contract.md and backend tests"],
        ],
        [2.0, 5.7],
    )

    add_h1(doc, "17. External References Used For Defense Framing")
    for text in [
        "WHO PCOS fact sheet: https://www.who.int/news-room/fact-sheets/detail/polycystic-ovary-syndrome",
        "Google OR-Tools CP-SAT documentation: https://developers.google.com/optimization/cp/cp_solver",
        "USDA National Agricultural Library Dietary Guidance and DRI resources: https://www.nal.usda.gov/human-nutrition-and-food-safety/dietary-guidance",
        "Dietary Guidelines for Americans portal: https://www.dietaryguidelines.gov/",
    ]:
        add_body(doc, text)

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
