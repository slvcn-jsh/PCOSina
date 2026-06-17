from __future__ import annotations

import re
import textwrap
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "backend" / "services" / "meal_planner.py"
OUT_DIR = ROOT / "docs" / "defense"
ASSET_DIR = OUT_DIR / "PCOSINA_MILP_Stage1_Stage2_Code_Evidence_2026-06-07.assets"
DOCX_PATH = OUT_DIR / "PCOSINA_MILP_Stage1_Stage2_Code_Evidence_2026-06-07.docx"


EVIDENCE = [
    {
        "category": "Solver and Stage Boundary",
        "title": "OR-Tools CP-SAT import",
        "line": 10,
        "end": 10,
        "what": "Shows the backend imports OR-Tools CP-SAT, the solver API used for the planner model.",
        "defense": "This proves we use a solver library, not only handcrafted if-else meal selection.",
    },
    {
        "category": "Solver and Stage Boundary",
        "title": "Main two-stage planner entry point",
        "line": 2468,
        "end": 2475,
        "what": "Defines the production meal-plan generation function.",
        "defense": "This is where the Stage 1 shortlist and Stage 2 solver process are connected.",
    },
    {
        "category": "Solver and Stage Boundary",
        "title": "Stage 1 starts inside planner",
        "line": 2601,
        "end": 2613,
        "what": "Marks the start of Stage 1 pruning, shortlist preparation, and feature context setup.",
        "defense": "Stage 1 prepares safe candidate recipes before optimization.",
    },
    {
        "category": "Solver and Stage Boundary",
        "title": "Stage 1 shortlist call",
        "line": 2643,
        "end": 2651,
        "what": "Calls shortlist_candidates using the user's profile, recipes, policy, pricing context, and ML feature context.",
        "defense": "The production planner does not send the whole catalog directly to the solver; it filters and ranks first.",
    },
    {
        "category": "Solver and Stage Boundary",
        "title": "Stage 2 model starts",
        "line": 2893,
        "end": 2898,
        "what": "Creates the CP-SAT model and the binary decision variables.",
        "defense": "This is the start of the MILP-style 0-1 optimization formulation.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Stage 1 function definition",
        "line": 1654,
        "end": 1662,
        "what": "Defines shortlist_candidates, the deterministic candidate filtering and shortlisting stage.",
        "defense": "Stage 1 is a real planner stage, not only UI-side filtering.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Read user constraints",
        "line": 1665,
        "end": 1669,
        "what": "Reads dietary restrictions, allergies, weekly budget, maximum cooking time, and pantry tokens.",
        "defense": "The user profile directly affects the planning constraints and candidate filtering.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Normalize ingredients",
        "line": 287,
        "end": 299,
        "what": "Converts recipe ingredients into normalized tokens for consistent matching.",
        "defense": "This supports allergy, pantry, tag, and feature matching using comparable tokens.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Normalize pantry",
        "line": 306,
        "end": 313,
        "what": "Converts user pantry items into normalized tokens.",
        "defense": "Pantry use is not guessed; pantry strings are processed into the same token format as recipe ingredients.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Normalize allergies",
        "line": 321,
        "end": 350,
        "what": "Maps allergy inputs into allergen families and custom allergy tokens.",
        "defense": "Allergy filtering is based on normalized allergy categories and ingredient-token matching.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Symptom and goal adjustments",
        "line": 394,
        "end": 444,
        "what": "Applies deterministic nutrition nudges for goals and selected symptoms.",
        "defense": "PCOS-related profile fields affect fiber, sugar, protein, calorie, and Stage 1 preference boosts.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Recipe feature extraction",
        "line": 680,
        "end": 697,
        "what": "Builds tags, ingredient tokens, protein group, vegetable tokens, and allowed meal labels.",
        "defense": "The planner extracts structured recipe features before scoring and solving.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Meal-slot eligibility inference",
        "line": 613,
        "end": 628,
        "what": "Maps recipe mealType values into allowed meal labels such as Breakfast, Lunch, Dinner, and snacks.",
        "defense": "Recipes are not randomly placed in slots; meal-type compatibility is computed first.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Hard restriction check",
        "line": 1036,
        "end": 1059,
        "what": "Detects allergy, pork, beef, vegetarian, pescatarian, and lactose-intolerant failures.",
        "defense": "The if-statements here are safety gates, which are normal and necessary in any decision-support app.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Exclude unsafe recipes",
        "line": 1726,
        "end": 1737,
        "what": "Applies static features and skips recipes that fail hard restriction checks.",
        "defense": "Unsafe recipes are removed before Stage 2 optimization.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Prep-time exclusion",
        "line": 1738,
        "end": 1742,
        "what": "Excludes recipes above the user's maximum cooking-time limit.",
        "defense": "This is a user feasibility rule, not a hard-coded meal decision.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Pantry overlap scoring",
        "line": 1757,
        "end": 1760,
        "what": "Counts how many recipe ingredients overlap with saved pantry items.",
        "defense": "Pantry data influences ranking and later solver reward, while still preserving hard safety rules.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Stage 1 score formula",
        "line": 805,
        "end": 810,
        "what": "Ranks candidates using protein, cost estimate, calorie distance, pantry match, and Stage 1 boost.",
        "defense": "Stage 1 ranking is deterministic and explainable.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Cost estimate",
        "line": 761,
        "end": 802,
        "what": "Estimates recipe cost using ingredient pricing context and serving multiplier, with a bounded fallback.",
        "defense": "Cost is computed before optimization so budget can affect planning.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Meal bucket assignment",
        "line": 1823,
        "end": 1831,
        "what": "Places safe recipes into Breakfast, Lunch, Dinner, or Universal buckets.",
        "defense": "Stage 1 organizes the candidates by meal type before Stage 2.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Optional ML scoring",
        "line": 1850,
        "end": 1869,
        "what": "Uses a ranker or deterministic shadow score only after hard filtering.",
        "defense": "ML is optional and assistive; it does not override hard constraints.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Sort, deduplicate, and trim",
        "line": 1875,
        "end": 1889,
        "what": "Sorts by base score, removes near-duplicates, applies limits, and does budget-aware trimming.",
        "defense": "Stage 1 narrows the search space while keeping safe and relevant options.",
    },
    {
        "category": "Stage 1 Evidence",
        "title": "Stage 1 returns buckets",
        "line": 1901,
        "end": 1912,
        "what": "Records Stage 1 diagnostics and returns the final buckets to the planner.",
        "defense": "Stage 1 produces explicit output consumed by Stage 2.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Binary decision variable x[s,i]",
        "line": 2893,
        "end": 2898,
        "what": "Creates x[s,i] as a Boolean variable for each meal slot and recipe candidate.",
        "defense": "x[s,i] = 1 means recipe i is selected for slot s; x[s,i] = 0 means it is not selected.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Exactly one recipe per slot",
        "line": 2901,
        "end": 2905,
        "what": "Adds a linear equality requiring one allowed recipe for each meal slot.",
        "defense": "This is a formal assignment constraint, not a manual if-else selection.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Disallow wrong meal type",
        "line": 2906,
        "end": 2910,
        "what": "Forces non-allowed recipe-slot pairs to zero.",
        "defense": "The solver is prevented from assigning Breakfast-only recipes to the wrong slot type.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Greedy warm-start hint",
        "line": 2913,
        "end": 2927,
        "what": "Adds a greedy starting hint for the solver search.",
        "defense": "This hint can guide the solver, but it does not replace the CP-SAT optimization.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "No adjacent duplicate",
        "line": 2928,
        "end": 2937,
        "what": "Prevents adjacent repeats and limits weekly repeats.",
        "defense": "Variety constraints are directly encoded in the solver model.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Repeat penalty variable",
        "line": 2947,
        "end": 2953,
        "what": "Creates repeat_over variables to penalize repeated recipe use.",
        "defense": "The model can discourage repetition without making every repeat impossible.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Protein group variety",
        "line": 2955,
        "end": 2968,
        "what": "Groups recipes by protein source and penalizes overuse of the same group.",
        "defense": "Variety is handled as a model term, not only by random shuffling.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Vegetable diversity",
        "line": 2970,
        "end": 2987,
        "what": "Tracks vegetable token coverage and adds a minimum diversity target with slack.",
        "defense": "Ingredient diversity is explicitly represented in the formulation.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Pantry reward",
        "line": 2989,
        "end": 2998,
        "what": "Builds a pantry-match expression from selected recipe variables.",
        "defense": "Pantry utilization affects the objective through selected recipes.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Total cost expression and optional rough cap",
        "line": 3000,
        "end": 3002,
        "what": "Computes total selected cost and optionally applies a rough budget cap.",
        "defense": "Budget is part of the mathematical model when the policy enables the rough cap.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Daily calorie constraints",
        "line": 3013,
        "end": 3021,
        "what": "Computes daily calories from selected meals and enforces calorie bounds.",
        "defense": "Daily calorie feasibility is enforced by linear constraints.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Macro, fiber, sodium, and sugar constraints",
        "line": 3022,
        "end": 3040,
        "what": "Computes daily protein, carbs, fats, fiber, sodium, and sugar from selected meals.",
        "defense": "Macros and fiber are hard bounded; sodium and sugar are tracked through overage variables.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Meal calorie distribution",
        "line": 3042,
        "end": 3049,
        "what": "Creates meal-level calorie deviation variables around meal distribution targets.",
        "defense": "The planner does not only hit daily totals; it also balances individual meals.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Objective component totals",
        "line": 3051,
        "end": 3067,
        "what": "Combines error, sodium/sugar overage, repeat pressure, group pressure, pantry reward, and diversity reward.",
        "defense": "The model has measurable optimization terms before the final objective is declared.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Objective weights",
        "line": 3068,
        "end": 3092,
        "what": "Applies user planning priority, variety preference, budget, pantry, prep time, and acceptance weights.",
        "defense": "User priorities affect objective weights, not hard-coded final meals.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Weighted minimization objective",
        "line": 3093,
        "end": 3100,
        "what": "Calls model.Minimize with the full weighted objective expression.",
        "defense": "This is the strongest proof that Stage 2 is an optimization model.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Solver created and configured",
        "line": 3111,
        "end": 3145,
        "what": "Creates the CP-SAT solver and sets time, gap, worker, memory, and solution-count parameters.",
        "defense": "The system uses an actual solver run with configured search limits.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Solver execution",
        "line": 3151,
        "end": 3165,
        "what": "Runs solver.Solve(model) and reads the solver status.",
        "defense": "The selected plan comes from the solver result.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Accept feasible or optimal",
        "line": 3197,
        "end": 3202,
        "what": "Stops only when the solver returns OPTIMAL, FEASIBLE, INFEASIBLE, or exits retry logic.",
        "defense": "The planner respects solver status instead of always forcing an output.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Read selected meals from variables",
        "line": 3215,
        "end": 3229,
        "what": "Uses solver.Value(x[idx, i]) to convert selected binary variables into planned meals.",
        "defense": "The final displayed meal plan is literally read from the solved decision variables.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Final grocery budget validation",
        "line": 3232,
        "end": 3268,
        "what": "Builds grocery output, checks final budget authority, and rejects/retries over-budget plans.",
        "defense": "The system verifies the selected plan against final grocery cost, not only the rough solver estimate.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Success return with explanation",
        "line": 3297,
        "end": 3349,
        "what": "Builds the explanation object and returns the successful plan.",
        "defense": "The system returns both a plan and explanation metadata for defense/debugging.",
    },
    {
        "category": "Stage 2 MILP-Style Formulation Evidence",
        "title": "Infeasible fallback",
        "line": 3378,
        "end": 3382,
        "what": "Returns budget failure, timeout, or infeasible instead of pretending a safe plan exists.",
        "defense": "The planner handles infeasibility gracefully.",
    },
]


def safe_name(value: str) -> str:
    value = value.lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_")[:90]


def load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    font_candidates = [
        Path("C:/Windows/Fonts/consola.ttf"),
        Path("C:/Windows/Fonts/cour.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"),
    ]
    for candidate in font_candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_code_image(
    source_lines: list[str],
    start_line: int,
    end_line: int,
    target_line: int,
    target_end: int,
    output_path: Path,
) -> None:
    font = load_font(21)
    title_font = load_font(18)
    line_height = 30
    padding_x = 28
    padding_y = 24
    gutter_width = 86
    max_code_px = 2050

    rendered: list[tuple[int, str, bool]] = []
    for number in range(start_line, end_line + 1):
        text = source_lines[number - 1].rstrip("\n")
        is_target = target_line <= number <= target_end
        rendered.append((number, text, is_target))

    measure_img = Image.new("RGB", (1, 1))
    measure = ImageDraw.Draw(measure_img)
    max_width = 0
    for number, text, _ in rendered:
        candidate = f"{number:>5}  {text}"
        bbox = measure.textbbox((0, 0), candidate, font=font)
        max_width = max(max_width, bbox[2] - bbox[0])
    width = min(max(1100, max_width + padding_x * 2 + 24), max_code_px)
    height = padding_y * 2 + len(rendered) * line_height + 34

    image = Image.new("RGB", (width, height), "#fbfbfb")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width - 1, height - 1), outline="#d0d7de", width=2)
    draw.rectangle((0, 0, width, 36), fill="#f0f3f6")
    draw.text((padding_x, 8), str(SOURCE.relative_to(ROOT)).replace("\\", "/"), fill="#57606a", font=title_font)

    y = padding_y + 30
    for number, text, is_target in rendered:
        if is_target:
            draw.rectangle((padding_x - 8, y - 3, width - padding_x + 8, y + line_height - 2), fill="#fff3bf")
        draw.text((padding_x, y), f"{number:>5}", fill="#6e7781", font=font)
        draw.text((padding_x + gutter_width, y), text, fill="#24292f", font=font)
        y += line_height

    image.save(output_path)


def add_run(paragraph, text: str, *, bold: bool = False, color: str | None = None):
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_evidence_table(document: Document, items: Iterable[dict]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Evidence", "Source line(s)", "What it proves", "Defense use"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for item in items:
        row = table.add_row().cells
        row[0].text = item["title"]
        line_text = f"backend/services/meal_planner.py:{item['line']}"
        if item["end"] != item["line"]:
            line_text += f"-{item['end']}"
        row[1].text = line_text
        row[2].text = item["what"]
        row[3].text = item["defense"]


def build_doc() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()

    for index, item in enumerate(EVIDENCE, start=1):
        context_before = 3
        context_after = 3
        start = max(1, int(item["line"]) - context_before)
        end = min(len(source_lines), int(item["end"]) + context_after)
        output = ASSET_DIR / f"{index:02d}_{safe_name(item['title'])}.png"
        item["image"] = output
        draw_code_image(source_lines, start, end, int(item["line"]), int(item["end"]), output)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.name = "Arial"
    styles["Heading 3"].font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PCOSina Code Evidence Dossier")
    run.bold = True
    run.font.size = Pt(22)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Stage 1 Filtering and Stage 2 MILP-Style OR-Tools CP-SAT Formulation").bold = True

    p = document.add_paragraph()
    p.add_run("Primary source: ").bold = True
    p.add_run(str(SOURCE.relative_to(ROOT)).replace("\\", "/"))

    p = document.add_paragraph()
    p.add_run("Defense wording: ").bold = True
    p.add_run(
        "Our Stage 2 is a MILP-style 0-1 linear optimization formulation implemented with OR-Tools CP-SAT. "
        "It uses binary decision variables, linear constraints, and a weighted minimization objective. "
        "Stage 1 uses deterministic rule-based filtering and scoring so unsafe or incompatible recipes are removed before optimization."
    )

    p = document.add_paragraph()
    p.add_run("Important clarification: ").bold = True
    p.add_run(
        "The code uses if-statements for normal control flow, validation, and safety checks. "
        "That does not mean the meal plan is hard-coded. The final meal combination is selected from solver decision variables."
    )

    document.add_heading("Quick Evidence Map", level=1)
    add_evidence_table(document, EVIDENCE)

    current_category = None
    for index, item in enumerate(EVIDENCE, start=1):
        if item["category"] != current_category:
            current_category = item["category"]
            document.add_page_break()
            document.add_heading(current_category, level=1)

        heading = document.add_heading(f"{index}. {item['title']}", level=2)
        heading.paragraph_format.keep_with_next = True

        p = document.add_paragraph()
        add_run(p, "Source: ", bold=True)
        source_ref = f"backend/services/meal_planner.py:{item['line']}"
        if item["end"] != item["line"]:
            source_ref += f"-{item['end']}"
        p.add_run(source_ref)

        p = document.add_paragraph()
        add_run(p, "What it does: ", bold=True)
        p.add_run(item["what"])

        p = document.add_paragraph()
        add_run(p, "Use in defense: ", bold=True)
        p.add_run(item["defense"])

        document.add_picture(str(item["image"]), width=Inches(9.75))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(f"Screenshot evidence: {source_ref}").italic = True

    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
    print(DOCX_PATH)
