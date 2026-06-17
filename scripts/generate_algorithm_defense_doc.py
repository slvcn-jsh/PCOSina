from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.generate_system_architecture_defense_doc import (
    add_body,
    add_bullets,
    add_callout,
    add_figure,
    add_heading,
    add_numbered,
    add_table,
    configure_document,
)


OUT = ROOT / "docs" / "defense" / "PCOSINA_Algorithm_Defense_Prep_2026-06-04.docx"
SLIDE_IMAGE = ROOT / "tmp" / "defense_prep" / "presentation_pages_png" / "page_3.png"
ASSET_DIR = ROOT / "docs" / "defense" / "PCOSINA_SOP_Evidence_Extraction_Dossier_FINAL_2026-05-24.assets"
ALGORITHM_FLOW = ASSET_DIR / "figure_12.png"
FALLBACK_FLOW = ASSET_DIR / "figure_13.png"
RELAXATION_LADDER = ASSET_DIR / "figure_14.png"
RUNTIME_CHART = ROOT / "docs" / "assets" / "chapter4_final" / "figure_20_planner_runtime_optimization_progression.png"
ALGORITHM_COMPARISON = ROOT / "docs" / "assets" / "algorithm_evaluation_defense_ready" / "figure_4_6_algorithm_comparison.png"
ML_UPLIFT = ROOT / "docs" / "assets" / "algorithm_evaluation_defense_ready" / "figure_4_5_lightgbm_ranking_uplift.png"


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = 1
    title.add_run("PCOSina Algorithm Defense Prep")

    subtitle = doc.add_paragraph()
    subtitle.alignment = 1
    subtitle.add_run(
        "Focused packet for thesis defense: slide guide, MILP-formulated algorithm flow, "
        "CP-SAT solver implementation, safe failure boundary, script, Q&A, evidence, and checklist"
    )

    add_callout(
        doc,
        "Defense identity",
        "PCOSina's algorithm is MILP-formulated and solver-first. The thesis-title phrase "
        "MILP-based refers to the mathematical formulation: a 0-1 weekly meal-assignment "
        "model with binary recipe-slot variables, linear constraints, and objective terms. "
        "OR-Tools CP-SAT is the practical implementation solver used to solve that integer "
        "and constraint model. Stage 1 constructs a safe candidate pool through deterministic "
        "filtering and optional ranking. Stage 2 solves the MILP-formulated assignment for the "
        "final 7-day, 21-slot plan. ML may assist ranking, but it cannot override hard "
        "constraints or choose the final plan.",
    )

    add_table(
        doc,
        ["Item", "Defense-ready answer"],
        [
            [
                "Main point",
                "The algorithm separates safety filtering from final MILP-formulated optimization: unsafe or incompatible recipes are removed first, then CP-SAT solves the final weekly assignment.",
            ],
            [
                "Presentation slide",
                "Slide 3: Algorithm from [Presentation] Quadrant_Presentation_Revised_S2526T3 (1).pdf.",
            ],
            [
                "Manuscript anchors",
                "Chapter 2 algorithmic model paragraphs 210-214; Chapter 3 Figure 12 and Figures 13-14; Chapter 4 Tables 30-33 and Figures 20, 26.",
            ],
            [
                "System anchors",
                "backend/services/meal_planner.py, production_contract.md, planner_contract_2026-05-26.md, ADR-001, ADR-004, ADR-005, and progress_ledger.md.",
            ],
            [
                "Title defense",
                "MILP-based means MILP in formulation; CP-SAT means solver in implementation. Defend the title through binary variables, linear constraints, and objective terms.",
            ],
            [
                "Boundary",
                "Do not say the app returns a greedy fallback as a successful plan. Current production contract returns structured no-safe-plan guidance when no safe full plan is possible.",
            ],
        ],
        [Inches(1.65), Inches(5.65)],
    )

    doc.add_page_break()
    add_heading(doc, "1. Slide And Manuscript Figures", 1)
    add_figure(doc, SLIDE_IMAGE, "Figure ALG-1. Presentation slide used for the Algorithm defense section.", width=7.35)

    doc.add_page_break()
    add_figure(
        doc,
        ALGORITHM_FLOW,
        "Figure ALG-2. Manuscript Figure 12. Two-stage MILP-based meal planning algorithm.",
        width=6.7,
    )

    doc.add_page_break()
    add_figure(
        doc,
        FALLBACK_FLOW,
        "Figure ALG-3. Manuscript Figure 13. Adaptive fallback flowchart for weekly meal-plan generation.",
        width=7.25,
    )
    add_figure(
        doc,
        RELAXATION_LADDER,
        "Figure ALG-4. Manuscript Figure 14. Relaxation ladder used by adaptive fallback.",
        width=7.25,
    )

    doc.add_page_break()
    add_figure(
        doc,
        RUNTIME_CHART,
        "Figure ALG-5. Manuscript Figure 20 / Chapter 4 asset. Planner runtime optimization progression.",
        width=7.25,
    )

    doc.add_page_break()
    add_figure(
        doc,
        ALGORITHM_COMPARISON,
        "Figure ALG-6. Chapter 4 algorithm choice comparison evidence.",
        width=7.25,
    )
    add_figure(
        doc,
        ML_UPLIFT,
        "Figure ALG-7. LightGBM Stage 1 ranking uplift evidence.",
        width=7.25,
    )

    doc.add_page_break()
    add_heading(doc, "2. Core Defense Explanation", 1)
    add_callout(
        doc,
        "30-second answer",
        "Our algorithm is a two-stage, constraint-first meal-planning pipeline. First, PCOSina validates the user profile and prepares solver-ready inputs such as allergies, exclusions, dietary restrictions, goal, budget, cooking-time limit, pantry tokens, and nutrition targets. Stage 1 removes unsafe or incompatible recipes and may rank the remaining safe candidates using deterministic scores and optional LightGBM support. Stage 2 is the MILP-formulated core of the title: the weekly plan is represented as binary recipe-slot decisions with linear constraints and objective penalties, then OR-Tools CP-SAT solves that integer/constraint model to assign breakfast, lunch, and dinner for 7 days. The solver enforces hard constraints such as allergies, restrictions, budget, meal-slot validity, repetition limits, and nutrition bounds. Soft preferences like variety, pantry overlap, cost behavior, and ranking scores influence the objective, but they do not override safety. If no complete safe plan can be found, the system returns no-safe-plan guidance instead of forcing an unsafe recommendation.",
        fill="EAF2F8",
    )

    doc.add_page_break()
    add_heading(doc, "3. Two-Stage Algorithm Flow", 1)
    add_table(
        doc,
        ["Step", "What happens", "Defense sentence"],
        [
            [
                "Input preparation",
                "The app captures profile values, restrictions, allergies, exclusions, budget, cooking-time limit, pantry items, and preference settings.",
                "Inputs are planning data, not diagnostic data.",
            ],
            [
                "Validation",
                "Invalid, missing, conflicting, or unsupported inputs are rejected or converted into safe no-safe-plan style feedback.",
                "The algorithm cannot optimize bad or unsafe inputs.",
            ],
            [
                "Nutrition target generation",
                "The backend computes BMR/TDEE-style calorie targets, macro targets, fiber bounds, and supported policy adjustments.",
                "The solver receives bounded numeric targets instead of vague health text.",
            ],
            [
                "Stage 1 filtering",
                "Recipes are removed when they violate allergies, recognized dietary restrictions, explicit exclusions, meal-slot suitability, or other hard rules.",
                "Unsafe recipes do not reach the final optimizer.",
            ],
            [
                "Stage 1 ranking",
                "The remaining safe recipes are scored using pantry overlap, cost, nutrition fit, preference signals, and optional LightGBM ranking.",
                "Ranking changes order, not safety authority.",
            ],
            [
                "Shortlisting",
                "The system caps and balances candidate pools so CP-SAT solves a tractable model under runtime limits.",
                "Shortlisting reduces search pressure while preserving safe candidates.",
            ],
            [
                "Stage 2 MILP-formulated CP-SAT solve",
                "The backend creates boolean recipe-slot variables, linear constraints, and objective terms, then solves the 7-day, 21-slot assignment model.",
                "The MILP-formulated model is the thesis core; CP-SAT is the solver implementation and final meal-plan authority.",
            ],
            [
                "Output or no-safe-plan",
                "If a complete safe plan exists, the backend returns meals, recipes, grocery guidance, explanations, and metadata. If not, it returns no-safe-plan reasons and safe adjustment guidance.",
                "No unsafe partial plan is accepted as a successful result.",
            ],
        ],
        [Inches(1.35), Inches(3.45), Inches(2.5)],
    )

    add_heading(doc, "4. Mathematical Model Quick Version", 1)
    add_callout(
        doc,
        "Plain formula",
        "This is the MILP-formulated part of the thesis title. Let x[s,i] = 1 if recipe i is assigned to meal slot s, otherwise 0. "
        "There are 21 slots: breakfast, lunch, and dinner across 7 days. "
        "The model uses linear constraints and objective penalties; OR-Tools CP-SAT chooses the x values that satisfy hard constraints and minimize weighted penalties for deviation, repetition, cost behavior, pantry mismatch, and ranking preference.",
    )
    add_table(
        doc,
        ["Model element", "Current defense wording", "What to avoid"],
        [
            [
                "Decision variable",
                "Binary x[slot, recipe] variables decide whether a recipe is selected for a specific meal slot.",
                "Do not describe it as a random recommender list.",
            ],
            [
                "Slot constraint",
                "Each meal slot must receive exactly one allowed recipe.",
                "Do not say the plan can leave blank meals as success.",
            ],
            [
                "Hard pre-solver filters",
                "Allergies, recognized restrictions, explicit exclusions, and invalid meal-slot matches are removed before solve.",
                "Do not say ML can reinsert filtered recipes.",
            ],
            [
                "Hard MILP/CP-SAT constraints",
                "The MILP-formulated model defines slot assignment, repetition limits, active weekly budget cap, calorie/macro/fiber bounds, and solver-valid meal structure; CP-SAT enforces them in implementation.",
                "Do not say soft preferences are equal to hard medical rules.",
            ],
            [
                "Advisory penalties",
                "Sodium and sugar are displayable and penalized as advisory overage targets in the current planner contract.",
                "Do not call sodium/sugar a full infeasibility gate in the current system.",
            ],
            [
                "Pantry behavior",
                "Pantry overlap steers scoring, objective rewards, and grocery-gap guidance.",
                "Do not claim the planner proves every pantry ingredient is physically available.",
            ],
            [
                "Objective",
                "The objective minimizes weighted deviations while rewarding feasible, practical, varied, and preference-aligned plans.",
                "Do not say the cheapest plan always wins.",
            ],
            [
                "Infeasibility",
                "If no safe complete assignment exists, the API returns no-safe-plan diagnostics and guidance.",
                "Do not call an unsafe greedy plan a fallback success.",
            ],
        ],
        [Inches(1.45), Inches(3.6), Inches(2.25)],
    )

    doc.add_page_break()
    add_heading(doc, "5. Repo-Truth Cross-Check", 1)
    add_table(
        doc,
        ["Claim", "Current source evidence", "Defense-safe wording"],
        [
            [
                "Backend planner authority",
                "backend/services/meal_planner.py defines shortlist_candidates() and solve_meal_plan(); solve_meal_plan builds cp_model.CpModel(), NewBoolVar variables, Add constraints, and Minimize objective.",
                "The backend solves the MILP-formulated 0-1 assignment with CP-SAT as the authoritative selector for new optimized plans.",
            ],
            [
                "Two-stage contract",
                "docs/architecture/production_contract.md defines Stage 1 filtering/ranking/candidate generation and Stage 2 as a formal 0-1 model solved by OR-Tools CP-SAT.",
                "Stage 1 prepares safe candidates; Stage 2 solves the MILP-formulated final weekly assignment.",
            ],
            [
                "ML boundary",
                "ADR-004 states ML scores are allowed only in Stage 1 candidate ranking; CP-SAT remains final authority.",
                "ML improves ordering but cannot override hard constraints or select unsafe plans.",
            ],
            [
                "No unsafe fallback",
                "ADR-005 requires status=no-safe-plan with reason codes, guidance, suggested relaxations, policy version, and diagnostics when authoritative solve is infeasible.",
                "PCOSina explains blockers instead of returning an unsafe weak plan.",
            ],
            [
                "Constraint classification",
                "planner_contract_2026-05-26 separates hard, soft, advisory, and tracking fields.",
                "Use hard/soft/advisory/tracking language when explaining tradeoffs.",
            ],
            [
                "Pantry boundary",
                "planner_contract_2026-05-26 states planning is pantry-aware, not hard pantry-feasible.",
                "Pantry supports scoring and grocery-gap guidance; it is not inventory proof.",
            ],
            [
                "Benchmark evidence",
                "progress_ledger.md records final 100/100 valid runs, 0 hard-rule violations, average 257 ms, P95 313 ms, max 357 ms.",
                "The tested planner was fast and constraint-preserving under the benchmark profile set.",
            ],
            [
                "ML evidence",
                "ml_progress_ledger.md and Chapter 4 assets report NDCG@10 0.9211 versus 0.8838 baseline.",
                "ML helped rank candidates, but final plan authority stayed deterministic.",
            ],
        ],
        [Inches(1.35), Inches(3.55), Inches(2.4)],
    )

    add_heading(doc, "6. Full Speaker Script", 1)
    script_paragraphs = [
        "Good day panel. This slide explains the algorithm used by PCOSina to generate a weekly meal plan. The safest way to describe it is a MILP-formulated, solver-first hybrid meal-planning algorithm. It is not a pure AI generator, and it is not a simple decision tree. It is a two-stage pipeline that combines deterministic filtering, optional ranking support, and CP-SAT optimization.",
        "The algorithm begins with user and system inputs. These include age, height, weight, activity level, goal, dietary restrictions, allergies, ingredient exclusions, weekly budget, cooking-time preference, pantry entries, and recipe or nutrition data. These values are validated and converted into structured tokens, numeric targets, and policy-controlled planning values. This is important because the optimizer can only enforce rules correctly when the inputs are clean and bounded.",
        "Stage 1 is candidate construction. In this stage, recipes that violate allergies, recognized dietary restrictions, explicit exclusions, meal-slot suitability, or basic feasibility are removed before optimization. The system may then score or rank the safe recipes using pantry overlap, calorie alignment, nutrition properties, cost behavior, preference signals, and optional LightGBM support. However, the machine learning part is non-authoritative. It only changes ordering or priority among already safe candidates.",
        "Stage 2 is the formal MILP-formulated optimization stage. This is the part that supports the MILP-based title. The backend models the weekly meal plan as binary decisions: whether a recipe is assigned to a specific meal slot. There are 21 meal slots because the plan covers breakfast, lunch, and dinner for 7 days. The formulation uses linear constraints and objective penalties, while OR-Tools CP-SAT is the practical solver implementation that searches for a complete assignment satisfying hard constraints and minimizing objective penalties related to nutrition deviation, variety, repetition, pantry use, cost behavior, and ranking preference.",
        "The important boundary is that hard constraints remain protected. Allergies, explicit exclusions, dietary restrictions, supported nutrition bounds, active budget caps, and final solver feasibility cannot be overridden by ML or by soft preferences. Pantry data is used as a planning signal and grocery guidance source, but the current system should be described as pantry-aware rather than a full physical inventory proof.",
        "If the model cannot find a safe complete weekly plan, PCOSina does not force a weak or unsafe plan. It returns a no-safe-plan response with reason codes, diagnostics, and suggested safe adjustments such as reviewing budget, cooking-time limit, variety preference, or overly strict restrictions. This is why the algorithm is defensible: it prioritizes feasibility, safety, explainability, and practical runtime over uncontrolled recommendation output.",
        "The final benchmark evidence supports this design. The tested planner completed 100 out of 100 valid benchmark runs with zero hard-rule violations, and the final optimized runtime was 257 milliseconds on average, with P95 at 313 milliseconds and maximum runtime at 357 milliseconds. LightGBM improved candidate ranking, with NDCG@10 of 0.9211 versus 0.8838 baseline, but CP-SAT remained the final selector.",
    ]
    for paragraph in script_paragraphs:
        add_body(doc, paragraph)

    doc.add_page_break()
    add_heading(doc, "7. Numbers And Terms To Memorize", 1)
    add_table(
        doc,
        ["Evidence point", "Value", "How to say it"],
        [
            ["Algorithm type", "Two-stage MILP-formulated solver-first hybrid", "Deterministic filtering plus optional ranking plus CP-SAT final assignment."],
            ["Stage 1", "Preference-aware candidate construction", "Removes unsafe recipes and ranks safe candidates."],
            ["Stage 2", "MILP-formulated 0-1 assignment solved by OR-Tools CP-SAT", "Binary recipe-slot variables, linear constraints, and objective terms select the final weekly plan."],
            ["Title defense", "MILP-based = formulation; CP-SAT = solver", "Do not call CP-SAT a pure MILP solver; say it solves the MILP-formulated integer/constraint model."],
            ["Planning horizon", "7 days / 21 meal slots", "Breakfast, lunch, and dinner for one week."],
            ["ML role", "Stage 1 ranking only", "ML can rank safe candidates; it cannot select unsafe meals."],
            ["Final benchmark", "100/100 runs; 0 hard-rule violations", "The tested planner preserved hard constraints."],
            ["Runtime", "Average 257 ms; P95 313 ms; max 357 ms", "The final optimized benchmark was sub-second."],
            ["ML ranking metric", "NDCG@10 0.9211 vs 0.8838 baseline", "Ranking improved while CP-SAT remained authoritative."],
            ["Failure behavior", "status=no-safe-plan", "No unsafe fallback is returned as a valid plan."],
            ["Pantry boundary", "Pantry-aware, not full pantry-feasible", "Pantry influences scoring and grocery guidance, not guaranteed inventory proof."],
        ],
        [Inches(1.7), Inches(2.0), Inches(3.6)],
    )

    doc.add_page_break()
    add_heading(doc, "8. Say This, Not That", 1)
    add_table(
        doc,
        ["Avoid saying", "Say instead", "Why"],
        [
            [
                "It is an AI meal planner.",
                "It is a deterministic two-stage planner with optional ML ranking support.",
                "ML is assistive only; the solver keeps final authority.",
            ],
            [
                "The slide says fallback greedy, so greedy returns a plan.",
                "The current safe wording is adaptive retry or safe relaxation of non-safety preferences, followed by no-safe-plan if needed.",
                "ADR-005 rejects unsafe fallback plans as successful output.",
            ],
            [
                "MILP is the actual solver.",
                "The model is MILP-formulated 0-1 assignment, solved in implementation by OR-Tools CP-SAT.",
                "This distinguishes mathematical formulation from implementation solver.",
            ],
            [
                "CP-SAT means the title is not MILP-based.",
                "MILP-based refers to the formulation; CP-SAT refers to the solver implementation.",
                "The defense should separate model identity from solver technology.",
            ],
            [
                "Pantry constraints guarantee available ingredients.",
                "The current planner is pantry-aware and produces grocery-gap guidance.",
                "The repo contract says pantry is not hard inventory proof.",
            ],
            [
                "ML optimizes the final meal plan.",
                "ML may rank safe candidates; CP-SAT optimizes the final plan.",
                "This prevents safety-authority confusion.",
            ],
            [
                "If constraints conflict, the app still gives a plan.",
                "If a complete safe plan is impossible, PCOSina returns no-safe-plan reasons and safe suggestions.",
                "Unsafe approximate outputs are not accepted as valid plans.",
            ],
            [
                "Sodium and sugar are always hard infeasibility gates.",
                "Current planner treats sodium and sugar as advisory overage penalties while displaying them in the app.",
                "This matches the 2026-05-26 planner contract.",
            ],
        ],
        [Inches(2.1), Inches(3.0), Inches(2.2)],
        header_fill="B94B61",
    )

    add_heading(doc, "9. Panel Q&A", 1)
    qas = [
        (
            "What is the main idea of your algorithm?",
            "It is a two-stage MILP-formulated, constraint-first planner. Stage 1 removes unsafe or incompatible recipes and ranks safe candidates; Stage 2 solves a binary recipe-slot assignment through OR-Tools CP-SAT to produce the final 7-day meal plan.",
        ),
        (
            "Why use two stages instead of one big solver?",
            "Stage 1 reduces the search space and protects safety early. Stage 2 then solves a smaller, cleaner assignment problem, which improves runtime while preserving hard constraints.",
        ),
        (
            "Why call it MILP-based if you use CP-SAT?",
            "The title is defensible because MILP-based refers to how the planning problem is formulated: binary decision variables, linear constraints, and objective terms. CP-SAT is not a pure MILP solver; it is the implementation solver used to solve the integer/constraint model.",
        ),
        (
            "What are the decision variables?",
            "The main decision variable is x[slot, recipe]. It equals 1 when a recipe is assigned to a meal slot and 0 otherwise.",
        ),
        (
            "What are your hard constraints?",
            "Hard constraints include allergy and restriction exclusion, explicit ingredient exclusions, meal-slot validity, active budget cap, repetition limits, calorie and macro/fiber bounds, and final solver feasibility.",
        ),
        (
            "What are soft constraints?",
            "Soft constraints guide ranking or objective weights but do not replace safety. Examples include variety preference, planning priority, pantry overlap, and preference/ranking signals.",
        ),
        (
            "Where does ML fit?",
            "ML sits only in Stage 1 ranking. It can improve the order of already safe candidates, but it cannot override filtering or choose the final meal plan.",
        ),
        (
            "What happens if no safe plan exists?",
            "The system returns a structured no-safe-plan result with reason codes, diagnostics, and suggested safe adjustments. It does not force an unsafe or incomplete plan.",
        ),
        (
            "Does pantry act as a hard constraint?",
            "No. Current backend planning is pantry-aware. Pantry overlap rewards likely useful recipes and supports grocery-gap guidance, but it does not prove every ingredient is physically available.",
        ),
        (
            "How did you prove it works?",
            "The final benchmark completed 100/100 valid runs with zero hard-rule violations, average runtime of 257 ms, P95 of 313 ms, and maximum of 357 ms. The ML ranking component also improved NDCG@10 from 0.8838 to 0.9211.",
        ),
        (
            "Why not let generative AI produce the plan?",
            "Because unconstrained generation can violate hard rules. PCOSina uses deterministic filtering and CP-SAT so allergy, budget, nutrition, and safety boundaries are auditable.",
        ),
        (
            "What is the limitation of the algorithm?",
            "It depends on recipe dataset quality, nutrition estimates, current policy settings, and backend availability for new optimized generation. It is a wellness decision-support planner, not a clinical decision engine.",
        ),
    ]
    for q, a in qas:
        add_body(doc, f"Q: {q}", bold_prefix="Q:")
        add_body(doc, f"A: {a}", bold_prefix="A:")

    add_heading(doc, "10. Quick Whiteboard Version", 1)
    add_body(doc, "If asked to explain without the slide, draw seven boxes:")
    add_numbered(
        doc,
        [
            "User inputs: profile, goals, restrictions, allergies, exclusions, budget, pantry, preferences.",
            "Validation and preparation: normalize text, compute calorie/macros, resolve policy values.",
            "Stage 1 hard filtering: remove unsafe or incompatible recipes.",
            "Stage 1 ranking and shortlist: score pantry/cost/nutrition/preference and optional ML score.",
            "Stage 2 MILP-formulated CP-SAT model: x[slot, recipe] binary assignment across 21 meal slots.",
            "Objective and constraints: enforce hard constraints, minimize weighted soft deviations.",
            "Result: full safe plan, or no-safe-plan diagnostics and safe adjustment guidance.",
        ],
    )

    add_heading(doc, "11. Source And Reference Anchors", 1)
    add_table(
        doc,
        ["Anchor", "Use"],
        [
            ["Presentation PDF page 3", "Main defense slide for Algorithm."],
            ["Manuscript paragraphs 210-214", "Algorithmic model: MILP-formulated 0-1 formulation solved by OR-Tools CP-SAT."],
            ["Manuscript paragraphs 270-271", "Stage 2 MILP-formulated CP-SAT weekly assignment explanation."],
            ["Manuscript Figure 12", "Two-stage MILP-based meal planning algorithm."],
            ["Manuscript Figures 13-14", "Adaptive retry/relaxation and safe no-safe-plan boundary."],
            ["Manuscript Tables 30-33", "Candidate filtering, CP-SAT model elements, benchmark evidence, safety outcomes."],
            ["Manuscript Figure 20", "Runtime optimization progression."],
            ["Manuscript Figure 26", "Algorithm choice comparison and hard-rule behavior."],
            ["backend/services/meal_planner.py", "Current planner implementation: Stage 1 shortlisting and Stage 2 MILP-formulated CP-SAT solve."],
            ["docs/architecture/production_contract.md", "Two-stage planning authority and no unsafe output contract."],
            ["docs/architecture/planner_contract_2026-05-26.md", "Hard, soft, advisory, tracking classifications and pantry boundary."],
            ["docs/adr/ADR-001-remote-cpsat-authority.md", "Remote CP-SAT is authoritative planner."],
            ["docs/adr/ADR-004-ml-shadow-canary.md", "ML Stage 1 ranking only."],
            ["docs/adr/ADR-005-no-safe-plan-contract.md", "Structured no-safe-plan instead of unsafe fallback."],
            ["docs/thesis_validation/02_FORMULAS_AND_COMPUTATIONS.md", "Manual computation evidence and implemented formulas."],
            ["docs/thesis_validation/04_DECISION_TREES.md", "Decision-flow diagrams based on actual code paths; not ML decision trees."],
            ["docs/roadmap/progress_ledger.md", "Current benchmark status and runtime evidence."],
        ],
        [Inches(2.65), Inches(4.65)],
    )

    add_heading(doc, "12. References To Keep Ready", 1)
    add_bullets(
        doc,
        [
            "Optimization and integer-programming diet-planning references in the manuscript: use them to justify binary assignment, constraints, and objective functions.",
            "Guevarra et al. (2022) and De Leon et al. (2022): use as local optimization-meal-planning support.",
            "Benvenuti et al. (2024): use for diet optimization and acceptability/cultural practicality framing.",
            "OR-Tools CP-SAT references in the manuscript and codebase: use to justify integer/combinatorial optimization implementation.",
            "LightGBM/personalized ranking references: use only for candidate prioritization, not final medical or meal-plan authority.",
            "ISO/IEC 25010 references: use for software quality evaluation, not clinical effectiveness.",
        ],
    )

    add_heading(doc, "13. Study Checklist", 1)
    add_bullets(
        doc,
        [
            "Memorize the 30-second answer.",
            "Practice saying: MILP-formulated/MILP-based 0-1 assignment solved by OR-Tools CP-SAT.",
            "Be ready to answer: CP-SAT is not a pure MILP solver; the title refers to the formulation.",
            "Practice the Stage 1 versus Stage 2 boundary.",
            "Memorize x[slot, recipe] as the decision variable explanation.",
            "Memorize hard versus soft versus advisory constraints.",
            "Practice the no-safe-plan answer and avoid saying unsafe greedy fallback.",
            "Memorize final benchmark: 100/100 runs, 0 hard-rule violations, 257 ms average, P95 313 ms, max 357 ms.",
            "Memorize ML metric: NDCG@10 0.9211 versus 0.8838 baseline.",
            "Avoid saying AI final planner, pantry hard feasibility, or clinical diagnosis engine.",
        ],
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendix A. Manuscript Anchor Summary", 1)
    add_table(
        doc,
        ["Paragraphs / Figures", "Key content"],
        [
            ["210-214", "MILP-formulated 0-1 weekly assignment formulation solved by OR-Tools CP-SAT; binary variables, constraints, and objective terms."],
            ["270-271", "Stage 2 MILP-formulated CP-SAT assignment: filtered candidates become binary recipe-slot variables over meal slots."],
            ["Figure 12", "Two-stage meal-planning algorithm: user input, validation, Stage 1 filtering/ranking, shortlist, Stage 2 optimization, output, feedback."],
            ["Figures 13-14", "Adaptive fallback and relaxation ladder: soft adjustment attempts are allowed, hard safety rules stay protected, no-safe-plan if needed."],
            ["Tables 30-33", "Candidate filtering, CP-SAT model, benchmark evidence, and final safety outcomes."],
            ["Figure 20", "Runtime progression from slow baseline to final optimized 257 ms average with 0 hard-rule violations."],
            ["Figure 26", "Algorithm comparison showing constraint-first CP-SAT behavior versus unconstrained proxy behavior."],
            ["1460-1478", "Conclusion: two-stage algorithm implemented; 21-slot schedules; 100/100 benchmark runs; no-safe-plan for impossible cases."],
        ],
        [Inches(1.65), Inches(5.65)],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
