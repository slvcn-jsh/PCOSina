"""Generate smaller Breakfast recipe-record DOCX parts.

The full Breakfast record volume is hard to upload to Google Docs. This script
regenerates the same Breakfast recipe records in original order, split into
smaller DOCX files. Data comes from the same source files and generator helpers
used by the full recipe-record packet.
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = ROOT / "backend"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from docx import Document

from scripts.generate_recipe_philfct_pricing_records import (
    GENERATED_DATE,
    PHILFCT_CACHE_PATH,
    PRICED_DRAFT_PATH,
    RUNTIME_PATH,
    TEMPLATE_PATH,
    add_recipe_record,
    add_table,
    add_title,
    configure_document,
    git_commit_short,
    meal_type,
    philfct_code_index,
    priced_meal_index,
    read_json,
    recipe_evidence_rows,
    recipe_id,
    recipe_name,
    sorted_recipes,
)


OUT_DIR = (
    ROOT
    / "docs"
    / "thesis_validation"
    / "03_ACTUAL_SYSTEM_DATA_EXPORTS"
    / "recipe_records"
    / "breakfast_split_90_page_target"
)
INDEX_CSV = OUT_DIR / "pcosina_breakfast_split_record_index.csv"
INDEX_JSON = OUT_DIR / "pcosina_breakfast_split_record_index.json"
RECORDS_PER_PART = 30


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    return " ".join(text.split())


def chunked(items: list[dict[str, Any]], size: int) -> list[list[dict[str, Any]]]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def part_filename(part_number: int, start_record: int, end_record: int) -> str:
    return (
        "PCOSINA_PhilFCT_Pricing_Recipe_Records_Breakfast_"
        f"Part_{part_number:02d}_Records_{start_record:03d}-{end_record:03d}.docx"
    )


def build_part_doc(
    part_number: int,
    total_parts: int,
    part_recipes: list[dict[str, Any]],
    absolute_start: int,
    rows_by_recipe: dict[str, list[dict[str, Any]]],
    git_short: str,
) -> Path:
    absolute_end = absolute_start + len(part_recipes) - 1
    output_path = OUT_DIR / part_filename(part_number, absolute_start, absolute_end)
    doc = Document()
    configure_document(doc)
    add_title(
        doc,
        f"PCOSina PhilFCT and Pricing Recipe Records - Breakfast Part {part_number:02d} of {total_parts:02d}",
        "Split Breakfast recipe-record packet for easier Google Docs upload.",
        (
            f"Generated: {GENERATED_DATE} | Git commit: {git_short} | "
            f"Template inspected: {TEMPLATE_PATH.relative_to(ROOT)} | "
            f"Original record range: {absolute_start:03d}-{absolute_end:03d}"
        ),
    )
    doc.add_heading("Split Volume Scope", level=1)
    add_table(
        doc,
        ["Question", "Answer"],
        [
            ["Which records are included?", f"Breakfast records {absolute_start:03d}-{absolute_end:03d} in the same order as the full Breakfast volume."],
            ["Why split this file?", "The original Breakfast volume is large for Google Docs upload. This part targets about 90 rendered pages, but exact page count depends on Word/Google Docs pagination."],
            ["Was any recipe data changed?", "No. Recipe sections are regenerated from the same source data and generator helpers as the full record packet."],
            ["What should reviewers edit?", "Use the red review tables: Reviewer, Date checked, and Correction notes. Edit data cells only when correcting a verified source value."],
        ],
        widths=[2.4, 7.2],
        font_size=7.6,
        header_fill="5B9BD5",
    )

    for offset, recipe in enumerate(part_recipes):
        absolute_record = absolute_start + offset
        add_recipe_record(
            doc,
            recipe,
            rows_by_recipe.get(recipe_id(recipe), []),
            absolute_record,
            first=offset == 0,
        )
    doc.save(output_path)
    return output_path


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for stale_docx in OUT_DIR.glob("PCOSINA_PhilFCT_Pricing_Recipe_Records_Breakfast_Part_*.docx"):
        stale_docx.unlink()
    for stale_index in (INDEX_CSV, INDEX_JSON):
        if stale_index.exists():
            stale_index.unlink()

    recipes = read_json(RUNTIME_PATH)
    priced = read_json(PRICED_DRAFT_PATH)
    philfct = read_json(PHILFCT_CACHE_PATH)
    priced_by_recipe = priced_meal_index(priced)
    philfct_by_code = philfct_code_index(philfct)

    breakfast_recipes = [recipe for recipe in sorted_recipes(recipes) if meal_type(recipe) == "Breakfast"]
    rows_by_recipe = {
        recipe_id(recipe): recipe_evidence_rows(recipe, priced_by_recipe, philfct_by_code)
        for recipe in breakfast_recipes
    }
    parts = chunked(breakfast_recipes, RECORDS_PER_PART)
    git_short = git_commit_short()

    output_records: list[dict[str, Any]] = []
    for part_number, part_recipes in enumerate(parts, start=1):
        absolute_start = ((part_number - 1) * RECORDS_PER_PART) + 1
        output_path = build_part_doc(
            part_number,
            len(parts),
            part_recipes,
            absolute_start,
            rows_by_recipe,
            git_short,
        )
        for offset, recipe in enumerate(part_recipes):
            record_number = absolute_start + offset
            output_records.append(
                {
                    "recordNumber": record_number,
                    "recipeId": recipe_id(recipe),
                    "recipeName": recipe_name(recipe),
                    "mealType": meal_type(recipe),
                    "partNumber": part_number,
                    "partFile": output_path.name,
                }
            )

    with INDEX_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(output_records[0].keys()))
        writer.writeheader()
        writer.writerows(output_records)
    INDEX_JSON.write_text(
        json.dumps(
            {
                "generatedDate": GENERATED_DATE,
                "gitCommitShort": git_short,
                "recordsPerPart": RECORDS_PER_PART,
                "breakfastRecordCount": len(breakfast_recipes),
                "partCount": len(parts),
                "records": output_records,
            },
            indent=2,
            ensure_ascii=True,
        ),
        encoding="utf-8",
    )

    summary = {
        "outputDir": str(OUT_DIR),
        "partCount": len(parts),
        "breakfastRecordCount": len(breakfast_recipes),
        "recordsPerPart": RECORDS_PER_PART,
        "files": sorted(path.name for path in OUT_DIR.glob("*.docx")),
        "indexCsv": str(INDEX_CSV),
        "indexJson": str(INDEX_JSON),
    }
    print(json.dumps(summary, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
