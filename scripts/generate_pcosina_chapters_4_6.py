from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date
from pathlib import Path
from statistics import mean

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


REPO_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = REPO_ROOT / "docs" / "thesis" / "QUADRANT_Final Chapters 1-3.docx"
OUTPUT_PATH = REPO_ROOT / "docs" / "thesis" / "PCOSINA_Chapter_4-6.docx"
DATED_OUTPUT_PATH = REPO_ROOT / "docs" / "thesis" / f"PCOSINA_Chapter_4-6_UPDATED_{date.today().isoformat()}.docx"
GOOGLE_DOCS_OUTPUT_PATH = REPO_ROOT / "docs" / "thesis" / "PCOSINA_Chapter_4-6_GOOGLE_DOCS_READY.docx"
GOOGLE_DOCS_JPEG_OUTPUT_PATH = REPO_ROOT / "docs" / "thesis" / "PCOSINA_Chapter_4-6_GOOGLE_DOCS_JPEG_SAFE.docx"
GOOGLE_DOCS_NATIVE_OUTPUT_PATH = REPO_ROOT / "docs" / "thesis" / "PCOSINA_Chapter_4-6_GOOGLE_DOCS_NATIVE_DIAGRAMS.docx"
GOOGLE_DOCS_PLACEHOLDER_OUTPUT_PATH = REPO_ROOT / "docs" / "thesis" / "PCOSINA_Chapter_4-6_GOOGLE_DOCS_PLACEHOLDERS.docx"
GOOGLE_DOCS_ASSET_DIR = REPO_ROOT / "output" / "doc" / "assets" / "chapter_4_6_google_docs_safe"

DATASET_MANIFEST_PATH = REPO_ROOT / "ml" / "offline_training" / "artifacts" / "dataset_v1" / "dataset_manifest.json"
TRAINING_METRICS_PATH = REPO_ROOT / "ml" / "offline_training" / "artifacts" / "model_v1" / "training_metrics.json"
ML_LEDGER_PATH = REPO_ROOT / "docs" / "roadmap" / "ml_progress_ledger.json"
PROGRESS_LEDGER_PATH = REPO_ROOT / "docs" / "roadmap" / "progress_ledger.md"
APP_GRADLE_PATH = REPO_ROOT / "app" / "build.gradle.kts"
VERSION_CATALOG_PATH = REPO_ROOT / "gradle" / "libs.versions.toml"
RESPONDENT_RELEASE_NOTES_PATH = REPO_ROOT / "release-notes" / "respondent-test-notes.txt"
RESPONDENT_DEPLOYMENT_DOC_PATH = REPO_ROOT / "docs" / "respondent-deployment.md"
PRE_RELEASE_CHECKLIST_PATH = REPO_ROOT / "docs" / "pre-release-checklist.md"

FIGURE_PATHS = {
    "Figure 4.1": REPO_ROOT / "output" / "doc" / "assets" / "comment11" / "figure_4_1_effectiveness_evaluation_flowchart.png",
    "Figure 4.2": REPO_ROOT / "output" / "doc" / "assets" / "comment11" / "figure_4_2_effectiveness_dimension_map.png",
    "Figure 4.3": REPO_ROOT / "output" / "doc" / "assets" / "comment11" / "figure_4_3_benchmark_evidence_snapshot.png",
    "Figure 4.4": REPO_ROOT / "output" / "doc" / "assets" / "comment12" / "figure_4_4_existing_vs_proposed_comparison_design.png",
    "Figure 4.5": REPO_ROOT / "output" / "doc" / "assets" / "comment12" / "figure_4_5_scenario_runtime_and_quality_comparison.png",
    "Figure 4.6": REPO_ROOT / "output" / "doc" / "assets" / "comment12" / "figure_4_6_acceptance_criteria_matrix.png",
}

SYSTEM_FLOW_ROWS = [
    ("Splash", "Initial branding, offline-first framing, and startup feedback."),
    ("Login", "Existing user sign-in and local account handoff."),
    ("Sign Up", "Account creation, validation, and verification guidance."),
    ("Onboarding", "First-time orientation and expectation setting."),
    ("Profile Setup Step 1", "Identity, anthropometrics, and health-profile entry."),
    ("Profile Setup Step 2", "Dietary restrictions, symptom inputs, and medical planning constraints."),
    ("Profile Setup Step 3", "Household size, pantry behavior, and planning preferences."),
    ("Goal Selection", "User chooses weight loss, symptom management, or general health focus."),
    ("Dashboard", "Status-first weekly landing page with next-step guidance."),
    ("Meal Plan", "Weekly plan generation, day browsing, and meal swap flow."),
    ("Grocery List", "Pantry-aware grocery generation with category grouping and sharing."),
    ("Progress", "Meal logging, reflection capture, and weekly/daily progress tracking."),
    ("Recipe Details", "Recipe rationale, ingredients, instructions, and meal logging handoff."),
    ("Settings", "Reminder controls, profile summary, and maintenance tools."),
]

GENERAL_OBJECTIVE = (
    "To develop and evaluate PCOSina, an offline-first, preference-aware weekly meal "
    "planning mobile application for Filipino individuals diagnosed with PCOS using a "
    "two-stage deterministic filtering and optimization approach with pantry-aware "
    "grocery support."
)

SPECIFIC_OBJECTIVES = [
    "Design and implement a two-stage meal-planning algorithm that filters recipes based on user preferences and optimizes weekly meal plans using deterministic OR-Tools CP-SAT optimization.",
    "Develop a personalization mechanism incorporating dietary restrictions, food preferences, and PCOS-related nutritional goals.",
    "Integrate a digital pantry inventory system that prioritizes available ingredients and generates efficient grocery recommendations to reduce food waste.",
    "Implement an offline-first mobile application architecture ensuring core functionalities operate without internet connectivity.",
    "Evaluate system performance in terms of nutritional adequacy, usability, cultural relevance, pantry utilization, and computational efficiency using algorithmic and user-centered metrics.",
]

SCENARIO_RUNTIME_ROWS = [
    ("S1", "Base", 0.2852, 1.9052, "100.00%", "100.00%", "n/a", "n/a", "0.00", "3.00", "0.00", "0.00", "0.00%"),
    ("S2", "BudgetNoPork", 0.2240, 1.1701, "100.00%", "100.00%", "5.60%", "6.20%", "0.00", "3.00", "0.00", "0.00", "0.00%"),
    ("S3", "VegLactose", 0.1366, 3.9361, "100.00%", "100.00%", "n/a", "n/a", "0.00", "3.19", "0.00", "0.00", "0.00%"),
    ("S4", "PantryConstrained", 0.2898, 2.4698, "100.00%", "100.00%", "n/a", "n/a", "0.00", "3.00", "0.00", "0.00", "0.00%"),
    ("S5", "ResilienceProxy300", 6.9306, 2.2258, "100.00%", "100.00%", "n/a", "n/a", "n/a", "n/a", "n/a", "n/a", "0.00%"),
]

SCENARIO_QUALITY_ROWS = [
    ("S1", 531.57, 530.14, 33.27, 33.10, 0, 0),
    ("S2", 540.43, 539.00, 33.27, 33.02, 0, 0),
    ("S3", 605.43, 601.86, 36.31, 36.28, 0, 0),
    ("S4", 565.71, 557.14, 33.97, 33.10, 0, 0),
]

LIVE_BEFORE_RUNS_MS = [207816, 188249]
LIVE_AFTER_RUNS_MS = [9803, 13674, 6890, 6486, 9593]


def set_rtl_font(style) -> None:
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def clear_document(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_rtl_font(normal)
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    for name, size in [("Heading 1", 12), ("Heading 2", 12), ("Heading 3", 12)]:
        style = doc.styles[name]
        set_rtl_font(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.first_line_indent = Inches(0)


def configure_page_setup(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def load_android_version(path: Path) -> tuple[str, str]:
    text = path.read_text(encoding="utf-8")
    version_code_match = re.search(r"versionCode\s*=\s*(\d+)", text)
    version_name_match = re.search(r'versionName\s*=\s*"([^"]+)"', text)
    version_code = version_code_match.group(1) if version_code_match else "not found"
    version_name = version_name_match.group(1) if version_name_match else "not found"
    return version_name, version_code


def load_catalog_version(path: Path, key: str) -> str:
    pattern = re.compile(rf'^{re.escape(key)}\s*=\s*"([^"]+)"\s*$', re.MULTILINE)
    match = pattern.search(path.read_text(encoding="utf-8"))
    return match.group(1) if match else "not found"


def load_first_nonblank_line(path: Path) -> str:
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if stripped:
            return stripped
    return "not found"


def file_status(path: Path) -> str:
    return "Present" if path.exists() else "Missing"


def load_progress_status_counts(path: Path) -> Counter:
    counts: Counter[str] = Counter()
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        if "Roadmap Item" in line or line.startswith("|---"):
            continue
        parts = [part.strip() for part in line.strip("|").split("|")]
        if len(parts) >= 2:
            counts[parts[1]] += 1
    return counts


def count_test_sources() -> dict[str, int]:
    return {
        "Backend Python tests": sum(1 for _ in (REPO_ROOT / "backend" / "tests").rglob("*.py")),
        "Android JVM unit test files": sum(
            1
            for path in (REPO_ROOT / "app" / "src" / "test" / "java").rglob("*")
            if path.suffix in {".kt", ".java"}
        ),
        "Android instrumentation test files": sum(
            1
            for path in (REPO_ROOT / "app" / "src" / "androidTest" / "java").rglob("*")
            if path.suffix in {".kt", ".java"}
        ),
    }


def add_paragraph(doc: Document, text: str = "", *, bold: bool = False, italic: bool = False,
                  centered: bool = False, first_line_indent: float | None = 0.5,
                  single_spaced: bool = False) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line_indent is None:
        paragraph.paragraph_format.first_line_indent = Inches(0)
    else:
        paragraph.paragraph_format.first_line_indent = Inches(first_line_indent)
    if single_spaced:
        paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    font = run.font
    font.name = "Times New Roman"
    font.size = Pt(12)


def add_blank_line(doc: Document) -> None:
    add_paragraph(doc, "", first_line_indent=None, single_spaced=True)


def add_chapter_title(doc: Document, chapter_no: int, title: str) -> None:
    if len(doc.paragraphs) > 0:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    add_paragraph(doc, f"CHAPTER {chapter_no}", bold=True, centered=True, first_line_indent=None)
    add_paragraph(doc, title.upper(), bold=True, centered=True, first_line_indent=None)
    add_blank_line(doc)


def add_section_heading(doc: Document, number: str, title: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 2")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Inches(0)
    run = paragraph.add_run(f"{number}. {title}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_table_caption(doc: Document, text: str) -> None:
    add_paragraph(doc, text, bold=True, centered=True, first_line_indent=None, single_spaced=True)


def add_figure_caption(doc: Document, text: str) -> None:
    add_paragraph(doc, text, italic=True, centered=True, first_line_indent=None, single_spaced=True)


def format_table_text(paragraph, text: str, *, bold: bool = False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]]) -> None:
    add_table_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    for style_name in ("Table Grid", "Normal Table", "Light Grid"):
        try:
            table.style = style_name
            break
        except KeyError:
            continue
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        header_cells[idx].text = ""
        format_table_text(header_cells[idx].paragraphs[0], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            format_table_text(cells[idx].paragraphs[0], value, bold=False)
    add_blank_line(doc)


def add_native_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]]) -> None:
    add_figure_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for style_name in ("Table Grid", "Normal Table", "Light Grid"):
        try:
            table.style = style_name
            break
        except KeyError:
            continue
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        format_table_text(cell.paragraphs[0], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[idx].text = ""
            format_table_text(cells[idx].paragraphs[0], value)
    add_blank_line(doc)


def add_figure(doc: Document, image_path: Path, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(5.8))
    add_figure_caption(doc, caption)
    add_blank_line(doc)


def add_native_figure(doc: Document, figure_key: str, caption: str) -> None:
    if figure_key == "Figure 4.1":
        add_native_table(
            doc,
            caption,
            ["Sequence", "Process / Branch", "Details"],
            [
                ["1", "Start", "Define evaluation scope: software effectiveness only."],
                ["2", "Evidence streams", "Alpha tests, beta survey, benchmark runs, and numeric checks."],
                ["3", "Compute metrics", "Weighted mean, Cronbach alpha, AE/PE, feasibility rate, runtime, deviation, repetition, and pantry-utilization U."],
                ["4A", "Quality branch", "ISO/IEC 25010 scoring: functional suitability, usability, reliability, and performance."],
                ["4B", "Planner branch", "Benchmark metrics: time, feasibility rate, compliance, pantry behavior, and repetition."],
                ["5", "Synthesis", "Synthesize effectiveness report and interpretation."],
                ["6", "Chapter 4 output", "Chapter 4 tables and conclusions; clinical outcomes are explicitly excluded."],
            ],
        )
    elif figure_key == "Figure 4.2":
        add_native_table(
            doc,
            caption,
            ["Effectiveness Dimension", "Metrics", "Output Artifact"],
            [
                ["Quality", "ISO/IEC 25010 ratings, weighted mean, frequency, percentage", "Quality results table and interpretation"],
                ["Instrument reliability", "Cronbach alpha coefficient", "Reliability table"],
                ["Numeric correctness", "Absolute error and percentage error against manual reference", "Error table per computed metric"],
                ["Planner performance", "Runtime, feasibility rate, compliance deviation, repetition, pantry-utilization U", "Benchmark table by scenario"],
                ["Resilience behavior", "Fallback trigger rate and fallback completion", "Fallback behavior evidence table"],
            ],
        )
    elif figure_key == "Figure 4.3":
        add_native_table(
            doc,
            caption,
            ["Evidence Set", "Feasibility Rate (%)", "p95 Runtime (s)", "Interpretation"],
            [
                ["sample_growth", "0.0", "1.96", "Stress set exposed infeasible/no-safe conditions."],
                ["sample_growth_fi", "93.3", "7.87", "Fallback and guidance branches improved feasibility."],
                ["smoke_after_fall", "100.0", "7.11", "Stabilized fallback-era smoke set reached full feasibility."],
                ["smoke_custom_sce", "66.7", "7.18", "Custom stress scenario confirms need for no-safe-plan handling."],
                ["Summary", "Current evidence", "Below configured runtime threshold", "Values feed planner-performance and resilience dimensions."],
            ],
        )
    elif figure_key == "Figure 4.4":
        add_native_table(
            doc,
            caption,
            ["Step", "Comparison Design Element", "Purpose"],
            [
                ["1", "Fix shared inputs: profile, constraints, dataset version, pantry snapshot", "Controls all non-method variables."],
                ["2", "Run existing baseline: rule-filter plus greedy slot fill", "Measures the simpler comparison method."],
                ["3", "Run proposed method: two-stage optimizer with fallback policy", "Measures PCOSina's deterministic planning method."],
                ["4", "Collect metric outputs: time, feasibility rate, budget delta, U, repetition, hard violations", "Creates aligned measurement fields."],
                ["5", "Compute scenario table: S1-S5 aligned rows with identical conditions", "Allows method effects to be interpreted fairly."],
                ["6", "Produce Chapter 4 evidence: figures, tables, acceptance checks, interpretation", "Supports software-behavior conclusions."],
            ],
        )
    elif figure_key == "Figure 4.5":
        add_native_table(
            doc,
            caption,
            ["Scenario", "Existing Time (s)", "Proposed Time (s)", "Existing Calorie Dev.", "Proposed Calorie Dev."],
            [
                ["S1", "0.29", "1.91", "532", "530"],
                ["S2", "0.22", "1.17", "540", "539"],
                ["S3", "0.14", "3.94", "605", "602"],
                ["S4", "0.29", "2.47", "566", "557"],
                ["S5", "6.93", "2.23", "n/a", "n/a"],
            ],
        )
    elif figure_key == "Figure 4.6":
        add_native_table(
            doc,
            caption,
            ["Status", "Acceptance Criterion", "Evidence"],
            [
                ["PASS", "FR_proposed >= FR_base in restrictive scenario S3", "S3 feasibility rate 100.00% -> 100.00%"],
                ["PASS", "Budget-stress delta parity in S2, difference <= 1.0 percentage point", "S2 budget delta 5.60% vs 6.20%"],
                ["PASS", "No hard-constraint violations in matched scenarios S1-S4", "hard_viol_total = 0 for both methods in S1-S4"],
                ["PASS", "Non-silent behavior under stress: message coverage and resilient completion in S5", "message coverage 100.00%/100.00%; S5 feasibility rate 100.00% -> 100.00%"],
            ],
        )
    else:
        raise ValueError(f"Unsupported native figure key: {figure_key}")


def add_figure_or_native(
    doc: Document,
    figure_key: str,
    image_path: Path,
    caption: str,
    *,
    use_native_diagrams: bool = False,
    use_image_placeholders: bool = False,
) -> None:
    if use_image_placeholders:
        add_paragraph(doc, f"[[{figure_key.upper().replace(' ', '_').replace('.', '')}_IMAGE]]", bold=True, centered=True, first_line_indent=None)
        add_figure_caption(doc, caption)
        add_blank_line(doc)
    elif use_native_diagrams:
        add_native_figure(doc, figure_key, caption)
    else:
        add_figure(doc, image_path, caption)


def google_docs_safe_image_path(image_path: Path) -> Path:
    GOOGLE_DOCS_ASSET_DIR.mkdir(parents=True, exist_ok=True)
    output_path = GOOGLE_DOCS_ASSET_DIR / f"{image_path.stem}_google_docs_safe.jpg"
    with Image.open(image_path) as source:
        source.load()
        flattened = Image.new("RGB", source.size, "white")
        if source.mode == "RGBA":
            flattened.paste(source, mask=source.getchannel("A"))
        else:
            flattened.paste(source.convert("RGB"))
        flattened.save(
            output_path,
            format="JPEG",
            quality=92,
            optimize=False,
            progressive=False,
            dpi=(150, 150),
        )
    return output_path


def convert_figures_for_google_docs() -> None:
    for path in FIGURE_PATHS.values():
        google_docs_safe_image_path(path)


def save_doc_if_possible(doc: Document, path: Path) -> bool:
    try:
        doc.save(str(path))
        return True
    except PermissionError:
        print(f"SKIPPED_LOCKED {path}")
        return False


def safe_format_float(value: float, places: int = 2) -> str:
    return f"{value:.{places}f}"


def build_document(
    *,
    use_google_docs_safe_images: bool = False,
    use_native_diagrams: bool = False,
    use_image_placeholders: bool = False,
) -> Document:
    dataset_manifest = load_json(DATASET_MANIFEST_PATH)
    training_metrics = load_json(TRAINING_METRICS_PATH)
    ml_ledger = load_json(ML_LEDGER_PATH)
    roadmap_counts = load_progress_status_counts(PROGRESS_LEDGER_PATH)
    ml_counts = Counter(item["status"] for item in ml_ledger["items"])
    test_counts = count_test_sources()
    android_version_name, android_version_code = load_android_version(APP_GRADLE_PATH)
    sentry_android_version = load_catalog_version(VERSION_CATALOG_PATH, "sentryAndroid")
    respondent_release_summary = load_first_nonblank_line(RESPONDENT_RELEASE_NOTES_PATH)
    respondent_build_label = respondent_release_summary.replace(
        "PCOSINA respondent testing build ",
        "",
    )

    matched_runtime_base_avg = mean(row[2] for row in SCENARIO_RUNTIME_ROWS[:4])
    matched_runtime_proposed_avg = mean(row[3] for row in SCENARIO_RUNTIME_ROWS[:4])
    matched_calorie_base_avg = mean(row[1] for row in SCENARIO_QUALITY_ROWS)
    matched_calorie_proposed_avg = mean(row[2] for row in SCENARIO_QUALITY_ROWS)
    matched_macro_base_avg = mean(row[3] for row in SCENARIO_QUALITY_ROWS)
    matched_macro_proposed_avg = mean(row[4] for row in SCENARIO_QUALITY_ROWS)

    before_avg = mean(LIVE_BEFORE_RUNS_MS)
    after_avg = mean(LIVE_AFTER_RUNS_MS)
    improvement_factor = before_avg / after_avg

    doc = Document()
    configure_page_setup(doc)
    configure_styles(doc)

    add_chapter_title(doc, 4, "Presentation, Analysis, and Interpretation of Data")
    add_paragraph(
        doc,
        "This chapter presents the implemented outputs of PCOSina and interprets the repository-backed "
        "evidence gathered from the live Android application, benchmark comparison packs, roadmap ledgers, "
        "operational reports, and machine-learning training artifacts. The discussion remains specific to "
        "PCOSina as an offline-first Filipino PCOS meal-planning system and does not reuse external study data."
    )

    add_section_heading(doc, "4.1", "Implemented System Output and Feature Coverage")
    add_paragraph(
        doc,
        "PCOSina is implemented as an offline-first mobile application whose core workflow moves from account "
        "entry to profile setup, goal selection, weekly meal-plan generation, pantry-aware grocery guidance, "
        "progress logging, and reminder or maintenance settings. The thesis flow record in "
        "docs/thesis/pcosina_thesis.md documents fourteen representative screens from the live application."
    )
    add_paragraph(
        doc,
        f"The current Android configuration identifies version {android_version_name} "
        f"({android_version_code}), while the respondent testing notes identify the staging build as "
        f"{respondent_build_label}. This updated release line exposes transparent BMI and kcal explainers, "
        "goal-specific insights, grocery categorization, progress logging, support reporting, and respondent "
        "testing safeguards. Therefore, the system described in this chapter is a functional workflow rather "
        "than a static prototype."
    )
    add_table(
        doc,
        "Table 4.1\nPCOSina Implemented Workflow and Functional Coverage",
        ["Module / Screen", "Role in the User Workflow"],
        [[name, role] for name, role in SYSTEM_FLOW_ROWS],
    )
    add_paragraph(
        doc,
        "Table 4.1 shows that PCOSina covers the full operational path expected from the study: user onboarding, "
        "constraint capture, plan generation, grocery planning, adherence monitoring, and reflection support. "
        "The workflow is consistent with the product contract stated in the repository: deterministic preference-aware "
        "filtering plus deterministic optimization, with ML used only as an assistive ranking layer. Pantry evidence "
        "is interpreted as pantry-aware prioritization and grocery support unless a later build adds and validates a "
        "true hard pantry-feasibility mode."
    )

    add_section_heading(doc, "4.2", "Evaluation Framework and Available Evidence")
    add_paragraph(
        doc,
        "Chapter 3 defined both system-centered and user-centered evaluation paths. For this repository-backed draft, "
        "the strongest verified evidence comes from automated tests, benchmark comparison documents, roadmap status records, "
        "operational measurements, and respondent deployment records. The respondent accounting framework and ISO/IEC 25010 "
        "scoring structure already exist in the project, but the final raw questionnaire matrix and computed weighted means "
        "were not included in the current repository export. For that reason, this chapter reports only verifiable PCOSina "
        "evidence and leaves the final respondent-weighted table for later insertion."
    )
    add_table(
        doc,
        "Table 4.2\nRepository-Backed Evidence Used in Chapter 4",
        ["Artifact Source", "Evidence Provided", "Use in This Chapter"],
        [
            ["docs/thesis/pcosina_thesis.md", "Implemented flow screens and release snapshot", "System output and feature coverage"],
            ["output/doc/PanelCommentOutput3_*.docx", "Controlled baseline-versus-proposed comparison tables", "Algorithm comparison and interpretation"],
            ["docs/roadmap/team_system_evolution_report_2026-04-14.md", "Operational latency regression and recovery narrative", "Live planner performance analysis"],
            ["docs/roadmap/progress_ledger.md and docs/roadmap/ml_progress_ledger.json", "Implementation completion status", "Readiness and remaining work summary"],
            ["docs/respondent-deployment.md and release-notes/respondent-test-notes.txt", "Firebase App Distribution staging workflow and respondent build notes", "Respondent testing readiness and release interpretation"],
            ["ml/offline_training/artifacts/*", "Dataset size and model metrics", "Stage-1 ML assist analysis"],
            [
                "backend/tests, app/src/test/java, app/src/androidTest/java",
                f"{sum(test_counts.values())} verified test source files across backend and Android layers",
                "Engineering validation coverage",
            ],
        ],
    )
    add_paragraph(
        doc,
        f"The current repository contains {test_counts['Backend Python tests']} backend Python test files, "
        f"{test_counts['Android JVM unit test files']} Android JVM unit test files, and "
        f"{test_counts['Android instrumentation test files']} Android instrumentation test files, for a total of "
        f"{sum(test_counts.values())} test source files. This does not replace user evaluation, but it does confirm that "
        "the thesis implementation has broad engineering coverage across API contracts, solver behavior, Android view models, "
        "and UI instrumentation."
    )
    figure_4_1 = google_docs_safe_image_path(FIGURE_PATHS["Figure 4.1"]) if use_google_docs_safe_images else FIGURE_PATHS["Figure 4.1"]
    figure_4_2 = google_docs_safe_image_path(FIGURE_PATHS["Figure 4.2"]) if use_google_docs_safe_images else FIGURE_PATHS["Figure 4.2"]
    figure_4_3 = google_docs_safe_image_path(FIGURE_PATHS["Figure 4.3"]) if use_google_docs_safe_images else FIGURE_PATHS["Figure 4.3"]
    add_figure_or_native(doc, "Figure 4.1", figure_4_1, "Figure 4.1. PCOSina effectiveness evaluation flowchart.", use_native_diagrams=use_native_diagrams, use_image_placeholders=use_image_placeholders)
    add_figure_or_native(doc, "Figure 4.2", figure_4_2, "Figure 4.2. PCOSina effectiveness dimensions and evidence map.", use_native_diagrams=use_native_diagrams, use_image_placeholders=use_image_placeholders)
    add_figure_or_native(doc, "Figure 4.3", figure_4_3, "Figure 4.3. Benchmark evidence snapshot used in the Chapter 4 analysis.", use_native_diagrams=use_native_diagrams, use_image_placeholders=use_image_placeholders)
    add_paragraph(
        doc,
        "Figures 4.1 to 4.3 summarize the evaluation path used in this draft: implementation evidence, benchmark evidence, "
        "operational evidence, and ML artifact evidence are interpreted together so that Chapter 4 remains grounded in actual "
        "PCOSina data even before the consolidated survey workbook is appended."
    )

    add_section_heading(doc, "4.3", "Controlled Comparison Between the Existing Baseline and the Proposed PCOSina System")
    add_paragraph(
        doc,
        "The controlled comparison pack prepared for panel revisions defines the existing system as a rule-filter plus greedy "
        "slot-by-slot planner, while the proposed system is PCOSina's two-stage optimizer using deterministic filtering, "
        "OR-Tools CP-SAT optimization, and a guarded fallback policy. Both methods were measured using matched scenarios, "
        "fixed inputs, and identical metrics."
    )
    figure_4_4 = google_docs_safe_image_path(FIGURE_PATHS["Figure 4.4"]) if use_google_docs_safe_images else FIGURE_PATHS["Figure 4.4"]
    add_figure_or_native(doc, "Figure 4.4", figure_4_4, "Figure 4.4. Existing-versus-proposed comparison design used for PCOSina.", use_native_diagrams=use_native_diagrams, use_image_placeholders=use_image_placeholders)
    add_table(
        doc,
        "Table 4.3\nScenario-Level Runtime, Feasibility, and Diversity Comparison",
        [
            "Scenario",
            "Label",
            "Time Base (s)",
            "Time Proposed (s)",
            "FR Base",
            "FR Proposed",
            "Delta Budget Base",
            "Delta Budget Proposed",
            "U Base",
            "U Proposed",
            "Fallback Proposed",
        ],
        [
            [
                row[0],
                row[1],
                safe_format_float(row[2], 4),
                safe_format_float(row[3], 4),
                row[4],
                row[5],
                row[6],
                row[7],
                row[8],
                row[9],
                row[12],
            ]
            for row in SCENARIO_RUNTIME_ROWS
        ],
    )
    add_table(
        doc,
        "Table 4.4\nScenario-Level Nutritional Quality and Safety Comparison",
        [
            "Scenario",
            "Calorie Dev Base",
            "Calorie Dev Proposed",
            "Macro Dev Base",
            "Macro Dev Proposed",
            "Hard Viol Base",
            "Hard Viol Proposed",
        ],
        [
            [
                row[0],
                safe_format_float(row[1], 2),
                safe_format_float(row[2], 2),
                safe_format_float(row[3], 2),
                safe_format_float(row[4], 2),
                str(row[5]),
                str(row[6]),
            ]
            for row in SCENARIO_QUALITY_ROWS
        ],
    )
    add_paragraph(
        doc,
        f"For the matched scenarios S1 to S4, the existing baseline averaged {safe_format_float(matched_runtime_base_avg, 4)} seconds "
        f"per plan while the proposed PCOSina method averaged {safe_format_float(matched_runtime_proposed_avg, 4)} seconds. "
        f"This shows that the deterministic optimizer is slower than the simpler greedy baseline under controlled conditions, "
        "which is expected because the proposed system solves a richer weekly planning problem with pantry awareness, variety, "
        "and more explicit constraints."
    )
    add_paragraph(
        doc,
        f"Despite the runtime tradeoff, the proposed system slightly improved the average calorie deviation "
        f"({safe_format_float(matched_calorie_base_avg, 2)} to {safe_format_float(matched_calorie_proposed_avg, 2)}) and the average "
        f"macro deviation ({safe_format_float(matched_macro_base_avg, 2)} to {safe_format_float(matched_macro_proposed_avg, 2)}) "
        "while preserving zero hard-constraint violations in every matched scenario. In the resilience proxy scenario S5, the proposed "
        "system was also faster than the baseline (2.2258 seconds versus 6.9306 seconds), indicating that the optimized architecture can "
        "perform competitively under stress conditions."
    )
    figure_4_5 = google_docs_safe_image_path(FIGURE_PATHS["Figure 4.5"]) if use_google_docs_safe_images else FIGURE_PATHS["Figure 4.5"]
    figure_4_6 = google_docs_safe_image_path(FIGURE_PATHS["Figure 4.6"]) if use_google_docs_safe_images else FIGURE_PATHS["Figure 4.6"]
    add_figure_or_native(doc, "Figure 4.5", figure_4_5, "Figure 4.5. Scenario runtime and quality comparison for the PCOSina benchmark pack.", use_native_diagrams=use_native_diagrams, use_image_placeholders=use_image_placeholders)
    add_figure_or_native(doc, "Figure 4.6", figure_4_6, "Figure 4.6. Acceptance criteria matrix for the controlled PCOSina comparison.", use_native_diagrams=use_native_diagrams, use_image_placeholders=use_image_placeholders)
    add_paragraph(
        doc,
        "Taken together, the controlled comparison supports the study's design choice: the proposed PCOSina system intentionally spends "
        "more computation to maintain feasibility, preserve hard constraints, and slightly improve nutritional fit. The evidence therefore "
        "supports the use of a deterministic two-stage optimizer as the defended system, rather than a purely greedy baseline."
    )

    add_section_heading(doc, "4.4", "Operational Planner Recovery and ML-Assist Results")
    add_paragraph(
        doc,
        "Beyond benchmark comparisons, the repository also contains live engineering evidence about operational performance. "
        "The April 14, 2026 system evolution report documents a real latency regression that emerged while the project was expanding "
        "toward ML-assisted shortlist ranking, richer telemetry, and more production-shaped planning behavior."
    )
    add_table(
        doc,
        "Table 4.5\nObserved Live Planner Latency Before and After the April 2026 Recovery Work",
        ["Metric", "Before Fix", "After Fix", "Interpretation"],
        [
            ["Observed slow runs", "207,816 ms and 188,249 ms", "6,486 ms to 13,674 ms", "Recent successful runs returned to practical waiting time"],
            ["Average observed generation time", f"{safe_format_float(before_avg, 2)} ms", f"{safe_format_float(after_avg, 2)} ms", "Average latency improved substantially"],
            ["Best successful generation", "Not below 188,249 ms in the cited log window", f"{min(LIVE_AFTER_RUNS_MS):,} ms", "Best recent run is under seven seconds"],
            ["Worst successful generation", "207,816 ms", f"{max(LIVE_AFTER_RUNS_MS):,} ms", "Upper bound dropped from minutes to seconds"],
            ["Improvement factor", "-", f"{safe_format_float(improvement_factor, 2)}x faster", "Operational bottleneck was materially reduced"],
        ],
    )
    add_paragraph(
        doc,
        f"Table 4.5 shows that the mean observed live planning time fell from {safe_format_float(before_avg / 1000, 2)} seconds "
        f"to {safe_format_float(after_avg / 1000, 2)} seconds, equivalent to roughly a {safe_format_float(improvement_factor, 2)}x speedup. "
        "This is important because it confirms that the project team did not only design the solver academically; it also diagnosed and corrected "
        "a real deployment-era performance issue."
    )
    add_table(
        doc,
        "Table 4.6\nStage-1 ML Assist Artifact Summary for PCOSina",
        ["Metric", "Observed Value", "Interpretation"],
        [
            ["Dataset version", dataset_manifest["datasetVersion"], "Current offline training dataset snapshot"],
            ["Rows / positive rate", f"{dataset_manifest['quality']['rows']:,} rows / {dataset_manifest['quality']['positive_rate']:.4f}", "Training set built from repository telemetry artifacts"],
            ["Unique requests / users", f"{dataset_manifest['quality']['unique_requests']} / {dataset_manifest['quality']['unique_users']}", "Coverage of request-level behavior in the dataset"],
            ["Train / val / test split", f"{dataset_manifest['splitRows']['train']:,} / {dataset_manifest['splitRows']['val']:,} / {dataset_manifest['splitRows']['test']:,}", "Leakage-aware split by request identifier"],
            ["Feature count", str(training_metrics["feature_count"]), "Structured stage-1 ranking feature space"],
            ["Test AUC", safe_format_float(training_metrics["test_binary_metrics"]["auc"], 4), "Binary discrimination quality"],
            ["Test log loss", safe_format_float(training_metrics["test_binary_metrics"]["logloss"], 4), "Probability calibration proxy"],
            ["Test NDCG@10", safe_format_float(training_metrics["test_ranking_metrics"]["ndcg@10"], 4), "Top-10 ranking quality"],
            ["Baseline NDCG@10", safe_format_float(training_metrics["test_baseline_ranking_metrics"]["ndcg@10"], 4), "Reference ranking without the trained model"],
            ["Test MAP@10", safe_format_float(training_metrics["test_ranking_metrics"]["map@10"], 4), "Precision-oriented ranking quality"],
            ["Baseline MAP@10", safe_format_float(training_metrics["test_baseline_ranking_metrics"]["map@10"], 4), "Reference precision before learned ranking"],
        ],
    )
    add_paragraph(
        doc,
        "The ML results should be interpreted carefully. The trained LightGBM ranker improves ranking metrics over the baseline "
        "heuristic, but the repository contract explicitly states that ML remains optional and assistive only. Hard constraints such as "
        "allergies, cost ceilings, nutritional bounds, explicit hard preferences, and offline-first behavior still belong to deterministic "
        "filtering and deterministic optimization. Pantry handling should be reported as pantry-aware prioritization and grocery support "
        "in the current implementation. In other words, the ML layer helps prioritize candidates, but it does not replace solver authority."
    )
    add_table(
        doc,
        "Table 4.7\nImplementation and Readiness Snapshot from the Current Repository",
        ["Area", "Current Status", "Interpretation"],
        [
            [
                "Main roadmap ledger",
                f"{sum(roadmap_counts.values())} workstreams: {roadmap_counts.get('Done', 0)} done, {roadmap_counts.get('In progress', 0)} in progress",
                "Core hardening work is active, with configuration externalization already complete",
            ],
            [
                "ML progress ledger",
                f"{len(ml_ledger['items'])} items: {ml_counts.get('done', 0)} done, {ml_counts.get('in_progress', 0)} in progress",
                "ML foundation is advanced, but live rollout evidence is still being accumulated",
            ],
            [
                "Automated test sources",
                f"{sum(test_counts.values())} verified files across backend and Android layers",
                "Engineering validation is broad even while some rollout evidence is still pending",
            ],
        ],
    )
    add_paragraph(
        doc,
        "Table 4.7 clarifies the maturity of PCOSina at the time of writing. The system is no longer a narrow thesis prototype, because it already "
        "contains planner diagnostics, async job handling, rollout guardrails, telemetry pipelines, and substantial automated coverage. At the same time, "
        "the roadmap also shows that some operational items such as staged sync validation, connected-device evidence, and wider live ML canary evidence "
        "are still in progress."
    )

    add_section_heading(doc, "4.5", "Respondent Testing Deployment and Compatibility Update")
    add_paragraph(
        doc,
        "The May 2026 respondent testing materials add a practical deployment layer to the Chapter 4 evidence. The project now documents a single "
        "Firebase App Distribution staging flow so respondents and teammates can install the same APK without Android Studio, ADB, emulator setup, "
        "wireless debugging, or manual Firebase App Check debug-token configuration. This improves the consistency of user testing because the test "
        "population receives a controlled build rather than locally compiled variants."
    )
    add_table(
        doc,
        "Table 4.8\nCurrent Respondent Testing and Compatibility Readiness Snapshot",
        ["Area", "Current Evidence", "Interpretation"],
        [
            [
                "Android application version",
                f"versionName {android_version_name}, versionCode {android_version_code}",
                "The thesis draft now reflects the current app build line rather than the earlier April release.",
            ],
            [
                "Respondent build label",
                respondent_build_label,
                "Staging APK is the intended distribution artifact for respondent testing.",
            ],
            [
                "Distribution path",
                "Firebase App Distribution staging upload",
                "Respondents do not need Android Studio, ADB, emulators, or App Check debug tokens.",
            ],
            [
                "Backend testing environment",
                "Render staging with Firebase Auth enabled and App Check enforcement disabled for respondent testing only",
                "The testing path keeps authentication while relaxing App Check only for controlled thesis/respondent evaluation.",
            ],
            [
                "Native-library compatibility",
                f"Sentry Android {sentry_android_version}; release notes cite updated native crash reporting libraries for 16 KB page-size compatibility",
                "Pixel 9a and Android 16 KB page-size testing should use a fresh staging APK built from this dependency line.",
            ],
            [
                "Deployment checklist files",
                f"{file_status(RESPONDENT_DEPLOYMENT_DOC_PATH)} respondent deployment guide; {file_status(PRE_RELEASE_CHECKLIST_PATH)} pre-release checklist",
                "The team has an explicit path for preparing, uploading, and monitoring respondent builds.",
            ],
        ],
    )
    add_paragraph(
        doc,
        "This deployment update does not change the scientific interpretation of the planner, but it improves the reliability of data collection. "
        "A consistent staging APK reduces environmental differences among respondents and makes issues such as sign-in failures, plan-generation "
        "failures, grocery-output errors, progress-logging issues, and native compatibility warnings easier to reproduce and triage."
    )

    add_section_heading(doc, "4.6", "Discussion and Interpretation")
    add_paragraph(
        doc,
        "The evidence in this chapter supports five major interpretations. First, PCOSina successfully implements the intended full workflow of the study, "
        "from account entry and profile capture to weekly planning, grocery assistance, and progress tracking. Second, the deterministic two-stage optimizer "
        "achieves its purpose: it preserves hard constraints while producing better nutritional fit than the simpler greedy comparator, even if it requires more "
        "runtime under matched scenarios. Third, the project has already moved beyond a purely academic build because operational bottlenecks were observed in live "
        "conditions, diagnosed, and reduced through engineering fixes."
    )
    add_paragraph(
        doc,
        "Fourth, the ML layer has measurable ranking value, but it remains appropriately limited by the repository's safety contract. This is important for a "
        "wellness decision-support application because ML should personalize candidate ordering without ever weakening allergy, budget, nutrient, pantry, or "
        "offline guarantees. Fifth, the current chapter is strongest in implementation and engineering evidence. The remaining gap is not the absence of a working "
        "system, but the absence of the final consolidated respondent questionnaire matrix in the repository export. That missing matrix prevents the honest reporting "
        "of weighted survey means and Cronbach alpha values in this draft."
    )

    add_chapter_title(doc, 5, "Summary of Findings")
    add_section_heading(doc, "5.1", "Summary of Findings")
    add_paragraph(
        doc,
        "Based on the evidence presented in Chapter 4, the following findings summarize the current outcome of the study:"
    )
    for idx, text in enumerate(
        [
            "PCOSina implements a complete fourteen-screen workflow that covers onboarding, planning, grocery support, progress logging, and settings management in a single offline-first Android application.",
            "The proposed system's deterministic two-stage planner maintains zero hard-constraint violations in the matched benchmark scenarios while slightly improving calorie and macro fit over the baseline comparator.",
            "The application already integrates pantry-aware grocery generation, goal-based personalization, transparent health metrics, and reflection capture, which means the study objectives were implemented as working features instead of conceptual mockups.",
            "Live operational evidence shows that the major April 2026 latency regression was reduced from roughly 198.03 seconds average observed generation time to roughly 9.29 seconds average across recent successful runs.",
            "The stage-1 ML assist pipeline shows stronger ranking metrics than the heuristic baseline, but it remains correctly bounded by deterministic filtering and optimization, preserving the study's solver-first contract.",
            f"The respondent testing release path has been updated to {respondent_build_label}, distributed through Firebase App Distribution with documented tester instructions, staging backend configuration, and 16 KB page-size native-library compatibility notes.",
        ],
        start=1,
    ):
        add_paragraph(doc, f"{idx}. {text}", first_line_indent=0.25)

    add_section_heading(doc, "5.2", "Synthesis Relative to the Objectives of the Study")
    add_table(
        doc,
        "Table 5.1\nSynthesis of Findings Relative to the Specific Objectives",
        ["Specific Objective", "Repository-Backed Finding", "Current Thesis Interpretation"],
        [
            [
                "1. Two-stage meal-planning algorithm",
                "Implemented and defended through deterministic filtering plus OR-Tools CP-SAT optimization and controlled comparison tables.",
                "Achieved in the current system build.",
            ],
            [
                "2. Personalization mechanism",
                "Profile setup, restrictions, goals, and symptom-aware planning are integrated in the Android workflow.",
                "Achieved with live product evidence.",
            ],
            [
                "3. Digital pantry inventory and grocery recommendations",
                "Pantry-aware grocery generation and grouped shopping flows are present in the live app; current evidence should not be worded as hard pantry feasibility.",
                "Achieved as pantry-aware support, with hard pantry-feasibility enforcement reserved for future validation if implemented.",
            ],
            [
                "4. Offline-first mobile architecture",
                "The repository contract and Android flows emphasize local-first behavior, queueing, and non-blocking app use.",
                "Achieved, while cross-device sync evidence remains an operational follow-up.",
            ],
            [
                "5. Performance and evaluation",
                "Benchmark, test, and operational evidence are available; final respondent-weighted survey matrix is not yet attached in this repository export.",
                "Substantially achieved, with final user-weighted evaluation tables pending archival.",
            ],
        ],
    )
    add_paragraph(
        doc,
        f"The general objective of the study was: {GENERAL_OBJECTIVE} The Chapter 4 evidence supports this objective because the application is already implemented, "
        "the planner contract is measurable, the pantry and personalization mechanisms are integrated, and the system has been examined through benchmark, operational, "
        "and artifact-based evaluation. The main missing addition for the final manuscript is the insertion of the consolidated respondent-weighted survey workbook."
    )

    add_chapter_title(doc, 6, "Conclusions and Recommendations")
    add_section_heading(doc, "6.1", "Conclusions")
    add_paragraph(
        doc,
        "This study concludes that PCOSina achieved the core design and development goals of the project. It successfully produced an offline-first mobile meal-planning "
        "application tailored for Filipinos with PCOS, supported by a deterministic two-stage planning architecture, pantry-aware grocery support, goal-based personalization, "
        "and adherence-focused tracking features."
    )
    add_paragraph(
        doc,
        "The comparison evidence indicates that the proposed PCOSina system is a justifiable improvement over the simpler greedy baseline because it preserves hard constraints "
        "and offers slightly better nutritional fit, even when it requires more computational work under matched scenarios. Operational evidence also shows that the team handled "
        "a real runtime regression responsibly by diagnosing and reducing latency, which strengthens confidence in the engineering maturity of the study output."
    )
    add_paragraph(
        doc,
        "The study also concludes that ML can be integrated responsibly into a wellness planning system when it is limited to assistive ranking and remains subordinate to deterministic "
        "constraint handling. PCOSina's repository contract preserves that boundary. Therefore, the current system can be defended as a decision-support application rather than a diagnosis "
        "engine or medical device."
    )

    add_section_heading(doc, "6.2", "Recommendations")
    for idx, text in enumerate(
        [
            "Consolidate and archive the final respondent questionnaire matrix so that weighted means, reliability statistics, and the final ISO/IEC 25010 summary table can be inserted without estimation.",
            "Capture the complete thesis screenshot set referenced in docs/thesis/pcosina_thesis.md so that future manuscript revisions can include the full live application walkthrough.",
            "Continue staged multi-device sync validation and connected-device Android testing to complement the current repository-backed engineering evidence.",
            "Use fresh Firebase App Distribution staging builds when testing on Pixel 9a or other 16 KB page-size devices, and verify that APK Analyzer or zipalign checks do not report stale unaligned native libraries.",
            "Maintain the guarded ML rollout policy and collect additional live canary evidence before expanding ML-assisted traffic beyond the current controlled readiness posture.",
            "Expand culturally relevant Filipino recipe coverage and continue collecting user feedback so that future versions can improve personalization while preserving deterministic safety constraints.",
        ],
        start=1,
    ):
        add_paragraph(doc, f"{idx}. {text}", first_line_indent=0.25)

    return doc


if __name__ == "__main__":
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    standard_doc = build_document()
    save_doc_if_possible(standard_doc, OUTPUT_PATH)
    save_doc_if_possible(standard_doc, DATED_OUTPUT_PATH)
    save_doc_if_possible(standard_doc, GOOGLE_DOCS_OUTPUT_PATH)

    convert_figures_for_google_docs()
    google_docs_doc = build_document(use_google_docs_safe_images=True)
    save_doc_if_possible(google_docs_doc, GOOGLE_DOCS_JPEG_OUTPUT_PATH)

    native_diagrams_doc = build_document(use_native_diagrams=True)
    save_doc_if_possible(native_diagrams_doc, GOOGLE_DOCS_NATIVE_OUTPUT_PATH)

    placeholder_doc = build_document(use_image_placeholders=True)
    save_doc_if_possible(placeholder_doc, GOOGLE_DOCS_PLACEHOLDER_OUTPUT_PATH)

    print(OUTPUT_PATH)
    print(DATED_OUTPUT_PATH)
    print(GOOGLE_DOCS_OUTPUT_PATH)
    print(GOOGLE_DOCS_JPEG_OUTPUT_PATH)
    print(GOOGLE_DOCS_NATIVE_OUTPUT_PATH)
    print(GOOGLE_DOCS_PLACEHOLDER_OUTPUT_PATH)
