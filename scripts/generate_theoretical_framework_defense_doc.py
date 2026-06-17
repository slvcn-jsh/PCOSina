from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "defense" / "PCOSINA_Theoretical_Framework_Defense_Prep_2026-06-04.docx"
SLIDE_IMAGE = ROOT / "tmp" / "defense_prep" / "presentation_pages_png" / "page_1.png"


ACCENT = "2F7D5A"
DARK = "1F2933"
MUTED = "5B6770"
PALE = "EAF5EF"
PALE_BLUE = "EAF2F8"
PALE_PINK = "FDECEF"
PALE_YELLOW = "FFF7DA"
BORDER = "B7C6BC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=DARK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05


def add_table(doc, headers, rows, widths=None, header_fill=ACCENT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8.8)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = widths[idx]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_text(cell, str(value), color=DARK, size=8.5)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[idx]
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Aptos Display"
        run.font.color.rgb = RGBColor.from_string(ACCENT if level == 1 else DARK)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(DARK)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.color.rgb = RGBColor.from_string(DARK)
    else:
        r = p.add_run(text)
        r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(item)
        run.font.color.rgb = RGBColor.from_string(DARK)


def add_numbered(doc, items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(f"{idx}.    {item}")
        run.font.color.rgb = RGBColor.from_string(DARK)


def add_callout(doc, title, body, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=ACCENT)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    cell.text = ""
    p = cell.paragraphs[0]
    rt = p.add_run(title)
    rt.bold = True
    rt.font.color.rgb = RGBColor.from_string(ACCENT)
    rt.font.size = Pt(10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    rb = p2.add_run(body)
    rb.font.color.rgb = RGBColor.from_string(DARK)
    rb.font.size = Pt(9)
    doc.add_paragraph()


def configure_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)

    for style_name, size, color in [
        ("Title", 24, ACCENT),
        ("Heading 1", 15, ACCENT),
        ("Heading 2", 12, DARK),
        ("Heading 3", 10.5, DARK),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("PCOSina Theoretical Framework Defense Prep")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(ACCENT)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Focused packet for thesis defense: script, slide guide, manuscript anchors, evidence, references, and Q&A")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "Defense identity",
        "PCOSina is an offline-first Filipino-PCOS wellness decision-support application. It uses deterministic preference-aware filtering and a MILP-formulated 0-1 optimization model solved with OR-Tools CP-SAT. ML is optional and assistive only.",
        fill=PALE,
    )

    add_table(
        doc,
        ["Item", "Defense-ready answer"],
        [
            ["Main point", "The theoretical framework explains why the system uses optimization, scheduling, a MILP-formulated assignment model, modular offline-first software design, rule-based guardrails, and assistive ML."],
            ["Presentation slide", "Slide 1: Theoretical Framework from [Presentation] Quadrant_Presentation_Revised_S2526T3 (1).pdf."],
            ["Manuscript anchor", "Chapter 1, Theoretical Framework section: paragraphs 197-239 in the extraction pass."],
            ["System anchor", "ARCHITECTURE.md, production_contract.md, planner_contract_2026-05-26.md, ADR-004, ADR-005, and backend/services/meal_planner.py."],
            ["MILP title defense", "MILP-based refers to the mathematical formulation: binary recipe-slot variables, linear constraints, and objective terms. CP-SAT is the implementation solver, not a pure MILP solver claim."],
            ["Boundary", "The app is not a diagnosis engine. The app does not treat PCOS. It supports meal planning decisions within declared constraints."],
        ],
        [Inches(1.55), Inches(5.75)],
    )

    add_heading(doc, "1. Slide To Present", 1)
    if SLIDE_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SLIDE_IMAGE), width=Inches(7.1))
        cap = doc.add_paragraph("Figure TF-1. Presentation slide used for the Theoretical Framework defense section.")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(8)
    else:
        add_callout(doc, "Slide image not embedded", f"Expected slide image was not found at {SLIDE_IMAGE}.", fill=PALE_PINK)

    add_body(
        doc,
        "Manuscript figure note: the Theoretical Framework section in Chapter 1 is text-based and has no dedicated figure under that heading. The slide above is the direct presentation visual for this topic. Conceptual Framework, System Architecture, and Algorithm figures should be handled in their own packets so topics do not get mixed."
    )

    add_heading(doc, "2. Core Defense Explanation", 1)
    add_callout(
        doc,
        "30-second answer",
        "Our theoretical framework is grounded in four computer-science and software-engineering foundations: optimization and scheduling theory, a MILP-formulated 0-1 assignment model solved with OR-Tools CP-SAT, offline-first modular design, and rule-based decision support with assistive ML. These theories explain why PCOSina filters unsafe recipes first, optimizes the complete 7-day meal schedule second, keeps saved data available locally, and limits ML to ranking support only. The title is MILP-based because the model formulation uses binary variables, linear constraints, and objective terms; CP-SAT is the practical implementation solver.",
        fill=PALE_BLUE,
    )

    add_heading(doc, "3. Four Foundations", 1)
    add_table(
        doc,
        ["Foundation", "Meaning in the manuscript", "How PCOSina implements it", "Defense sentence"],
        [
            [
                "Optimization and Scheduling Theory",
                "Meal planning is treated as selecting the best feasible meal combination while assigning meals to fixed slots.",
                "The system plans 7 days with breakfast, lunch, and dinner slots, while balancing nutrition, budget, pantry overlap, preference, and variety.",
                "PCOSina treats meal planning as a constrained scheduling problem, not a static recipe list.",
            ],
            [
                "MILP-Formulated 0-1 Model with CP-SAT Solver",
                "Each recipe-slot assignment can be represented as a binary decision: selected or not selected.",
                "The backend builds binary variables for slot assignments and solves them through OR-Tools CP-SAT as the final authority.",
                "This is why the title can say MILP-based: the model is MILP-formulated, while CP-SAT is the solver used in implementation.",
            ],
            [
                "Offline-first and Modular Design",
                "Essential user data and generated outputs should remain available without continuous connectivity.",
                "Android uses local persistence through DataStore, encrypted shared preferences, and local artifacts; backend optimization is separated from UI and local state.",
                "Offline-first means saved data continuity; new optimized generation still requires the backend.",
            ],
            [
                "Rule-based Systems and Assistive ML",
                "Explicit rules are used for explainable decisions; ML can assist ranking but cannot replace safety rules.",
                "Stage 1 removes recipes that violate allergies, restrictions, meal-slot rules, and hard constraints; LightGBM ranks only already-safe candidates.",
                "ML improves ordering, but deterministic filtering and CP-SAT decide validity.",
            ],
        ],
        [Inches(1.35), Inches(2.0), Inches(2.2), Inches(1.75)],
    )

    add_heading(doc, "4. Slide Narration Script", 1)
    add_body(doc, "Use this for the actual presentation. Do not read every bullet. Explain the logic behind the four boxes.")
    add_numbered(
        doc,
        [
            "Start with the purpose: This slide explains the theories behind the technical design of PCOSina, not the medical diagnosis of PCOS.",
            "For optimization and scheduling theory: say that a weekly meal plan is a scheduling problem because meals must be assigned across 7 days and 21 meal slots while satisfying nutrition and practical constraints.",
            "For the MILP-formulated CP-SAT model: explain that the system represents choices as yes-or-no decisions, such as whether a recipe is assigned to a specific slot. The formulation uses linear constraints and objective terms; CP-SAT then searches for a complete valid plan.",
            "For rule-based systems and assistive ML: emphasize that hard rules are checked first. Allergies, exclusions, budget ceilings, and nutrition boundaries cannot be overridden by ML.",
            "For offline-first and modular design: clarify that saved plans, grocery data, pantry records, and progress logs remain available locally, while new optimized plan generation uses the backend service.",
            "Close by connecting the four foundations: The system is safe because hard rules are deterministic, useful because optimization balances competing goals, practical because saved data remains local, and scalable because modules are separated.",
        ],
    )

    add_heading(doc, "5. Full Speaker Script", 1)
    add_body(
        doc,
        "Good day panel. This section explains the theoretical foundation of PCOSina. Since PCOSina is a computer science and wellness decision-support project, our theoretical framework focuses on how the system transforms user needs into safe and practical meal-planning decisions."
    )
    add_body(
        doc,
        "First, the study uses optimization and scheduling theory. A weekly meal plan is not just a list of recipes. It is a schedule where meals must be assigned across seven days and three major meal slots per day. Each assignment must respect nutrition targets, allergies, dietary restrictions, budget behavior, pantry overlap, and variety requirements."
    )
    add_body(
        doc,
        "Second, PCOSina uses a MILP-formulated 0-1 assignment model solved with OR-Tools CP-SAT. This is the technical center of the MILP-based title. Each candidate recipe can be represented as a binary decision variable for a meal slot: selected or not selected. The formulation uses linear constraints and objective terms for feasibility and optimization. In the actual implementation, CP-SAT is the solver that handles the integer and constraint-based search."
    )
    add_body(
        doc,
        "Third, the system follows offline-first and modular software design. Saved user profiles, generated plans, grocery guidance, pantry records, and progress-related data are stored locally so users can still review important information without continuous internet access. At the same time, backend optimization, local persistence, UI, data access, and ML ranking remain separated so each part can be tested and maintained independently."
    )
    add_body(
        doc,
        "Fourth, the system uses rule-based decision support with assistive machine learning. Rule-based filtering handles non-negotiable safety rules before optimization begins. LightGBM can help rank already-safe candidate recipes, but it cannot reintroduce a rejected recipe and cannot override allergies, dietary restrictions, nutrition bounds, budget caps, or the final CP-SAT assignment."
    )
    add_body(
        doc,
        "In short, the theoretical framework justifies our design: deterministic rules protect safety, the MILP-formulated CP-SAT solve handles the weekly optimization problem, offline-first design protects access to saved data, and ML remains supportive rather than authoritative."
    )

    add_heading(doc, "6. Repo-Truth Cross-Check", 1)
    add_table(
        doc,
        ["Claim", "Current source evidence", "Defense-safe wording"],
        [
            [
                "Offline-first system",
                "ARCHITECTURE.md states local app storage is the primary source of truth for user-facing continuity; Android uses DataStore, encrypted shared preferences, and local artifacts.",
                "Saved data remains accessible offline. New optimized generation requires backend connectivity.",
            ],
            [
                "Two-stage planner",
                "production_contract.md and backend/services/meal_planner.py define Stage 1 filtering/ranking/candidate generation and Stage 2 as a formal 0-1 model solved by OR-Tools CP-SAT.",
                "Stage 1 prepares safe candidates; Stage 2 solves the MILP-formulated final weekly plan.",
            ],
            [
                "ML is assistive",
                "ADR-004 states ML scores are allowed only in Stage 1 candidate ranking and CP-SAT remains final authority.",
                "LightGBM can prioritize safe candidates but cannot make unsafe recipes valid.",
            ],
            [
                "No unsafe fallback",
                "ADR-005 states infeasible solves return status=no-safe-plan with diagnostics and guidance; no unsafe approximation is returned as successful.",
                "If constraints are impossible, PCOSina explains the blocker instead of forcing a plan.",
            ],
            [
                "Pantry boundary",
                "planner_contract_2026-05-26 states backend planning remains pantry-aware, not fully hard pantry-feasible.",
                "Pantry data rewards overlap and supports grocery guidance; full pantry-consumption modeling is future work.",
            ],
        ],
        [Inches(1.4), Inches(3.05), Inches(2.85)],
    )

    doc.add_page_break()
    add_heading(doc, "7. Numbers To Memorize", 1)
    add_table(
        doc,
        ["Evidence point", "Value", "How to say it"],
        [
            ["Final benchmark runs", "100/100 passed", "Across the final local benchmark, all tested runs generated valid complete plans for valid profiles."],
            ["MILP title defense", "MILP-formulated 0-1 model; CP-SAT solver implementation", "The title is defended by the model formulation, not by calling CP-SAT a pure MILP solver."],
            ["Hard-rule violations", "0", "No tested final run violated hard rules such as allergy/restriction/budget/nutrition boundaries."],
            ["Final runtime", "Average 257 ms, P95 313 ms, max 357 ms", "The final optimized local benchmark was sub-second."],
            ["ML ranking metric", "NDCG@10 0.9211 vs 0.8838 baseline", "ML improved candidate ordering, but CP-SAT remained the final authority."],
            ["No-safe-plan behavior", "Intentional impossible budget returned no-safe-plan", "The system fails safely when a safe complete plan cannot be generated."],
            ["Verification", "Full local verification passed on 2026-05-26", "Android unit tests and backend pytest passed in the latest recorded full local verification."],
        ],
        [Inches(1.55), Inches(1.4), Inches(4.35)],
    )

    add_heading(doc, "8. Say This, Not That", 1)
    add_table(
        doc,
        ["Avoid saying", "Say instead", "Why"],
        [
            ["The app diagnoses PCOS.", "The app is a wellness decision-support tool for meal planning.", "Diagnosis and treatment claims are outside the system boundary."],
            ["ML chooses the final meal plan.", "ML only ranks already-safe candidates in Stage 1.", "CP-SAT is the final authority."],
            ["CP-SAT means it is not MILP-based.", "MILP-based refers to the mathematical formulation; CP-SAT is the practical solver implementation.", "This is the clean answer when panelists ask about the title."],
            ["The planner is a decision tree.", "It is a decision flow with deterministic filtering plus optimization.", "The implementation is not a tree model."],
            ["Everything works offline.", "Saved data works offline; new optimized generation requires backend connectivity.", "This matches the actual local-first architecture."],
            ["The pantry is always a hard feasibility constraint.", "The planner is pantry-aware and grocery-gap aware.", "Current backend uses pantry as scoring/reward/guidance, not complete hard inventory consumption."],
            ["Fallback returns a greedy meal plan.", "If no safe plan exists, PCOSina returns no-safe-plan guidance.", "The accepted safety contract rejects unsafe approximate output."],
        ],
        [Inches(2.0), Inches(3.0), Inches(2.3)],
        header_fill="B94B61",
    )

    add_heading(doc, "9. Panel Q&A", 1)
    qas = [
        ("Why is your theoretical framework mostly computer-science based and not purely medical?",
         "Because PCOSina is not a clinical diagnosis system. The medical and nutrition literature defines the wellness boundaries, but the system contribution is how software engineering and optimization turn those boundaries into a safe meal-planning workflow."),
        ("Why use optimization theory?",
         "Because the system must balance many constraints at once: nutrition, allergies, budget, pantry overlap, meal-slot validity, repetition, and variety. Optimization gives a structured way to select the best feasible plan instead of manually choosing recipes."),
        ("Why does your title say MILP-based if CP-SAT is the solver?",
         "The title is defensible because MILP-based refers to the formulation: binary recipe-slot decisions, linear constraints, and an objective function. CP-SAT is not a pure MILP solver; it is the practical implementation solver used to solve the integer/constraint model."),
        ("What is a binary decision variable in your system?",
         "It is a yes-or-no variable. For example, x[slot, recipe] equals 1 if that recipe is assigned to that meal slot, and 0 if it is not."),
        ("What is the difference between Stage 1 and Stage 2?",
         "Stage 1 decides which recipes are safe and strong enough to enter the pool. Stage 2 decides the final weekly assignment across all meal slots."),
        ("Why is ML not the final decision-maker?",
         "Because health-adjacent meal planning needs deterministic safety boundaries. ML may improve ranking, but hard constraints must remain auditable and non-negotiable."),
        ("What happens if the model cannot find a safe plan?",
         "The system returns a structured no-safe-plan response with reason codes and suggested non-safety adjustments. It does not force a partial or unsafe plan."),
        ("What does offline-first mean here?",
         "It means saved profiles, plans, groceries, pantry records, and progress logs remain accessible locally. It does not mean the phone performs every new backend optimization while offline."),
        ("How does modular design help the study?",
         "It separates UI, local persistence, backend optimization, recipe data, ML ranking, and evaluation. This improves maintainability, testing, traceability, and defense clarity."),
        ("What is the limitation of the theoretical framework?",
         "It supports safe software decision-making, not clinical outcome proof. Long-term adherence, symptom improvement, and full pantry-consumption modeling are future research directions."),
    ]
    for q, a in qas:
        add_body(doc, f"Q: {q}", bold_prefix="Q:")
        add_body(doc, f"A: {a}", bold_prefix="A:")

    add_heading(doc, "10. Quick Whiteboard Version", 1)
    add_body(doc, "If asked to explain without slides, draw this sequence:")
    add_numbered(
        doc,
        [
            "User inputs: profile, goal, restrictions, allergies, pantry, budget.",
            "Rule-based gate: remove unsafe or incompatible recipes.",
            "Optional ML ranking: reorder only safe candidates.",
            "MILP-formulated CP-SAT optimizer: assign recipes into 21 weekly meal slots.",
            "Output: safe complete weekly plan, grocery guidance, explanation, or no-safe-plan.",
        ],
    )

    add_heading(doc, "11. Source And Reference Anchors", 1)
    add_table(
        doc,
        ["Anchor", "Use"],
        [
            ["Manuscript paragraphs 197-239", "Main theoretical framework section."],
            ["Presentation PDF page 1", "Defense slide for the theoretical framework."],
            ["ARCHITECTURE.md", "Current system shape: Android local persistence, backend authority, ML guardrails, offline-first boundary."],
            ["docs/architecture/production_contract.md", "Two-stage MILP-formulated planner and no unsafe output contract."],
            ["docs/architecture/planner_contract_2026-05-26.md", "Hard/soft/advisory/tracking classifications, MILP-formulated planner boundary, and pantry-aware boundary."],
            ["docs/adr/ADR-004-ml-shadow-canary.md", "ML remains Stage 1 ranking support only."],
            ["docs/adr/ADR-005-no-safe-plan-contract.md", "No-safe-plan behavior when constraints are infeasible."],
            ["backend/services/meal_planner.py", "Implemented Stage 1 filtering, MILP-formulated binary variables, constraints, objective terms, and CP-SAT solve."],
            ["docs/roadmap/progress_ledger.md", "Final benchmark and verification evidence."],
        ],
        [Inches(2.5), Inches(4.8)],
    )

    add_heading(doc, "12. References To Keep Ready", 1)
    add_bullets(
        doc,
        [
            "Benvenuti et al. (2024): integer programming for healthy and sustainable diet planning; supports MILP-formulated optimization-based meal planning.",
            "De Leon et al. (2022): Philippine MILP dietary planning example; supports local relevance of mathematical diet planning.",
            "Guevarra et al. (2022): LP/ILP meal-cost optimization; supports whole-food/integer planning rather than impractical fractional solutions.",
            "Paluyo et al. (2023): Philippine digital health model; supports local continuity and accessibility concerns.",
            "Agrawal et al. (2025): AI in personalized nutrition; useful for explaining why ML must remain interpretable and bounded.",
            "Teede et al. and PCOS nutrition guidance in the manuscript: supports the wellness/nutrition boundary, not diagnostic claims.",
        ],
    )

    add_heading(doc, "13. Study Checklist", 1)
    add_bullets(
        doc,
        [
            "Memorize the 30-second answer.",
            "Explain the four foundations without reading the slide word-for-word.",
            "Practice the difference between MILP-formulated model and CP-SAT solver.",
            "Practice the Stage 1 vs Stage 2 explanation.",
            "Memorize the no-safe-plan answer.",
            "Memorize the offline-first boundary.",
            "Memorize the key evidence numbers: 100/100, 0 violations, 257 ms, P95 313 ms, max 357 ms, NDCG@10 0.9211.",
            "Be ready to say what the framework does not prove: diagnosis, treatment, long-term clinical outcomes, full pantry consumption, or fully offline new optimization.",
        ],
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendix A. Manuscript Anchor Summary", 1)
    add_table(
        doc,
        ["Paragraphs", "Key content"],
        [
            ["197-200", "The framework is grounded in computer science, operations research, and software engineering."],
            ["202-207", "Optimization and scheduling theory model meal planning as constrained allocation over 7 days and meal slots."],
            ["209-221", "MILP-formulated binary assignment with CP-SAT, two-stage candidate construction, assistive LightGBM, and final weekly assignment."],
            ["223-230", "Offline-first, modular design, privacy-conscious handling, and user-centered design."],
            ["232-239", "Rule-based systems and assistive ML preserve explainability and safety boundaries."],
        ],
        [Inches(1.4), Inches(5.9)],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
