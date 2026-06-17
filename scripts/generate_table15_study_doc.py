from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_NAME = "PCOSINA_Table15_Sets_Parameters_Decision_Variables_Study_2026-06-10.docx"
REPO_OUTPUT = ROOT / "docs" / "defense" / OUTPUT_NAME
DOWNLOADS_OUTPUT = Path.home() / "Downloads" / OUTPUT_NAME


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def build_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PCOSINA Table 15 Defense Study Notes")
    run.bold = True
    run.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Sets, Parameters, and Decision Variables").italic = True

    doc.add_paragraph(
        "Purpose: This file explains Table 15 in plain language for final defense study. "
        "It also clarifies the grocery gap wording so the explanation matches the current system."
    )

    add_heading(doc, "Exact Table 15 Content", 1)
    add_table(
        doc,
        ["Symbol", "Meaning", "Type"],
        [
            ["R", "Candidate recipes after Stage 1 filtering", "Set"],
            ["S", "Weekly meal slots (D x M = 7 x 3 = 21)", "Set"],
            ["K", "Nutrition dimensions, such as calories, protein, carbohydrates, and fat", "Set"],
            ["I", "Ingredient set used for pantry and grocery guidance", "Set"],
            ["x_(r,s)", "1 if recipe r assigned to slot s, else 0", "Binary decision variable"],
            ["y_r", "1 if recipe r is used at least once during the week, otherwise 0", "Binary helper variable"],
            ["u_k, v_k", "Under-target and over-target nutrition deviation for nutrition dimension k", "Non-negative deviation variables"],
            ["g_i", "Estimated grocery gap for ingredient i, when ingredient quantity estimates are available", "Non-negative guidance variable"],
        ],
    )

    add_heading(doc, "Plain-Language Explanation", 1)
    add_table(
        doc,
        ["Symbol", "Layman Meaning", "Defense Point"],
        [
            [
                "R",
                "The safe recipe pool after Stage 1 filtering. It is not the whole recipe database.",
                "Stage 1 removes unsafe or incompatible recipes before ranking and solving.",
            ],
            [
                "S",
                "The meal positions the system must fill, usually 7 days x 3 meals = 21 slots.",
                "The solver assigns one allowed recipe to each required slot.",
            ],
            [
                "K",
                "The nutrition categories checked by the system, such as calories, protein, carbs, fats, fiber, sodium, and sugar.",
                "These are used to keep the plan close to nutrition targets and policy limits.",
            ],
            [
                "I",
                "The ingredient names used for pantry matching and grocery guidance.",
                "Ingredients connect the meal plan to pantry overlap and grocery output.",
            ],
            [
                "x_(r,s)",
                "A yes/no decision: should recipe r be placed in slot s?",
                "This is the main optimization decision variable.",
            ],
            [
                "y_r",
                "A helper idea: did recipe r appear at least once in the weekly plan?",
                "The current code mostly tracks this through sums of x variables for repeat and usage rules.",
            ],
            [
                "u_k, v_k",
                "How far the plan is below or above a nutrition target.",
                "Some nutrition dimensions are enforced by bounds, while sodium/sugar overages can be tracked as penalties.",
            ],
            [
                "g_i",
                "The estimated missing/needed grocery amount for an ingredient when quantity data are available.",
                "This supports grocery guidance. It is not the main CP-SAT decision variable.",
            ],
        ],
    )

    add_heading(doc, "Grocery Gap Clarification", 1)
    doc.add_paragraph(
        "Yes, in layman terms, a grocery gap is related to the ingredients that appear in the grocery list. "
        "If the selected meal plan needs eggs and the user does not already have eggs in the pantry, eggs are part of the practical grocery need."
    )
    doc.add_paragraph(
        "More precise defense wording: g_i represents an estimated grocery gap or needed ingredient quantity for ingredient i, when quantity estimates are available. "
        "In the current implementation, the backend builds the displayed grocery output after the final plan is selected by aggregating recipe ingredients and pricing them. "
        "The Android side can also rebuild grocery items from saved meal ingredient sources. Therefore, grocery gap is best explained as grocery guidance rather than the main optimizer decision variable."
    )

    add_heading(doc, "Defense Script", 1)
    doc.add_paragraph(
        "Table 15 defines the symbols used in our optimization model. R is the safe candidate recipe set after Stage 1 filtering, "
        "S is the set of weekly meal slots, K is the nutrition dimensions, and I is the ingredient set for pantry and grocery guidance. "
        "The main decision variable is x(r,s), which tells whether a recipe is assigned to a specific meal slot. "
        "Helper variables or equivalent expressions track recipe usage, nutrition deviation, and grocery guidance."
    )
    doc.add_paragraph(
        "For grocery gap, we explain it as the estimated ingredient need after the plan is selected. "
        "It supports the grocery list and budget guidance, but it is not the same as the main CP-SAT assignment variable."
    )

    add_heading(doc, "Common Panel Questions", 1)
    add_table(
        doc,
        ["Question", "Best Answer"],
        [
            [
                "Is R all recipes?",
                "No. R is the candidate set after Stage 1 filtering. Unsafe or incompatible recipes are removed before ranking and solving.",
            ],
            [
                "What does x(r,s) prove?",
                "It proves that the system uses binary decision variables for meal assignment, not random selection.",
            ],
            [
                "Is y_r literally present in code?",
                "The thesis uses y_r as a helper concept. In the current code, recipe usage is mainly represented by sums of x variables, especially for repeat rules.",
            ],
            [
                "Is g_i the grocery list?",
                "It is closely related. g_i means estimated grocery need/gap for an ingredient. The actual displayed grocery list is built after selecting the final plan by aggregating and pricing recipe ingredients.",
            ],
        ],
    )

    add_heading(doc, "Code Evidence Map", 1)
    add_table(
        doc,
        ["Evidence", "File and Line", "What It Proves"],
        [
            [
                "Stage 1 shortlist candidates",
                "backend/services/meal_planner.py:1654",
                "Defines the safe candidate shortlisting process that creates R.",
            ],
            [
                "Binary x variables",
                "backend/services/meal_planner.py:2898",
                "Creates x[s,i] binary variables for recipe-slot assignment.",
            ],
            [
                "Exactly one recipe per slot",
                "backend/services/meal_planner.py:2905",
                "Each meal slot receives one allowed recipe.",
            ],
            [
                "Recipe usage/repeat expression",
                "backend/services/meal_planner.py:2949",
                "Counts selected uses of each recipe through sums of x variables.",
            ],
            [
                "Nutrition constraints",
                "backend/services/meal_planner.py:3017 and 3032",
                "Applies calorie, macro, and fiber bounds.",
            ],
            [
                "Sodium/sugar overage handling",
                "backend/services/meal_planner.py:3039",
                "Tracks advisory overage behavior rather than hiding it.",
            ],
            [
                "Backend grocery aggregation",
                "backend/services/meal_planner.py:1125 and backend/services/grocery_aggregator.py:702",
                "Builds grocery output by aggregating selected recipe ingredients.",
            ],
            [
                "Backend grocery pricing",
                "backend/services/grocery_aggregator.py:738",
                "Prices aggregated grocery items and returns the displayed estimate.",
            ],
            [
                "Android grocery rebuild",
                "app/src/main/java/com/pcosina/app/domain/GroceryRebuildUseCase.kt:8",
                "Rebuilds grocery items from meal ingredient sources on Android.",
            ],
        ],
    )

    add_heading(doc, "One-Sentence Memory Aid", 1)
    doc.add_paragraph(
        "Table 15 is the planner's math vocabulary: R is safe recipes, S is meal slots, K is nutrition, I is ingredients, "
        "x chooses recipe-slot assignments, y tracks recipe use, u/v track nutrition deviation, and g supports grocery needs."
    )

    return doc


def main() -> None:
    REPO_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(REPO_OUTPUT)
    DOWNLOADS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPO_OUTPUT, DOWNLOADS_OUTPUT)
    print(REPO_OUTPUT)
    print(DOWNLOADS_OUTPUT)


if __name__ == "__main__":
    main()
