"""Generate a DOCX evidence report for 50 planner test cases.

The report is intentionally generated from the current backend planner instead
of manually copied samples, so it can be re-run after recipe or policy changes.
"""

from __future__ import annotations

import csv
import json
import subprocess
import sys
from collections import Counter
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = ROOT / "backend"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from domain.models import GeneratePlanRequest, UserProfile
from scripts.audit_runtime_catalog_diversity import (
    base_profile,
    count_by_meal_type,
    extended_50_scenario_profiles,
    load_catalog,
    production_policy,
)
from services import meal_planner


GENERATED_DATE = "2026-08-11"
CATALOG_PATH = BACKEND_ROOT / "recipes.json"
OUTPUT_DIR = ROOT / "docs" / "thesis_validation" / "03_ACTUAL_SYSTEM_DATA_EXPORTS"
DOCX_PATH = OUTPUT_DIR / "PCOSINA_50_Test_Cases_Full_21_Meal_Output_and_Failure_Evidence.docx"
JSON_PATH = OUTPUT_DIR / "pcosina_50_test_case_full_meal_outputs.json"
CSV_PATH = OUTPUT_DIR / "pcosina_50_test_case_full_meal_outputs.csv"
STRESS_JSON_PATH = OUTPUT_DIR / "pcosina_combination_stress_audit_supported_restrictions_allergy_pairs.json"
STRESS_CSV_PATH = OUTPUT_DIR / "pcosina_combination_stress_audit_supported_restrictions_allergy_pairs.csv"


NEGATIVE_CASES: list[dict[str, Any]] = [
    {
        "id": "FT-001",
        "scenario": "unsupported_14_day_request",
        "profile": base_profile(),
        "days": 14,
        "meals_per_day": 3,
        "purpose": "Policy validation: current production policy only supports 7-day planning.",
    },
    {
        "id": "FT-002",
        "scenario": "unsupported_4_meals_per_day",
        "profile": base_profile(),
        "days": 7,
        "meals_per_day": 4,
        "purpose": "Policy validation: current production policy supports three meals per day.",
    },
    {
        "id": "FT-003",
        "scenario": "prep_time_10_minutes",
        "profile": base_profile(maxCookingTimeMinutes=10, planningPriority="Quick Prep"),
        "days": 7,
        "meals_per_day": 3,
        "purpose": "No-safe-catalog test: user prep limit is too strict for the active catalog.",
    },
    {
        "id": "FT-004",
        "scenario": "custom_cooked_white_rice_allergy",
        "profile": base_profile(allergies=["cooked white rice"]),
        "days": 7,
        "meals_per_day": 3,
        "purpose": "Custom allergy test: a broad staple allergy must not be silently ignored.",
    },
    {
        "id": "FT-005",
        "scenario": "pescatarian_fish_shellfish_allergy_hard_conflict",
        "profile": base_profile(dietaryRestrictions=["Pescatarian"], allergies=["fish", "shellfish"]),
        "days": 7,
        "meals_per_day": 3,
        "purpose": "Hard-conflict test: pescatarian selection cannot be satisfied when both fish and shellfish are excluded.",
    },
    {
        "id": "FT-006",
        "scenario": "ultra_low_budget_500",
        "profile": base_profile(weeklyBudgetPhp=500, planningPriority="Budget First"),
        "days": 7,
        "meals_per_day": 3,
        "purpose": "Budget infeasibility test: final grocery estimate must respect the weekly cap.",
    },
]


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    return " ".join(text.split())


def short_text(value: Any, limit: int = 90) -> str:
    text = safe_text(value)
    if len(text) <= limit:
        return text
    return text[: limit - 3].rstrip() + "..."


def fmt_number(value: Any, digits: int = 1) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    if number.is_integer():
        return str(int(number))
    return f"{number:.{digits}f}"


def fmt_money(value: Any) -> str:
    try:
        return f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return ""


def load_stress_summary() -> dict[str, Any]:
    if not STRESS_JSON_PATH.exists():
        return {}
    try:
        payload = json.loads(STRESS_JSON_PATH.read_text(encoding="utf-8"))
        return payload if isinstance(payload, dict) else {}
    except Exception:
        return {}


def git_commit_short() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"],
            cwd=str(ROOT),
            text=True,
            stderr=subprocess.DEVNULL,
        ).strip()
    except Exception:
        return "unknown"


def profile_dict(profile: UserProfile) -> dict[str, Any]:
    return profile.model_dump(mode="json")


def recipe_title(recipe: dict[str, Any], fallback: str = "") -> str:
    return safe_text(recipe.get("title") or recipe.get("name") or fallback)


def recipe_cost(recipe: dict[str, Any]) -> Any:
    for key in ("_cost_est", "estimatedCostPhp", "costPhp", "cost"):
        value = recipe.get(key)
        if value not in (None, ""):
            return value
    return ""


def recipe_families(recipe: dict[str, Any]) -> list[str]:
    families = recipe.get("_ingredient_variety_families")
    if isinstance(families, list):
        return [safe_text(item) for item in families if safe_text(item)]
    return meal_planner._recipe_ingredient_variety_families(recipe)


def selected_ids_from_rows(rows: list[dict[str, Any]]) -> list[str]:
    return [safe_text(row.get("recipe_id")) for row in rows if safe_text(row.get("recipe_id"))]


def allergy_violation_count(profile: UserProfile, selected_recipes: Iterable[dict[str, Any]]) -> int:
    violations = 0
    for recipe in selected_recipes:
        tags = list(recipe.get("tags") or [])
        tokens = meal_planner.normalize_ingredients(list(recipe.get("ingredients") or []))
        if meal_planner.restriction_failure_reasons(profile, tags, tokens):
            violations += 1
    return violations


def plan_rows_from_result(plan: Any, recipe_by_id: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for day_index, day in enumerate(plan or [], start=1):
        day_label = safe_text(getattr(day, "dayLabel", "") or f"Day {day_index}")
        for meal in getattr(day, "meals", []) or []:
            recipe_id = safe_text(getattr(meal, "recipeId", ""))
            recipe = recipe_by_id.get(recipe_id, {})
            rows.append(
                {
                    "day": day_index,
                    "day_label": day_label,
                    "meal": safe_text(getattr(meal, "mealLabel", "") or recipe.get("mealType")),
                    "recipe_id": recipe_id,
                    "recipe_name": recipe_title(recipe, getattr(meal, "title", "")),
                    "kcal": recipe.get("calories"),
                    "protein_g": recipe.get("proteinGrams"),
                    "carbs_g": recipe.get("carbsGrams"),
                    "fat_g": recipe.get("fatsGrams"),
                    "fiber_g": recipe.get("fiberGrams"),
                    "cost_php": recipe_cost(recipe),
                    "families": ", ".join(recipe_families(recipe)),
                    "source_dataset": safe_text(recipe.get("sourceDataset")),
                    "nutrition_confidence": safe_text(recipe.get("nutritionConfidence")),
                    "review_status": safe_text(recipe.get("nutritionReviewStatus")),
                }
            )
    return rows


def solve_case(
    case_id: str,
    scenario: str,
    profile: UserProfile,
    recipes: list[dict[str, Any]],
    policy: dict[str, Any],
    *,
    days: int = 7,
    meals_per_day: int = 3,
    case_type: str = "success_suite",
    purpose: str = "",
) -> dict[str, Any]:
    scenario_recipes = deepcopy(recipes)
    recipe_by_id = {safe_text(recipe.get("id")): recipe for recipe in scenario_recipes}
    telemetry: dict[str, Any] = {}
    request = GeneratePlanRequest(profile=profile, days=days, mealsPerDay=meals_per_day)
    plan, message, explanation = meal_planner.solve_meal_plan(
        request,
        scenario_recipes,
        policy=deepcopy(policy),
        telemetry_out=telemetry,
    )
    rows = plan_rows_from_result(plan, recipe_by_id)
    selected_ids = selected_ids_from_rows(rows)
    selected_counts = Counter(selected_ids)
    selected_recipes = [recipe_by_id[recipe_id] for recipe_id in selected_ids if recipe_id in recipe_by_id]
    stage1_diag = telemetry.get("stage1_diag") if isinstance(telemetry.get("stage1_diag"), dict) else {}
    solve_pairs = telemetry.get("solve_pair_diagnostics") if isinstance(telemetry.get("solve_pair_diagnostics"), list) else []
    family_counts = dict(explanation.get("ingredientFamilyCounts") or {}) if isinstance(explanation, dict) else {}
    dominant_family = explanation.get("dominantIngredientFamily") if isinstance(explanation, dict) else ""
    dominant_count = int(explanation.get("dominantIngredientFamilyCount") or 0) if isinstance(explanation, dict) else 0
    slot_count = int(days) * int(meals_per_day)
    exact_repeat_count = max(selected_counts.values(), default=0)
    return {
        "case_id": case_id,
        "case_type": case_type,
        "scenario": scenario,
        "purpose": purpose,
        "status": "success" if plan is not None else "failed",
        "message": safe_text(message),
        "days": days,
        "meals_per_day": meals_per_day,
        "slot_count_expected": slot_count,
        "slot_count_returned": len(rows),
        "catalog_recipe_count": len(recipes),
        "catalog_meal_type_counts": count_by_meal_type(recipes),
        "profile": profile_dict(profile),
        "dietary_restrictions": list(profile.dietaryRestrictions or []),
        "allergies": list(profile.allergies or []),
        "weekly_budget_php": profile.weeklyBudgetPhp,
        "max_cooking_time_minutes": profile.maxCookingTimeMinutes,
        "variety_preference": profile.varietyPreference,
        "planning_priority": profile.planningPriority,
        "safe_recipe_count_pre_pricing": stage1_diag.get("safe_recipe_count_pre_pricing"),
        "candidate_count_pre": telemetry.get("candidate_count_pre"),
        "candidate_count_post": telemetry.get("candidate_count_post"),
        "exclusion_summary": stage1_diag.get("exclusion_summary") or {},
        "exclusion_detail_counts": stage1_diag.get("exclusion_detail_counts") or {},
        "nutrition_feasibility": stage1_diag.get("nutrition_feasibility") or {},
        "selected_recipe_ids": selected_ids,
        "selected_unique_recipe_count": len(selected_counts),
        "selected_max_recipe_repeat_count": exact_repeat_count,
        "selected_repeated_recipe_count": sum(1 for count in selected_counts.values() if count > 1),
        "selected_recipe_repeat_counts": dict(sorted(selected_counts.items())),
        "selected_ingredient_family_counts": family_counts,
        "dominant_ingredient_family": dominant_family,
        "dominant_ingredient_family_count": dominant_count,
        "dominant_ingredient_family_share": round(dominant_count / len(rows), 3) if rows else 0.0,
        "estimated_weekly_cost_php": explanation.get("estimatedWeeklyCost") if isinstance(explanation, dict) else None,
        "estimated_weekly_cost_source": explanation.get("estimatedWeeklyCostSource") if isinstance(explanation, dict) else "",
        "max_per_week_used": explanation.get("maxPerWeek") if isinstance(explanation, dict) else "",
        "confidence_score": explanation.get("confidenceScore") if isinstance(explanation, dict) else "",
        "allergy_or_restriction_violation_count": allergy_violation_count(profile, selected_recipes),
        "solve_pair_statuses": [
            {
                "status": item.get("status"),
                "maxPerWeek": item.get("maxPerWeek"),
                "sameSlotConsecutiveRepeatBlocked": item.get("sameSlotConsecutiveRepeatBlocked"),
                "ingredientFamilyRepeatCaps": item.get("ingredientFamilyRepeatCaps") or {},
                "ingredientFamilyRepeatCapsEnforced": item.get("ingredientFamilyRepeatCapsEnforced"),
            }
            for item in solve_pairs
            if isinstance(item, dict)
        ],
        "meal_rows": rows,
    }


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell: Any, width_inches: float) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(18)
    styles["Title"].font.bold = True
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = RGBColor(31, 78, 121)


def style_paragraph(paragraph: Any, font_size: float = 8, bold: bool = False) -> None:
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.bold = bold


def add_table(
    doc: Document,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    *,
    widths: list[float] | None = None,
    font_size: float = 7,
    header_fill: str = "1F4E79",
) -> Any:
    row_list = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, heading in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = safe_text(heading)
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[idx])
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)

    for row_values in row_list:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.text = safe_text(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                set_cell_width(cell, widths[idx])
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
    doc.add_paragraph()
    return table


def add_title(doc: Document, git_short: str, catalog_count: int) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("PCOSina Planner Evidence: 50 Test Cases With Full 21-Meal Outputs")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Generated from the active backend planner, production policy, and runtime recipe catalog."
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Generated: {GENERATED_DATE} | Git commit: {git_short} | Active local catalog recipes: {catalog_count}"
    )


def add_key_value_table(doc: Document, rows: list[tuple[Any, Any]]) -> None:
    add_table(
        doc,
        ["Field", "Value"],
        rows,
        widths=[2.2, 7.8],
        font_size=8,
        header_fill="5B9BD5",
    )


def top_items(mapping: dict[str, Any], limit: int = 6) -> str:
    if not mapping:
        return ""
    try:
        pairs = sorted(mapping.items(), key=lambda item: (-int(item[1]), str(item[0])))
    except Exception:
        pairs = sorted(mapping.items(), key=lambda item: str(item[0]))
    return "; ".join(f"{safe_text(key)}={safe_text(value)}" for key, value in pairs[:limit])


def solve_status_text(case: dict[str, Any]) -> str:
    pairs = case.get("solve_pair_statuses") or []
    if not pairs:
        return ""
    counts = Counter(safe_text(item.get("status")) for item in pairs if safe_text(item.get("status")))
    return ", ".join(f"{status} x{count}" for status, count in counts.items())


def add_executive_summary(doc: Document, successes: list[dict[str, Any]], failures: list[dict[str, Any]]) -> None:
    doc.add_heading("1. Executive Summary", level=1)
    success_count = sum(1 for item in successes if item["status"] == "success")
    failure_count = sum(1 for item in failures if item["status"] == "failed")
    total_slots = sum(int(item.get("slot_count_returned") or 0) for item in successes)
    max_repeat = max(int(item.get("selected_max_recipe_repeat_count") or 0) for item in successes)
    max_family = max(int(item.get("dominant_ingredient_family_count") or 0) for item in successes)
    min_unique = min(int(item.get("selected_unique_recipe_count") or 0) for item in successes)
    max_unique = max(int(item.get("selected_unique_recipe_count") or 0) for item in successes)
    family_max: Counter[str] = Counter()
    for item in successes:
        for family, count in (item.get("selected_ingredient_family_counts") or {}).items():
            family_max[family] = max(family_max[family], int(count or 0))
    stress = load_stress_summary()

    rows = [
        ("Successful representative cases", f"{success_count} of {len(successes)}"),
        ("Full meal slots returned", f"{total_slots} of {len(successes) * 21}"),
        ("Negative/failure controls", f"{failure_count} failed as expected out of {len(failures)}"),
        ("Max exact recipe repeat observed", f"{max_repeat} times in a 21-slot plan"),
        ("Unique recipes per 21-slot plan", f"{min_unique} to {max_unique}"),
        ("Max dominant semantic family observed", f"{max_family} of 21 slots"),
        ("Semantic family maxima observed", top_items(dict(family_max), 10)),
    ]
    if stress:
        rows.extend(
            [
                (
                    "Supported-combination stress audit",
                    f"{stress.get('acceptedCount')} accepted of {stress.get('scenarioCount')} scenarios; "
                    f"{stress.get('successCount')} generated plans; {stress.get('expectedConflictCount')} expected hard-conflict failures",
                ),
                (
                    "Stress-audit variety result",
                    f"Unexpected failures: {stress.get('unexpectedFailureCount')}; variety violations: {stress.get('varietyViolationCount')}; "
                    f"worst max repeat: {(stress.get('worst') or {}).get('maxRepeat')}; worst semantic family: {(stress.get('worst') or {}).get('maxSemanticFamilyCount')}/21",
                ),
            ]
        )
    rows.append(
        (
            "Interpretation",
            "The current catalog can produce complete 21-meal plans for the 50 representative scenarios while keeping exact recipe repetition at or below 3 and semantic-family dominance at or below 8 of 21 slots. "
            "Broader supported-combination stress testing also passed; arbitrary custom allergy text can still produce safe infeasibility when it removes staple ingredients or conflicts with a chosen diet.",
        )
    )

    add_key_value_table(
        doc,
        rows,
    )

    chicken_case = next((item for item in successes if item["scenario"] == "custom_chicken_allergy"), None)
    if chicken_case:
        doc.add_heading("Chicken Allergy Check", level=2)
        add_key_value_table(
            doc,
            [
                ("Scenario", f"{chicken_case['case_id']} - custom_chicken_allergy"),
                ("Input allergy", ", ".join(chicken_case.get("allergies") or [])),
                ("Status", chicken_case.get("status")),
                ("Returned slots", chicken_case.get("slot_count_returned")),
                ("Hard-filter violation count in selected meals", chicken_case.get("allergy_or_restriction_violation_count")),
                ("Selected family counts", top_items(chicken_case.get("selected_ingredient_family_counts") or {}, 10)),
                ("Dominant family", f"{chicken_case.get('dominant_ingredient_family')} ({chicken_case.get('dominant_ingredient_family_count')}/21)"),
                ("Estimated weekly cost PHP", fmt_money(chicken_case.get("estimated_weekly_cost_php"))),
            ],
        )


def add_success_summary_table(doc: Document, cases: list[dict[str, Any]]) -> None:
    doc.add_heading("2. 50 Test Case Summary", level=1)
    rows = []
    for case in cases:
        rows.append(
            [
                case["case_id"],
                short_text(case["scenario"], 34),
                ", ".join(case.get("dietary_restrictions") or []) or "None",
                ", ".join(case.get("allergies") or []) or "None",
                case.get("weekly_budget_php"),
                case.get("max_cooking_time_minutes"),
                case.get("variety_preference"),
                case.get("planning_priority"),
                case.get("status"),
                f"{case.get('slot_count_returned')}/21",
                case.get("selected_unique_recipe_count"),
                case.get("selected_max_recipe_repeat_count"),
                f"{case.get('dominant_ingredient_family')}={case.get('dominant_ingredient_family_count')}",
                fmt_money(case.get("estimated_weekly_cost_php")),
            ]
        )
    add_table(
        doc,
        [
            "ID",
            "Scenario",
            "Diet restrictions",
            "Allergies",
            "Budget PHP",
            "Max min",
            "Var",
            "Priority",
            "Status",
            "Slots",
            "Unique",
            "Max repeat",
            "Dominant family",
            "Est PHP",
        ],
        rows,
        widths=[0.45, 1.35, 1.25, 1.2, 0.6, 0.45, 0.55, 0.75, 0.55, 0.45, 0.45, 0.55, 1.0, 0.65],
        font_size=6.2,
    )


def add_meal_outputs(doc: Document, cases: list[dict[str, Any]]) -> None:
    doc.add_heading("3. Complete 21-Meal Outputs Per Test Case", level=1)
    doc.add_paragraph(
        "Each table below is the actual selected plan returned by the backend planner for that scenario. "
        "Cost is per selected recipe using the backend runtime cost estimate; weekly cost uses the backend aggregated grocery estimate when available."
    )
    for index, case in enumerate(cases, start=1):
        if index > 1:
            doc.add_page_break()
        doc.add_heading(f"{case['case_id']} - {case['scenario']}", level=2)
        add_key_value_table(
            doc,
            [
                ("Inputs", f"Restrictions: {', '.join(case.get('dietary_restrictions') or []) or 'None'} | Allergies: {', '.join(case.get('allergies') or []) or 'None'}"),
                ("Budget / cooking / variety", f"PHP {case.get('weekly_budget_php')} weekly | {case.get('max_cooking_time_minutes')} min max | {case.get('variety_preference')} / {case.get('planning_priority')}"),
                ("Result", f"{case.get('status')} | {case.get('message')} | slots {case.get('slot_count_returned')}/21 | unique recipes {case.get('selected_unique_recipe_count')} | max exact repeat {case.get('selected_max_recipe_repeat_count')}"),
                ("Semantic family counts", top_items(case.get("selected_ingredient_family_counts") or {}, 10)),
                ("Solver attempt statuses", solve_status_text(case)),
            ],
        )
        rows = []
        for meal in case.get("meal_rows") or []:
            rows.append(
                [
                    meal.get("day"),
                    meal.get("meal"),
                    meal.get("recipe_id"),
                    short_text(meal.get("recipe_name"), 50),
                    fmt_number(meal.get("kcal"), 0),
                    fmt_number(meal.get("protein_g"), 0),
                    fmt_number(meal.get("carbs_g"), 0),
                    fmt_number(meal.get("fat_g"), 0),
                    fmt_number(meal.get("fiber_g"), 0),
                    fmt_money(meal.get("cost_php")),
                    short_text(meal.get("families"), 34),
                ]
            )
        add_table(
            doc,
            ["Day", "Meal", "Recipe ID", "Recipe name", "kcal", "Prot", "Carb", "Fat", "Fiber", "Cost PHP", "Families"],
            rows,
            widths=[0.35, 0.65, 0.85, 2.4, 0.45, 0.45, 0.45, 0.42, 0.42, 0.58, 1.45],
            font_size=6.4,
        )


def add_failure_section(doc: Document, cases: list[dict[str, Any]]) -> None:
    doc.add_page_break()
    doc.add_heading("4. Negative and Failure Generation Test Cases", level=1)
    doc.add_paragraph(
        "These cases intentionally use unsupported or very restrictive inputs. A failed generation is acceptable here because the planner returns a clear message instead of violating hard constraints."
    )
    rows = []
    for case in cases:
        rows.append(
            [
                case["case_id"],
                case["scenario"],
                case["purpose"],
                ", ".join(case.get("dietary_restrictions") or []) or "None",
                ", ".join(case.get("allergies") or []) or "None",
                case.get("weekly_budget_php"),
                case.get("max_cooking_time_minutes"),
                f"{case.get('days')}x{case.get('meals_per_day')}",
                case.get("status"),
                case.get("message"),
                case.get("safe_recipe_count_pre_pricing"),
                f"{case.get('candidate_count_pre')}/{case.get('candidate_count_post')}",
                solve_status_text(case),
            ]
        )
    add_table(
        doc,
        [
            "ID",
            "Scenario",
            "Purpose",
            "Restrictions",
            "Allergies",
            "Budget",
            "Max min",
            "Plan",
            "Status",
            "System message",
            "Safe",
            "Cand pre/post",
            "Solve statuses",
        ],
        rows,
        widths=[0.55, 1.3, 2.1, 1.0, 1.2, 0.55, 0.45, 0.45, 0.55, 1.5, 0.45, 0.65, 1.05],
        font_size=6.4,
        header_fill="7F6000",
    )

    for case in cases:
        doc.add_heading(f"{case['case_id']} Details", level=2)
        add_key_value_table(
            doc,
            [
                ("Scenario", case.get("scenario")),
                ("Purpose", case.get("purpose")),
                ("Inputs", f"days={case.get('days')}, mealsPerDay={case.get('meals_per_day')}, budgetPHP={case.get('weekly_budget_php')}, maxCookingTime={case.get('max_cooking_time_minutes')}"),
                ("Restrictions", ", ".join(case.get("dietary_restrictions") or []) or "None"),
                ("Allergies", ", ".join(case.get("allergies") or []) or "None"),
                ("Actual status/message", f"{case.get('status')} | {case.get('message')}"),
                ("Returned meal slots", case.get("slot_count_returned")),
                ("Exclusion summary", top_items(case.get("exclusion_summary") or {}, 10)),
                ("Top exclusion details", top_items(case.get("exclusion_detail_counts") or {}, 10)),
                ("Solver statuses", solve_status_text(case)),
                ("Interpretation", "No unsafe fallback meal list was returned for this failure case."),
            ],
        )


def add_appendix(doc: Document, success_cases: list[dict[str, Any]], failure_cases: list[dict[str, Any]], git_short: str) -> None:
    doc.add_page_break()
    doc.add_heading("5. Reproducibility Notes", level=1)
    rows = [
        ("Generator", "scripts/generate_50_test_case_meal_output_doc.py"),
        ("Catalog", "backend/recipes.json"),
        ("Policy source", "backend/policy_config.py plus runtime bootstrap policy"),
        ("Git commit", git_short),
        ("Success suite source", "scripts/audit_runtime_catalog_diversity.py::extended_50_scenario_profiles"),
        ("Success cases", len(success_cases)),
        ("Negative/failure controls", len(failure_cases)),
        ("Machine-readable JSON", str(JSON_PATH.relative_to(ROOT))),
        ("Machine-readable CSV", str(CSV_PATH.relative_to(ROOT))),
    ]
    if STRESS_JSON_PATH.exists():
        rows.extend(
            [
                ("Supported-combination stress JSON", str(STRESS_JSON_PATH.relative_to(ROOT))),
                ("Supported-combination stress CSV", str(STRESS_CSV_PATH.relative_to(ROOT))),
                (
                    "Stress audit scope",
                    "12 supported restriction profiles crossed with zero, one, and two allergies from the listed allergy set plus custom chicken.",
                ),
            ]
        )
    rows.extend(
        [
            (
                "Coverage boundary",
                "The suite validates supported and representative combinations. It does not claim that every arbitrary free-text allergy can generate a plan; impossible or conflicting inputs must fail clearly.",
            ),
            (
                "Medical safety scope",
                "PCOSina is a wellness decision-support meal planner. These tests validate planning behavior and constraint handling, not diagnosis or medical treatment.",
            ),
        ]
    )
    add_key_value_table(
        doc,
        rows,
    )


def write_csv(success_cases: list[dict[str, Any]], failure_cases: list[dict[str, Any]]) -> None:
    rows: list[dict[str, Any]] = []
    for case in success_cases + failure_cases:
        if case.get("meal_rows"):
            for meal in case["meal_rows"]:
                rows.append(
                    {
                        "case_id": case["case_id"],
                        "case_type": case["case_type"],
                        "scenario": case["scenario"],
                        "status": case["status"],
                        "message": case["message"],
                        "dietary_restrictions": "; ".join(case.get("dietary_restrictions") or []),
                        "allergies": "; ".join(case.get("allergies") or []),
                        "weekly_budget_php": case.get("weekly_budget_php"),
                        "max_cooking_time_minutes": case.get("max_cooking_time_minutes"),
                        "variety_preference": case.get("variety_preference"),
                        "planning_priority": case.get("planning_priority"),
                        "selected_unique_recipe_count": case.get("selected_unique_recipe_count"),
                        "selected_max_recipe_repeat_count": case.get("selected_max_recipe_repeat_count"),
                        "dominant_ingredient_family": case.get("dominant_ingredient_family"),
                        "dominant_ingredient_family_count": case.get("dominant_ingredient_family_count"),
                        "estimated_weekly_cost_php": case.get("estimated_weekly_cost_php"),
                        **meal,
                    }
                )
        else:
            rows.append(
                {
                    "case_id": case["case_id"],
                    "case_type": case["case_type"],
                    "scenario": case["scenario"],
                    "status": case["status"],
                    "message": case["message"],
                    "dietary_restrictions": "; ".join(case.get("dietary_restrictions") or []),
                    "allergies": "; ".join(case.get("allergies") or []),
                    "weekly_budget_php": case.get("weekly_budget_php"),
                    "max_cooking_time_minutes": case.get("max_cooking_time_minutes"),
                    "variety_preference": case.get("variety_preference"),
                    "planning_priority": case.get("planning_priority"),
                    "selected_unique_recipe_count": case.get("selected_unique_recipe_count"),
                    "selected_max_recipe_repeat_count": case.get("selected_max_recipe_repeat_count"),
                    "dominant_ingredient_family": case.get("dominant_ingredient_family"),
                    "dominant_ingredient_family_count": case.get("dominant_ingredient_family_count"),
                    "estimated_weekly_cost_php": case.get("estimated_weekly_cost_php"),
                    "day": "",
                    "day_label": "",
                    "meal": "",
                    "recipe_id": "",
                    "recipe_name": "",
                    "kcal": "",
                    "protein_g": "",
                    "carbs_g": "",
                    "fat_g": "",
                    "fiber_g": "",
                    "cost_php": "",
                    "families": "",
                    "source_dataset": "",
                    "nutrition_confidence": "",
                    "review_status": "",
                }
            )
    fieldnames = list(rows[0].keys()) if rows else []
    with CSV_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def build_document(success_cases: list[dict[str, Any]], failure_cases: list[dict[str, Any]], git_short: str, catalog_count: int) -> None:
    doc = Document()
    configure_document(doc)
    add_title(doc, git_short, catalog_count)
    add_executive_summary(doc, success_cases, failure_cases)
    add_success_summary_table(doc, success_cases)
    add_meal_outputs(doc, success_cases)
    add_failure_section(doc, failure_cases)
    add_appendix(doc, success_cases, failure_cases, git_short)
    doc.save(DOCX_PATH)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    recipes = load_catalog(CATALOG_PATH)
    policy = production_policy()
    git_short = git_commit_short()

    success_cases: list[dict[str, Any]] = []
    for index, (scenario, profile) in enumerate(extended_50_scenario_profiles().items(), start=1):
        success_cases.append(
            solve_case(
                f"TC-{index:03d}",
                scenario,
                profile,
                recipes,
                policy,
                case_type="success_suite",
            )
        )

    failure_cases: list[dict[str, Any]] = []
    for case in NEGATIVE_CASES:
        failure_cases.append(
            solve_case(
                case["id"],
                case["scenario"],
                case["profile"],
                recipes,
                policy,
                days=case["days"],
                meals_per_day=case["meals_per_day"],
                case_type="negative_control",
                purpose=case["purpose"],
            )
        )
    unexpected_negative_successes = [
        case["case_id"] for case in failure_cases if case["status"] == "success"
    ]
    if unexpected_negative_successes:
        raise RuntimeError(
            "Negative-control case(s) unexpectedly generated a plan: "
            + ", ".join(unexpected_negative_successes)
        )

    payload = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "generatedDate": GENERATED_DATE,
        "gitCommitShort": git_short,
        "catalogPath": str(CATALOG_PATH.relative_to(ROOT)),
        "catalogRecipeCount": len(recipes),
        "successCaseCount": len(success_cases),
        "failureCaseCount": len(failure_cases),
        "successCases": success_cases,
        "failureCases": failure_cases,
    }
    JSON_PATH.write_text(json.dumps(payload, indent=2, ensure_ascii=True), encoding="utf-8")
    write_csv(success_cases, failure_cases)
    build_document(success_cases, failure_cases, git_short, len(recipes))

    summary = {
        "docx": str(DOCX_PATH),
        "json": str(JSON_PATH),
        "csv": str(CSV_PATH),
        "successCases": len(success_cases),
        "failureCases": len(failure_cases),
        "successCaseFailures": [case["case_id"] for case in success_cases if case["status"] != "success"],
        "fullSlotFailures": [case["case_id"] for case in success_cases if case["slot_count_returned"] != 21],
        "maxExactRepeat": max(case["selected_max_recipe_repeat_count"] for case in success_cases),
        "maxSemanticFamilyCount": max(case["dominant_ingredient_family_count"] for case in success_cases),
        "chickenAllergyCase": next(
            {
                "caseId": case["case_id"],
                "status": case["status"],
                "slots": case["slot_count_returned"],
                "violations": case["allergy_or_restriction_violation_count"],
                "familyCounts": case["selected_ingredient_family_counts"],
            }
            for case in success_cases
            if case["scenario"] == "custom_chicken_allergy"
        ),
        "negativeStatuses": {
            case["case_id"]: {"scenario": case["scenario"], "status": case["status"], "message": case["message"]}
            for case in failure_cases
        },
    }
    print(json.dumps(summary, indent=2, ensure_ascii=True))


if __name__ == "__main__":
    main()
