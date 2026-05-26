from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "docs" / "PCOSINA MANUSCRIPT.docx"
OUT = ROOT / "docs" / "defense" / "PCOSINA_Manuscript_SOP_Coverage_Audit_2026-05-23.docx"


def manuscript_stats() -> dict[str, int]:
    doc = Document(MANUSCRIPT)
    text = "\n".join(p.text for p in doc.paragraphs)
    lowered = text.lower()
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "insulin_resistance": lowered.count("insulin resistance"),
        "insulin_resistance_related": lowered.count("insulin-resistance-related"),
        "pantry_constrained": lowered.count("pantry-constrained"),
        "milp_based": lowered.count("milp-based"),
        "cp_sat": lowered.count("cp-sat"),
    }


def set_cell_text(cell, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.font.size = Pt(9.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        r.font.size = Pt(9)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, size=8)
        if widths:
            table.rows[0].cells[idx].width = Inches(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=8)
            if widths:
                cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def build_doc() -> None:
    stats = manuscript_stats()

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(9.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PCOSINA Manuscript SOP Coverage and Evidence Audit")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Audit of whether the latest manuscript contains the required How, Inputs, Outputs/Evidence, and Metrics for each SOP")
    r.font.size = Pt(10)

    add_body(doc, f"Manuscript scanned: {MANUSCRIPT}")
    add_body(doc, "Date prepared: 2026-05-23")
    add_body(
        doc,
        "Scope note: This audit assumes the SOP wording will be updated separately. It focuses on what the current manuscript already contains, what is controllable now, and what should be added or rewritten before defense. Paragraph references such as P0615 are DOCX extraction indexes for the current manuscript version; use the section title and quoted wording with Ctrl+F if paragraph numbering changes after edits.",
    )

    add_heading(doc, "Executive Readout", 1)
    add_bullets(
        doc,
        [
            "Overall status: mostly aligned, but not yet defense-clean. The manuscript already contains the core algorithm pipeline, input lists, output descriptions, evaluation results, Cronbach alpha results, ISO/IEC 25010 survey results, pantry/offline boundaries, and benchmark results.",
            "Main weakness: the needed SOP evidence exists, but it is scattered across Chapters 1, 3, 4, and 5. Add one explicit SOP-to-evidence bridge table so panelists can immediately see how each question was answered.",
            "Highest-risk wording to revise: insulin resistance still appears in the SOP and planning inputs; pantry is sometimes described as pantry-constrained even though the manuscript later says pantry overlap is a scoring/guidance signal; food-waste reduction is sometimes stated as an achieved effect instead of a supported/proxy outcome.",
            "Strongest existing evidence: Table 25 Cronbach alpha, Tables 33-41 user evaluation, Table 45 functional testing, Table 47 final planner safety results, Table 49 pantry/grocery/budget behavior, Table 50 offline-first boundaries, and Table 51 final runtime progression.",
        ],
    )

    add_heading(doc, "Scan Snapshot", 1)
    add_table(
        doc,
        ["Scan Item", "Observed Count / Value", "Defense Meaning"],
        [
            ["Paragraphs extracted", str(stats["paragraphs"]), "The audit scanned the full text body of the manuscript."],
            ["Tables extracted", str(stats["tables"]), "The manuscript has enough tabular evidence to support a structured defense, but it needs clearer SOP mapping."],
            ["'insulin resistance' occurrences", str(stats["insulin_resistance"]), "Still needs cleanup for the revised SOP direction."],
            ["'insulin-resistance-related' occurrences", str(stats["insulin_resistance_related"]), "Still appears in planning input descriptions and should be removed or reframed."],
            ["'pantry-constrained' occurrences", str(stats["pantry_constrained"]), "Risky unless pantry is truly hard-constrained; safer wording is pantry-aware."],
            ["'MILP-based' occurrences", str(stats["milp_based"]), "Acceptable only if clearly explained as formulation style, because CP-SAT is the implementation solver."],
            ["'CP-SAT' occurrences", str(stats["cp_sat"]), "Good: the actual implementation solver is strongly represented."],
        ],
        [2.2, 1.4, 5.4],
    )

    add_heading(doc, "Overall Coverage Matrix", 1)
    add_table(
        doc,
        ["SOP", "How PCOSina Solves This", "Inputs Required", "Outputs / Evidence Required", "Metrics / Results Needed", "Audit Status"],
        [
            [
                "1. Weekly nutrition and weight-management meal planning",
                "Present in IPO framework, algorithm pipeline, formal computations, Stage 1/Stage 2 planner, and Chapter 4 results.",
                "Present: profile, anthropometrics, activity, goal, restrictions, allergies, budget, pantry, recipe/nutrition data. Revise insulin-related inputs.",
                "Present: 21-slot weekly plan, nutrition summary, cost, pantry overlap, grocery guidance, no-safe-plan handling.",
                "Mostly present: 100/100 final runs, zero hard-rule violations, nutrition feasibility, runtime. Add clearer nutrient-target acceptance thresholds.",
                "Mostly included; revise SOP wording and add a compact SOP 1 evidence table.",
            ],
            [
                "2. Preferences, culture, and restrictions",
                "Present in deterministic filtering, LightGBM-assisted ranking, CP-SAT assignment, and cultural relevance protocol.",
                "Present: dietary restrictions, allergies, explicit exclusions, cooking time, budget, Filipino recipe data, pantry entries, preferences.",
                "Present: filtered candidates, Filipino-context meal plans, changed outputs by profile, survey support.",
                "Partial: user ratings are present; cultural CTS/CIS formulas are present; actual CTS/CIS result table is not clearly shown.",
                "Partially included; add computed cultural relevance results by scenario.",
            ],
            [
                "3. Pantry inventory and unnecessary purchase reduction",
                "Present as pantry-aware ranking and grocery-deficit guidance, not as a hard pantry-feasibility guarantee.",
                "Present: manual pantry entries, normalized pantry/ingredient tokens, available quantities when known, recipe ingredient requirements.",
                "Present: grocery guidance, pantry-covered item flags, pantry-overlap behavior, budget enforcement.",
                "Partial: Table 49 has proxy evidence. Missing actual food-waste/adherence measurement, so claim should be softened.",
                "Included but wording risk remains; reframe as helps promote reduction or supports reduction.",
            ],
            [
                "4. Offline-first mobile support",
                "Present as local storage continuity for profiles, saved plans, pantry data, grocery guides, and progress records.",
                "Present: locally stored profiles, pantry records, saved meal plans, grocery guides, progress records; backend required for new optimized generation.",
                "Present: offline access to saved records, portability survey item, functional testing, Table 50.",
                "Partial: user ratings and pass/fail evidence are present. Add a clearer offline scenario test matrix if not already in Appendix L.",
                "Mostly included; ensure all claims say saved-data continuity, not full offline generation.",
            ],
            [
                "5. Nutritional quality, usability, cultural relevance, computational efficiency",
                "Present through evaluation methodology, ISO/IEC 25010 survey, Cronbach alpha, functional tests, planner benchmarks, ML ranking evaluation.",
                "Present: respondent survey data, expert validation, test cases, benchmark profiles, planner logs, ML dataset metrics.",
                "Present: tables for Cronbach alpha, ISO results, functional tests, planner safety, ML metrics, runtime, pantry/offline evidence.",
                "Mostly present: Cronbach alpha all above 0.700, overall ISO mean 4.02, final 100/100 runs, avg 257 ms, P95 313 ms, max 357 ms. Cultural metric results still need explicit values.",
                "Strongly included; add traceability from each metric to each SOP and show actual cultural CTS/CIS results.",
            ],
        ],
        [0.95, 1.75, 1.75, 1.75, 1.75, 1.35],
    )

    add_heading(doc, "SOP 1 Detailed Audit", 1)
    add_body(doc, "Question focus: generate nutritionally adequate weekly meal plans for Filipino users with PCOS while supporting weight-management constraints. Revised direction: remove insulin resistance from the SOP and system framing.")
    add_table(
        doc,
        ["Bucket", "Where It Is Found", "What Is Already Included", "Still Needed / Fix"],
        [
            ["How", "Chapter 1 Conceptual Framework P0251-P0293; Chapter 3 Algorithm P0615-P0652; Stage 1/Stage 2 P0691-P0759; Chapter 4 P1108-P1121 and P1158-P1161.", "Two-stage workflow: deterministic filtering, non-authoritative LightGBM ranking, CP-SAT weekly assignment, fallback/no-safe-plan handling.", "Add a short SOP 1 bridge paragraph after the revised SOP: profile -> targets -> candidate filtering -> CP-SAT -> validation -> stored output."],
            ["Inputs", "P0264, P0620, P0652-P0657, Table 14 and Table 15.", "Profile, anthropometrics, activity level, wellness goal, restrictions, allergies, budget, pantry, cooking time, variety, planning priority, recipe/nutrition/price data.", "Remove insulin-resistance-related setting from the input list unless the app still intentionally collects it. Keep weight management, calorie target, fiber/carbohydrate boundary, and general PCOS-supportive nutrition."],
            ["Outputs / Evidence", "P0282-P0286, P0632, P0780-P0782, Table 45, Table 47.", "7-day/21-slot plan, recipe details, nutrition summaries, cost, pantry indicators, grocery guidance, reason codes when no safe plan exists.", "Add one table showing accepted output artifacts for SOP 1: weekly plan, nutrition summary, constraint report, no-safe-plan diagnostics, and screenshots/log files."],
            ["Metrics / Results", "P1080, P1119-P1124, Table 47, Table 51.", "Final benchmark: 100/100 runs completed, 21 meal slots, zero hard-rule violations, nutrition feasibility passed, average 257 ms, P95 313 ms, max 357 ms.", "Clarify exact nutrition acceptance thresholds: calorie band, macro band, protein/fiber lower bounds, advisory sodium/sugar handling, and how deviation was computed."],
        ],
        [1.0, 2.0, 3.0, 3.0],
    )

    add_heading(doc, "SOP 2 Detailed Audit", 1)
    add_body(doc, "Question focus: integrate user preferences, Filipino cultural food practices, and dietary restrictions into personalized planning.")
    add_table(
        doc,
        ["Bucket", "Where It Is Found", "What Is Already Included", "Still Needed / Fix"],
        [
            ["How", "P0274, P0293, P0630, P0691-P0719, P0729, P0787-P0795, P0918-P0930, P1163-P1166.", "Restrictions and allergies are hard filters; preferences, pantry overlap, cost, Filipino food context, and LightGBM ranking influence candidate ordering; CP-SAT makes the final assignment.", "Add one simple example showing two different user profiles producing different candidate pools or meal selections."],
            ["Inputs", "P0264, P0620, P0694-P0708, Table 12, Table 14.", "Preferences, allergies, restrictions, explicit exclusions, cooking time, budget, pantry, nutrition values, recipe attributes, meal-slot suitability.", "Separate hard preferences from soft preferences. Mark which preference type is never relaxed."],
            ["Outputs / Evidence", "Table 33, Table 36, P1041, P1050, P1164-P1166.", "Functional Suitability includes familiar Filipino food options; Usability supports app clarity; the manuscript claims outputs change depending on profile configuration.", "Add actual test evidence that a banned ingredient/allergen never appears, and that Filipino recipe selections are traceable to cultural tags or recipe names."],
            ["Metrics / Results", "P0787-P0795, P0918-P0930, Table 19, Table 20, Table 33.", "CTS/CIS formulas and acceptance rules are defined; user ratings for familiar Filipino food options are present.", "Missing: final CTS/CIS values by evaluation scenario. Add a small table with scenario, selected meals, CTS, CIS, safety violations, and pass/fail."],
        ],
        [1.0, 2.0, 3.0, 3.0],
    )

    add_heading(doc, "SOP 3 Detailed Audit", 1)
    add_body(doc, "Question focus: use pantry inventory data to support grocery guidance and help promote reduced unnecessary purchases without claiming guaranteed food-waste reduction.")
    add_table(
        doc,
        ["Bucket", "Where It Is Found", "What Is Already Included", "Still Needed / Fix"],
        [
            ["How", "P0330-P0334, P0644, P0680-P0682, P1130-P1135, P1168-P1171, P1216-P1217.", "Manual pantry entries are normalized; pantry overlap influences ranking; grocery gaps are computed; pantry-covered items are flagged; impossible budgets return no-safe-plan.", "Replace guaranteed wording with: helps promote ingredient reuse, supports reduced unnecessary purchases, or provides pantry-aware guidance."],
            ["Inputs", "P0264, P0330-P0332, P0620, P0680-P0682, Table 49.", "User-entered pantry items, normalized ingredient tokens, recipe ingredient requirements, quantities when available, budget cap, price estimates.", "Add explicit reminder that pantry data quality depends on user updates, ingredient quantity, freshness, and adherence."],
            ["Outputs / Evidence", "P0282, P0334, P0780, Table 45, Table 49, P1171.", "Grocery list, grocery deficit guidance, pantry-covered flags, scenario evidence with 38 pantry-overlap matches, PHP 1,000 budget pass, PHP 20 no-safe-plan boundary.", "Avoid saying the system proved actual food waste reduction unless you have pre/post household waste or purchase logs."],
            ["Metrics / Results", "Table 31, Table 33, Table 49.", "Proxy evidence exists: pantry-checking behavior, user perception of reduced purchases, pantry-overlap matches, grocery deficit behavior, budget enforcement.", "Missing: actual food-waste measurement. Add proxy-only label and metrics such as pantry overlap count, grocery deficit count, duplicate-purchase avoidance prompts, and user-reported perceived purchase reduction."],
        ],
        [1.0, 2.0, 3.0, 3.0],
    )

    add_heading(doc, "SOP 4 Detailed Audit", 1)
    add_body(doc, "Question focus: offline-first support for meal planning, recipe access, and pantry management under limited connectivity.")
    add_table(
        doc,
        ["Bucket", "Where It Is Found", "What Is Already Included", "Still Needed / Fix"],
        [
            ["How", "P0232, P0262, P0284, P0350, P0540-P0546, P0572-P0575, P1137-P1142, P1173-P1176, P1218-P1219.", "Local persistence keeps existing profiles, saved plans, pantry data, grocery guidance, and progress records accessible offline; new optimized generation may require backend CP-SAT service.", "Keep this boundary visible wherever offline-first is mentioned: saved-data continuity is offline, new optimized generation is hybrid/backend-dependent."],
            ["Inputs", "P0232, P0262, P0284, P0572, Table 42, Table 50.", "Locally stored profile, pantry records, generated meal plans, grocery guides, progress logs, optional cloud/sync support.", "Add offline test inputs: airplane mode, app restart, no backend, existing saved plan, unsynced pantry update, and cached recipe details."],
            ["Outputs / Evidence", "Table 45, Table 50, P1138-P1142, P1174-P1176.", "User ratings support offline access; Table 50 explains saved-data continuity and backend dependency; functional tests passed offline saved-data behavior.", "Add screenshots or logs showing saved plan/pantry/grocery views opening while offline."],
            ["Metrics / Results", "Table 50, Table 51, Table 45, Table 47.", "Portability overall mean 4.04 Agree; Reliability overall mean 4.00 Agree; offline reliability functional test passed.", "Add an offline scenario result matrix with pass/fail, observed behavior, data preserved, and limitation message for new generation."],
        ],
        [1.0, 2.0, 3.0, 3.0],
    )

    add_heading(doc, "SOP 5 Detailed Audit", 1)
    add_body(doc, "Question focus: evaluate nutritional quality, usability, cultural relevance, and computational efficiency using algorithmic and user-centered metrics.")
    add_table(
        doc,
        ["Bucket", "Where It Is Found", "What Is Already Included", "Still Needed / Fix"],
        [
            ["How", "P0784-P0959, P0985-P1149, P1157-P1181, P1206-P1220.", "Evaluation methodology covers ISO/IEC 25010, Cronbach alpha, functional testing, planner benchmarks, ML ranking evaluation, pantry/offline evaluation, and results discussion.", "Add a direct 'Metric-to-SOP Traceability Table' before or after Chapter 4 summary."],
            ["Inputs", "P0830-P0877, P0944-P0959, P1102-P1113, Table 25, Tables 33-41, Tables 45-51.", "50 respondents, 25 primary and 25 secondary; expert validators; benchmark profiles; functional test cases; ML dataset; survey instrument and Cronbach alpha.", "Keep raw survey spreadsheet, Excel Cronbach computation, benchmark logs, and functional test records ready for panel checking."],
            ["Outputs / Evidence", "Table 25, Tables 33-41, Table 45, Table 47, Table 48, Table 49, Table 50, Table 51.", "Cronbach alpha all subscales exceed 0.700; ISO overall total 4.02 Agree; functional tests passed; final benchmark 100/100; ML NDCG@10 reported; offline/pantry results reported.", "Make sure each table is introduced with which SOP it answers. Avoid leaving Chapter 4 as a collection of unrelated results."],
            ["Metrics / Results", "P0997-P0999, P1038-P1066, P1080, P1119-P1128, P1130-P1146.", "Usability 4.06 Agree, Functional Suitability 3.93 Agree, Security 4.24 Strongly Agree, final runtime avg 257 ms, P95 313 ms, max 357 ms, zero hard-rule violations.", "Missing explicit actual cultural CTS/CIS result values. If unavailable, state cultural relevance is supported by user ratings and protocol, not completed numeric CTS/CIS results."],
        ],
        [1.0, 2.0, 3.0, 3.0],
    )

    add_heading(doc, "High-Risk Manuscript Wording to Revise", 1)
    add_table(
        doc,
        ["Issue", "Exact Locations", "Why It Matters", "Recommended Defense-Safe Wording"],
        [
            ["Insulin resistance still appears in SOP and inputs", "P0185, P0264, P0315, P0367, P0620, P1158 plus related background paragraphs.", "The revised SOP removes insulin resistance. Leaving it in the system input and SOP framing invites panel questions about clinical claims and current app fields.", "Use weight management, calorie-appropriate planning, balanced meal composition, fiber-aware and carbohydrate-conscious guidance. Remove insulin-resistance-related setting unless still implemented and defended."],
            ["Food waste is stated as an outcome instead of a supported/proxy outcome", "P0187, P0198, P1168-P1171, P1186, P1216-P1217.", "Actual waste reduction depends on user adherence, pantry accuracy, quantity, freshness, and purchase behavior. The system can support or promote reduction, not guarantee it.", "How can pantry inventory data be integrated to help promote reduced unnecessary purchases and support food-waste-conscious meal planning while maintaining nutritional adequacy?"],
            ["Pantry-constrained wording conflicts with pantry-aware implementation", "P0002, P0032, P0193, P0196, P0497, P1212.", "The manuscript correctly says pantry overlap is a scoring/guidance signal in P0644 and P0680-P0682. Calling the whole optimizer pantry-constrained may overclaim hard pantry feasibility.", "Use pantry-aware optimization or pantry-informed planning unless the solver strictly requires selected meals to be feasible from pantry inventory."],
            ["MILP-based title needs CP-SAT clarification", "P0002, P0032, P0623; mitigated by P0215, P0219, P0397, P0398, P0724.", "The actual solver is OR-Tools CP-SAT. The manuscript already explains MILP as formulation style, but title/objective wording must remain consistent.", "Use MILP-formulated 0-1 assignment model solved using OR-Tools CP-SAT, or simply constraint-based CP-SAT weekly assignment optimization."],
            ["Offline-first can sound like full offline generation", "P0199, P1188, P1218-P1219.", "The manuscript itself says new optimized meal generation may depend on backend availability. Overclaiming full offline behavior creates a contradiction.", "PCOSina supports offline access to locally stored profiles, saved meal plans, pantry records, grocery guides, and progress logs; new optimized plan generation may require the backend optimizer."],
            ["Cultural relevance metrics are defined but final numeric values are not clearly reported", "P0787-P0795, P0918-P0930; results rely more on Table 33 and narrative.", "SOP 5 asks for cultural relevance performance. A panelist may ask for actual CTS/CIS results because the formulas are introduced.", "Add a scenario-level cultural relevance results table or explicitly state that cultural relevance was evaluated through user ratings and expert review while CTS/CIS remains a proposed metric."],
        ],
        [1.3, 1.7, 2.7, 3.4],
    )

    add_heading(doc, "Recommended Insertions", 1)
    add_table(
        doc,
        ["Priority", "Where to Insert", "What to Add", "Purpose"],
        [
            ["1", "Immediately after revised Statement of the Problem in Chapter 1", "One SOP operationalization table with columns: SOP, system process, required inputs, output artifacts, evaluation metrics, chapter/table evidence.", "This directly answers all 'How' questions and prevents the manuscript from looking scattered."],
            ["2", "Chapter 3 after Algorithm section P0615-P0652", "A concise pipeline diagram/list: collect profile -> validate inputs -> compute targets -> filter hard constraints -> rank safe candidates -> CP-SAT optimize -> validate output -> store locally -> show grocery guidance.", "This makes the algorithm process defense-ready."],
            ["3", "Chapter 4 before Presentation and Analysis of System Results P1157", "A results traceability table mapping each SOP to Tables 25, 33-41, 45, 47-51 and key paragraph findings.", "This makes the results chapter answer the SOPs explicitly."],
            ["4", "Pantry/Grocery section P1130-P1135 and SOP 3 answer P1168-P1171", "Replace guarantee wording with proxy wording and add pantry metrics: overlap matches, grocery deficits, pantry-covered item count, budget outcome, user perception score.", "This keeps the claim honest and defensible."],
            ["5", "Cultural relevance evaluation P0787-P0795 and Chapter 4 results", "Add actual CTS/CIS table, or state clearly that CTS/CIS is a defined protocol and not a completed numeric result if no computed data exists.", "This closes the biggest metric gap for cultural relevance."],
            ["6", "Offline-first section P1137-P1142 and Chapter 5 P1218-P1219", "Add offline scenario test matrix: saved plan view, pantry CRUD, grocery list access, progress logs, app restart, backend unavailable, new generation limitation.", "This turns offline-first from a claim into testable evidence."],
        ],
        [0.8, 2.1, 4.2, 2.0],
    )

    add_heading(doc, "Defense Answer Scripts", 1)
    add_table(
        doc,
        ["Possible Panel Question", "Best Short Answer"],
        [
            ["Is the system treating insulin resistance?", "No. The revised framing should focus on weight-management-oriented, calorie-appropriate, balanced meal planning. Low-GI/fiber-aware nutrition can remain as general PCOS-supportive guidance, but PCOSina is not diagnosing or treating insulin resistance."],
            ["Does pantry integration reduce food waste?", "The system does not guarantee real household food-waste reduction. It uses pantry records to prioritize ingredient reuse and generate grocery guidance, which can help promote reduced unnecessary purchases if the user keeps pantry data accurate and follows the plan."],
            ["Is the planner MILP or CP-SAT?", "The meal assignment is described with a 0-1 optimization formulation using binary variables and constraints, while the actual implementation solver is OR-Tools CP-SAT. Deterministic filtering and CP-SAT remain the authority; ML only helps ranking."],
            ["What works offline?", "Saved profiles, generated meal plans, pantry records, grocery guides, and progress records remain accessible offline after they are created. Generating a new optimized weekly plan may still require backend optimizer availability."],
            ["Are the survey results reliable?", "Based on Table 25, all Cronbach alpha values for primary and secondary user subscales are above 0.700, which supports acceptable to excellent internal consistency for the survey instrument."],
            ["What is still missing in the manuscript?", "The manuscript needs clearer SOP-to-evidence mapping, insulin cleanup, softer pantry/food-waste wording, explicit CTS/CIS cultural result values or a limitation statement, and a more concrete offline scenario matrix."],
        ],
        [2.3, 6.5],
    )

    add_heading(doc, "Final Action Checklist", 1)
    add_bullets(
        doc,
        [
            "Revise SOP 1 and all input lists to remove insulin resistance and insulin-resistance-related setting unless it is intentionally retained and fully defended.",
            "Change pantry-constrained wording to pantry-aware wherever pantry is not a strict hard feasibility constraint.",
            "Change reduce food waste to help promote reduced unnecessary purchases or support food-waste-conscious planning.",
            "Add one SOP-to-evidence matrix in Chapter 1 or Chapter 4.",
            "Add actual CTS/CIS cultural relevance results or label CTS/CIS as a defined protocol rather than a completed result.",
            "Add or point to offline scenario evidence showing saved-data continuity under no internet.",
            "Keep raw survey, Cronbach computation, expert validation, functional test matrix, benchmark logs, and generated meal-plan examples ready for final defense.",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
