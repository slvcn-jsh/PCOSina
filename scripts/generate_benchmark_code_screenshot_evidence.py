from __future__ import annotations

import json
import shutil
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()
OUT_DIR = ROOT / "docs" / "defense" / f"benchmark_code_screenshots_{TODAY}"
DOCX_PATH = ROOT / "docs" / "defense" / f"PCOSINA_Benchmark_Code_Screenshot_Evidence_{TODAY}.docx"
ZIP_PATH = ROOT / "docs" / "defense" / f"PCOSINA_Benchmark_Code_Screenshot_Evidence_{TODAY}.zip"
DOWNLOADS = Path.home() / "Downloads"


@dataclass(frozen=True)
class SourceShot:
    title: str
    source: Path
    start: int
    end: int
    caption: str
    filename: str


SHOTS = [
    SourceShot(
        title="Benchmark Fixture and CLI Controls",
        source=ROOT / "scripts" / "benchmark_planner_profiles.py",
        start=29,
        end=50,
        caption=(
            "Shows that the benchmark runner uses the fixed T01-T20 fixture path, "
            "supports repeated runs per profile, and can require the ML ranker to be ready."
        ),
        filename="01_benchmark_fixture_and_cli.png",
    ),
    SourceShot(
        title="P95 Runtime Formula",
        source=ROOT / "scripts" / "benchmark_planner_profiles.py",
        start=77,
        end=89,
        caption="Shows the script's interpolated percentile computation used for P95 runtime.",
        filename="02_p95_percentile_formula.png",
    ),
    SourceShot(
        title="Per-Run Validation Metrics",
        source=ROOT / "scripts" / "benchmark_planner_profiles.py",
        start=445,
        end=472,
        caption=(
            "Shows that each run records slot count, solver status, nutrition feasibility, "
            "constraint validation, and hard-rule violation count."
        ),
        filename="03_per_run_validation_metrics.png",
    ),
    SourceShot(
        title="Final Suite Summary Formulas",
        source=ROOT / "scripts" / "benchmark_planner_profiles.py",
        start=792,
        end=800,
        caption=(
            "Shows how total runs, passed runs, average runtime, P95 runtime, and maximum "
            "runtime are computed for the benchmark suite."
        ),
        filename="04_suite_summary_formulas.png",
    ),
    SourceShot(
        title="ML-Ready Benchmark Guard",
        source=ROOT / "scripts" / "benchmark_planner_profiles.py",
        start=724,
        end=730,
        caption="Shows that --require-ml-ready fails the benchmark if the LightGBM ranker is not ready.",
        filename="05_ml_ready_guard.png",
    ),
    SourceShot(
        title="Archived Final Benchmark Result",
        source=ROOT / "benchmarks" / "reports" / "planner_realistic_profiles.local.lightgbm.final.json",
        start=12951,
        end=12975,
        caption=(
            "Shows the archived final result: 100 total runs, 100 passed, 5 runs per case, "
            "257.13 ms average, 313.05 ms P95, 357 ms max, and LightGBM ready."
        ),
        filename="06_archived_final_benchmark_result.png",
    ),
    SourceShot(
        title="PHP 20 No-Safe-Plan Boundary Test",
        source=ROOT / "backend" / "tests" / "test_meal_planner.py",
        start=2488,
        end=2532,
        caption=(
            "Shows the separate automated budget-infeasibility test where weeklyBudgetPhp=20 "
            "must return no plan instead of an unsafe or impossible recommendation."
        ),
        filename="07_php20_no_safe_plan_test.png",
    ),
]


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts/consolab.ttf" if bold else "C:/Windows/Fonts/consola.ttf"),
        Path("C:/Windows/Fonts/lucon.ttf"),
        Path("C:/Windows/Fonts/cour.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def _read_lines(path: Path, start: int, end: int) -> list[tuple[int, str]]:
    raw = path.read_text(encoding="utf-8", errors="replace").splitlines()
    if start < 1 or end > len(raw) or start > end:
        raise ValueError(f"Invalid line range {start}-{end} for {path}")
    return [(idx, raw[idx - 1]) for idx in range(start, end + 1)]


def _wrap_code_line(line: str, width: int) -> list[str]:
    if len(line) <= width:
        return [line]
    chunks = []
    remaining = line
    while len(remaining) > width:
        cut = remaining.rfind(" ", 0, width)
        if cut < max(24, width // 2):
            cut = width
        chunks.append(remaining[:cut])
        remaining = "    " + remaining[cut:].lstrip()
    if remaining:
        chunks.append(remaining)
    return chunks


def render_source_shot(shot: SourceShot) -> Path:
    rel = shot.source.relative_to(ROOT).as_posix()
    lines = _read_lines(shot.source, shot.start, shot.end)
    body_font = _font(18)
    title_font = _font(22, bold=True)
    meta_font = _font(14)
    line_height = 25
    max_code_chars = 118
    gutter_width = 70
    left = 24
    top = 24
    header_height = 86
    body_rows: list[tuple[int | None, str]] = []
    for number, text in lines:
        wrapped = _wrap_code_line(text.replace("\t", "    "), max_code_chars)
        body_rows.append((number, wrapped[0]))
        for continuation in wrapped[1:]:
            body_rows.append((None, continuation))
    width = 1700
    height = top + header_height + (len(body_rows) * line_height) + 32
    image = Image.new("RGB", (width, height), "#f7f7f8")
    draw = ImageDraw.Draw(image)
    draw.rectangle([0, 0, width, 78], fill="#1f2937")
    draw.text((left, 16), shot.title, font=title_font, fill="#ffffff")
    draw.text((left, 52), f"{rel}:{shot.start}-{shot.end}", font=meta_font, fill="#d1d5db")
    draw.rectangle([0, 78, width, height], fill="#ffffff")
    y = top + header_height
    for row_idx, (line_number, text) in enumerate(body_rows):
        if row_idx % 2 == 0:
            draw.rectangle([0, y - 2, width, y + line_height - 2], fill="#fbfbfc")
        if line_number is not None:
            draw.text((left, y), str(line_number).rjust(4), font=body_font, fill="#6b7280")
        draw.text((left + gutter_width, y), text, font=body_font, fill="#111827")
        y += line_height
    out = OUT_DIR / shot.filename
    image.save(out, "PNG")
    return out


def _build_fixture_summary_image() -> Path:
    path = ROOT / "benchmarks" / "canonical_scenarios" / "planner_realistic_profiles_20.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    cases = data.get("cases", [])
    rows = [("ID", "Case label", "Goal", "Budget / restrictions / allergies")]
    for case in cases:
        profile = case.get("profile", {})
        restrictions = ", ".join(profile.get("dietaryRestrictions") or []) or "none"
        allergies = ", ".join(profile.get("allergies") or []) or "none"
        budget = profile.get("weeklyBudgetPhp") or profile.get("budgetWeekly") or "none"
        detail = f"budget={budget}; restrictions={restrictions}; allergies={allergies}"
        rows.append((
            str(case.get("id", "")),
            str(case.get("userType", "")),
            str(profile.get("goal", "")),
            detail,
        ))
    title_font = _font(22, bold=True)
    body_font = _font(16)
    meta_font = _font(14)
    widths = [80, 340, 260, 850]
    row_h = 31
    width = sum(widths) + 60
    height = 120 + len(rows) * row_h + 28
    image = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.rectangle([0, 0, width, 82], fill="#1f2937")
    draw.text((24, 16), "T01-T20 Fixture Summary", font=title_font, fill="#ffffff")
    draw.text((24, 52), path.relative_to(ROOT).as_posix(), font=meta_font, fill="#d1d5db")
    y = 104
    x0 = 24
    for idx, row in enumerate(rows):
        fill = "#e5e7eb" if idx == 0 else ("#fbfbfc" if idx % 2 == 0 else "#ffffff")
        draw.rectangle([x0, y - 4, x0 + sum(widths), y + row_h - 4], fill=fill)
        x = x0
        for col, value in enumerate(row):
            draw.text((x + 8, y), str(value)[:120], font=body_font, fill="#111827")
            x += widths[col]
        y += row_h
    out = OUT_DIR / "08_t01_t20_fixture_summary.png"
    image.save(out, "PNG")
    return out


def _add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9)


def _add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.italic = True


def build_doc(image_paths: Iterable[tuple[str, Path, str]]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PCOSINA Benchmark Code Screenshot Evidence")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(f"Generated {TODAY} from repository source files")
    sub_run.italic = True
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(9)

    _add_body(
        doc,
        "Purpose: This pack provides printable code-level evidence for how the T01-T20 planner benchmark was defined, "
        "executed, measured, summarized, and validated. It does not invent benchmark values; screenshots are rendered "
        "from the exact local source files and archived report.",
    )
    _add_body(
        doc,
        "Defense boundary: The final 100-run benchmark is archived as machine-readable JSON/CSV. The earliest two "
        "Table 32 speed rows are documented development benchmark values, but their raw JSON/CSV artifacts were not retained.",
    )

    for idx, (title_text, image_path, caption) in enumerate(image_paths, start=1):
        doc.add_heading(f"{idx}. {title_text}", level=1)
        doc.add_picture(str(image_path), width=Inches(7.35))
        _add_caption(doc, caption)

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    image_entries: list[tuple[str, Path, str]] = []
    for shot in SHOTS:
        image_entries.append((shot.title, render_source_shot(shot), shot.caption))
    fixture_summary = _build_fixture_summary_image()
    image_entries.append((
        "T01-T20 Fixture Summary",
        fixture_summary,
        "Shows the 20 fixed benchmark cases read from the canonical fixture.",
    ))
    build_doc(image_entries)

    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(DOCX_PATH, DOCX_PATH.name)
        for _, image_path, _ in image_entries:
            archive.write(image_path, f"screenshots/{image_path.name}")

    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOCX_PATH, DOWNLOADS / DOCX_PATH.name)
    shutil.copy2(ZIP_PATH, DOWNLOADS / ZIP_PATH.name)

    print(DOCX_PATH)
    print(ZIP_PATH)
    print(DOWNLOADS / DOCX_PATH.name)
    print(DOWNLOADS / ZIP_PATH.name)


if __name__ == "__main__":
    main()
