from __future__ import annotations

import hashlib
import json
import subprocess
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
HISTORICAL_FIXTURE_REF = (
    "997534e:benchmarks/canonical_scenarios/planner_realistic_profiles_20.json"
)
RESULT_PATH = (
    ROOT
    / "benchmarks"
    / "reports"
    / "planner_realistic_profiles.local.lightgbm.final.json"
)
OUTPUT_PATH = (
    ROOT
    / "docs"
    / "defense"
    / "PCOSINA_T01_T20_Actual_Planner_Benchmark_Appendix_2026-06-10.docx"
)


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def load_historical_fixture() -> tuple[dict[str, Any], bytes]:
    raw = subprocess.check_output(
        ["git", "show", HISTORICAL_FIXTURE_REF],
        cwd=ROOT,
    )
    return json.loads(raw.decode("utf-8")), raw


def sha256(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 7.5,
    color: str | None = None,
) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 7.2,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_text(cell, label, bold=True, size=font_size, color="FFFFFF")
        shade_cell(cell, "1F4E78")
        cell.width = Inches(widths[index])

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            set_cell_text(cells[column_index], value, size=font_size)
            cells[column_index].width = Inches(widths[column_index])
            if row_index % 2:
                shade_cell(cells[column_index], "EAF2F8")
    doc.add_paragraph()


def join_values(values: list[Any] | None) -> str:
    if not values:
        return "None"
    return ", ".join(str(value) for value in values)


def format_budget(value: Any) -> str:
    amount = float(value or 0)
    return "No configured budget" if amount <= 0 else f"PHP {amount:,.0f}"


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        prefix, remainder = text[: len(bold_prefix)], text[len(bold_prefix) :]
        prefix_run = paragraph.add_run(prefix)
        prefix_run.bold = True
        prefix_run.font.name = "Arial"
        prefix_run.font.size = Pt(9)
        remainder_run = paragraph.add_run(remainder)
        remainder_run.font.name = "Arial"
        remainder_run.font.size = Pt(9)
    else:
        run = paragraph.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)


def build_case_rows(
    fixture: dict[str, Any],
    report: dict[str, Any],
) -> tuple[list[list[str]], list[list[str]]]:
    cases = fixture["cases"]
    summaries = {item["caseId"]: item for item in report["caseSummaries"]}
    runs_by_case: dict[str, list[dict[str, Any]]] = {}
    for run in report["runs"]:
        runs_by_case.setdefault(run["caseId"], []).append(run)

    definition_rows: list[list[str]] = []
    result_rows: list[list[str]] = []

    if len(cases) != 20 or len(summaries) != 20:
        raise ValueError("Expected exactly 20 fixture cases and 20 case summaries.")

    for case in cases:
        case_id = case["id"]
        profile = case["profile"]
        summary = summaries[case_id]
        runs = sorted(runs_by_case[case_id], key=lambda item: item["runIndex"])
        if summary["userType"] != case["userType"]:
            raise ValueError(
                f"{case_id} label mismatch: fixture={case['userType']!r}, "
                f"report={summary['userType']!r}"
            )
        if len(runs) != 5:
            raise ValueError(f"{case_id} expected 5 runs, found {len(runs)}.")

        hard_violations = sum(
            int(run["hardConstraintViolationCount"]) for run in runs
        )
        advisory_warnings = sum(
            int(run["advisoryConstraintWarningCount"]) for run in runs
        )
        costs = [int(run["estimatedWeeklyCostPhp"]) for run in runs]
        if not all(run["slotCount"] == 21 for run in runs):
            raise ValueError(f"{case_id} contains a non-21-slot run.")
        if not all(run["constraintValidationOk"] for run in runs):
            raise ValueError(f"{case_id} contains a failed constraint validation.")
        if not all(run["nutritionFeasibilityOk"] for run in runs):
            raise ValueError(f"{case_id} contains a nutrition-infeasible run.")

        definition_rows.append(
            [
                case_id,
                case["userType"],
                str(profile.get("goal", "")),
                (
                    f"Age {profile.get('age')}; "
                    f"{profile.get('heightCm')} cm; {profile.get('weightKg')} kg; "
                    f"{profile.get('activityLevel')}"
                ),
                str(profile.get("insulinResistanceLevel", "Not recorded")),
                join_values(profile.get("dietaryRestrictions")),
                join_values(profile.get("allergies")),
                format_budget(profile.get("weeklyBudgetPhp")),
                (
                    f"{profile.get('maxCookingTimeMinutes')} min; "
                    f"{profile.get('varietyPreference')}; "
                    f"{profile.get('planningPriority')}; "
                    f"household {profile.get('householdSize', 'not recorded')}; "
                    f"pantry: {join_values(profile.get('pantryItems'))}"
                ),
            ]
        )

        result_rows.append(
            [
                case_id,
                summary["userType"],
                f"{summary['passed']}/{summary['runs']}",
                "21 in every run",
                str(summary["terminalSolverStatuses"][0]),
                f"{summary['avgRuntimeMs']:.1f}",
                f"{summary['p95RuntimeMs']:.1f}",
                f"{summary['minRuntimeMs']:.0f}-{summary['maxRuntimeMs']:.0f}",
                f"PHP {min(costs):,}-{max(costs):,}",
                "5/5",
                str(hard_violations),
                str(advisory_warnings),
            ]
        )

    return definition_rows, result_rows


def build_document() -> None:
    fixture, fixture_raw = load_historical_fixture()
    report_raw = RESULT_PATH.read_bytes()
    report = json.loads(report_raw.decode("utf-8"))
    definition_rows, result_rows = build_case_rows(fixture, report)
    suite = report["suite"]

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(
        "Appendix: Actual T01-T20 Planner Benchmark Cases and Results"
    )
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        "Evidence-faithful reconstruction of the production-shaped local benchmark "
        "started May 19, 2026"
    )
    subtitle_run.italic = True
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(9)

    add_body(
        doc,
        "Purpose: This appendix exposes the complete T01-T20 case set behind the "
        "manuscript's 20-profile, 100-run summary. It reports archived measurements; "
        "it does not claim that every device, server, profile, or future dataset will "
        "produce the same runtime or selected meals.",
        bold_prefix="Purpose:",
    )
    add_body(
        doc,
        "Source integrity: Inputs were recovered from Git commit 997534e because its "
        "case labels exactly match the archived result report. This avoids combining "
        "the May 19 measurements with the subsequently revised fixture labels.",
        bold_prefix="Source integrity:",
    )

    doc.add_heading("Suite-Level Result", level=1)
    suite_rows = [
        ["Benchmark mode", f"{suite['mode']} / {suite['environment']}"],
        ["Cases and repetitions", f"{suite['caseCount']} cases x {suite['runsPerCase']} runs"],
        ["Total result", f"{suite['passedRuns']}/{suite['totalRuns']} runs passed"],
        ["Runtime", (
            f"Average {suite['avgRuntimeMs']:.2f} ms; "
            f"P95 {suite['p95RuntimeMs']:.2f} ms; "
            f"minimum {suite['minRuntimeMs']:.0f} ms; "
            f"maximum {suite['maxRuntimeMs']:.0f} ms"
        )],
        ["Planner evidence", (
            f"LightGBM ready: {suite['mlRankerReady']}; "
            f"model: {suite['mlModelVersion']}; recipes: {suite['recipeCount']}"
        )],
        ["Validation outcome", (
            "100/100 runs had 21 slots, nutrition feasibility passed, "
            "constraint validation passed, and zero recorded hard-rule violations."
        )],
    ]
    add_table(doc, ["Measure", "Archived result"], suite_rows, [2.0, 8.0], font_size=8)

    doc.add_heading("Table A. Historical T01-T20 Benchmark Input Definitions", level=1)
    add_body(
        doc,
        "These are the inputs associated with the archived result labels. The "
        "insulin-resistance field appears because it existed in this historical "
        "benchmark version; its presence should not be interpreted as a current "
        "product or clinical claim.",
    )
    add_table(
        doc,
        [
            "ID",
            "Case label",
            "Goal",
            "Anthropometric / activity profile",
            "Historical insulin setting",
            "Dietary restrictions",
            "Allergies",
            "Weekly budget",
            "Practical settings",
        ],
        definition_rows,
        [0.4, 1.15, 1.0, 1.45, 0.9, 1.0, 0.85, 0.85, 2.4],
        font_size=6.7,
    )

    doc.add_heading("Table B. Actual Five-Run Results for T01-T20", level=1)
    add_table(
        doc,
        [
            "ID",
            "Archived case label",
            "Passed",
            "Meal slots",
            "Solver",
            "Avg ms",
            "P95 ms",
            "Min-max ms",
            "Estimated weekly cost range",
            "Nutrition feasible",
            "Hard violations",
            "Advisory warnings",
        ],
        result_rows,
        [0.4, 1.35, 0.55, 0.75, 0.65, 0.6, 0.6, 0.75, 1.25, 0.8, 0.75, 0.8],
        font_size=6.8,
    )

    doc.add_heading("Interpretation and Limits", level=1)
    for text in [
        "A pass means the archived benchmark runner recorded success, 21 meal slots, "
        "nutritionFeasibilityOk=true, constraintValidationOk=true, and no hard-rule "
        "violation for that run.",
        "FEASIBLE means CP-SAT found a valid solution under the tested settings. It "
        "does not necessarily mean that global mathematical optimality was proven.",
        "Advisory warnings are non-hard sodium or sugar quality warnings. They do not "
        "represent allergy, dietary-restriction, slot-validity, or other hard-rule failures.",
        "Estimated weekly cost varied between repeated runs because the solver could "
        "return different feasible meal combinations. The amounts are planner estimates, "
        "not guaranteed market prices.",
        "The impossible PHP 20 no-safe-plan boundary test is separate from T01-T20 and "
        "is therefore not inserted as a twenty-first profile.",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(8.5)

    doc.add_heading("Evidence Provenance", level=1)
    provenance_rows = [
        [
            "Historical fixture",
            HISTORICAL_FIXTURE_REF,
            sha256(fixture_raw),
        ],
        [
            "Archived result report",
            str(RESULT_PATH.relative_to(ROOT)).replace("\\", "/"),
            sha256(report_raw),
        ],
        [
            "Benchmark start",
            suite["startedAt"],
            "Timestamp stored inside the archived report",
        ],
        [
            "Generated appendix",
            str(OUTPUT_PATH.relative_to(ROOT)).replace("\\", "/"),
            "Generated June 10, 2026; values validated programmatically",
        ],
    ]
    add_table(
        doc,
        ["Artifact", "Repository reference", "Integrity / note"],
        provenance_rows,
        [1.4, 4.3, 4.3],
        font_size=7.2,
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
