"""Compute consumed-portion prices for the PhilFCT portioned catalog."""

from __future__ import annotations

import csv
import json
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = REPO_ROOT / "backend"
PORTIONED = BACKEND_ROOT / "seed_data" / "pcosina_philfct_portioned_catalog_v2_draft.json"
PRICE_RULES = BACKEND_ROOT / "seed_data" / "reviewed_market_price_rules.csv"
OUT_JSON = BACKEND_ROOT / "seed_data" / "pcosina_philfct_portioned_priced_catalog_v2_draft.json"
OUT_CSV = REPO_ROOT / "docs" / "thesis_validation" / "03_ACTUAL_SYSTEM_DATA_EXPORTS" / "pcosina_philfct_portioned_pricing_review_queue.csv"
OUT_XLSX = OUT_CSV.with_suffix(".xlsx")


LIQUID_HINTS = {
    "oil", "cooking oil", "corn oil", "canola oil", "vinegar", "soy sauce",
    "fish sauce", "patis", "coconut milk", "coconut cream", "milk", "broth",
    "sauce", "juice",
}


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9\s]+", " ", str(text or "").lower())).strip()


def load_price_rules() -> list[dict[str, Any]]:
    rows = []
    with PRICE_RULES.open("r", encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            if str(row.get("active") or "").lower() != "true":
                continue
            try:
                price = float(row.get("price_php") or 0)
            except Exception:
                continue
            if price <= 0 and str(row.get("zero_price") or "").lower() != "true":
                continue
            keywords = [norm(part) for part in str(row.get("keywords") or "").split(",") if norm(part)]
            if not keywords:
                keywords = [norm(row.get("id") or "")]
            rows.append({
                **row,
                "_keywords": keywords,
                "_keyword_tokens": [set(keyword.split()) for keyword in keywords],
                "price_php_float": price,
                "unit_norm": norm(row.get("unit") or "kg") or "kg",
            })
    return rows


def best_rule(name: str, rules: list[dict[str, Any]], cache: dict[str, dict[str, Any] | None]) -> dict[str, Any] | None:
    query = norm(name)
    if not query:
        return None
    if query in cache:
        return cache[query]
    qtokens = set(query.split())
    best = None
    best_score = 0.0
    for rule in rules:
        for keyword, ktokens in zip(rule["_keywords"], rule["_keyword_tokens"]):
            if not (qtokens & ktokens):
                continue
            token_score = len(qtokens & ktokens) / max(len(ktokens), 1)
            seq_score = SequenceMatcher(None, query, keyword).ratio()
            score = (0.75 * token_score) + (0.25 * seq_score)
            if score > best_score:
                best = rule
                best_score = score
    if best is None or best_score < 0.35:
        cache[query] = None
        return None
    result = {**best, "match_score": round(best_score, 3)}
    cache[query] = result
    return result


def consumed_cost(portion_g: float | None, rule: dict[str, Any] | None, name: str) -> tuple[float | None, str]:
    if portion_g is None or portion_g <= 0:
        return None, "no_portion"
    if not rule:
        return None, "no_price_rule"
    unit = str(rule.get("unit_norm") or "kg").lower()
    price = float(rule.get("price_php_float") or 0.0)
    lname = norm(name)
    adjusted_portion_g = float(portion_g)
    if "rice" in lname and any(token in lname for token in {"cooked", "boiled", "sinaing"}):
        adjusted_portion_g *= 0.34
    if any(token in lname for token in {"munggo", "monggo", "mung bean"}) and any(token in lname for token in {"cooked", "boiled"}):
        adjusted_portion_g *= 0.40
    if str(rule.get("zero_price") or "").lower() == "true":
        return 0.0, "zero_price_rule"
    if unit == "kg":
        return round((adjusted_portion_g / 1000.0) * price, 2), "portion_g_to_kg"
    if unit == "l":
        # For cooking liquids, the portion parser stores grams as a practical
        # mL-equivalent estimate. This is acceptable for pricing drafts.
        return round((adjusted_portion_g / 1000.0) * price, 2), "portion_g_to_liter_estimate"
    if unit == "piece":
        # Use 100g as conservative default piece weight when no exact market
        # piece conversion is available. Keep this reviewable.
        return round((adjusted_portion_g / 100.0) * price, 2), "portion_g_to_piece_estimate"
    return round((adjusted_portion_g / 1000.0) * price, 2), f"fallback_unit_{unit}"


def export_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    wb = Workbook()
    ws = wb.active
    ws.title = "Portion Pricing"
    for row in rows:
        ws.append(row)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    numeric = {
        "portion_g", "price_php", "price_min_php", "price_max_php", "price_match_score",
        "cost_php", "meal_total_cost_php",
    }
    headers = [cell.value for cell in ws[1]]
    for name in numeric:
        if name in headers:
            idx = headers.index(name) + 1
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row_idx, idx)
                try:
                    if cell.value not in (None, ""):
                        cell.value = float(cell.value)
                        cell.number_format = "0.00"
                except Exception:
                    pass
    for idx, col in enumerate(ws.columns, start=1):
        max_len = max(min(len(str(cell.value or "")), 70) for cell in col)
        ws.column_dimensions[get_column_letter(idx)].width = max(12, min(max_len + 2, 70))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx_path)
    load_workbook(xlsx_path, read_only=True).close()


def main() -> int:
    catalog = json.loads(PORTIONED.read_text(encoding="utf-8"))
    rules = load_price_rules()
    cache: dict[str, dict[str, Any] | None] = {}
    rows = []
    priced_meals = []
    matched = 0
    priced = 0
    total_ingredients = 0
    for meal in catalog["meals"]:
        meal_cost = 0.0
        priced_ingredients = []
        for ing in meal.get("portionIngredients") or []:
            total_ingredients += 1
            if ing.get("parse_status") == "ignored_non_nutritive":
                priced_ingredients.append({**ing, "pricing": {"cost_php": 0.0, "status": "ignored_non_nutritive"}})
                continue
            name = ing.get("clean_name") or ing.get("source_text") or ""
            rule = best_rule(name, rules, cache)
            if rule:
                matched += 1
            cost, status = consumed_cost(ing.get("portion_g"), rule, name)
            if cost is not None:
                priced += 1
                meal_cost += cost
            pricing = {
                "cost_php": cost,
                "status": status,
                "price_rule_id": rule.get("id") if rule else None,
                "price_keywords": rule.get("keywords") if rule else None,
                "price_php": rule.get("price_php") if rule else None,
                "market_unit": rule.get("unit") if rule else None,
                "source": rule.get("source") if rule else None,
                "market_source": rule.get("market_source") if rule else None,
                "source_url": rule.get("source_url") if rule else None,
                "confidence": rule.get("confidence") if rule else None,
                "match_score": rule.get("match_score") if rule else None,
            }
            priced_ingredients.append({**ing, "pricing": pricing})
            rows.append({
                "recipe_id": meal.get("sourceRecipeId"),
                "recipe_name": meal.get("sourceRecipeName"),
                "source_text": ing.get("source_text"),
                "clean_name": name,
                "portion_g": ing.get("portion_g"),
                "parse_status": ing.get("parse_status"),
                "philfct_name": (ing.get("philfct") or {}).get("philfct_name"),
                "price_rule_id": pricing["price_rule_id"],
                "price_keywords": pricing["price_keywords"],
                "market_unit": pricing["market_unit"],
                "price_php": pricing["price_php"],
                "cost_php": pricing["cost_php"],
                "pricing_status": pricing["status"],
                "price_confidence": pricing["confidence"],
                "price_match_score": pricing["match_score"],
                "market_source": pricing["market_source"],
                "source_url": pricing["source_url"],
            })
        priced_meals.append({
            **meal,
            "portionIngredients": priced_ingredients,
            "portionPricing": {
                "estimatedMealCostPhp": round(meal_cost, 2),
                "pricedIngredientCount": sum(1 for item in priced_ingredients if (item.get("pricing") or {}).get("cost_php") is not None),
                "source": "reviewed_market_price_rules.csv",
                "formula": "kg/l cost = portion_g / 1000 * market_unit_price; piece cost uses portion_g / 100 * price_php and remains reviewable",
            },
        })
    payload = {
        **catalog,
        "version": "pcosina-philfct-portioned-priced-catalog-v2-draft",
        "pricingSummary": {
            "priceRuleCount": len(rules),
            "sourceIngredientRows": total_ingredients,
            "priceMatchedRows": matched,
            "pricedRows": priced,
            "uniquePriceQueries": len(cache),
        },
        "meals": priced_meals,
    }
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    export_xlsx(OUT_CSV, OUT_XLSX)
    print(json.dumps(payload["pricingSummary"], indent=2))
    print(OUT_JSON)
    print(OUT_XLSX)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
