from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "defense" / "PCOSINA_System_Architecture_Defense_Prep_2026-06-04.docx"
SLIDE_IMAGE = ROOT / "tmp" / "defense_prep" / "presentation_pages_png" / "page_4.png"
ASSET_DIR = ROOT / "docs" / "defense" / "PCOSINA_SOP_Evidence_Extraction_Dossier_FINAL_2026-05-24.assets"
ARCH_FIGURE = ASSET_DIR / "figure_08.png"
OFFLINE_FLOW_FIGURE = ASSET_DIR / "figure_09.png"
FEATURE_MAP_FIGURE = ASSET_DIR / "figure_10.png"


ACCENT = "2F7D5A"
DARK = "1F2933"
MUTED = "5B6770"
PALE = "EAF5EF"
PALE_BLUE = "EAF2F8"
PALE_PINK = "FDECEF"
BORDER = "B7C6BC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=DARK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05


def add_table(doc, headers, rows, widths=None, header_fill=ACCENT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8.8)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = widths[idx]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_text(cell, str(value), color=DARK, size=8.5)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[idx]
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Aptos Display"
        run.font.color.rgb = RGBColor.from_string(ACCENT if level == 1 else DARK)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(DARK)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.color.rgb = RGBColor.from_string(DARK)
    else:
        r = p.add_run(text)
        r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(item)
        run.font.color.rgb = RGBColor.from_string(DARK)


def add_numbered(doc, items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(f"{idx}.    {item}")
        run.font.color.rgb = RGBColor.from_string(DARK)


def add_callout(doc, title, body, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=ACCENT)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    cell.text = ""
    p = cell.paragraphs[0]
    rt = p.add_run(title)
    rt.bold = True
    rt.font.color.rgb = RGBColor.from_string(ACCENT)
    rt.font.size = Pt(10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    rb = p2.add_run(body)
    rb.font.color.rgb = RGBColor.from_string(DARK)
    rb.font.size = Pt(9)
    doc.add_paragraph()


def add_figure(doc, path, caption, width=7.1):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(8)
        cap.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    else:
        add_callout(doc, "Figure not embedded", f"Expected image was not found at {path}.", fill=PALE_PINK)


def configure_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)

    for style_name, size, color in [
        ("Title", 24, ACCENT),
        ("Heading 1", 15, ACCENT),
        ("Heading 2", 12, DARK),
        ("Heading 3", 10.5, DARK),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("PCOSina System Architecture Defense Prep")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(ACCENT)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Focused packet for thesis defense: slide guide, manuscript figures, architecture zones, script, evidence, boundaries, and Q&A")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "Defense identity",
        "The system architecture explains how PCOSina separates the Android mobile client, local persistence, secure backend control plane, remote optimizer, data sources, optional sync/deployment services, and evaluation loop while preserving local-first continuity and deterministic planning authority. The planning authority is a MILP-formulated 0-1 weekly assignment model solved in implementation through OR-Tools CP-SAT.",
        fill=PALE,
    )

    add_table(
        doc,
        ["Item", "Defense-ready answer"],
        [
            ["Main point", "PCOSina is an offline-first modular hybrid system: saved user data remains available locally, while new optimized weekly plan generation is handled by a backend optimizer with deterministic filtering and MILP-formulated CP-SAT planning authority."],
            ["Presentation slide", "Slide 4: System Architecture from [Presentation] Quadrant_Presentation_Revised_S2526T3 (1).pdf."],
            ["Manuscript anchor", "Chapter 3, System Architecture section: extracted manuscript paragraphs 588-615, Figure 8, and adjacent offline-first Figures 9-10."],
            ["System anchor", "ARCHITECTURE.md, production_contract.md, optimizer_queue_worker.md, planner_contract_2026-05-26.md, ADR-004, ADR-005, and render.yaml."],
            ["MILP title defense", "MILP-based refers to the planner formulation: binary recipe-slot variables, linear constraints, and objective terms. CP-SAT is the implementation solver, not a claim that CP-SAT itself is a pure MILP solver."],
            ["Boundary", "Do not describe PCOSina as a pure offline solver or a medical system. Saved data works offline; new optimized plan generation may need the backend."],
        ],
        [Inches(1.55), Inches(5.75)],
    )

    doc.add_page_break()
    add_heading(doc, "1. Slide And Manuscript Figures", 1)
    add_figure(doc, SLIDE_IMAGE, "Figure SA-1. Presentation slide used for the System Architecture defense section.", width=7.1)
    doc.add_page_break()
    add_figure(doc, ARCH_FIGURE, "Figure SA-2. Manuscript Figure 8. PCOSina interactive system architecture map.", width=7.1)
    add_figure(doc, OFFLINE_FLOW_FIGURE, "Figure SA-3. Manuscript Figure 9. Offline-first planning workflow.", width=6.5)
    doc.add_page_break()
    add_figure(doc, FEATURE_MAP_FIGURE, "Figure SA-4. Manuscript Figure 10. Offline and online feature availability map.", width=3.85)

    doc.add_page_break()
    add_heading(doc, "2. Core Defense Explanation", 1)
    add_callout(
        doc,
        "30-second answer",
        "Our system architecture is a local-first hybrid architecture. The Android client handles user-facing flows such as onboarding, profile, pantry, grocery, meal plan, progress, and saved offline access. Local persistence uses DataStore, encrypted shared preferences, and local artifacts. When a new optimized plan is requested, the app sends a validated request to the FastAPI backend. The backend protects the request with schema checks, Firebase/App Check support, readiness checks, and an async job path. The meal-planning brain performs Stage 1 deterministic filtering and optional non-authoritative LightGBM ranking, then Stage 2 MILP-formulated 0-1 weekly assignment solved with OR-Tools CP-SAT as the final authority. Data sources include recipe/nutrition data, price/pantry references, Postgres production persistence, Redis queue support, and optional Firebase sync. The evaluation loop feeds testing and expert/user findings back into UI, rules, constraints, dataset, and deployment improvements.",
        fill=PALE_BLUE,
    )

    add_heading(doc, "3. Architecture Zones", 1)
    add_table(
        doc,
        ["Zone", "Purpose", "Main components", "Defense sentence"],
        [
            [
                "Zone A: User Side",
                "Everything the user sees and controls on the phone.",
                "Onboarding, profile setup, symptoms/goals, allergies, pantry, budget, meal plan, grocery, progress, settings, and saved local records.",
                "This zone keeps the app usable and lets users open saved profiles, plans, groceries, pantry records, and progress data offline.",
            ],
            [
                "Zone B: Secure System Side",
                "Protects backend resources before planning begins.",
                "FastAPI request checks, Firebase App Check/Auth support, schema version checks, production safety settings, planning job manager, and async job status.",
                "This zone makes sure only valid, supported planning requests reach the optimizer.",
            ],
            [
                "Zone C: Meal Planning Brain",
                "Chooses the weekly plan under safety and optimization rules.",
                "Stage 1 safety filter, optional smart sorting, LightGBM ranking support, Stage 2 MILP-formulated 0-1 CP-SAT optimizer, and no-safe-plan guide.",
                "This is the decision authority: filtering protects safety, the MILP formulation defines the assignment problem, CP-SAT selects the final plan, and ML only assists ranking.",
            ],
            [
                "Zone D: Data Sources",
                "Supplies recipe, nutrition, price, policy, and operational records.",
                "Recipe/nutrition data, ingredient tokens, nutrition tags, price and pantry data, Postgres backend records, Redis queue, and optional cloud sync.",
                "Data sources feed the planner but do not replace deterministic safety checks.",
            ],
            [
                "Zone E: Output and Improvement",
                "Shows results and feeds evaluation back into system refinement.",
                "User result, recipe and grocery guidance, nutrition/budget summary, evaluation, findings, app design, safety rules, prices, reliability, and logs.",
                "This zone proves the system is evaluated and improved, not just drawn as a static architecture.",
            ],
        ],
        [Inches(1.25), Inches(1.75), Inches(2.55), Inches(1.75)],
    )

    add_heading(doc, "4. Request Flow To Explain", 1)
    add_numbered(
        doc,
        [
            "User opens PCOSina. The phone loads saved local data: profile, pantry, previous plans, grocery guidance, progress, and reflections.",
            "If the user only views saved information, the user can continue offline.",
            "If the user requests a new optimized weekly plan, the app sends a planning request to the backend when connectivity is available.",
            "The FastAPI backend validates the request, checks schema version, applies security/readiness controls, and may queue the job for the planning worker.",
            "Stage 1 constructs a safe candidate pool by filtering unsafe or incompatible recipes and optionally ranking safe candidates.",
            "Stage 2 builds the MILP-formulated 0-1 weekly assignment and uses OR-Tools CP-SAT to select the final plan across the week.",
            "If a complete safe plan is found, the backend returns it to the phone and the app saves it locally for later offline access.",
            "If no safe plan is possible, the backend returns no-safe-plan reasons and safe adjustment guidance instead of forcing an unsafe result.",
            "Evaluation findings, logs, and user/expert feedback feed back into UI, rules, constraints, dataset, and deployment improvements.",
        ],
    )

    add_heading(doc, "5. Full Speaker Script", 1)
    add_body(
        doc,
        "Good day panel. This slide presents the system architecture of PCOSina. The most important idea is that PCOSina is an offline-first, modular hybrid system. It is not a single mobile screen doing everything, and it is not a fully cloud-dependent app. The architecture separates the user-facing Android client, local persistence, secure backend validation, remote optimization, data sources, optional synchronization, and evaluation feedback."
    )
    add_body(
        doc,
        "On the user side, the Android application handles onboarding, profile setup, pantry records, grocery guidance, meal-plan viewing, progress records, settings, and saved information. The mobile side uses local persistence for continuity, specifically DataStore-backed preferences, encrypted shared preferences for sensitive local reflection content, and local artifacts for saved outputs. This is why users can still open saved profiles, previous meal plans, grocery lists, pantry data, and progress records even in low-signal or offline conditions."
    )
    add_body(
        doc,
        "When the user needs a new optimized weekly plan, the architecture uses the backend. The FastAPI backend acts as the secure control plane. It validates incoming requests, checks schema versions, supports Firebase Authentication and App Check verification when enabled, applies readiness and safety checks, and can route planning work through an async job manager and worker path."
    )
    add_body(
        doc,
        "The meal-planning brain is separated from the mobile UI. Stage 1 performs deterministic safety filtering and candidate construction. It removes recipes that violate allergies, dietary restrictions, explicit exclusions, meal-slot suitability, and basic feasibility. LightGBM may help rank already-safe candidates, but it is optional and non-authoritative. Stage 2 uses a MILP-formulated 0-1 assignment model solved through OR-Tools CP-SAT to assign recipes to the weekly meal slots. That stage is the final authority for the generated meal plan."
    )
    add_body(
        doc,
        "The architecture also includes data-source and deployment components. Recipe and nutrition data, ingredient tokens, price references, pantry-related signals, backend records, policy settings, and job records support planning. In production, the backend persistence target is PostgreSQL, while Redis supports the queue path and Render hosts the backend and worker services. Optional Firebase sync can support account recovery or continuity, but it is supplementary and does not replace local continuity."
    )
    add_body(
        doc,
        "The output and improvement zone closes the loop. The user receives the generated plan, recipe information, grocery guidance, and nutrition or budget explanation. The system can also produce metadata, diagnostics, no-safe-plan reasons, and evidence used for evaluation. Findings from testing, ISO/IEC 25010 evaluation, CS/IT review, OB-GYN and RND review, and user feedback are used to refine the UI, rules, constraints, dataset, reliability, and deployment."
    )
    add_body(
        doc,
        "So the architecture supports the study goal in three ways: it protects offline access to saved user data, it keeps optimization authority separate from the UI, and it preserves deterministic safety boundaries even when ML or cloud services are available.",
    )

    add_heading(doc, "6. Repo-Truth Cross-Check", 1)
    add_table(
        doc,
        ["Architecture claim", "Current source evidence", "Defense-safe wording"],
        [
            [
                "Mobile local persistence",
                "ARCHITECTURE.md states Android currently uses DataStore-backed preferences, encrypted shared preferences, and local artifacts. It does not use Room or SQLite as the mobile database.",
                "Say local persistence layer, DataStore, encrypted shared preferences, and local artifacts. Do not say Room/SQLite for Android.",
            ],
            [
                "Backend authority",
                "ARCHITECTURE.md identifies backend/main.py as FastAPI and backend/services/meal_planner.py as deterministic filtering plus final optimization authority.",
                "The backend performs new optimized generation and protects the planning contract.",
            ],
            [
                "Production database",
                "ARCHITECTURE.md and render.yaml indicate production backend persistence is PostgreSQL; SQLite is only local/test fallback.",
                "Say Postgres for production backend persistence. Do not present SQLite as production storage.",
            ],
            [
                "Queue and worker",
                "optimizer_queue_worker.md and render.yaml define Redis queue support and a dedicated worker running backend/worker_plan_jobs.py.",
                "Queued planning supports backend scale and durability; it does not make the mobile client dependent on constant connectivity for saved data.",
            ],
            [
                "Planning contract",
                "production_contract.md defines Stage 1 filtering/ranking/candidate generation and Stage 2 as a formal 0-1 model solved by OR-Tools CP-SAT.",
                "Stage 1 prepares safe candidates; Stage 2 solves the MILP-formulated final weekly assignment.",
            ],
            [
                "ML boundary",
                "ADR-004 allows ML only in Stage 1 ranking; CP-SAT remains final authority.",
                "ML can rank safe candidates but cannot override hard constraints or select unsafe plans.",
            ],
            [
                "No unsafe fallback",
                "ADR-005 requires status=no-safe-plan with diagnostics and guidance when authoritative solve is infeasible.",
                "If constraints conflict, PCOSina explains the blocker instead of returning an unsafe plan.",
            ],
            [
                "Pantry boundary",
                "planner_contract_2026-05-26 states backend planning is pantry-aware, not hard pantry-feasible.",
                "Pantry supports scoring and grocery-gap guidance; it is not proof that every ingredient is physically available.",
            ],
        ],
        [Inches(1.45), Inches(3.25), Inches(2.6)],
    )

    add_heading(doc, "7. Numbers And Terms To Memorize", 1)
    add_table(
        doc,
        ["Evidence point", "Value", "How to say it"],
        [
            ["Architecture type", "Offline-first modular hybrid", "Mobile client preserves saved local data; backend handles new optimized planning."],
            ["Mobile persistence", "DataStore, encrypted shared preferences, local artifacts", "This is the correct Android storage wording."],
            ["Backend runtime", "FastAPI, Python 3.12+", "The backend is the secure control plane and remote optimizer entry point."],
            ["Production services", "Render web service, dedicated worker, Postgres, Redis", "Production shape separates web requests, queue support, worker execution, and database persistence."],
            ["Planner horizon", "7 days / 21 meal slots", "The optimizer selects breakfast, lunch, and dinner assignments across a week."],
            ["MILP title defense", "MILP-formulated 0-1 model; CP-SAT solver implementation", "The architecture supports the title because the model is formulated with binary variables, linear constraints, and objective terms."],
            ["Final benchmark", "100/100 passed, 0 hard-rule violations", "The planner preserved hard constraints in the final benchmark evidence."],
            ["Runtime evidence", "Average 257 ms, P95 313 ms, max 357 ms", "The final benchmark showed sub-second planning performance."],
            ["ML evidence", "NDCG@10 0.9211 vs 0.8838 baseline", "ML improved ranking, but CP-SAT remained authoritative."],
            ["Offline boundary", "Saved data offline; new optimized generation needs backend", "Offline-first means continuity for saved records, not fully offline new optimization."],
        ],
        [Inches(1.65), Inches(2.05), Inches(3.6)],
    )

    doc.add_page_break()
    add_heading(doc, "8. Say This, Not That", 1)
    add_table(
        doc,
        ["Avoid saying", "Say instead", "Why"],
        [
            ["The app is fully offline.", "Saved data works offline; new optimized plan generation may need backend connectivity.", "This matches the actual architecture and avoids overclaiming."],
            ["The phone runs the whole optimizer.", "The mobile app prepares and saves data; the backend optimizer generates new plans.", "The remote optimizer service is the planning authority."],
            ["Android uses SQLite/Room.", "Android uses DataStore, encrypted shared preferences, and local artifacts.", "Current repo explicitly says no Room/SQLite mobile database."],
            ["ML is the planner.", "ML is optional Stage 1 ranking support only.", "CP-SAT remains final authority."],
            ["CP-SAT means the title is not MILP-based.", "MILP-based refers to the formulation; CP-SAT is the solver implementation.", "This separates mathematical model identity from solver technology."],
            ["Cloud sync is required.", "Cloud sync is optional and supplementary.", "Local continuity is the primary user-facing source of truth."],
            ["A fallback returns a weaker plan.", "If no safe complete plan exists, PCOSina returns no-safe-plan guidance.", "Unsafe approximate plans are not accepted as valid output."],
            ["Pantry data guarantees inventory feasibility.", "Pantry is used for planning support and grocery-gap guidance.", "Current implementation is pantry-aware, not full inventory proof."],
        ],
        [Inches(2.05), Inches(3.0), Inches(2.25)],
        header_fill="B94B61",
    )

    add_heading(doc, "9. Panel Q&A", 1)
    qas = [
        (
            "What is the main idea of your system architecture?",
            "It is an offline-first modular hybrid architecture. The mobile app preserves saved user data locally, while the backend handles secure validation and new optimized plan generation.",
        ),
        (
            "Why not run the optimizer only on the phone?",
            "The weekly assignment problem is constraint-heavy and benefits from backend execution, queue control, and centralized recipe/policy data. The phone still keeps saved outputs available offline.",
        ),
        (
            "What exactly works offline?",
            "Saved profiles, pantry records, previously generated meal plans, grocery guidance, progress records, and local reflections. Brand-new optimized plan generation may need backend connectivity.",
        ),
        (
            "What is the role of the backend?",
            "The backend validates requests, applies security and readiness checks, executes deterministic candidate filtering, runs the MILP-formulated CP-SAT optimization, returns plans or no-safe-plan diagnostics, and supports operational records.",
        ),
        (
            "Where is the source of truth?",
            "For user-facing continuity, local app storage is the source of truth. For backend planning records and production operations, Postgres is the production persistence target. Sync is supplementary.",
        ),
        (
            "What does Redis do?",
            "Redis supports queue signaling and runtime backend infrastructure for async planning and rate-limiting paths in the production shape.",
        ),
        (
            "What does PostgreSQL store?",
            "It stores production backend records such as backend data, policy/admin records, planning job records, and related service data used by the backend.",
        ),
        (
            "Where does Firebase fit?",
            "Firebase is optional support for authentication, App Check verification, distribution, and supplementary sync. It is not the final planning authority.",
        ),
        (
            "Where does ML fit in the architecture?",
            "ML sits in Stage 1 as optional smart sorting or ranking support. It can help prioritize safe candidates but cannot override deterministic filters or CP-SAT.",
        ),
        (
            "Why does the title say MILP-based if the architecture uses CP-SAT?",
            "Because MILP-based describes the planning formulation: binary recipe-slot variables, linear constraints, and objective terms. CP-SAT is the implementation solver used by the backend to solve that integer/constraint model, so it should not be described as a pure MILP solver.",
        ),
        (
            "How does the architecture handle unsafe or impossible requests?",
            "Unsafe candidates are removed before optimization. If no complete safe plan can be generated, the backend returns no-safe-plan reasons and safe suggestions instead of forcing an unsafe plan.",
        ),
        (
            "Why separate UI, backend, solver, and data layers?",
            "Separation improves maintainability, testing, privacy boundaries, performance tuning, and defense traceability. It also prevents the UI from being tightly coupled to solver internals.",
        ),
        (
            "What is the limitation of this architecture?",
            "It supports offline access to saved information, not fully offline generation of new optimized plans. It also remains a wellness decision-support system, not a clinical decision engine.",
        ),
    ]
    for q, a in qas:
        add_body(doc, f"Q: {q}", bold_prefix="Q:")
        add_body(doc, f"A: {a}", bold_prefix="A:")

    add_heading(doc, "10. Quick Whiteboard Version", 1)
    add_body(doc, "If asked to explain without the slide, draw five boxes:")
    add_numbered(
        doc,
        [
            "Android client: onboarding, profile, pantry, meal plan, grocery, progress, saved records.",
            "Local persistence: DataStore, encrypted shared preferences, local artifacts, cached plans.",
            "Secure backend: FastAPI validation, Firebase/App Check support, schema/readiness checks, async jobs.",
            "Planner brain: Stage 1 deterministic filtering plus optional ML ranking, then Stage 2 MILP-formulated CP-SAT final assignment.",
            "Data/deployment/evaluation: recipe and price data, Postgres, Redis, Render worker, optional sync, tests and expert feedback.",
        ],
    )

    add_heading(doc, "11. Source And Reference Anchors", 1)
    add_table(
        doc,
        ["Anchor", "Use"],
        [
            ["Manuscript paragraphs 588-615", "Main system architecture explanation and offline-first flow wording."],
            ["Manuscript Figure 8", "Primary system architecture figure."],
            ["Manuscript Figures 9-10", "Offline-first workflow and feature availability map."],
            ["Presentation PDF page 4", "Defense slide for the system architecture."],
            ["ARCHITECTURE.md", "Current implementation truth for mobile storage, backend authority, production persistence, ML guardrails, and deployment shape."],
            ["docs/architecture/production_contract.md", "Offline-first modular hybrid architecture and planning authority."],
            ["docs/architecture/optimizer_queue_worker.md", "Async queue/worker architecture and diagnostics."],
            ["docs/architecture/planner_contract_2026-05-26.md", "Hard/soft/advisory/tracking constraints, MILP-formulated planner boundary, and pantry-aware boundary."],
            ["docs/adr/ADR-004-ml-shadow-canary.md", "ML remains Stage 1 ranking only."],
            ["docs/adr/ADR-005-no-safe-plan-contract.md", "No unsafe fallback output."],
            ["render.yaml", "Production-target shape: Postgres, Redis, FastAPI web service, worker service."],
        ],
        [Inches(2.65), Inches(4.65)],
    )

    add_heading(doc, "12. References To Keep Ready", 1)
    add_bullets(
        doc,
        [
            "Offline-first and local-first software design references in the manuscript: use them to justify saved-data continuity in low connectivity.",
            "Health informatics and privacy references in the manuscript: use them to justify consent, data minimization, and non-diagnostic boundaries.",
            "Optimization, MILP, and OR-Tools references in the manuscript: use them to justify a MILP-formulated backend model solved through CP-SAT planning authority.",
            "ISO/IEC 25010 references in the manuscript: use them to justify maintainability, reliability, security, performance, and portability evaluation.",
            "AI/personalized nutrition references in the manuscript: use them only for bounded Stage 1 ranking support, not autonomous clinical decision-making.",
        ],
    )

    add_heading(doc, "13. Study Checklist", 1)
    add_bullets(
        doc,
        [
            "Memorize the 30-second answer.",
            "Practice the five zones: User Side, Secure System Side, Meal Planning Brain, Data Sources, Output and Improvement.",
            "Practice the offline-first boundary: saved data offline, new optimized generation needs backend.",
            "Memorize correct Android storage wording: DataStore, encrypted shared preferences, local artifacts.",
            "Memorize production wording: Postgres production persistence, Redis queue support, Render backend and worker.",
            "Practice the MILP title boundary: MILP-based refers to formulation; CP-SAT is solver implementation.",
            "Practice the ML boundary: optional Stage 1 ranking only; CP-SAT is final authority.",
            "Practice the no-safe-plan answer.",
            "Avoid saying Room/SQLite for Android, full offline optimization, or ML final authority.",
        ],
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendix A. Manuscript Anchor Summary", 1)
    add_table(
        doc,
        ["Paragraphs", "Key content"],
        [
            ["588-589", "Architecture interaction among Android app, local persistence, remote optimizer, reference data, optional sync, and evaluation; Figure 8."],
            ["591-593", "Two-stage planning: Stage 1 filters unsafe/infeasible recipes; Stage 2 performs MILP-formulated OR-Tools CP-SAT weekly assignment."],
            ["595-606", "Role delineation: expert validators/research evaluators, user with PCOS, admin/operator, and software engineering team/proponents."],
            ["613-615", "Offline-first workflow: saved local data remains available; new optimized generation uses backend; infeasible planning returns guidance."],
            ["649-651", "Offline and online feature map: offline saved data, backend-required new optimization, and optional online services."],
        ],
        [Inches(1.35), Inches(5.95)],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
