from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import re
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "[Manuscript] Quadrant_Revised_Ch 1-6_S2526T3 (3).docx"
OUT_DIR = ROOT / "docs" / "defense"
ASSET_DIR = OUT_DIR / "PCOSINA_Manuscript_All_Figures_Study_Copy_2026-06-09.assets"
TABLES_DOCX = OUT_DIR / "PCOSINA_Manuscript_All_Tables_Study_Copy_2026-06-09.docx"
FIGURES_DOCX = OUT_DIR / "PCOSINA_Manuscript_All_Figures_Study_Copy_2026-06-09.docx"
DOWNLOADS = Path.home() / "Downloads"


TABLE_CAPTION_RE = re.compile(r"^\s*Table\s+\d+\.", re.IGNORECASE)
FIGURE_CAPTION_RE = re.compile(r"^\s*Figure\s+\d+\.", re.IGNORECASE)


def paragraph_text_from_element(element: Any) -> str:
    text = "".join(element.xpath(".//w:t/text()"))
    return " ".join(text.split())


def source_blocks(document: Document) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph_index += 1
            blocks.append(
                {
                    "type": "paragraph",
                    "paragraph_index": paragraph_index,
                    "element": child,
                    "text": paragraph_text_from_element(child),
                    "blips": child.xpath('.//*[local-name()="blip"]'),
                }
            )
        elif child.tag == qn("w:tbl"):
            table_index += 1
            blocks.append(
                {
                    "type": "table",
                    "table_index": table_index,
                    "element": child,
                    "text": "",
                    "blips": child.xpath('.//*[local-name()="blip"]'),
                }
            )
    return blocks


def nearest_nonempty_paragraph(
    blocks: list[dict[str, Any]],
    block_index: int,
    *,
    direction: int,
    limit: int = 8,
) -> tuple[int, str] | None:
    seen = 0
    index = block_index + direction
    while 0 <= index < len(blocks) and seen < limit:
        block = blocks[index]
        if block["type"] == "paragraph" and block.get("text"):
            seen += 1
            return index, str(block["text"])
        index += direction
    return None


def find_caption_for_table(blocks: list[dict[str, Any]], block_index: int) -> str | None:
    previous = nearest_nonempty_paragraph(blocks, block_index, direction=-1)
    if previous and TABLE_CAPTION_RE.match(previous[1]):
        return previous[1]
    following = nearest_nonempty_paragraph(blocks, block_index, direction=1)
    if following and TABLE_CAPTION_RE.match(following[1]):
        return following[1]
    return None


def find_caption_for_figure(blocks: list[dict[str, Any]], block_index: int) -> str | None:
    same_text = str(blocks[block_index].get("text") or "")
    if FIGURE_CAPTION_RE.match(same_text):
        return same_text
    for direction in (1, -1):
        index = block_index + direction
        seen = 0
        while 0 <= index < len(blocks) and seen < 8:
            block = blocks[index]
            if block["type"] == "paragraph" and block.get("text"):
                seen += 1
                text = str(block["text"])
                if FIGURE_CAPTION_RE.match(text):
                    return text
            index += direction
    return None


def visible_context_for_block(blocks: list[dict[str, Any]], block_index: int) -> str:
    same_text = str(blocks[block_index].get("text") or "").strip()
    if same_text:
        return same_text
    previous = nearest_nonempty_paragraph(blocks, block_index, direction=-1)
    following = nearest_nonempty_paragraph(blocks, block_index, direction=1)
    parts = []
    if previous:
        parts.append(f"Previous text: {previous[1]}")
    if following:
        parts.append(f"Next text: {following[1]}")
    return " | ".join(parts) if parts else "No nearby visible text."


def set_landscape(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)


def setup_styles(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(12)


def add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(subtitle)
    run.bold = True
    run.font.size = Pt(11)
    p = document.add_paragraph()
    p.add_run("Source manuscript: ").bold = True
    p.add_run(str(SOURCE.relative_to(ROOT)).replace("\\", "/"))
    p = document.add_paragraph()
    p.add_run("Extraction rule: ").bold = True
    p.add_run(
        "Content was copied from the source manuscript for study organization only. "
        "Table structures are copied from the original Word XML; figure images are copied from the original embedded JPG/PNG media."
    )


def copy_image_relationships(source_document: Document, target_document: Document, cloned_element: Any) -> None:
    for blip in cloned_element.xpath('.//*[local-name()="blip"]'):
        old_rid = blip.get(qn("r:embed"))
        if not old_rid:
            continue
        old_part = source_document.part.related_parts.get(old_rid)
        if old_part is None:
            continue
        new_rid, _ = target_document.part.get_or_add_image(BytesIO(old_part.blob))
        blip.set(qn("r:embed"), new_rid)


def add_list_paragraph(document: Document, label: str, value: str) -> None:
    p = document.add_paragraph()
    p.add_run(label).bold = True
    p.add_run(value)


def build_tables_doc(source_document: Document, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    table_entries: list[dict[str, Any]] = []
    table_doc = Document()
    set_landscape(table_doc)
    setup_styles(table_doc)
    add_title(
        table_doc,
        "PCOSina Manuscript Tables",
        "All Word table objects extracted from the revised Chapter 1-6 manuscript",
    )

    table_blocks = [
        (index, block)
        for index, block in enumerate(blocks)
        if block["type"] == "table"
    ]
    add_list_paragraph(table_doc, "Total manuscript Word tables extracted: ", str(len(table_blocks)))
    table_doc.add_heading("Table Index", level=1)

    for block_index, block in table_blocks:
        caption = find_caption_for_table(blocks, block_index)
        label = caption or f"Uncaptioned Word table object {block['table_index']}"
        context = visible_context_for_block(blocks, block_index)
        entry = {
            "table_index": block["table_index"],
            "label": label,
            "caption": caption,
            "context": context,
        }
        table_entries.append(entry)
        p = table_doc.add_paragraph()
        p.add_run(f"{entry['table_index']}. ").bold = True
        p.add_run(label)
        if not caption:
            p.add_run(f" | Context: {context[:180]}")

    for entry, (block_index, block) in zip(table_entries, table_blocks):
        table_doc.add_page_break()
        table_doc.add_heading(f"Table Object {entry['table_index']}: {entry['label']}", level=1)
        if not entry["caption"]:
            p = table_doc.add_paragraph()
            p.add_run("Nearby context: ").bold = True
            p.add_run(entry["context"])
        cloned_table = deepcopy(block["element"])
        copy_image_relationships(source_document, table_doc, cloned_table)
        table_doc._body._element.append(cloned_table)
        table_doc.add_paragraph()

    table_doc.save(TABLES_DOCX)
    return table_entries


def media_extension(part: Any, fallback_index: int) -> str:
    suffix = Path(str(getattr(part, "partname", ""))).suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg"}:
        return ".jpg" if suffix == ".jpeg" else suffix
    content_type = str(getattr(part, "content_type", "")).lower()
    if "jpeg" in content_type or "jpg" in content_type:
        return ".jpg"
    if "png" in content_type:
        return ".png"
    return f".image{fallback_index}"


def build_figures_doc(source_document: Document, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    for existing in ASSET_DIR.glob("*"):
        if existing.is_file():
            existing.unlink()

    figure_entries: list[dict[str, Any]] = []
    occurrence = 0
    for block_index, block in enumerate(blocks):
        if block["type"] != "paragraph":
            continue
        for blip in block.get("blips") or []:
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            part = source_document.part.related_parts.get(rid)
            if part is None:
                continue
            occurrence += 1
            caption = find_caption_for_figure(blocks, block_index)
            context = visible_context_for_block(blocks, block_index)
            ext = media_extension(part, occurrence)
            image_path = ASSET_DIR / f"embedded_image_{occurrence:03d}{ext}"
            image_path.write_bytes(part.blob)
            label = caption or f"Embedded image occurrence {occurrence}"
            figure_entries.append(
                {
                    "occurrence": occurrence,
                    "label": label,
                    "caption": caption,
                    "context": context,
                    "paragraph_index": block.get("paragraph_index"),
                    "image_path": image_path,
                }
            )

    figure_doc = Document()
    set_landscape(figure_doc)
    setup_styles(figure_doc)
    add_title(
        figure_doc,
        "PCOSina Manuscript Figures and Embedded Images",
        "All embedded image occurrences extracted from the revised Chapter 1-6 manuscript",
    )
    numbered_captions = []
    seen_captions = set()
    for entry in figure_entries:
        caption = entry.get("caption")
        if caption and caption not in seen_captions:
            seen_captions.add(caption)
            numbered_captions.append(caption)
    add_list_paragraph(figure_doc, "Total embedded image occurrences extracted: ", str(len(figure_entries)))
    add_list_paragraph(figure_doc, "Numbered manuscript figure captions detected: ", str(len(numbered_captions)))

    figure_doc.add_heading("Numbered Figure Index", level=1)
    for caption in numbered_captions:
        figure_doc.add_paragraph(caption)

    figure_doc.add_heading("All Embedded Image Occurrences Index", level=1)
    for entry in figure_entries:
        p = figure_doc.add_paragraph()
        p.add_run(f"{entry['occurrence']}. ").bold = True
        p.add_run(str(entry["label"]))
        if not entry["caption"]:
            p.add_run(f" | Context: {entry['context'][:180]}")

    for entry in figure_entries:
        figure_doc.add_page_break()
        figure_doc.add_heading(f"Image Occurrence {entry['occurrence']}: {entry['label']}", level=1)
        p = figure_doc.add_paragraph()
        p.add_run("Source paragraph index: ").bold = True
        p.add_run(str(entry["paragraph_index"]))
        if not entry["caption"]:
            p = figure_doc.add_paragraph()
            p.add_run("Nearby context: ").bold = True
            p.add_run(str(entry["context"]))
        try:
            figure_doc.add_picture(str(entry["image_path"]), width=Inches(9.2))
        except Exception:
            p = figure_doc.add_paragraph()
            p.add_run("Image extraction note: ").bold = True
            p.add_run(f"Saved original media at {entry['image_path']}, but python-docx could not embed this media type.")
        caption = figure_doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(str(entry["label"])).italic = True

    figure_doc.save(FIGURES_DOCX)
    return figure_entries


def copy_to_downloads(paths: list[Path]) -> list[Path]:
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    copied = []
    for path in paths:
        destination = DOWNLOADS / path.name
        shutil.copy2(path, destination)
        copied.append(destination)
    return copied


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    source_document = Document(SOURCE)
    blocks = source_blocks(source_document)
    table_entries = build_tables_doc(source_document, blocks)
    figure_entries = build_figures_doc(source_document, blocks)
    downloads = copy_to_downloads([TABLES_DOCX, FIGURES_DOCX])
    print(f"source={SOURCE}")
    print(f"tables_doc={TABLES_DOCX}")
    print(f"figures_doc={FIGURES_DOCX}")
    print(f"tables_extracted={len(table_entries)}")
    print(f"embedded_images_extracted={len(figure_entries)}")
    print(f"downloads={';'.join(str(path) for path in downloads)}")


if __name__ == "__main__":
    main()
