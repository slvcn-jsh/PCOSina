from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "defense" / "PCOSINA_Conceptual_Framework_Defense_Prep_2026-06-04.docx"
SLIDE_IMAGE = ROOT / "tmp" / "defense_prep" / "presentation_pages_png" / "page_2.png"
MANUSCRIPT_FIGURE = (
    ROOT
    / "docs"
    / "defense"
    / "PCOSINA_SOP_Evidence_Extraction_Dossier_FINAL_2026-05-24.assets"
    / "figure_01.png"
)


ACCENT = "2F7D5A"
DARK = "1F2933"
MUTED = "5B6770"
PALE = "EAF5EF"
PALE_BLUE = "EAF2F8"
PALE_PINK = "FDECEF"
PALE_YELLOW = "FFF7DA"
BORDER = "B7C6BC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=DARK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05


def add_table(doc, headers, rows, widths=None, header_fill=ACCENT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8.8)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = widths[idx]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_text(cell, str(value), color=DARK, size=8.5)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[idx]
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Aptos Display"
        run.font.color.rgb = RGBColor.from_string(ACCENT if level == 1 else DARK)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(DARK)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.color.rgb = RGBColor.from_string(DARK)
    else:
        r = p.add_run(text)
        r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(item)
        run.font.color.rgb = RGBColor.from_string(DARK)


def add_numbered(doc, items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(f"{idx}.    {item}")
        run.font.color.rgb = RGBColor.from_string(DARK)


def add_callout(doc, title, body, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=ACCENT)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    cell.text = ""
    p = cell.paragraphs[0]
    rt = p.add_run(title)
    rt.bold = True
    rt.font.color.rgb = RGBColor.from_string(ACCENT)
    rt.font.size = Pt(10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    rb = p2.add_run(body)
    rb.font.color.rgb = RGBColor.from_string(DARK)
    rb.font.size = Pt(9)
    doc.add_paragraph()


def add_figure(doc, path, caption, width=7.1):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(8)
        cap.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    else:
        add_callout(doc, "Figure not embedded", f"Expected image was not found at {path}.", fill=PALE_PINK)


def configure_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)

    for style_name, size, color in [
        ("Title", 24, ACCENT),
        ("Heading 1", 15, ACCENT),
        ("Heading 2", 12, DARK),
        ("Heading 3", 10.5, DARK),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("PCOSina Conceptual Framework Defense Prep")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(ACCENT)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Focused packet for thesis defense: IPO model, slide guide, manuscript figure, script, evidence, references, and Q&A")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "Defense identity",
        "The conceptual framework explains how PCOSina transforms knowledge resources, system requirements, user inputs, and MILP-formulated optimization logic into evaluated 7-day Filipino-context meal-planning outputs through an IPO model with an evaluation feedback loop.",
        fill=PALE,
    )

    add_table(
        doc,
        ["Item", "Defense-ready answer"],
        [
            ["Main point", "The conceptual framework is the traceability map: inputs enter the system, the local-first client and backend planner process them, outputs are generated, and evaluation feedback improves the UI, rules, constraints, dataset, and deployment."],
            ["Presentation slide", "Slide 2: Conceptual Framework from [Presentation] Quadrant_Presentation_Revised_S2526T3 (1).pdf."],
            ["Manuscript anchor", "Chapter 1, Conceptual Framework section: extracted manuscript paragraphs 245-289 and Figure 1."],
            ["System anchor", "ARCHITECTURE.md, production_contract.md, planner_contract_2026-05-26.md, ADR-004, ADR-005, and backend/services/meal_planner.py."],
            ["MILP title defense", "In the Process block, MILP-based refers to the 0-1 weekly assignment formulation with binary recipe-slot variables, linear constraints, and objective terms; CP-SAT is the implementation solver."],
            ["Boundary", "PCOSina is a wellness decision-support system. It does not diagnose PCOS, prescribe treatment, or replace professional healthcare or nutrition consultation."],
        ],
        [Inches(1.55), Inches(5.75)],
    )

    add_heading(doc, "1. Slide And Manuscript Figure", 1)
    add_figure(doc, SLIDE_IMAGE, "Figure CF-1. Presentation slide used for the Conceptual Framework defense section.", width=7.1)
    add_figure(doc, MANUSCRIPT_FIGURE, "Figure CF-2. Manuscript Figure 1. PCOSina updated Input-Process-Output-Evaluation framework.", width=7.1)
    add_body(
        doc,
        "Figure use: the slide is what you present to the panel; the manuscript figure is your cleaner source figure. The figure connects to the next defense topics, but the conceptual framework should stay focused on the IPO transformation and evaluation loop."
    )

    add_heading(doc, "2. Core Defense Explanation", 1)
    add_callout(
        doc,
        "30-second answer",
        "Our conceptual framework uses an Input-Process-Output model with an evaluation feedback loop. The inputs are nutrition knowledge, Filipino food context, system requirements, and user planning data. The process is the local-first Android client, FastAPI backend control plane, deterministic Stage 1 filtering, Stage 2 MILP-formulated 0-1 weekly assignment solved with OR-Tools CP-SAT, optional non-authoritative ML ranking, and deployment/sync support. The outputs are the 7-day Filipino meal plan, recipe and grocery guidance, nutrition and budget explanations, local cached artifacts, solver metadata, and research evidence. The feedback loop uses testing, ISO/IEC 25010 evaluation, user feedback, and expert review to refine the UI, rules, constraints, dataset, and deployment.",
        fill=PALE_BLUE,
    )

    doc.add_page_break()
    add_heading(doc, "3. IPO Plus Evaluation Map", 1)
    add_table(
        doc,
        ["Framework block", "What it contains", "PCOSina implementation", "Defense sentence"],
        [
            [
                "Input",
                "Knowledge requirements, hardware, software, and user planning inputs.",
                "PCOS nutrition principles, Filipino food context, Android device, Android Studio/Kotlin/Compose, FastAPI, OR-Tools CP-SAT, LightGBM, profile, goals, allergies, restrictions, pantry, budget, recipe, ingredient, price, and feedback records.",
                "Inputs define what the system must know and what the user needs the planner to respect.",
            ],
            [
                "Process",
                "Development workflow and core meal-planning workflow.",
                "Local-first client, backend validation, Firebase Auth/App Check support, Stage 1 deterministic filtering/ranking, Stage 2 MILP-formulated 0-1 assignment solved by CP-SAT, optional ML ranking, and sync/deployment support.",
                "The process is where user data becomes safe candidates and then a MILP-formulated complete weekly assignment.",
            ],
            [
                "Output",
                "User-facing, system, and research outputs.",
                "7-day meal plan, daily meals, recipe details, grocery guidance, nutrition/budget explanation, local cached artifacts, solver metadata, no-safe-plan reason codes, usage evidence, and validation evidence.",
                "Outputs are not only meal plans; they also include explainability, safety diagnostics, and evaluation evidence.",
            ],
            [
                "Evaluation and feedback",
                "Functional, system quality, research, and user evaluation.",
                "Functional testing, ISO/IEC 25010, alpha/beta testing, respondent feedback, CS/IT review, OB-GYN/RND validation, and benchmark evidence.",
                "The feedback loop turns test results and panel/user findings into refinement of the UI, constraints, rules, dataset, and deployment.",
            ],
        ],
        [Inches(1.15), Inches(1.75), Inches(2.65), Inches(1.75)],
    )

    add_heading(doc, "4. Slide Narration Script", 1)
    add_body(doc, "Use this during the presentation. Keep your eyes on the panel and explain the flow left to right.")
    add_numbered(
        doc,
        [
            "Start with the whole idea: This framework uses IPO with evaluation feedback, so it shows how PCOSina transforms resources and user data into evaluated meal-planning outputs.",
            "Explain Input: The inputs include knowledge requirements like PCOS-supportive nutrition and Filipino food context, technical requirements like Android, Kotlin, FastAPI, OR-Tools, and user data like profile, restrictions, allergies, pantry, budget, recipes, prices, and feedback.",
            "Explain Process: The process starts in the local-first Android app, then the FastAPI backend validates the request and runs the planner. Stage 1 removes unsafe or incompatible recipes and scores safe candidates. Stage 2 uses a MILP-formulated 0-1 assignment solved through OR-Tools CP-SAT to select the 7-day plan. ML is optional and only supports ranking.",
            "Explain Output: The outputs include the weekly Filipino meal plan, recipe details, grocery guidance, nutrition and budget explanations, local cached records, solver metadata, infeasibility reason codes, and research validation evidence.",
            "Explain Evaluation and Feedback: The system is evaluated through functional testing, ISO/IEC 25010 quality criteria, alpha/beta testing, respondent feedback, and expert review. The results refine the UI, rules, constraints, dataset, and deployment.",
            "Close with the boundary: The framework supports meal-planning decisions only. PCOSina is not a diagnosis engine or treatment system.",
        ],
    )

    add_heading(doc, "5. Full Speaker Script", 1)
    add_body(
        doc,
        "Good day panel. This slide presents the conceptual framework of PCOSina. We used an Input-Process-Output model with an evaluation feedback loop because the project is not only about creating an app screen. It is about tracing how knowledge, technical requirements, user data, planning logic, and evaluation results work together to produce a safe and practical meal-planning output."
    )
    add_body(
        doc,
        "On the left side are the inputs. These include knowledge requirements such as PCOS-supportive nutrition principles, Filipino food context, allergy and dietary restriction rules, budget, pantry, nutrition, and variety constraints, and the ethical boundary that the system is for wellness support only. The inputs also include hardware and software requirements, such as an Android device, Android Studio, Kotlin, Jetpack Compose, FastAPI, OR-Tools CP-SAT, Firebase support, GitHub, Gradle, and testing tools. Finally, the inputs include the actual planning data: user profile, goals, symptoms, allergies, restrictions, pantry items, weekly budget, recipe data, ingredient data, price data, progress records, and feedback."
    )
    add_body(
        doc,
        "The middle part is the process. The process begins with the local-first Android client, where users complete onboarding, manage their profile, pantry, groceries, meal plan, and progress data. Local storage supports continuity by keeping saved records accessible. The request then goes to the FastAPI backend control plane, where the system validates the request, checks schema versions, verifies security controls when enabled, and runs readiness and safety checks."
    )
    add_body(
        doc,
        "The core planning logic has two stages. Stage 1 is deterministic candidate filtering and construction. It removes unsafe or incompatible recipes, applies allergy and restriction filters, checks meal-slot suitability, and scores candidates based on pantry match, cost, preparation time, and nutrition fit. LightGBM may help rank candidates in this stage, but it is non-authoritative. Stage 2 is the MILP-formulated optimization block: the system represents the weekly plan with binary recipe-slot variables, linear constraints, and objective penalties, then uses OR-Tools CP-SAT to select the final 7-day meal plan while enforcing calories, macros, budget behavior, repetition, variety, and other constraints."
    )
    add_body(
        doc,
        "On the right side are the outputs. User-facing outputs include the personalized 7-day Filipino meal plan, daily meals, recipe details, grocery guidance, and nutrition and budget explanations. System outputs include solver metadata, infeasibility reason codes, suggested relaxations, cached local plan and grocery artifacts, progress logs, and admin or operator data. Research outputs include usability evidence, functional testing results, planner correctness evidence, offline-first behavior evidence, and security or privacy validation evidence."
    )
    add_body(
        doc,
        "The last block is evaluation and feedback. This is important because the framework does not stop after output generation. Functional testing checks whether meal generation, allergy and restriction enforcement, budget and macro constraints, and grocery or pantry outputs work as intended. ISO/IEC 25010 evaluates software quality such as usability, reliability, performance, maintainability, and security. Alpha and beta testing, respondent feedback, and expert review then feed back into improvements for the UI, rules, constraints, dataset, and deployment."
    )
    add_body(
        doc,
        "In summary, the conceptual framework shows the complete transformation: inputs become processed through local-first software, deterministic filtering, MILP-formulated CP-SAT optimization, and bounded ML support; the result is a 7-day Filipino meal-planning output, and evaluation feedback continuously improves the system while keeping the project within its wellness decision-support boundary."
    )

    add_heading(doc, "6. Repo-Truth Cross-Check", 1)
    add_table(
        doc,
        ["Conceptual claim", "Current source evidence", "Defense-safe wording"],
        [
            [
                "Local-first client",
                "ARCHITECTURE.md states Android uses DataStore-backed preferences, encrypted shared preferences, and local artifacts. It explicitly says Android does not currently use Room or SQLite.",
                "Saved data remains available locally through DataStore, encrypted shared preferences, and local artifacts.",
            ],
            [
                "Backend process",
                "ARCHITECTURE.md and production_contract.md describe a FastAPI backend and remote optimizer service.",
                "New optimized plan generation is handled by the backend planner, while the mobile app preserves saved outputs for offline continuity.",
            ],
            [
                "Stage 1 planner",
                "production_contract.md defines Stage 1 as filtering, ranking, and candidate generation.",
                "Stage 1 prepares the safe candidate pool by removing unsafe or incompatible recipes before final optimization.",
            ],
            [
                "Stage 2 planner",
                "production_contract.md defines Stage 2 as a formal 0-1 model solved by OR-Tools CP-SAT.",
                "Stage 2 is the MILP-formulated final authority that assigns recipes to the weekly meal slots through CP-SAT.",
            ],
            [
                "ML boundary",
                "ADR-004 says ML scores are allowed only in Stage 1 candidate ranking and CP-SAT Stage 2 remains final authority.",
                "ML can reorder safe candidates, but it cannot override hard constraints or select the final plan by itself.",
            ],
            [
                "No unsafe fallback",
                "ADR-005 says infeasible solves return a structured status=no-safe-plan response.",
                "If constraints are impossible, PCOSina explains the blockers instead of forcing an unsafe plan.",
            ],
            [
                "Pantry boundary",
                "planner_contract_2026-05-26 states backend planning remains pantry-aware, not hard pantry-feasible.",
                "Pantry supports scoring and grocery-gap guidance; it should not be overclaimed as complete inventory consumption proof.",
            ],
        ],
        [Inches(1.5), Inches(3.1), Inches(2.7)],
    )

    add_heading(doc, "7. Numbers And Data To Memorize", 1)
    add_table(
        doc,
        ["Evidence point", "Value", "How to say it"],
        [
            ["Planning horizon", "7 days / 21 meal slots", "The conceptual output is a weekly Filipino meal plan with breakfast, lunch, and dinner slots."],
            ["MILP title defense", "MILP-formulated 0-1 assignment; CP-SAT solver implementation", "The Process block supports the title because it uses binary variables, linear constraints, and objective terms."],
            ["Minimum Android environment", "Android 8.0 minimum; Android 11+ recommended", "The manuscript and slide describe Android 8.0 as minimum and Android 11 or later as recommended for smoother testing."],
            ["Respondents", "50 total: 25 primary, 25 secondary", "The evaluation feedback loop includes actual user groups, not only developer testing."],
            ["Expert reviewers", "2 CS/IT professionals, 1 OB-GYN, 1 RND", "Technical and health-related validators reviewed system quality and safety boundaries."],
            ["Final benchmark runs", "100/100 passed", "Across the final benchmark, valid profiles produced complete plans."],
            ["Hard-rule violations", "0", "The planner preserved hard constraints across the final benchmark evidence."],
            ["Final runtime", "Average 257 ms, P95 313 ms, max 357 ms", "The final benchmark showed sub-second local planning performance."],
            ["ML ranking evidence", "NDCG@10 0.9211 vs 0.8838 baseline", "ML improved candidate ordering, but CP-SAT remained the final authority."],
            ["ISO/IEC 25010 overall", "4.02 interpreted as Agree", "The manuscript reports positive system quality evaluation across selected ISO/IEC 25010 characteristics."],
        ],
        [Inches(1.65), Inches(1.85), Inches(3.8)],
    )

    add_heading(doc, "8. Say This, Not That", 1)
    add_table(
        doc,
        ["Avoid saying", "Say instead", "Why"],
        [
            ["The framework proves clinical effectiveness.", "The framework supports software-based wellness decision support.", "Clinical outcomes are outside the study scope."],
            ["The app diagnoses PCOS.", "The app supports meal-planning decisions for users managing PCOS-related nutrition needs.", "Diagnosis and treatment are outside the boundary."],
            ["Everything works offline.", "Saved plans and local records are available offline; new optimized generation requires backend connectivity.", "This matches the actual local-first implementation."],
            ["ML personalizes the final plan by itself.", "ML only helps rank already-safe candidates in Stage 1.", "CP-SAT remains the final authority."],
            ["CP-SAT means the study is not MILP-based.", "MILP-based describes the formulation; CP-SAT describes the solver implementation.", "This keeps the title aligned with the actual model and code."],
            ["Pantry guarantees full feasibility.", "The system is pantry-aware and grocery-gap aware.", "Current backend planning does not model complete hard inventory consumption."],
            ["The IPO is just a drawing.", "The IPO is the traceability map from requirements to processing, output, and evaluated refinement.", "This explains why every block exists in the study."],
        ],
        [Inches(2.0), Inches(3.0), Inches(2.3)],
        header_fill="B94B61",
    )

    add_heading(doc, "9. Panel Q&A", 1)
    qas = [
        (
            "What is the conceptual framework in simple terms?",
            "It is the map of how the study works. It shows the inputs we need, the process that transforms those inputs, the outputs produced by the system, and the evaluation feedback used to improve the system.",
        ),
        (
            "Why use an IPO model?",
            "Because PCOSina is a system that transforms data and requirements into outputs. IPO makes that transformation clear: user and technical inputs go through the planner process and become meal plans, explanations, local records, and evaluation evidence.",
        ),
        (
            "What is the difference between the theoretical and conceptual framework?",
            "The theoretical framework explains the theories behind the system, such as optimization, MILP formulation, CP-SAT solving, offline-first design, and rule-based decision support. The conceptual framework explains how those ideas are organized in this actual study from input to process to output and feedback.",
        ),
        (
            "What are the main inputs?",
            "The inputs are knowledge requirements, technical resources, and user planning data. In practical terms, these include PCOS nutrition principles, Filipino food context, Android and backend tools, user profile, goals, allergies, restrictions, pantry items, budget, recipes, prices, progress records, and feedback.",
        ),
        (
            "What happens in the process block?",
            "The Android app collects and stores user data locally, the backend validates the request, Stage 1 filters and ranks safe candidate recipes, Stage 2 solves the MILP-formulated weekly assignment through CP-SAT, and optional sync or deployment components support operation.",
        ),
        (
            "Why mention MILP in the conceptual framework if the implementation uses CP-SAT?",
            "Because the conceptual Process block describes the model formulation: binary recipe-slot variables, linear constraints, and objective terms. CP-SAT is the solver implementation used to solve that integer/constraint model, so the title remains MILP-based in formulation.",
        ),
        (
            "Where does machine learning fit?",
            "ML fits only in Stage 1 as optional assistive ranking. It can help prioritize candidate recipes, but it cannot override allergies, restrictions, budget, nutrition boundaries, or the CP-SAT final decision.",
        ),
        (
            "What outputs does PCOSina produce?",
            "It produces a 7-day Filipino-context meal plan, recipe details, grocery guidance, nutrition and budget explanations, cached local records, solver metadata, no-safe-plan reason codes when needed, and research evidence from testing and evaluation.",
        ),
        (
            "Why include evaluation and feedback in the framework?",
            "Because the study must prove more than output generation. Evaluation checks functional correctness, software quality, usability, reliability, safety boundaries, and expert acceptability, then feeds improvements back into the UI, rules, constraints, dataset, and deployment.",
        ),
        (
            "What is offline-first in this framework?",
            "Offline-first means the app preserves saved profiles, plans, grocery guidance, pantry records, and progress data locally so users can still access them without continuous internet. New optimized plan generation still depends on backend connectivity.",
        ),
        (
            "How does the framework handle infeasibility?",
            "If constraints cannot produce a safe complete plan, the planner returns no-safe-plan diagnostics and suggested safe adjustments. It does not force an unsafe plan just to show an output.",
        ),
        (
            "What is the main limitation of the conceptual framework?",
            "It supports software planning and evaluation. It does not prove long-term clinical improvement, direct symptom reduction, medical treatment effectiveness, or complete pantry inventory accuracy.",
        ),
        (
            "Why emphasize Filipino context?",
            "Because the system is designed for Filipino users. The input block includes Filipino food context, affordability, local ingredients, and cultural food familiarity so generated plans are practical and not only theoretically healthy.",
        ),
    ]
    for q, a in qas:
        add_body(doc, f"Q: {q}", bold_prefix="Q:")
        add_body(doc, f"A: {a}", bold_prefix="A:")

    add_heading(doc, "10. Quick Whiteboard Version", 1)
    add_body(doc, "If asked to explain without slides, draw this left-to-right chain:")
    add_numbered(
        doc,
        [
            "Input: PCOS nutrition knowledge, Filipino food context, technical tools, user profile, allergies, restrictions, pantry, budget, recipe and price data.",
            "Process: local-first Android client, backend validation, Stage 1 deterministic filtering and ranking, Stage 2 MILP-formulated CP-SAT weekly assignment, optional ML support.",
            "Output: 7-day meal plan, recipes, grocery guidance, nutrition/budget explanation, local cached records, solver metadata, no-safe-plan diagnostics, research evidence.",
            "Evaluation: functional tests, ISO/IEC 25010, alpha/beta testing, respondent feedback, CS/IT review, OB-GYN/RND review.",
            "Feedback loop: refine UI, rules, constraints, dataset, planner settings, and deployment.",
        ],
    )

    add_heading(doc, "11. Source And Reference Anchors", 1)
    add_table(
        doc,
        ["Anchor", "Use"],
        [
            ["Manuscript paragraphs 245-289", "Main conceptual framework explanation: IPO model, input, process, output, evaluation, and boundaries."],
            ["Manuscript Figure 1", "Clean figure for the PCOSina Input-Process-Output-Evaluation framework."],
            ["Presentation PDF page 2", "Defense slide for the conceptual framework."],
            ["ARCHITECTURE.md", "Current implementation truth: local persistence, backend authority, Stage 1/Stage 2, ML guardrails."],
            ["docs/architecture/production_contract.md", "Two-stage planner, MILP-formulated CP-SAT authority, and no unsafe output contract."],
            ["docs/architecture/planner_contract_2026-05-26.md", "Hard/soft/advisory categories and pantry-aware boundary."],
            ["docs/adr/ADR-004-ml-shadow-canary.md", "ML remains Stage 1 ranking support only."],
            ["docs/adr/ADR-005-no-safe-plan-contract.md", "No-safe-plan behavior when constraints are infeasible."],
            ["backend/services/meal_planner.py", "Implemented planner process for filtering, ranking, binary variables, constraints, objective terms, and CP-SAT assignment."],
            ["docs/roadmap/progress_ledger.md", "Benchmark, hard-violation, and local verification evidence."],
        ],
        [Inches(2.6), Inches(4.7)],
    )

    add_heading(doc, "12. References To Keep Ready", 1)
    add_bullets(
        doc,
        [
            "Operations Research, MILP, and optimization references in the manuscript: use them to justify the constrained weekly assignment formulation.",
            "Filipino nutrition and local context references in the manuscript: use them to justify why the input block includes Filipino food context, affordability, and cultural fit.",
            "PCOS nutrition guidance in the manuscript: use it only for wellness and nutrition boundaries, not diagnosis or treatment claims.",
            "Personalized nutrition and AI references in the manuscript: use them to explain why ranking support can be useful but must remain bounded by deterministic rules.",
            "ISO/IEC 25010 references in the manuscript: use them to justify the system quality evaluation block.",
        ],
    )

    add_heading(doc, "13. Study Checklist", 1)
    add_bullets(
        doc,
        [
            "Memorize the 30-second answer.",
            "Practice the left-to-right slide narration: input, process, output, evaluation, feedback.",
            "Be able to explain the difference between theoretical framework and conceptual framework.",
            "Memorize the offline-first boundary: saved data offline, new optimized generation needs backend connectivity.",
            "Memorize the MILP title boundary: MILP-based refers to formulation; CP-SAT refers to solver implementation.",
            "Memorize the ML boundary: Stage 1 ranking only; CP-SAT final authority.",
            "Memorize the no-safe-plan answer.",
            "Memorize the key numbers: 7 days, 21 slots, Android 8.0 minimum, Android 11+ recommended, 50 respondents, 100/100 benchmark runs, 0 hard-rule violations, 257 ms average, 4.02 ISO/IEC 25010 overall.",
            "Be ready to state what the conceptual framework does not prove: diagnosis, treatment, clinical outcomes, full pantry-consumption proof, or fully offline new optimization.",
        ],
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendix A. Manuscript Anchor Summary", 1)
    add_table(
        doc,
        ["Paragraphs", "Key content"],
        [
            ["245-247", "Conceptual framework uses IPO with evaluation feedback and integrates operations research, nutrition science, mobile systems/HCI, and assistive ML with rule-based decision support."],
            ["249-258", "Input phase covers knowledge, hardware, software, and user planning inputs."],
            ["260-273", "Process phase covers software development and the two-stage meal-planning workflow: Stage 1 candidate construction and Stage 2 MILP-formulated CP-SAT weekly assignment."],
            ["275-280", "Output phase covers the 7-day Filipino-context plan, recipe/nutrition/grocery guidance, local storage for offline-first access, and non-diagnostic boundary."],
            ["282-289", "Evaluation phase covers ISO/IEC 25010, user and expert review, feedback, and the system boundary as wellness decision support only."],
        ],
        [Inches(1.4), Inches(5.9)],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
