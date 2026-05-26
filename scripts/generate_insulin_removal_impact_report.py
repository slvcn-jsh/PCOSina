from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "defense" / "PCOSINA_Insulin_Removal_System_Impact_Audit_2026-05-23.docx"
TMP = ROOT / "tmp" / "insulin_resistance_impact_audit"
CURRENT_BENCH = TMP / "current_with_insulin.json"
NO_INSULIN_BENCH = TMP / "what_if_no_insulin_field.json"


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def pct(delta: float, base: float) -> str:
    if not base:
        return "n/a"
    return f"{(delta / base) * 100:.1f}%"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.font.size = Pt(9.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        r.font.size = Pt(9)


def set_cell(cell, text: str, *, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True)
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()


def suite_summary_rows() -> tuple[list[list[str]], dict[str, Any], dict[str, Any]]:
    current = load_json(CURRENT_BENCH)
    no_insulin = load_json(NO_INSULIN_BENCH)
    c = current["suite"]
    n = no_insulin["suite"]
    rows = []
    for label, key in [
        ("Total runs", "totalRuns"),
        ("Passed runs", "passedRuns"),
        ("Failed runs", "failedRuns"),
        ("Success rate", "successRate"),
        ("Average runtime", "avgRuntimeMs"),
        ("P95 runtime", "p95RuntimeMs"),
        ("Max runtime", "maxRuntimeMs"),
    ]:
        cv = c[key]
        nv = n[key]
        if isinstance(cv, float):
            if key == "successRate":
                ctext = f"{cv * 100:.1f}%"
                ntext = f"{nv * 100:.1f}%"
                dtext = f"{(nv - cv) * 100:.1f} percentage points"
            else:
                ctext = f"{cv:.1f} ms"
                ntext = f"{nv:.1f} ms"
                dtext = f"{nv - cv:+.1f} ms ({pct(nv - cv, cv)})"
        else:
            ctext = str(cv)
            ntext = str(nv)
            dtext = str(nv - cv) if isinstance(cv, int) else ""
        rows.append([label, ctext, ntext, dtext, "Pass" if n["thresholdStatus"] == "pass" else "Review"])
    rows.extend(
        [
            ["Hard constraint violations", "0", "0", "No change", "Pass"],
            ["Constraint validation OK", "100/100", "100/100", "No change", "Pass"],
            ["Nutrition feasibility OK", "100/100", "100/100", "No change", "Pass"],
            ["Advisory sodium warnings", "12", "38", "+26 warnings", "Review wording and nutrition advisory table"],
            ["ML ranker ready", str(c["mlRankerReady"]), str(n["mlRankerReady"]), "No change", "Pass"],
            ["ML model version", str(c["mlModelVersion"]), str(n["mlModelVersion"]), "No change", "Pass"],
        ]
    )
    return rows, current, no_insulin


def build_doc() -> None:
    bench_rows, current, no_insulin = suite_summary_rows()

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("PCOSINA Insulin-Resistance Removal System Impact Audit")
    r.bold = True
    r.font.size = Pt(18)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Repo scan, current test run, and what-if planner benchmark before changing the system or manuscript")
    sr.font.size = Pt(10)

    add_body(doc, "Date prepared: 2026-05-23")
    add_body(doc, "Scope: This report does not modify application code. It scans where insulin-resistance logic currently exists and compares the current 20-profile planner benchmark against a temporary what-if fixture where insulinResistanceLevel is removed from every profile.")

    add_heading(doc, "Executive Conclusion", 1)
    add_bullets(
        doc,
        [
            "Removing insulin resistance is not only a manuscript edit. The field is active in Android profile capture/storage, backend schema, backend planner formulas, backend tests, benchmark fixtures, and thesis-validation documents.",
            "The core architecture does not change: deterministic filtering, optional LightGBM ranking, and OR-Tools CP-SAT final assignment remain the planning pipeline.",
            "The nutrition target computation does change for profiles that were Moderate or Severe. In the 20-profile benchmark, 8 of 20 profiles, or 40%, had different macro targets when the insulin field was removed.",
            "The fresh what-if benchmark still passed all 100 runs with zero hard-constraint violations, but runtime and advisory nutrition values changed. Therefore, final Chapter 4 algorithmic results should be rerun after the real implementation change.",
            "User survey/Cronbach results do not automatically change unless the survey item set or evaluated app build changes. If the survey included insulin-specific questions or users tested the old insulin-field flow, disclose the evaluated build or rerun user testing on the revised build.",
        ],
    )

    add_heading(doc, "What Was Scanned", 1)
    add_table(
        doc,
        ["Area", "Evidence Found", "Impact"],
        [
            ["Android app main", "UserProfile.kt stores insulinResistanceLevel. UserProfileScreen.kt requires an insulin resistance dropdown in Step 2. UserPreferencesRepository.kt persists insulin keys. UserViewModel.kt updates insulin details.", "UI flow, validation, local persistence, and API payload preparation need cleanup."],
            ["Backend runtime/schema", "domain/models.py defines insulinResistanceLevel. pcosina_contract.json exposes it. policy_config.py has cold-start default Mild. meal_planner.py uses it in macro_ratios(), profile explanations, and nutrition-pressure retry selection.", "Backend contract and solver target computation need a deliberate replacement policy."],
            ["Tests and benchmarks", "test_meal_planner.py has macro ratio tests for Moderate and Severe. planner_realistic_profiles_20.json has 20 insulin values: 11 Mild, 7 Moderate, 1 Severe, 1 None.", "Tests and benchmark fixtures must be updated, then benchmark values must be regenerated."],
            ["ML pipeline", "No direct insulin feature was found in LightGBM feature definitions. However, macro_distance_score depends on macro targets, and those targets are affected by insulin level.", "No direct feature removal appears required, but retraining or at least ML metric refresh is recommended after solver labels change."],
            ["Docs/manuscript support", "Thesis-validation docs, manuscript support scripts, and defense docs contain insulin-resistance language.", "Paper and defense materials need synchronized edits after the system contract is changed."],
        ],
        [1.4, 4.2, 3.2],
    )

    add_heading(doc, "Active Runtime Dependencies", 1)
    add_table(
        doc,
        ["File / Location", "Current Behavior", "Why It Matters If Removed"],
        [
            ["backend/services/meal_planner.py:1030-1036", "macro_ratios(insulin_level) returns Mild/default 25/40/35, Moderate 28/35/37, Severe 30/30/40.", "This directly changes targetProtein, targetCarbs, and targetFats for Moderate/Severe profiles."],
            ["backend/services/meal_planner.py:2416-2428", "Solver target macros are computed from macro ratios, then symptom/goal adjustments are added and clamped.", "Removing insulin requires a new fixed or policy-driven macro rule."],
            ["backend/services/meal_planner.py:1181-1188", "Severe insulin level contributes to high_nutrition_pressure retry ordering.", "Removing the field changes retry ordering for Severe-only cases unless replaced by another condition."],
            ["backend/services/meal_planner.py:462-463", "Plan explanation says insulin resistance changes macro targets.", "This must be removed from user-facing explanation output."],
            ["app/src/main/java/.../UserProfileScreen.kt:249,279,1110", "Step 2 requires insulin selection and labels the dropdown Insulin Resistance.", "The onboarding/profile flow must be redesigned so users are not blocked by a removed field."],
            ["app/src/main/java/.../UserPreferencesRepository.kt:117,364,473,503,575", "The insulin field is persisted and exported in profile maps.", "Migration/backward compatibility is needed for existing local profiles."],
        ],
        [2.8, 3.3, 3.0],
    )

    add_heading(doc, "Formula Impact", 1)
    add_body(doc, "The what-if assumption used here is: remove insulinResistanceLevel from benchmark profiles and allow the current backend default to apply, which is Mild/default macro ratios of 25% protein, 40% carbohydrates, and 35% fat. This is a simulation, not a code change.")
    add_table(
        doc,
        ["Case", "Original Level", "Target kcal", "Current Protein/Carbs/Fats", "No-Insulin Protein/Carbs/Fats", "Delta"],
        [
            ["T02", "Moderate", "1200", "84 / 120 / 49", "75 / 120 / 46", "-9 / 0 / -3"],
            ["T03", "Moderate", "1940", "135 / 159 / 79", "121 / 184 / 75", "-14 / +25 / -4"],
            ["T07", "Moderate", "1702", "119 / 138 / 69", "106 / 160 / 66", "-13 / +22 / -3"],
            ["T09", "Moderate", "1398", "97 / 122 / 57", "87 / 139 / 54", "-10 / +17 / -3"],
            ["T13", "Moderate", "1489", "104 / 120 / 61", "93 / 138 / 57", "-11 / +18 / -4"],
            ["T15", "Severe", "1200", "90 / 120 / 53", "75 / 120 / 46", "-15 / 0 / -7"],
            ["T16", "Moderate", "1963", "137 / 171 / 80", "122 / 196 / 76", "-15 / +25 / -4"],
            ["T20", "Moderate", "1522", "106 / 133 / 62", "95 / 152 / 59", "-11 / +19 / -3"],
        ],
        [0.7, 1.0, 0.8, 1.7, 1.7, 1.2],
    )
    add_body(doc, "Interpretation: calorie targets do not directly depend on insulin level. Macro targets do. Removing the field mostly lowers protein targets and raises carbohydrate targets for former Moderate cases because they fall back to the Mild/default macro split.")

    add_heading(doc, "Test Results", 1)
    add_table(
        doc,
        ["Test / Check", "Command", "Result", "Meaning"],
        [
            ["Backend meal planner tests", "python -m pytest backend/tests/test_meal_planner.py -q", "70 passed", "Current insulin-related planner tests and broader planner behavior pass before any change."],
            ["Backend schema/planner/policy tests", "python -m pytest backend/tests/test_schema_contract.py backend/tests/test_planner_response_contract.py backend/tests/test_planner_input_hardening.py backend/tests/test_policy_config.py -q", "11 passed", "Current contract and policy assumptions pass, but these tests will need updates after field removal."],
            ["Backend ML tests", "python -m pytest backend/tests/test_ml_ranker.py backend/tests/test_ml_dataset_builder.py backend/tests/test_lightgbm_training.py -q", "8 passed", "Current ML ranker/dataset/training tests pass. No direct insulin feature was found, but macro-derived features may shift."],
            ["Android targeted unit tests", "gradlew.bat testDebugUnitTest --tests HealthMetricsTest --tests PlannerProfilePreparationUseCaseTest --tests ProfileConstraintSemanticsTest", "Blocked: JAVA_HOME/java not configured", "Android-side validation still needs to be run in the local Android/Java environment after the change."],
        ],
        [1.5, 3.2, 1.2, 3.0],
    )

    add_heading(doc, "Benchmark Comparison", 1)
    add_body(doc, "Benchmark method: current fixture and temporary no-insulin fixture were both run with 20 profiles x 5 runs = 100 runs, require-ml-ready enabled, LightGBM model active, and fail-on-regression enabled. The temporary fixture removed insulinResistanceLevel from all 20 profiles and did not alter production code.")
    add_table(
        doc,
        ["Metric", "Current With Insulin Field", "What-If No Insulin Field", "Change", "Status"],
        bench_rows,
        [1.8, 1.9, 1.9, 1.8, 1.8],
    )
    add_body(doc, "Interpretation: the what-if run stayed valid and within thresholds, but exact performance and advisory nutrition values changed. Average runtime increased by 7.8 ms (+3.4%), P95 increased by 34.3 ms (+12.0%), max runtime increased by 64 ms (+20.1%), and sodium advisory warnings increased from 12 to 38. These are not catastrophic changes, but they are enough that final performance tables should be regenerated after the real code change.")

    add_heading(doc, "Answer to the Main Defense Question", 1)
    add_table(
        doc,
        ["Question", "Answer"],
        [
            ["Does removing insulin resistance require system changes?", "Yes. It touches Android UI/profile storage, backend schema, backend planner formulas, tests, benchmark fixtures, and documentation."],
            ["Will algorithm speed and percentages change?", "Likely yes in exact values. The what-if benchmark changed average, P95, max runtime, and advisory warning counts. Success rate and hard-rule violations stayed unchanged in this run."],
            ["Will meal-plan outputs change?", "Likely yes for some profiles. Macro targets changed for 8/20 canonical profiles. Solver-selected dishes can vary across valid solutions, so final output examples should be regenerated."],
            ["Will Cronbach alpha/user survey results change?", "Not automatically. Cronbach alpha measures the survey instrument, not the planner formula. It changes only if survey items are removed/rewritten or if you rerun survey data for the revised app."],
            ["Will ML accuracy metrics change?", "Possibly. There is no direct insulin feature, but macro_distance_score and selected-by-solver labels can change after macro target policy changes. Refreshing ML metrics is recommended if the final paper reports ML ranking performance."],
        ],
        [2.1, 6.7],
    )

    add_heading(doc, "Action Plan Before Manuscript Update", 1)
    add_table(
        doc,
        ["Priority", "Action", "Files / Evidence", "Retest Required"],
        [
            ["1", "Decide the replacement nutrition policy: fixed PCOS-supportive macro split, policy-configurable macro split, or weight-management-only macro range.", "backend/services/meal_planner.py, policy_config.py, manuscript Chapter 3 formulas", "Unit tests for macro target computation."],
            ["2", "Remove or neutralize insulinResistanceLevel from Android UI, profile model, persistence, API payload, backend domain model, schema, and cold-start defaults.", "UserProfile.kt, UserProfileScreen.kt, UserPreferencesRepository.kt, UserViewModel.kt, backend/domain/models.py, pcosina_contract.json", "Android unit tests, backend schema tests, migration/backward-compatibility check."],
            ["3", "Update planner explanations so they no longer say insulin resistance changes macro targets.", "backend/services/meal_planner.py profile_rule_summary and explanation payloads", "Planner response contract tests."],
            ["4", "Update benchmark fixtures and rerun 20 profiles x 5 runs with ML ready.", "benchmarks/canonical_scenarios/planner_realistic_profiles_20.json", "Fresh benchmark values for Chapter 4."],
            ["5", "Update ML artifacts only if final labels/feature distributions are materially changed or if ML metrics are reported in the thesis.", "ml/offline_training, docs/ml", "ML dataset build/training/evaluation tests."],
            ["6", "Update manuscript language after system results are refreshed.", "PCOSINA MANUSCRIPT.docx, thesis-validation docs, SOP defense report", "Document review and final table traceability check."],
        ],
        [0.7, 3.6, 2.7, 1.8],
    )

    add_heading(doc, "Recommended Thesis Wording", 1)
    add_body(doc, "Use this after the actual system change and retest:")
    add_bullets(
        doc,
        [
            "PCOSina computes calorie and macronutrient targets from anthropometric data, activity level, wellness goal, and policy-defined nutrition bounds.",
            "The system supports weight-management-oriented and PCOS-supportive meal planning through balanced calories, dietary restrictions, allergy filtering, fiber-aware choices, pantry-aware scoring, and CP-SAT weekly assignment optimization.",
            "The planner does not diagnose, classify, or treat insulin resistance. It provides wellness decision support and should not replace professional medical or nutrition consultation.",
            "Because the profile input contract changed, algorithmic benchmark results were regenerated using the revised fixture and updated system implementation.",
        ],
    )

    add_heading(doc, "Bottom Line", 1)
    add_body(doc, "The removal is manageable, but it is real. It is not a major architecture rewrite, but it changes the input contract and macro-target computation. The current what-if evidence suggests feasibility and hard-rule safety can remain stable, but final benchmark values, nutrition advisory counts, tests, and Chapter 4 tables must be refreshed after the implementation change.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
