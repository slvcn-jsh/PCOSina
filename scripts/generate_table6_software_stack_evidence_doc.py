from __future__ import annotations

import re
import subprocess
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "defense"
ASSET_DIR = OUT_DIR / "PCOSINA_Table6_Software_Stack_Code_Evidence_2026-06-09.assets"
DOCX_PATH = OUT_DIR / "PCOSINA_Table6_Software_Stack_Code_Evidence_2026-06-09.docx"
DOWNLOADS = Path.home() / "Downloads"


EVIDENCE = [
    {
        "software": "Backend dependency proof",
        "title": "Pinned backend software dependencies",
        "path": "backend/requirements.txt",
        "start": 1,
        "end": 15,
        "what": "Shows FastAPI, OR-Tools, Firebase Admin, PostgreSQL driver, Redis, LightGBM, and test dependencies used by the backend.",
        "defense": "This proves the backend stack is declared as actual dependencies, not only mentioned in the manuscript.",
    },
    {
        "software": "Android Studio / Gradle",
        "title": "Android project plugins",
        "path": "app/build.gradle.kts",
        "start": 5,
        "end": 13,
        "what": "Shows the Android application, Kotlin Android, Kotlin Compose, Firebase App Distribution, and Crashlytics Gradle plugins.",
        "defense": "Android Studio uses this Gradle configuration to build and manage the PCOSina Android application.",
    },
    {
        "software": "Android Studio / Android app configuration",
        "title": "Android app SDK and package configuration",
        "path": "app/build.gradle.kts",
        "start": 237,
        "end": 263,
        "what": "Shows compile SDK, application ID, minimum SDK, target SDK, version code, and version name.",
        "defense": "This is concrete evidence that the mobile client is a native Android application.",
    },
    {
        "software": "Kotlin and Jetpack Compose",
        "title": "Kotlin and Compose build configuration",
        "path": "app/build.gradle.kts",
        "start": 379,
        "end": 400,
        "what": "Shows Kotlin JVM target, Compose enabled in build features, and Compose UI/Material dependencies.",
        "defense": "This supports the claim that the mobile app is built in Kotlin using Jetpack Compose.",
    },
    {
        "software": "Jetpack Compose",
        "title": "Compose entry point in MainActivity",
        "path": "app/src/main/java/com/pcosina/app/MainActivity.kt",
        "start": 1,
        "end": 24,
        "what": "Shows setContent and the Compose app entry point.",
        "defense": "This is where Android launches the PCOSina Compose UI.",
    },
    {
        "software": "Jetpack Compose",
        "title": "Actual Compose screen function",
        "path": "app/src/main/java/com/pcosina/app/ui/screens/MealPlanRefinedScreen.kt",
        "start": 211,
        "end": 231,
        "what": "Shows a real @Composable screen for the meal-plan UI.",
        "defense": "Compose is not only installed; it is used directly in the app screens.",
    },
    {
        "software": "Android DataStore",
        "title": "DataStore declaration for local preferences",
        "path": "app/src/main/java/com/pcosina/app/data/repository/UserPreferencesRepository.kt",
        "start": 1,
        "end": 40,
        "what": "Shows DataStore imports and the user_prefs DataStore declaration.",
        "defense": "This is the local-first storage foundation for saved user app data.",
    },
    {
        "software": "Android DataStore",
        "title": "Saving user profile fields locally",
        "path": "app/src/main/java/com/pcosina/app/data/repository/UserPreferencesRepository.kt",
        "start": 455,
        "end": 489,
        "what": "Shows profile fields written to local DataStore preferences: name, age, height, weight, activity, goal, symptoms, restrictions, allergies, pantry, budget, cooking time, variety, and planning priority.",
        "defense": "This proves profile inputs used by planning are stored locally on the Android device.",
    },
    {
        "software": "Android DataStore / Local artifacts",
        "title": "Saving generated plans and groceries locally",
        "path": "app/src/main/java/com/pcosina/app/data/repository/UserPreferencesRepository.kt",
        "start": 1097,
        "end": 1131,
        "what": "Shows saved plan, plan history, active plan, and grocery JSON persistence.",
        "defense": "This supports the offline-first claim for saved plans and grocery data after they already exist locally.",
    },
    {
        "software": "Encrypted Shared Preferences",
        "title": "Encrypted local storage setup",
        "path": "app/src/main/java/com/pcosina/app/data/repository/ReflectionStore.kt",
        "start": 1,
        "end": 40,
        "what": "Shows EncryptedSharedPreferences with AES-based key/value encryption schemes.",
        "defense": "This proves sensitive reflection/progress-style local artifacts are not stored only as plain shared preferences.",
    },
    {
        "software": "Encrypted Shared Preferences",
        "title": "Safe encrypted storage read/write wrappers",
        "path": "app/src/main/java/com/pcosina/app/data/repository/ReflectionStore.kt",
        "start": 68,
        "end": 84,
        "what": "Shows safe read/write wrappers around encrypted local storage.",
        "defense": "This supports reliability and graceful handling of unreadable local encrypted records.",
    },
    {
        "software": "Firebase",
        "title": "Firebase dependencies in Android app",
        "path": "app/build.gradle.kts",
        "start": 409,
        "end": 416,
        "what": "Shows Firebase Analytics, Auth, Firestore, Crashlytics, and App Check dependencies.",
        "defense": "Firebase is used for authentication, optional cloud sync/support services, crash reporting, and App Check protection.",
    },
    {
        "software": "Firebase / Firestore",
        "title": "Profile cloud sync using Firestore",
        "path": "app/src/main/java/com/pcosina/app/data/repository/UserPreferencesRepository.kt",
        "start": 493,
        "end": 528,
        "what": "Shows profile payload fields synced to Firestore when cloud sync is available.",
        "defense": "This proves Firebase is not just a login label; it is also used for optional profile persistence/sync.",
    },
    {
        "software": "Firebase Admin / Backend auth",
        "title": "Backend Firebase initialization and auth guard",
        "path": "backend/main.py",
        "start": 581,
        "end": 627,
        "what": "Shows backend Firebase Admin initialization and authentication dependency.",
        "defense": "This is how the backend verifies authenticated requests when Firebase is configured.",
    },
    {
        "software": "FastAPI",
        "title": "FastAPI app initialization",
        "path": "backend/main.py",
        "start": 1,
        "end": 20,
        "what": "Shows FastAPI imported alongside request/response utilities.",
        "defense": "This proves the backend API layer is built with FastAPI.",
    },
    {
        "software": "FastAPI",
        "title": "FastAPI application object",
        "path": "backend/main.py",
        "start": 349,
        "end": 357,
        "what": "Shows app = FastAPI(...) with metadata and lifespan configuration.",
        "defense": "This is the actual FastAPI application object served by the backend.",
    },
    {
        "software": "FastAPI REST API",
        "title": "GET and POST recipe endpoints",
        "path": "backend/main.py",
        "start": 3942,
        "end": 3984,
        "what": "Shows GET list/status/detail endpoints and POST create/seed recipe endpoints.",
        "defense": "This is concrete evidence of FastAPI route decorators and admin content API operations.",
    },
    {
        "software": "FastAPI REST API",
        "title": "PUT and DELETE recipe endpoints",
        "path": "backend/main.py",
        "start": 3996,
        "end": 4026,
        "what": "Shows PUT update and DELETE recipe endpoints.",
        "defense": "This proves the backend has update/delete API methods, not only a single planner endpoint.",
    },
    {
        "software": "FastAPI planner API",
        "title": "Generate-plan API endpoint",
        "path": "backend/main.py",
        "start": 5416,
        "end": 5436,
        "what": "Shows POST /generate-plan with request model, Firebase auth, App Check, schema version, and planner policy setup.",
        "defense": "This is the backend doorway used by the Android app to request meal-plan generation.",
    },
    {
        "software": "PostgreSQL",
        "title": "PostgreSQL backend mode and connection",
        "path": "backend/database.py",
        "start": 27,
        "end": 35,
        "what": "Shows DATABASE_URL as the backend database connection setting.",
        "defense": "PostgreSQL is selected through a real database URL, not hardcoded text in the manuscript.",
    },
    {
        "software": "PostgreSQL",
        "title": "PostgreSQL connection and production guard",
        "path": "backend/database.py",
        "start": 199,
        "end": 305,
        "what": "Shows postgres mode detection, connection pooling, psycopg connection, and production guard requiring Postgres.",
        "defense": "This proves PostgreSQL is the backend production database path, while SQLite fallback is for non-production/local cases.",
    },
    {
        "software": "OR-Tools CP-SAT",
        "title": "CP-SAT import and binary decision variables",
        "path": "backend/services/meal_planner.py",
        "start": 1,
        "end": 16,
        "what": "Shows OR-Tools CP-SAT imported by the planner.",
        "defense": "This proves the optimization solver is used in the planner module.",
    },
    {
        "software": "OR-Tools CP-SAT",
        "title": "CP-SAT model and x[s,i] binary variables",
        "path": "backend/services/meal_planner.py",
        "start": 2893,
        "end": 2910,
        "what": "Shows CpModel creation, NewBoolVar variables, exactly-one slot assignment, and disallowed meal-type constraints.",
        "defense": "This is the code-level proof of the 0-1 meal assignment formulation.",
    },
    {
        "software": "OR-Tools CP-SAT",
        "title": "Objective and solver execution",
        "path": "backend/services/meal_planner.py",
        "start": 3093,
        "end": 3161,
        "what": "Shows model.Minimize, CpSolver creation, solver parameter setup, and solver.Solve(model).",
        "defense": "This proves the final plan is solved through an optimization model, not manually selected by if-else rules.",
    },
    {
        "software": "LightGBM",
        "title": "LightGBM model loading",
        "path": "backend/services/ml_ranker.py",
        "start": 71,
        "end": 108,
        "what": "Shows import lightgbm as lgb and loading a Booster model file.",
        "defense": "This proves LightGBM is implemented as an optional Stage 1 ranking artifact, when available.",
    },
    {
        "software": "LightGBM",
        "title": "LightGBM scoring functions",
        "path": "backend/services/ml_ranker.py",
        "start": 127,
        "end": 160,
        "what": "Shows score and score_many predictions using feature vectors.",
        "defense": "This is the actual ML scoring logic used for assistive candidate ranking.",
    },
    {
        "software": "Python",
        "title": "Python backend implementation",
        "path": "backend/main.py",
        "start": 1,
        "end": 20,
        "what": "Shows the backend is implemented in Python and imports FastAPI, Firebase Admin, and backend modules.",
        "defense": "Python is the backend language used to implement the API, solver bridge, and backend services.",
    },
    {
        "software": "GitHub Actions",
        "title": "GitHub workflow evidence",
        "path": ".github/workflows/ci.yml",
        "start": 1,
        "end": 45,
        "what": "Shows repository CI workflow configuration.",
        "defense": "This supports the Git/GitHub purpose: version tracking, collaboration, and validation automation.",
    },
]


TABLE6_ROWS = [
    ("Android Studio", "IDE", "Latest", "App development", "Windows/macOS"),
    ("Kotlin", "Mobile Programming Language", "Project-compatible version", "Android Application Logic", "Android"),
    ("Jetpack Compose", "Android UI toolkit", "Project-compatible version", "Declarative mobile user interface", "Android"),
    (
        "Android DataStore / Encrypted Shared Preferences / Local Artifacts",
        "Local persistence technologies",
        "Built-in / library-supported",
        "Offline-first storage of saved profiles, pantry records, generated plans, settings, and progress-related data",
        "Android",
    ),
    ("Firebase", "Optional cloud service", "Optional", "Authentication, distribution, backup, or supplementary synchronization when enabled", "Cloud"),
    ("FastAPI", "Backend web framework", "Project-compatible version", "API layer for plan-generation requests and optimizer service endpoints", "Backend"),
    ("PostgreSQL", "Relational Database", "Project-compatible version", "Backend data storage and persistent server-side records", "Backend / Cloud"),
    ("OR-Tools CP-SAT", "Optimization Solver", "Project-compatible version", "Solving the MILP-style 0-1 weekly assignment formulation", "Backend"),
    ("LightGBM", "Machine-learning ranking library", "Project-compatible version", "Assistive candidate ranking during Stage 1 candidate construction", "Backend"),
    ("Git / GitHub", "Version control and repository platform", "Latest available", "Collaboration, version tracking, issue tracking, and implementation traceability", "Web/Desktop"),
    ("Draw.io", "Diagram Tool", "Web-based", "System modeling", "Web"),
    ("Python", "Backend programming language", "Project-compatible", "FastAPI backend, optimizer service, CP-SAT planning logic, LightGBM-assisted ranking, backend utilities", "Backend"),
]


def safe_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")[:90]


def font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in [
        Path("C:/Windows/Fonts/consola.ttf"),
        Path("C:/Windows/Fonts/cour.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"),
    ]:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def render_lines(title: str, lines: Iterable[str], output_path: Path, *, target_line_count: int | None = None) -> None:
    code_font = font(19)
    title_font = font(17)
    line_height = 28
    padding_x = 24
    padding_y = 22
    max_width = 2050
    line_list = [line.rstrip("\n") for line in lines]
    if target_line_count is not None:
        line_list = line_list[:target_line_count]
    measure_img = Image.new("RGB", (1, 1), "#ffffff")
    measure = ImageDraw.Draw(measure_img)
    width = 1100
    for line in line_list + [title]:
        bbox = measure.textbbox((0, 0), line, font=code_font)
        width = max(width, min(max_width, bbox[2] - bbox[0] + padding_x * 2 + 20))
    height = padding_y * 2 + 34 + max(1, len(line_list)) * line_height
    image = Image.new("RGB", (width, height), "#fbfbfb")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width - 1, height - 1), outline="#d0d7de", width=2)
    draw.rectangle((0, 0, width, 36), fill="#f0f3f6")
    draw.text((padding_x, 8), title, fill="#57606a", font=title_font)
    y = padding_y + 30
    for line in line_list:
        draw.text((padding_x, y), line, fill="#24292f", font=code_font)
        y += line_height
    image.save(output_path)


def render_source_image(item: dict, index: int) -> Path:
    path = ROOT / item["path"]
    raw_lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    start = int(item["start"])
    end = int(item["end"])
    selected = []
    for number in range(start, min(end, len(raw_lines)) + 1):
        selected.append(f"{number:>5}  {raw_lines[number - 1]}")
    output = ASSET_DIR / f"{index:02d}_{safe_name(item['software'] + '_' + item['title'])}.png"
    render_lines(str(item["path"]).replace("\\", "/"), selected, output)
    return output


def render_git_evidence() -> Path:
    commands = [
        ["git", "rev-parse", "--show-toplevel"],
        ["git", "branch", "--show-current"],
        ["git", "remote", "-v"],
    ]
    lines = []
    for command in commands:
        proc = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)
        lines.append(f"> {' '.join(command)}")
        if proc.stdout.strip():
            lines.extend(proc.stdout.strip().splitlines())
        if proc.stderr.strip():
            lines.extend(proc.stderr.strip().splitlines())
        lines.append("")
    output = ASSET_DIR / "99_git_repository_evidence.png"
    render_lines("Git / GitHub repository evidence", lines, output)
    return output


def setup_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9)
    document.styles["Heading 1"].font.name = "Arial"
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.name = "Arial"
    document.styles["Heading 2"].font.size = Pt(12)


def add_table6(document: Document) -> None:
    document.add_heading("Manuscript Table 6 Software Specifications", level=1)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Software", "Description", "Version", "Purpose", "Platform"]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
        for run in table.rows[0].cells[idx].paragraphs[0].runs:
            run.bold = True
    for row in TABLE6_ROWS:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def add_evidence_index(document: Document) -> None:
    document.add_heading("Evidence Index", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Software", "Evidence", "Source", "Defense use"]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
        for run in table.rows[0].cells[idx].paragraphs[0].runs:
            run.bold = True
    for item in EVIDENCE:
        cells = table.add_row().cells
        cells[0].text = item["software"]
        cells[1].text = item["title"]
        cells[2].text = f"{item['path']}:{item['start']}-{item['end']}"
        cells[3].text = item["defense"]
    row = table.add_row().cells
    row[0].text = "Git / GitHub"
    row[1].text = "Repository remote and branch metadata"
    row[2].text = "git command output"
    row[3].text = "Shows the local project is connected to the GitHub repository for version control."
    row = table.add_row().cells
    row[0].text = "Draw.io"
    row[1].text = "Diagram tool clarification"
    row[2].text = "Manuscript/system diagram artifacts"
    row[3].text = "Draw.io is not runtime source code; it is used for diagrams such as architecture, workflow, use case, and sequence figures."


def build_doc() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    for image in ASSET_DIR.glob("*.png"):
        image.unlink()

    for idx, item in enumerate(EVIDENCE, start=1):
        item["image"] = render_source_image(item, idx)
    git_image = render_git_evidence()

    document = Document()
    setup_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PCOSina Table 6 Software Stack Code Evidence")
    run.bold = True
    run.font.size = Pt(20)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Concrete implementation proof for software specifications").bold = True

    p = document.add_paragraph()
    p.add_run("Purpose: ").bold = True
    p.add_run(
        "This document maps each Table 6 software item to concrete source-code or repository evidence. "
        "It is intended for final-defense study and panel questioning."
    )
    p = document.add_paragraph()
    p.add_run("Important scope note: ").bold = True
    p.add_run(
        "Some Table 6 items are runtime technologies with code evidence, such as FastAPI, DataStore, PostgreSQL, CP-SAT, and LightGBM. "
        "Draw.io is a diagramming tool, so its evidence is the manuscript/system diagrams rather than runtime source code."
    )

    add_table6(document)
    add_evidence_index(document)

    for idx, item in enumerate(EVIDENCE, start=1):
        document.add_page_break()
        document.add_heading(f"{idx}. {item['software']} - {item['title']}", level=1)
        p = document.add_paragraph()
        p.add_run("Source: ").bold = True
        p.add_run(f"{item['path']}:{item['start']}-{item['end']}")
        p = document.add_paragraph()
        p.add_run("What it proves: ").bold = True
        p.add_run(item["what"])
        p = document.add_paragraph()
        p.add_run("Defense use: ").bold = True
        p.add_run(item["defense"])
        document.add_picture(str(item["image"]), width=Inches(9.6))

    document.add_page_break()
    document.add_heading("Git / GitHub - Repository Evidence", level=1)
    document.add_paragraph(
        "Git/GitHub is not a runtime framework. The evidence is the repository metadata showing this workspace is a Git repository connected to the GitHub remote."
    )
    document.add_picture(str(git_image), width=Inches(9.6))

    document.add_page_break()
    document.add_heading("Draw.io - Diagram Tool Clarification", level=1)
    document.add_paragraph(
        "Draw.io is used for system modeling and diagrams, not as runtime application code. "
        "The evidence for this item is the manuscript's diagram figures such as conceptual framework, activity diagram, sequence diagram, use case diagram, class diagram, system architecture, and workflow diagrams."
    )
    document.add_paragraph(
        "Defense wording: Draw.io helped prepare visual system-modeling artifacts; it does not execute inside PCOSina."
    )

    document.save(DOCX_PATH)

    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    destination = DOWNLOADS / DOCX_PATH.name
    destination.write_bytes(DOCX_PATH.read_bytes())
    print(DOCX_PATH)
    print(destination)


if __name__ == "__main__":
    build_doc()
