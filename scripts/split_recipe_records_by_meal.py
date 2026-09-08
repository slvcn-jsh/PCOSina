"""Split meal-type recipe-record DOCX volumes into smaller uploadable parts."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = ROOT / "backend"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from docx import Document

from scripts.generate_recipe_philfct_pricing_records import (
    GENERATED_DATE,
    PHILFCT_CACHE_PATH,
    PRICED_DRAFT_PATH,
    RUNTIME_PATH,
    TEMPLATE_PATH,
    add_recipe_record,
    add_table,
    add_title,
    configure_document,
    git_commit_short,
    meal_type,
    philfct_code_index,
    priced_meal_index,
    read_json,
    recipe_evidence_rows,
    recipe_id,
    recipe_name,
    sorted_recipes,
)


OUTPUT_ROOT = ROOT / "docs" / "thesis_validation" / "03_ACTUAL_SYSTEM_DATA_EXPORTS" / "recipe_records"
DEFAULT_RECORDS_PER_PART = 30


def chunked(items: list[dict[str, Any]], size: int) -> list[list[dict[str, Any]]]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def safe_slug(value: str) -> str:
    return "".join(ch.lower() if ch.isalnum() else "_" for ch in value).strip("_")


def part_filename(meal: str, part_number: int, start_record: int, end_record: int) -> str:
    return (
        "PCOSINA_PhilFCT_Pricing_Recipe_Records_"
        f"{meal}_Part_{part_number:02d}_Records_{start_record:03d}-{end_record:03d}.docx"
    )


def split_output_dir(meal: str) -> Path:
    return OUTPUT_ROOT / f"{safe_slug(meal)}_split_90_page_target"


def build_part_doc(
    *,
    meal: str,
    part_number: int,
    total_parts: int,
    part_recipes: list[dict[str, Any]],
    absolute_start: int,
    rows_by_recipe: dict[str, list[dict[str, Any]]],
    git_short: str,
) -> Path:
    output_dir = split_output_dir(meal)
    absolute_end = absolute_start + len(part_recipes) - 1
    output_path = output_dir / part_filename(meal, part_number, absolute_start, absolute_end)
    doc = Document()
    configure_document(doc)
    add_title(
        doc,
        f"PCOSina PhilFCT and Pricing Recipe Records - {meal} Part {part_number:02d} of {total_parts:02d}",
        f"Split {meal} recipe-record packet for easier Google Docs upload.",
        (
            f"Generated: {GENERATED_DATE} | Git commit: {git_short} | "
            f"Template inspected: {TEMPLATE_PATH.relative_to(ROOT)} | "
            f"Original record range: {absolute_start:03d}-{absolute_end:03d}"
        ),
    )
    doc.add_heading("Split Volume Scope", level=1)
    add_table(
        doc,
        ["Question", "Answer"],
        [
            ["Which records are included?", f"{meal} records {absolute_start:03d}-{absolute_end:03d} in the same order as the full {meal} volume."],
            ["Why split this file?", "The original volume is large for Google Docs upload. This part targets about 90 rendered pages, but exact page count depends on Word/Google Docs pagination."],
            ["Was any recipe data changed?", "No. Recipe sections are regenerated from the same source data and generator helpers as the full record packet."],
            ["What should reviewers edit?", "Use the red review tables: Reviewer, Date checked, and Correction notes. Edit data cells only when correcting a verified source value."],
        ],
        widths=[2.4, 7.2],
        font_size=7.6,
        header_fill="5B9BD5",
    )

    for offset, recipe in enumerate(part_recipes):
        absolute_record = absolute_start + offset
        add_recipe_record(
            doc,
            recipe,
            rows_by_recipe.get(recipe_id(recipe), []),
            absolute_record,
            first=offset == 0,
        )
    doc.save(output_path)
    return output_path


def split_meal(meal: str, records_per_part: int) -> dict[str, Any]:
    output_dir = split_output_dir(meal)
    output_dir.mkdir(parents=True, exist_ok=True)
    for stale_docx in output_dir.glob(f"PCOSINA_PhilFCT_Pricing_Recipe_Records_{meal}_Part_*.docx"):
        stale_docx.unlink()

    index_csv = output_dir / f"pcosina_{safe_slug(meal)}_split_record_index.csv"
    index_json = output_dir / f"pcosina_{safe_slug(meal)}_split_record_index.json"
    for stale_index in (index_csv, index_json):
        if stale_index.exists():
            stale_index.unlink()

    recipes = read_json(RUNTIME_PATH)
    priced = read_json(PRICED_DRAFT_PATH)
    philfct = read_json(PHILFCT_CACHE_PATH)
    priced_by_recipe = priced_meal_index(priced)
    philfct_by_code = philfct_code_index(philfct)

    meal_recipes = [recipe for recipe in sorted_recipes(recipes) if meal_type(recipe) == meal]
    rows_by_recipe = {
        recipe_id(recipe): recipe_evidence_rows(recipe, priced_by_recipe, philfct_by_code)
        for recipe in meal_recipes
    }
    parts = chunked(meal_recipes, records_per_part)
    git_short = git_commit_short()

    output_records: list[dict[str, Any]] = []
    output_files: list[Path] = []
    for part_number, part_recipes in enumerate(parts, start=1):
        absolute_start = ((part_number - 1) * records_per_part) + 1
        output_path = build_part_doc(
            meal=meal,
            part_number=part_number,
            total_parts=len(parts),
            part_recipes=part_recipes,
            absolute_start=absolute_start,
            rows_by_recipe=rows_by_recipe,
            git_short=git_short,
        )
        output_files.append(output_path)
        for offset, recipe in enumerate(part_recipes):
            record_number = absolute_start + offset
            output_records.append(
                {
                    "recordNumber": record_number,
                    "recipeId": recipe_id(recipe),
                    "recipeName": recipe_name(recipe),
                    "mealType": meal_type(recipe),
                    "partNumber": part_number,
                    "partFile": output_path.name,
                }
            )

    with index_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(output_records[0].keys()))
        writer.writeheader()
        writer.writerows(output_records)
    index_json.write_text(
        json.dumps(
            {
                "generatedDate": GENERATED_DATE,
                "gitCommitShort": git_short,
                "recordsPerPart": records_per_part,
                "mealType": meal,
                "recordCount": len(meal_recipes),
                "partCount": len(parts),
                "records": output_records,
            },
            indent=2,
            ensure_ascii=True,
        ),
        encoding="utf-8",
    )

    return {
        "mealType": meal,
        "outputDir": str(output_dir),
        "partCount": len(parts),
        "recordCount": len(meal_recipes),
        "recordsPerPart": records_per_part,
        "files": [path.name for path in output_files],
        "indexCsv": str(index_csv),
        "indexJson": str(index_json),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Split recipe-record DOCX volumes by meal type.")
    parser.add_argument("--meal", action="append", choices=["Breakfast", "Lunch", "Dinner", "Universal"], required=True)
    parser.add_argument("--records-per-part", type=int, default=DEFAULT_RECORDS_PER_PART)
    args = parser.parse_args()
    if args.records_per_part <= 0:
        raise ValueError("--records-per-part must be positive")
    results = [split_meal(meal, args.records_per_part) for meal in args.meal]
    print(json.dumps({"results": results}, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
