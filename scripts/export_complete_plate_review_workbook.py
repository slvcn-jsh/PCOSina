"""Export the v2 complete-plate catalog into a detailed Excel workbook."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CATALOG = REPO_ROOT / "backend" / "seed_data" / "pcosina_philfct_complete_plate_catalog_v2_draft.json"
DEFAULT_OUTPUT = (
    REPO_ROOT
    / "docs"
    / "thesis_validation"
    / "03_ACTUAL_SYSTEM_DATA_EXPORTS"
    / "pcosina_philfct_complete_plate_catalog_v2_detailed_review.xlsx"
)


def _join(values) -> str:
    if isinstance(values, dict):
        return " | ".join(f"{key}: {', '.join(map(str, value)) if isinstance(value, list) else value}" for key, value in values.items())
    return " | ".join(str(value) for value in values or [])


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 70))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(max_len + 2, 70))


def _append_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_sheet(ws)


def main() -> int:
    catalog = json.loads(DEFAULT_CATALOG.read_text(encoding="utf-8"))
    meals = catalog["meals"]

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    summary_rows = []
    source_ingredient_rows = []
    adapted_ingredient_rows = []
    procedure_rows = []
    review_rows = []

    for meal in meals:
        nutrition = meal.get("nutrition") or {}
        summary_rows.append([
            meal.get("id"),
            meal.get("sourceRecipeId") or "",
            meal.get("name"),
            meal.get("baseRecipeName"),
            meal.get("mealType"),
            nutrition.get("calories"),
            nutrition.get("protein_g"),
            nutrition.get("carbs_g"),
            nutrition.get("fat_g"),
            nutrition.get("fiber_g"),
            _join(meal.get("mainComponents")),
            _join(meal.get("pinggangPinoyGroups")),
            meal.get("reviewStatus"),
            _join(meal.get("reviewFlags")),
        ])

        for idx, ingredient in enumerate(meal.get("sourceIngredients") or [], start=1):
            source_ingredient_rows.append([
                meal.get("id"),
                meal.get("name"),
                idx,
                ingredient,
                "Original source recipe ingredient. Needs gram conversion and PhilFCT match before final approval.",
            ])

        for idx, ingredient in enumerate(meal.get("adaptedIngredients") or [], start=1):
            per_100g = ingredient.get("per_100g") or {}
            computed = ingredient.get("computed") or {}
            adapted_ingredient_rows.append([
                meal.get("id"),
                meal.get("name"),
                idx,
                ingredient.get("name"),
                ingredient.get("portion_g"),
                ingredient.get("philfct_name"),
                ingredient.get("philfct_code") or "",
                per_100g.get("calories"),
                per_100g.get("protein_g"),
                per_100g.get("carbs_g"),
                per_100g.get("fat_g"),
                per_100g.get("fiber_g"),
                computed.get("calories"),
                computed.get("protein_g"),
                computed.get("carbs_g"),
                computed.get("fat_g"),
                computed.get("fiber_g"),
                ingredient.get("match_status"),
            ])

        for idx, step in enumerate(meal.get("sourceInstructions") or [], start=1):
            procedure_rows.append([
                meal.get("id"),
                meal.get("name"),
                idx,
                step,
            ])

        review_rows.append([
            meal.get("id"),
            meal.get("name"),
            meal.get("reviewStatus"),
            _join(meal.get("reviewFlags")),
            meal.get("nutritionProvenance", {}).get("baseDish", {}).get("source"),
            meal.get("nutritionProvenance", {}).get("baseDish", {}).get("confidence"),
            meal.get("nutritionProvenance", {}).get("baseDish", {}).get("reviewStatus"),
            meal.get("nutritionProvenance", {}).get("baseDish", {}).get("notes"),
        ])

    _append_sheet(
        wb,
        "Meal Summary",
        [
            "id", "sourceRecipeId", "name", "baseRecipeName", "mealType",
            "calories", "protein_g", "carbs_g", "fat_g", "fiber_g",
            "mainComponents", "pinggangPinoyGroups", "reviewStatus", "reviewFlags",
        ],
        summary_rows,
    )
    _append_sheet(
        wb,
        "Source Ingredients",
        ["mealId", "mealName", "ingredientOrder", "sourceIngredient", "remarks"],
        source_ingredient_rows,
    )
    _append_sheet(
        wb,
        "PhilFCT Added Components",
        [
            "mealId", "mealName", "ingredientOrder", "ingredient", "portion_g",
            "philfct_name", "philfct_code", "kcal_100g", "protein_100g",
            "carbs_100g", "fat_100g", "fiber_100g", "computed_kcal",
            "computed_protein", "computed_carbs", "computed_fat",
            "computed_fiber", "matchStatus",
        ],
        adapted_ingredient_rows,
    )
    _append_sheet(
        wb,
        "Procedures",
        ["mealId", "mealName", "stepNumber", "instruction"],
        procedure_rows,
    )
    _append_sheet(
        wb,
        "Review Notes",
        [
            "mealId", "mealName", "reviewStatus", "reviewFlags",
            "baseNutritionSource", "baseNutritionConfidence",
            "baseNutritionReviewStatus", "baseNutritionNotes",
        ],
        review_rows,
    )

    DEFAULT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DEFAULT_OUTPUT)

    check = load_workbook(DEFAULT_OUTPUT, read_only=True)
    print(DEFAULT_OUTPUT.resolve())
    print(f"sheets={check.sheetnames}")
    print(f"meals={len(meals)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
